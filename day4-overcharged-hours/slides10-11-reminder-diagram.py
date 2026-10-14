"""Day 4 deck, Overcharged hours Branch B.

Slide 10: drop the reminder-triggers sub-flow and the Penalty check box, move the
Timed Reminder format card under the "Request docs from CP" branch, and update
the recommended macro tile.
Slide 11: rebuild as a decision diagram (Reminder triggers -> Yes/No) that keeps
the four existing tiles, plus a recommended macros tile.

Usage: python3 slides10-11-reminder-diagram.py in.pptx out.pptx
"""
import copy, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)
s10, s11 = prs.slides[9], prs.slides[10]


def by_id(slide):
    return {sh.shape_id: sh for sh in slide.shapes}


def geom(sh, x, y, w, h):
    sh.left, sh.top, sh.width, sh.height = Pt(x), Pt(y), Pt(w), Pt(h)


def strip_style(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


S10 = by_id(s10)
assert S10[301].text_frame.text.startswith('Branch B: no clear evidence')
S11 = by_id(s11)
assert S11[343].text_frame.text.startswith('Branch B: when the reminder')

# ---------- slide 10 ----------
line_tpl = copy.deepcopy(S10[318]._element)  # grey connector with arrow head
for sid in range(314, 329):  # reminder node, yes/no lines + labels, outcomes, penalty box
    el = S10[sid]._element
    el.getparent().remove(el)

# connector from "Request docs" down into the timed reminder card
geom(S10[313], 447, 213, 0, 15)
ln = S10[313]._element.find('.//' + qn('a:ln'))
te = ln.find(qn('a:tailEnd'))
te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')

dx, dy = 322 - 32, 228 - 225
for sid in (329, 330, 331, 332):
    sh = S10[sid]
    sh.left, sh.top = sh.left + Pt(dx), sh.top + Pt(dy)
S10[329].width, S10[329].height = Pt(250), Pt(82)
S10[330].width = Pt(150)
S10[331].left = Pt(322 + 250 - 92)
S10[332].width = Pt(228)

# keep the card visible after its click-in: drop the click-2 exit group
timing = s10._element.find(qn('p:timing'))
if timing is not None:
    for ctn in timing.iter(qn('p:cTn')):
        if ctn.get('presetClass') == 'exit' and ctn.get('nodeType') == 'clickEffect':
            outer = ctn.getparent().getparent().getparent().getparent().getparent()
            outer.getparent().remove(outer)
            break

# recommended macro tile: three lines
tile, icon, txt = S10[333], S10[334], S10[335]
geom(tile, 32, 318, 262, 52)
geom(icon, 41, 325, 12, 12)
geom(txt, 58, 323, 230, 42)
paras = txt.text_frame._txBody.findall(qn('a:p'))
p1, p2 = paras[0], paras[1]
r_label, r_name = p1.findall(qn('a:r'))
r_label.find(qn('a:t')).text = 'Recommended macro:'
p1.remove(r_name)
new_p = copy.deepcopy(p1)
nr = new_p.findall(qn('a:r'))[0]
nr.find(qn('a:t')).text = 'Overcharged hrs (investigation needed)'
nr.find(qn('a:rPr')).find(qn('a:solidFill')).find(qn('a:srgbClr')).set('val', '1D1D1F')
p1.addnext(new_p)
for extra in paras[2:]:
    if not ''.join(t.text or '' for t in extra.iter(qn('a:t'))).strip():
        extra.getparent().remove(extra)

# ---------- slide 11 ----------
ORANGE, GREY, RED = RGBColor(0xE6, 0x91, 0x38), RGBColor(0x66, 0x66, 0x66), RGBColor(0xE0, 0x66, 0x66)


def box(slide, x, y, w, h, fill, lines, size=10):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    strip_style(sh)
    sh.adjustments[0] = 0.18
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ('margin_left', 'margin_right'):
        setattr(tf, m, Pt(6))
    tf.margin_top = tf.margin_bottom = Pt(2)
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = t
        f = r.font; f.size = Pt(size); f.bold = True; f.name = 'Play'
        f.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    return sh


def line(slide, x, y, w, h, arrow):
    el = copy.deepcopy(line_tpl)
    slide.shapes._spTree.append(el)
    sh = slide.shapes[-1]
    sh._element.find(qn('p:nvCxnSpPr')).find(qn('p:cNvPr')).set('id', str(slide.shapes._next_shape_id))
    sh._element.find(qn('p:nvCxnSpPr')).find(qn('p:cNvPr')).set('name', 'diagram line')
    geom(sh, x, y, w, h)
    if not arrow:
        sh._element.find('.//' + qn('a:tailEnd')).set('type', 'none')
    return sh


def label(slide, x, y, w, text):
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(14))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(10); r.font.name = 'Play'; r.font.color.rgb = RGBColor(0x1D, 0x1D, 0x1F)
    return tb


def set_size(sh, pt):
    for r in sh._element.iter(qn('a:rPr')):
        r.set('sz', str(int(pt * 100)))
    for r in sh._element.iter(qn('a:endParaRPr')):
        r.set('sz', str(int(pt * 100)))


def card(ids, x, y, w, h, border=None, title_pt=11, body_pt=9):
    bg, badge, badge_inner, title, body = (S11[i] for i in ids)
    geom(bg, x, y, w, h)
    geom(badge, x + 10, y + 9, 22, 22)
    if badge_inner.shape_type == 13:  # picture icon inside the badge
        geom(badge_inner, x + 14.5, y + 13.5, 13, 13)
    else:
        geom(badge_inner, x + 10, y + 9, 22, 22)
    geom(title, x + 40, y + 9, w - 50, 22)
    title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_size(title, title_pt)
    geom(body, x + 10, y + 38, w - 20, h - 46)
    set_size(body, body_pt)
    if border is not None:
        bl = bg._element.find(qn('p:spPr')).find(qn('a:ln'))
        bl.find(qn('a:solidFill')).find(qn('a:srgbClr')).set('val', border)
        bl.set('w', '15875')


# left column: the two questions
card((345, 346, 346, 348, 349), 32, 104, 236, 66)
geom(S11[347], 42, 113, 22, 22)
card((350, 351, 351, 353, 354), 32, 178, 236, 66)
geom(S11[352], 42, 187, 22, 22)

# diagram on the right
CX, LX, RX = 486, 381, 591
box(s11, CX - 115, 104, 230, 42, ORANGE, ['Reminder triggers', 'CP provided valid proof?'], 10.5)
line(s11, CX, 146, 0, 16, False)
line(s11, LX, 162, RX - LX, 0, False)
line(s11, LX, 162, 0, 18, True)
line(s11, RX, 162, 0, 18, True)
label(s11, CX - 46, 147, 40, 'Yes')
label(s11, CX + 6, 147, 40, 'No')
box(s11, LX - 92, 182, 184, 32, GREY, ['No refund'])
box(s11, RX - 92, 182, 184, 32, RED, ['Refund, then penalty check'])
card((355, 356, 357, 358, 359), LX - 92, 222, 184, 96, border='666666', title_pt=10.5, body_pt=9.5)
card((360, 361, 362, 363, 364), RX - 92, 222, 184, 96, border='E06666', title_pt=10.5, body_pt=9.5)

# recommended macros tile, styled like slide 10's
mt = copy.deepcopy(tile._element); mi = copy.deepcopy(icon._element); mx = copy.deepcopy(txt._element)
img_part = s10.part.related_part(icon._element.find('.//' + qn('a:blip')).get(qn('r:embed')))
rid = s11.part.relate_to(img_part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image')
mi.find('.//' + qn('a:blip')).set(qn('r:embed'), rid)
for el in (mt, mi, mx):
    s11.shapes._spTree.append(el)
new = list(s11.shapes)[-3:]
for sh in new:
    cnv = sh._element.find('.//' + qn('p:cNvPr'))
    cnv.set('id', str(s11.shapes._next_shape_id)); cnv.set('name', 'macros tile')
mtile, micon, mtxt = new
geom(mtile, 32, 254, 236, 64)
mtile.adjustments[0] = 0.14
geom(micon, 41, 262, 12, 12)
geom(mtxt, 58, 260, 204, 54)
mtxt.text_frame.vertical_anchor = MSO_ANCHOR.TOP
ps = mtxt.text_frame._txBody.findall(qn('a:p'))
ps[0].findall(qn('a:r'))[0].find(qn('a:t')).text = 'Recommended macros:'
ps[1].findall(qn('a:r'))[0].find(qn('a:t')).text = 'Overcharged hrs (valid claim, refund)'
p3 = copy.deepcopy(ps[1])
p3.findall(qn('a:r'))[0].find(qn('a:t')).text = 'Overcharged hrs (valid proof from CP, no refund)'
ps[1].addnext(p3)
for p in (ps[1], p3):
    p.find(qn('a:pPr')).find(qn('a:spcBef')).find(qn('a:spcPts')).set('val', '300')
for p in ps[2:]:
    p.getparent().remove(p)

prs.save(dst)
print('saved', dst)
