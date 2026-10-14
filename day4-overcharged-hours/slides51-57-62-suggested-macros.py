"""Day 4 deck: add suggested macros to slides 51, 57 and 62.

51 (Cleaner didn't show / Changing CP status): macro tile under the Never bar.
57 (Cleaner cancellation / handling process): extra line in the bottom tile.
62 (Unauthorized reschedule / handling guide): shorten the outcome boxes and
    add a full-width macro tile with both samples.
The tile is cloned from slide 10's recommended-macro tile so it matches the deck.

Usage: python3 slides51-57-62-suggested-macros.py in.pptx out.pptx
"""
import copy, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn

IMG = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'
prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''

s10 = prs.slides[9]
t10 = list(s10.shapes)
k = next(i for i, sh in enumerate(t10) if txt(sh).startswith('Recommended macro'))
TILE, ICON, TEXT = t10[k - 2], t10[k - 1], t10[k]
assert ICON.shape_type == 13
img_part = s10.part.related_part(ICON._element.find('.//' + qn('a:blip')).get(qn('r:embed')))
label_rpr = copy.deepcopy(TEXT._element.findall('.//' + qn('a:r'))[0].find(qn('a:rPr')))   # amber bold
name_rpr = copy.deepcopy(TEXT._element.findall('.//' + qn('a:r'))[1].find(qn('a:rPr')))    # ink bold
for r in (label_rpr, name_rpr):
    for h in r.findall(qn('a:hlinkClick')):
        r.remove(h)


def geom(sh, x, y, w, h):
    sh.left, sh.top, sh.width, sh.height = Pt(x), Pt(y), Pt(w), Pt(h)


def run(p, text, rpr, size=None):
    r = p.makeelement(qn('a:r'), {})
    rp = copy.deepcopy(rpr)
    if size:
        rp.set('sz', str(int(size * 100)))
    r.append(rp)
    t = r.makeelement(qn('a:t'), {}); t.text = text; r.append(t)
    end = p.find(qn('a:endParaRPr'))
    (end.addprevious(r) if end is not None else p.append(r))


def add_tile(slide, x, y, w, h, lines, size=9.5):
    """lines: list of [(text, 'label'|'name'|'plain')]."""
    els = [copy.deepcopy(e._element) for e in (TILE, ICON, TEXT)]
    rid = slide.part.relate_to(img_part, IMG)
    els[1].find('.//' + qn('a:blip')).set(qn('r:embed'), rid)
    for el in els:
        slide.shapes._spTree.append(el)
    tile, icon, tb = list(slide.shapes)[-3:]
    for sh in (tile, icon, tb):
        c = sh._element.find('.//' + qn('p:cNvPr'))
        c.set('id', str(slide.shapes._next_shape_id)); c.set('name', 'suggested macro')
    geom(tile, x, y, w, h)
    geom(icon, x + 9, y + (h - 12) / 2 if len(lines) == 1 else y + 8, 12, 12)
    geom(tb, x + 26, y + 4, w - 34, h - 8)
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    body = tb.text_frame._txBody
    ps = body.findall(qn('a:p'))
    tmpl = copy.deepcopy(ps[0])
    for p in ps:
        body.remove(p)
    plain = copy.deepcopy(name_rpr); plain.set('b', '0')
    plain.find('.//' + qn('a:srgbClr')).set('val', '6E6E73')
    for parts in lines:
        p = copy.deepcopy(tmpl)
        for r in p.findall(qn('a:r')):
            p.remove(r)
        for text, kind in parts:
            run(p, text, {'label': label_rpr, 'name': name_rpr, 'plain': plain}[kind], size)
        body.append(p)
    return tile


# ---- slide 51 ----
s51 = prs.slides[50]
assert txt(s51.shapes[1]) == 'Changing CP status'
add_tile(s51, 32, 324, 250, 26, [[('Suggested macro:  ', 'label'), ('No Show', 'name')]])

# ---- slide 57 ----
s57 = prs.slides[56]
assert txt(s57.shapes[1]) == 'The handling process'
last = next(sh for sh in s57.shapes if txt(sh).startswith('See the quick reference'))
ps = last.text_frame._txBody.findall(qn('a:p'))
assert ps[-1].findall(qn('a:r'))[-1].find(qn('a:t')).text.strip().startswith('Send comms')
newp = copy.deepcopy(ps[-1])
base = newp.findall(qn('a:r'))[0].find(qn('a:rPr'))
for r in newp.findall(qn('a:r')):
    newp.remove(r)
lab = copy.deepcopy(base); lab.set('b', '1')
f = lab.find(qn('a:solidFill'))
if f is not None:
    lab.remove(f)
f = lab.makeelement(qn('a:solidFill'), {}); c = f.makeelement(qn('a:srgbClr'), {'val': '9A6410'}); f.append(c)
lab.insert(0, f)
nm = copy.deepcopy(base); nm.set('b', '1')
run(newp, 'Suggested macro: ', lab)
run(newp, 'CP Cancel → Job Cancelled', nm)
ps[-1].addnext(newp)

# ---- slide 62 ----
s62 = prs.slides[61]
assert txt(s62.shapes[1]) == 'Handling guide'
for sh in s62.shapes:
    if sh.top == Pt(263) or abs(sh.top - Pt(263)) < Pt(1):
        if sh.height > Pt(80):
            sh.height = Pt(56)
add_tile(s62, 32, 327, 655, 44, [
    [('Suggested macros:', 'label')],
    [('Unapproved reschedule (sample 1)', 'name'), ('  →  customer-facing, includes a $20 credit', 'plain')],
    [('Unapproved reschedule (sample 2)', 'name'),
     ('  →  customer-facing, no credit: apology + coordination note only', 'plain')],
], size=9)

prs.save(sys.argv[2])
print('saved', sys.argv[2])
