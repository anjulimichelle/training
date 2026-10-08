"""Platform & Operating Mechanics, slide 42 (Max CP Rate: What the cap does): replace the three
legacy-CRM screenshots with a linked video thumbnail, update the steps to the new CRM flow shown in
the video, and reword what the cap does. Patches the live Google export so nothing else changes."""
import copy, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR

src, dst, thumb, url = sys.argv[1:5]
prs = Presentation(src)
s = prs.slides[41]
sh = list(s.shapes)
assert sh[1].text_frame.text == 'What the cap does' and sh[7].shape_type == 13

def set_para(shape, i, text):
    p = shape.text_frame.paragraphs[i]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]: r._r.getparent().remove(r._r)

set_para(sh[2], 0, 'Set in the CRM: Customer Information › Max rate › Set.')
set_para(sh[5], 0, 'Under Customer Information, find Max rate and click Set')
set_para(sh[13], 0, 'Enter the preferred hourly rate and click Save')
set_para(sh[21], 0, 'Max rate now shows the amount with a Hard cap tag')
set_para(sh[27], 1, 'Cleaning requests only go to cleaners whose rates don’t exceed the cap.')

# One wide panel (same style as the old screenshot boxes) holding the video.
panel = copy.deepcopy(sh[6]._element)
sh[6]._element.addnext(panel)
for i in (6, 7, 8, 9, 10, 14, 15, 16, 17, 18, 22, 23):
    sh[i]._element.getparent().remove(sh[i]._element)
off = panel.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}off')
ext = panel.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}ext')
off.set('x', str(Inches(0.45))); off.set('y', str(Inches(1.98)))
ext.set('cx', str(Inches(9.1))); ext.set('cy', str(Inches(1.75)))
panel.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr').set('id', '9001')

h = 1.55; w = h * 1228 / 800
pic = s.shapes.add_picture(thumb, Inches(0.55), Inches(2.08), Inches(w), Inches(h))
pic.name = 'maxRateVideo'; pic.click_action.hyperlink.address = url
pic.line.color.rgb = RGBColor.from_string('DCE1E7'); pic.line.width = Pt(0.75)

tb = s.shapes.add_textbox(Inches(0.55 + w + 0.3), Inches(2.08), Inches(9.45 - 0.55 - w - 0.3), Inches(h))
tb.name = 'maxRateVideoText'; tb.click_action.hyperlink.address = url
tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
lines = [('▶  Watch & Learn', 14, True, '0B7A6F'), ('How to set the Max CP Rate in the CRM (29 sec)', 11.5, True, '1D1D1F'),
         ('Click the video to play it.', 9.5, False, '6E6E73'),
         ('The hard cap applies to future submitted and claimed jobs.', 9.5, False, '6E6E73')]
for i, (t, size, bold, col) in enumerate(lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(4)
    r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'

# Speaker notes: new CRM steps.
notes = s.notes_slide.notes_text_frame
old = notes.text
new = ('An agent can set a maximum hourly rate on a customer’s behalf in the CRM. It caps the rate of the cleaners who can '
       'receive and claim the customer’s jobs, going forward: cleaning requests only go to cleaners whose rates don’t exceed the cap. '
       'Steps on the slide (also in the video): 1) In the CRM, under Customer Information, find Max rate (it shows "Not set" when there is no cap) and click Set. '
       '2) In the "Set max rate cap" box, enter the customer’s preferred hourly rate and click Save. '
       '3) Max rate now shows the amount with a Hard cap tag; click Edit to change it. The hard cap applies to future submitted and claimed jobs. '
       'TRAINER: in slideshow mode, click the video thumbnail to play the tutorial. ')
tail = old[old.index('Trade-off'):] if 'Trade-off' in old else ''
notes.text = new + tail
prs.save(dst)
print('notes tail kept:', bool(tail))
