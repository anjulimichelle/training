"""Day 1 deck (14YN4Hcekv...): slide 5 "How this training works" becomes "Let's get you set up!",
an access and tools checklist the trainer walks through with trainees, confirming each login.
The old slide's text moves to the speaker notes.

Usage: python3 slide5-lets-get-you-set-up.py in.pptx out.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
s = prs.slides[4]
assert txt(s.shapes[0]) == 'How this training works', txt(s.shapes[0])
old_text = ' '.join(txt(sh) for sh in list(s.shapes)[4:] if txt(sh))
old_notes = s.notes_slide.notes_text_frame.text
sl = list(s.shapes)
keep = {id(x._element) for x in sl[:4]}
for sh in sl:
    if id(sh._element) not in keep:
        sh._element.getparent().remove(sh._element)


def set_text(shape, t):
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = t
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


set_text(s.shapes[0], "Let's get you set up!")
set_text(s.shapes[1], 'GETTING STARTED')


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def box(x, y, w, h, fill, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=6, lw=0.75):
    sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return sh


def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(1)
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tb


text(32, 72, 655, 18, [[('Log in to each tool with your trainer, one at a time. Tick it off once you are in.', 11, False, MUTED)]])

tools = [('HG email', 'Your Homeaglow Google account'),
         ('Slack', 'Team chat, updates and questions'),
         ('Hubstaff', 'Time tracking for your shifts'),
         ('Wise Account', 'Account set up and you can log in'),
         ('New hire form', 'Filled in and submitted'),
         ('Legacy CRM live (Django)', 'The live Legacy CRM'),
         ('Legacy CRM staging', 'Practice copy of the Legacy CRM'),
         ('New CRM live', 'The live New CRM'),
         ('New CRM staging', 'Practice copy of the New CRM'),
         ('Notion', 'Knowledge articles and guides')]
X, W, Y0, G = 32.4, 655.2, 98, 6
cw = (W - 10) / 2; ch = 42
for i, (name, desc) in enumerate(tools):
    col, row = divmod(i, 5)
    x = X + col * (cw + 10); y = Y0 + row * (ch + G)
    box(x, y, cw, ch, WHITE, LINE)
    box(x + 10, y + (ch - 22) / 2, 22, 22, TEAL_PALE, paras=[[(str(i + 1), 9.5, True, TEAL)]], align=PP_ALIGN.CENTER,
        shape=MSO_SHAPE.OVAL)
    text(x + 42, y + 4, cw - 130, ch - 8, [[(name, 10.5, True, INK)], [(desc, 8, False, MUTED)]])
    box(x + cw - 82, y + (ch - 14) / 2, 14, 14, WHITE, TEAL, radius=2, lw=1.25)
    text(x + cw - 62, y, 54, ch, [[('Submitted' if name == 'New hire form' else 'Logged in', 8, False, MUTED)]])
box(X, 342, W, 30, AMBER_PALE, paras=[[('💡  Can\'t log in?  ', 10, True, AMBER),
                                       ('Tell your trainer right away so your access can be fixed before the next session.', 10, False, INK)]],
    margin=12)

s.notes_slide.notes_text_frame.text = (
    "Let's get you set up! Guide trainees through logging in to each tool, one at a time, and make sure every trainee can log in "
    'before moving on. Ask them to confirm out loud or in the chat ("I\'m in") for each item:\n'
    '1) HG email (Homeaglow Google account). 2) Slack. 3) Hubstaff. 4) Wise account. 5) New hire form (completed and submitted). '
    '6) Legacy CRM live (Django). 7) Legacy CRM staging. 8) New CRM live. 9) New CRM staging. 10) Notion.\n'
    'Note anyone who cannot log in and the tool, and follow up so access is fixed before the next session.\n'
    'Previous version of this slide (How this training works): ' + old_text + ' ' + old_notes)
prs.save(sys.argv[2]); print('saved', sys.argv[2])
