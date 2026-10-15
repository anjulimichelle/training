"""Day 7 deck, Membership Retention Scenario 1 (unaware, wants a one-time cleaning).

Slide 23: overview. Check whether the ETF applies, plus two reminders: FC Sales signups
          get a sales call review first; gift-voucher (dhj_gift) customers redeemed, not purchased.
Slide 24: flow chart for ETF applies (Legacy DHJ within MCT, FCF within MCT with a completed job).
New 25:   flow chart for ETF does not apply (Legacy DHJ beyond MCT, FCF with no completed job),
          with the "When the TR is triggered" pop-up (click 1 shows, click 2 hides).
Helpers and the slide 25 layout come from slide16-sr-no-etf-flowchart.py.

Usage: python3 slides23-25-membership-scenario1.py in.pptx out.pptx
"""
import copy, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
S = list(prs.slides)
assert txt(S[22].shapes[1]) == 'Scenario 1: unaware, wants a one-time cleaning', txt(S[22].shapes[1])
assert txt(S[23].shapes[1]) == 'Scenario 1: MF reduction and exit', txt(S[23].shapes[1])
s = None
s8_notes = "The same Timed Reminder rules apply in every beyond-MCT branch. If the next MF is more than 5 days away, set the reminder at least 3 days before it. If it is 5 days or less, give 1 free month first and set the reminder 3 days before the new MF date. When it triggers with no reply and no engagement, cancel FC; the auto comms are enough unless an unresolved pain point needs a manual email. If there is no reply but the customer is using the service, such as a new booking, don't cancel. Confirm with them, push the due date out 1 to 2 weeks, and set a new reminder 3 days before it."


def clear(slide):
    sl = list(slide.shapes)
    keep = {id(e._element) for e in sl[:3]} | {id(e._element) for e in sl if e.is_placeholder}
    for sh in sl:
        if id(sh._element) not in keep:
            sh._element.getparent().remove(sh._element)


def set_text(sh, t, size=None):
    runs = sh.text_frame.paragraphs[0].runs
    runs[0].text = t
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in sh.text_frame.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    if size:
        runs[0].font.size = Pt(size)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill_paras(tf, paras, align):
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, font, size, bold, color in runs:
            r = p.add_run(); r.text = t
            f = r.font; f.name = font; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tf


def hang(tf):
    for p in tf.paragraphs:
        lead = p.runs[0].text if p.runs else ''
        if lead.startswith('•'):
            p._p.get_or_add_pPr().set('marL', str(11 * 12700)); p._p.get_or_add_pPr().set('indent', str(-11 * 12700))
        elif lead.startswith('–'):
            p._p.get_or_add_pPr().set('marL', str(24 * 12700)); p._p.get_or_add_pPr().set('indent', str(-10 * 12700))


def box(x, y, wd, h, fill, line, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, lw=1):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(wd), Pt(h))
    strip(sh)
    sh.adjustments[0] = min(0.3, 8 / h)
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(7); tf.margin_top = tf.margin_bottom = Pt(3)
    fill_paras(tf, paras, align)
    hang(tf)
    return sh


def text(x, y, wd, h, paras, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(wd), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_paras(tf, paras, align)
    return tb


def arrow(x1, y1, x2, y2, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    strip(c)
    c.line.color.rgb = RGBColor.from_string(MUTED); c.line.width = Pt(1)
    if head:
        ln = c._element.find('.//' + qn('a:ln'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
    return c


P = lambda t, size=9, bold=True, color=WHITE: (t, 'Play', size, bold, color)
A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)
B = lambda t, size=8.5: [('•  ', 'Arial', size, True, TEAL), A(t, size)]
B2 = lambda t, size=8.5: [('–  ', 'Arial', size, False, MUTED), A(t, size)]


A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)
B = lambda t, size=8.5: [('•  ', 'Arial', size, True, TEAL), A(t, size)]
B2 = lambda t, size=8.5: [('–  ', 'Arial', size, False, MUTED), A(t, size)]

def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def leave(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="1" presetClass="exit" presetSubtype="0" fill="hold" grpId="1" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>'
            f'</p:childTnLst></p:cTn></p:par>')


def click(cid, inner):
    return (f'<p:par><p:cTn id="{cid}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{cid+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


def apply_timing(pop):
    ids = [sh.shape_id for sh in pop]
    cid = 5
    c1 = ''
    for i, spid in enumerate(ids):
        c1 += effect(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
    start2 = cid; cid += 2
    c2 = ''
    for i, spid in enumerate(ids):
        c2 += leave(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 2
    bld = ''.join(f'<p:bldP spid="{i}" grpId="0"/><p:bldP spid="{i}" grpId="1"/>' for i in ids
                  if s.shapes._spTree.xpath(f'.//p:sp[p:nvSpPr/p:cNvPr/@id="{i}"]'))
    NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           + click(3, c1) + click(start2, c2) +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
    old = s._element.find(qn('p:timing'))
    if old is not None:
        s._element.remove(old)
    t_el = etree.fromstring(tim)
    ext = s._element.find(qn('p:extLst'))
    (ext.addprevious(t_el) if ext is not None else s._element.append(t_el))



def add_note(slide, text_):
    nt = slide.notes_slide.notes_text_frame
    if text_[:40] in nt.text:
        return
    p = copy.deepcopy(nt.paragraphs[-1]._p)
    for r in p.findall(qn('a:r'))[1:]:
        p.remove(r)
    p.findall(qn('a:r'))[0].find(qn('a:t')).text = text_
    nt.paragraphs[-1]._p.addnext(p)


# ======================= slide 23: overview =======================
s = S[22]
clear(s)
set_text(s.shapes[1], 'Scenario 1: Unaware | Wants a One-Time Cleaning', 22)
set_text(s.shapes[2], 'Explain and probe before any offer. First, check whether the ETF applies.')
box(32, 102, 322, 50, TEAL_PALE, TEAL, [
    [P('ETF applies  →  next slide', 10, True, TEAL)],
    B('Legacy DHJ within MCT'), B('FCF within MCT with a completed job')], align=PP_ALIGN.LEFT)
box(365, 102, 322, 50, WHITE, LINE, [
    [P('ETF does not apply  →  slide after', 10, True, INK)],
    B('Legacy DHJ beyond MCT'), B('FCF (within or beyond MCT) with no completed job')], align=PP_ALIGN.LEFT)
text(32, 161, 300, 14, [[P('REMINDERS BEFORE YOU RESPOND', 8.5, True, AMBER)]])
box(32, 178, 322, 108, AMBER_PALE, AMBER, [
    [P('Signed up through FC Sales?', 10, True, AMBER)],
    [A('Raise the ticket for a sales call review before you respond with policy. It checks whether anything was misrepresented on the call.', 8)],
    B('Misrepresentation found: follow General Guidance: Incorrect Information Provided.', 8),
    B("No misrepresentation: continue the steps, and include the FC Sales agent's post-call email in your reply.", 8),
    [A("Why: if a rep misstated the terms, that's ours to fix, not a retention conversation.", 7.5, False, MUTED)],
], anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT)
box(365, 178, 322, 108, AMBER_PALE, AMBER, [
    [P('Redeemed a gift voucher (dhj_gift)?', 10, True, AMBER)],
    [A('A friend (the Host) gave them the voucher. Show them where the terms were shared:', 8)],
    B('Their customer_welcome_forever_clean_guest email, if they got it.', 8),
    B("The Host's customer_free_voucher_to_gift_a_friend email: screenshot it, crop out internal details.", 8),
    B('New new_booking_flow cohort: terms were also shown at their first booking.', 8),
    [A('Don\'t say they "purchased" the voucher. Say they redeemed it.', 8, True, INK)],
], anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT)
box(32, 296, 655, 26, WHITE, LINE, [
    [P('Note:  ', 8.5, True, TEAL), A('no offer is tied to "unawareness" itself. Clarifying the terms addresses it. The credit only gets C to book.', 8.5)]],
    align=PP_ALIGN.LEFT)
add_note(s, 'Before responding with policy: if the signup came through the FC Sales team, raise the ticket for a sales call review, which checks whether anything was misrepresented on the call. '
            'If misrepresentation is found, follow General Guidance: Incorrect Information Provided. If not, continue with the steps and include the FC Sales agent\'s post-call email in your reply so the customer can see the terms they were sent. '
            'If they redeemed a gift voucher (dhj_gift), a friend (the Host) gave them a free voucher, so they never went through a purchase. Show them where the terms were shared: '
            'the customer_welcome_forever_clean_guest email if they received it; the Host\'s customer_free_voucher_to_gift_a_friend email (tell them when their Host was told the voucher includes a membership, and include a screenshot with internal details like the Host\'s customer ID and the template name cropped out); '
            'and, if their new_booking_flow cohort is new, mention the terms were also shown when they booked their first cleaning (skip this for the old cohort). Don\'t say they "purchased" the voucher; say they redeemed it. '
            'Then check whether the ETF applies: Legacy DHJ within MCT or FCF within MCT with a completed job (next slide); or Legacy DHJ beyond MCT or FCF with no completed job, within or beyond MCT (the slide after).')

# ======================= slide 24: ETF applies =======================
s = S[23]
clear(s)
set_text(s.shapes[1], 'Scenario 1: ETF applies')
set_text(s.shapes[2], 'Legacy DHJ within MCT  |  FCF within MCT with a completed job')
CX = 345
box(CX - 200, 99, 400, 22, TEAL, None,
    [[P('FC retention – Unaware of the FC  |  C only wants a one-time cleaning', 8.5)], [P('(ETF applies)', 7.5, False, TEAL_PALE)]])
arrow(CX, 121, CX, 126)
box(CX - 150, 126, 300, 30, TEAL_PALE, TEAL,
    [[A('Explain FC terms (MF + ETF) and probe for their service experience. Apply free weeks/months to move the billing date as necessary.', 8)]])
arrow(CX, 156, CX, 162)
box(CX - 130, 162, 260, 22, AMBER_PALE, AMBER, [[P('Is unawareness the only reason for cancelling?', 8.5, True, AMBER)]])
arrow(CX + 130, 173, 500, 173)
text(CX + 136, 162, 30, 10, [[P('NO', 7.5, True, INK)]])
box(500, 158, 187, 30, TEAL_PALE, TEAL, [[A('Follow the playbook for the actual reason C gave.', 8, False, TEAL)]])
arrow(CX, 184, CX, 192)
text(CX + 5, 184, 30, 8, [[P('YES', 7.5, True, INK)]])
box(CX - 55, 192, 110, 18, AMBER_PALE, AMBER, [[P('Retainable?', 9, True, AMBER)]])
arrow(CX - 55, 201, 192, 201)
text(CX - 82, 190, 24, 9, [[P('NO', 7.5, True, INK)]])
box(32, 186, 160, 80, WHITE, LINE, [
    B('Reinforce the FC value', 7.5), B('Explain why the ETF applies', 7.5), B('Offer $20 credits, as necessary', 7.5),
    B('Apply free weeks/months to move the billing date as a one-time courtesy', 7.5)], align=PP_ALIGN.LEFT)
arrow(112, 266, 112, 285, head=False)
arrow(112, 285, CX - 130, 285)
arrow(CX, 210, CX, 217)
text(CX + 5, 209, 30, 8, [[P('YES', 7.5, True, INK)]])
box(CX - 150, 217, 300, 52, TEAL_PALE, TEAL, [
    B('Re-explain the FC value', 8),
    B('Offer one of the following + free months (if applicable):', 8),
    B2('$20 credit, or an MF reduction', 8),
    B('Apply free weeks/months to move the billing date as a one-time courtesy', 8)], align=PP_ALIGN.LEFT)
arrow(CX, 269, CX, 275)
box(CX - 130, 275, 260, 20, INK, None, [[P('C rejects the resolution and insists on cancelling', 8.5, False)]])
arrow(CX, 295, CX, 301)
box(CX - 100, 301, 200, 20, AMBER_PALE, AMBER, [[P('Any hint they\'d pay a reduced ETF?', 8.5, True, AMBER)]])
arrow(CX + 100, 311, 500, 311)
text(CX + 106, 300, 30, 9, [[P('YES', 7.5, True, INK)]])
box(500, 298, 187, 26, TEAL, None, [[P('EXIT Strategy', 9)], [P('ETF reduction, max 50%', 7.5, False, TEAL_PALE)]])
arrow(593, 324, 593, 331)
box(500, 331, 187, 22, INK, None, [[P('C declines and insists on cancelling', 8, False)]])
arrow(500, 342, CX + 70, 342)
arrow(CX, 321, CX, 331)
text(CX + 5, 321, 30, 9, [[P('NO', 7.5, True, INK)]])
box(CX - 70, 331, 140, 22, AMBER_PALE, AMBER, [[P('GRACEFUL CLOSURE', 9, True, AMBER)]])
add_note(s, 'Flow (ETF applies): explain the FC terms (MF and ETF) and probe for their service experience. If unawareness is not the only reason, follow the playbook for the reason they gave. '
            'Not retainable: reinforce the FC value, explain why the ETF applies, offer $20 credits as necessary, and move the billing date with free weeks/months as a one-time courtesy. '
            'Retainable: re-explain the value and offer a $20 credit or an MF reduction, plus free months if applicable. If C still rejects and insists on cancelling, check for any hint they would pay a reduced ETF: '
            'if yes, Exit Strategy with ETF reduction up to 50%; if they decline that too, or there is no hint, Graceful Closure.')

# ======================= new slide 25: ETF does not apply =======================
ref = S[23]
s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
num = [x for x in ref.shapes if x.is_placeholder][0]
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); list(lst)[23].addnext(el)
set_text(s.shapes[1], 'Scenario 1: ETF does not apply')
set_text(s.shapes[2], 'Legacy DHJ beyond MCT  |  FCF (within or beyond MCT) with no completed job')
s.notes_slide.notes_text_frame.text = (
    'Flow (ETF does not apply): is the customer still retainable? If not, address any pain points and cancel the FC. '
    'If retainable, attempt to retain and offer an MF reduction (optional) plus free months if applicable, explicitly confirm no further MF will be charged unless they agree to continue, and set a TR. '
    'Beyond MCT, MF reduction rarely moves the needle because the customer usually has not been using the service. If C rejects the resolution and insists on cancelling, cancel the FC.')
L, LW = 32, 400
CX = L + LW / 2
R, RW = 456, 231
RX = R + RW / 2

box(L, 98, LW, 22, TEAL, None,
    [[P('FC retention – Unaware of the FC  |  C only wants a one-time cleaning', 8.5)], [P('(ETF does not apply)', 8, False, TEAL_PALE)]])
arrow(CX, 120, CX, 127)
box(CX - 110, 127, 220, 26, AMBER_PALE, AMBER,
    [[P('Is the C still retainable?', 9.5, True, AMBER)],
     [P('Use the Five Key Factors', 7.5, False, AMBER)]])
arrow(CX + 110, 140, R, 140)
text(CX + 120, 129, 30, 10, [[P('NO', 8, True, INK)]])
box(R, 125, RW, 30, TEAL_PALE, TEAL,
    [[P('Address any pain points and cancel the FC', 9, False, TEAL)]])
box(R, 164, RW, 50, WHITE, LINE,
    [[P('Assessing retainability:  ', 8.5, True, TEAL),
      A('use the General Retention Playbook Five Key Factors. Weigh what is actually present, not a factor count.', 8)]],
    align=PP_ALIGN.LEFT)

box(R, 222, RW, 40, AMBER_PALE, AMBER,
    [[P('Beyond MCT:  ', 8.5, True, AMBER), A("MF reduction rarely moves the needle, because C usually hasn't been using the service.", 8)]],
    align=PP_ALIGN.LEFT)
arrow(CX, 153, CX, 163)
text(CX + 5, 153, 30, 10, [[P('YES', 8, True, INK)]])
box(L, 163, LW, 88, TEAL_PALE, TEAL, [
    B('Attempt to retain C and offer an MF reduction (optional) + free months (if applicable: unused paid MFs, not more than asked, not more than 3)'),
    B('Explicitly confirm no further MF will be charged unless they agree to continue'),
    B('Set a TR 3 days before the next MF date. If the next MF is within 5 days, issue a free month first'),
], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT)

# split into two branches
LB, RB = L + 100, L + LW - 100
arrow(CX, 251, CX, 260, head=False)
arrow(LB, 260, RB, 260, head=False)
arrow(LB, 260, LB, 269)
arrow(RB, 260, RB, 269)
box(LB - 92, 269, 184, 26, INK, None, [[P('C rejects the resolution', 9, False)]])
arrow(LB, 295, LB, 304)
box(LB - 70, 304, 140, 24, AMBER_PALE, AMBER, [[P('Cancel the FC', 9.5, True, AMBER)]])
box(RB - 92, 269, 184, 26, INK, None, [[P('No reply, but the TR was triggered', 9, False)]])
arrow(RB, 295, RB, 304)
box(RB - 70, 304, 140, 24, TEAL, None, [[P('What to do next  ▸', 9, True)]])
text(RB - 92, 331, 184, 10, [[A('click to see both outcomes', 7, False, MUTED)]], PP_ALIGN.CENTER)

# ---- pop-up: when the TR is triggered (from slide 8) ----
pop = []
card = box(110, 102, 500, 218, WHITE, TEAL, [[A('')]], lw=1.5)
sp = card._element.find(qn('p:spPr'))
eff = etree.SubElement(sp, qn('a:effectLst'))
sh_ = etree.SubElement(eff, qn('a:outerShdw'), blurRad='76200', dist='25400', dir='5400000', algn='t', rotWithShape='0')
clr = etree.SubElement(sh_, qn('a:srgbClr'), val='000000'); etree.SubElement(clr, qn('a:alpha'), val='25000')
pop.append(card)
pop.append(box(250, 114, 220, 30, AMBER_PALE, AMBER, [[P('When the TR is triggered', 11, True, AMBER)]]))
PL, PR = 245, 475
pop.append(arrow(360, 144, 360, 154, head=False))
pop.append(arrow(PL, 154, PR, 154, head=False))
pop.append(arrow(PL, 154, PL, 164))
pop.append(arrow(PR, 154, PR, 164))
pop.append(box(PL - 110, 164, 220, 30, AMBER, None, [[P('No reply and no new bookings', 9.5)]]))
pop.append(box(PL - 110, 200, 220, 84, AMBER_PALE, AMBER,
               [[A('Cancel the FC.', 9, True)],
                [A("Send comms only if there's an unresolved pain point the auto comms wouldn't cover.", 9)]]))
pop.append(box(PR - 110, 164, 220, 30, TEAL, None, [[P('No reply, but C is using the service', 9.5)]]))
pop.append(box(PR - 110, 200, 220, 84, TEAL_PALE, TEAL, [
    B("Don't cancel the FC", 9), B('Contact C to confirm the cancellation', 9),
    B('Add 1–2 weeks to move the due date, and set a new TR', 9)], align=PP_ALIGN.LEFT))
pop.append(text(130, 296, 460, 12, [[A('Click to close', 7.5, False, MUTED)]], PP_ALIGN.RIGHT))



apply_timing(pop)
add_note(s, 'TRAINER: click once to open the "When the TR is triggered" pop-up, click again to close it. ' + s8_notes)

prs.save(sys.argv[2]); print('saved', sys.argv[2])
