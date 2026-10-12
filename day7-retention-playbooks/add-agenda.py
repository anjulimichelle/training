"""Day 7 Retention Playbooks deck: add a clickable agenda slide right after the cover. Each row jumps to its
topic's title slide. Header shapes are copied from an existing content slide so fonts and colours match."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, INK, SOFT, BORDER, WHITE = '0B7A6F', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF'
prs = Presentation(src)
S = list(prs.slides)
assert not any(s.shapes[1].has_text_frame and s.shapes[1].text_frame.text.strip() == 'Agenda' for s in S)
ref = S[2]
topics = [s for s in S if s.shapes[0].has_text_frame and s.shapes[0].text_frame.text.strip().upper().startswith('TOPIC 0')]
assert len(topics) == 3

s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
num = [x for x in ref.shapes if x.is_placeholder][0]
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
for shape, t in ((s.shapes[0], 'OVERVIEW'), (s.shapes[1], 'Agenda'), (s.shapes[2], 'Three playbooks. Click a topic to jump to it.')):
    r = shape.text_frame.paragraphs[0].runs; r[0].text = t
    for x in r[1:]: x._r.getparent().remove(x._r)
lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); lst.insert(1, el)

def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
rows = [(topics[0], 'PCQ Retention Playbook', 'A poor cleaning leads C to say they want to cancel.'),
        (topics[1], 'Service Reliability Retention Playbook', 'A CP failure, not a bad cleaning, leads C to want to cancel.'),
        (topics[2], 'Membership Retention Playbook', 'The membership terms themselves drive the cancellation.')]
H, GAP = 1.0, 0.15
for i, (target, title, line) in enumerate(rows):
    y = Inches(1.45 + i * (H + GAP))
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), y, Inches(9.1), Inches(H)); strip(card)
    card.adjustments[0] = 0.06; card.fill.solid(); card.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    card.line.color.rgb = RGBColor.from_string(BORDER); card.line.width = Pt(0.75)
    tf = card.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.95); tf.margin_right = Inches(0.6); tf.margin_top = tf.margin_bottom = 0
    for k, (t, font, size, bold, col) in enumerate(((title, 'Play', 14, True, INK), (line, 'Arial', 10.5, False, SOFT))):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.LEFT; p.space_after = Pt(2)
        r = p.add_run(); r.text = t; r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(col)
    for x, w, t, size in ((0.55, 0.6, f'{i + 1:02d}', 16), (9.1, 0.35, '›', 20)):
        b = s.shapes.add_textbox(Inches(x), y, Inches(w), Inches(H))
        tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = t; r.font.name = 'Play'; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(TEAL)
        b.click_action.target_slide = target
    card.click_action.target_slide = target

s.notes_slide.notes_text_frame.text = (
    'Agenda. Topic 1: PCQ Retention Playbook (a poor cleaning leads the customer to say they want to cancel). '
    'Topic 2: Service Reliability Retention Playbook (a CP failure, not a bad cleaning, leads the customer to want to cancel). '
    'Topic 3: Membership Retention Playbook (the membership terms themselves, not a service failure, drive the cancellation). '
    'In slideshow mode, click any topic to jump straight to it.')
prs.save(dst)
