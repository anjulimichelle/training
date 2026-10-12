"""Day 7 (original deck): slide 4 keeps the trigger tile and gets the PCQ decision tree plus a
baseline reminder; the Baseline slide (5) is removed and its list becomes a pop-up button on the
Offers slide. Patches the live Google export so all other slides stay exactly as they are."""
import re, sys, io, zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

src, dst = sys.argv[1], sys.argv[2]
C = lambda h: RGBColor.from_string(h)
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE = C('0B7A6F'), C('E3F4F1'), C('1D1D1F'), C('6E6E73'), C('DCE1E7'), C('FFFFFF')
GOLD, GOLDSOFT = C('9A6410'), C('FFF3D6')
ORANGE, BLUE = C('E69138'), C('3D85C6')
RED, PINK, REDLINE = C('E06666'), C('F4CCCC'), C('CC4125')
DARK, GREY, LINE = C('666666'), C('D9D9D9'), C('595959')

prs = Presentation(src)
s4, s5, s6 = prs.slides[3], prs.slides[4], prs.slides[5]
txt = lambda s: ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
assert 'The trigger' in txt(s4) and 'Baseline for every PCQ' in txt(s5) and 'voucher and free month' in txt(s6)

def strip_style(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None: sh._element.remove(st)

def box(s, x, y, w, h, fill, line=None, name=None, kind=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.12):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h)); strip_style(sh)
    if name: sh.name = name
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = adj
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    return sh

def write(sh, parts, size, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font='Arial', margin=0.05, paras=None):
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right'): setattr(tf, m, Inches(margin))
    for m in ('margin_top', 'margin_bottom'): setattr(tf, m, Inches(0.02))
    for k, pp in enumerate(paras or [parts]):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align
        for t, bold, *rest in pp:
            r = p.add_run(); r.text = t; f = r.font; f.size = Pt(rest[1] if len(rest) > 1 else size); f.bold = bold
            f.name = rest[2] if len(rest) > 2 else font; f.color.rgb = rest[0] if rest else color

def tbox(s, x, y, w, h, name=None):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: t.name = name
    return t

def seg(s, x1, y1, x2, y2, arrow=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = LINE; c.line.width = Pt(1)
    if arrow:
        ln = c.line._get_or_add_ln(); etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
    return c

# ---------------- Slide 4: keep header + trigger tile, add the decision tree ----------------
for sh in list(s4.shapes)[8:21]: sh._element.getparent().remove(sh._element)

BASE, U, W = 3.62, 1.03, 1.78           # card column k sits at BASE + k*U, width W
HH, BH = 0.22, 0.42                      # Yes/No header and body heights
colx = lambda k: BASE + k * U
cen = lambda k: colx(k) + W / 2

def card(k, y, yes, body):
    hdr = box(s4, colx(k), y, W, HH, RED if yes else DARK, adj=0.2)
    write(hdr, [('Yes' if yes else 'No', True)], 9, WHITE, font='Play')
    b = box(s4, colx(k), y + HH + 0.02, W, BH, PINK if yes else GREY, line=REDLINE if yes else C('999999'), adj=0.14)
    write(b, [(body, False)], 8.5, INK)
    return y + HH + 0.02 + BH / 2        # body mid-line, where branches leave

QW = 2.55; qx = cen(1) - QW / 2
start = box(s4, qx, 1.32, QW, 0.32, ORANGE, adj=0.15)
write(start, [('Customer wants to cancel due to PCQ', True)], 9.5, WHITE, font='Play')
q1 = box(s4, qx, 1.82, QW, 0.4, BLUE, adj=0.15)
write(q1, [('Is the customer beyond their 6-month MCT?', True)], 9.5, WHITE, font='Play')
seg(s4, cen(1), 1.64, cen(1), 1.82, arrow=True)

R3, R4, R5 = 2.4, 3.27, 4.14
def branch(from_left, from_right, mid_y, ytop, kyes, kno):
    seg(s4, from_left, mid_y, cen(kyes), mid_y); seg(s4, cen(kyes), mid_y, cen(kyes), ytop, arrow=True)
    seg(s4, from_right, mid_y, cen(kno), mid_y); seg(s4, cen(kno), mid_y, cen(kno), ytop, arrow=True)

branch(qx, qx + QW, 2.02, R3, 0, 2)
card(0, R3, True, 'Use “Beyond MCT” scenario')
m = card(2, R3, False, 'Is this the customer’s 1st job?')
branch(colx(2), colx(2) + W, m, R4, 1, 3)
card(1, R4, True, 'Use “1st Job” scenario')
m = card(3, R4, False, 'Was the PREVIOUS job an issue job (back-to-back)?')
branch(colx(3), colx(3) + W, m, R5, 2, 4)
card(2, R5, True, 'Use “1st Job” scenario')
card(4, R5, False, 'Use “Any Job Except 1st Job” scenario')

# Baseline reminder under the trigger tile, linked to the Offers slide (where the list now lives).
note = box(s4, 0.45, 3.45, 2.93, 0.95, GOLDSOFT, adj=0.08)
write(note, [], 9, paras=[
    [('Check the baseline responses first', True, GOLD, 9.5, 'Play')],
    [('Every PCQ gets them, whatever the scenario. Open the list with the Baseline button on the Offers slide.', False, GOLD, 8.5)]],
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.12)
note.click_action.target_slide = s6

s4.notes_slide.notes_text_frame.text = (
    "The trigger: the CP showed up and did the job, but the result was poor, and now the customer wants to cancel. "
    "This is distinct from Service Reliability (late, no-show, cancelled) and from objections about the membership terms.\n"
    "Walk the decision tree. First: is the customer beyond their 6-month MCT? If yes, use the Beyond MCT scenario (no ETF exposure). "
    "If no: is this the customer's first job? If yes, use the 1st Job scenario. If not, was the previous job also an issue job (back-to-back)? "
    "If yes, use the 1st Job scenario; a second consecutive issue job is a broken promise on top of a broken promise, which is what the ETF Waivers policy addresses. "
    "If no, use the Any Job Except 1st Job scenario.\n"
    "Before any scenario, check the baseline responses. Every PCQ gets them; they are behind the Baseline button on the Offers slide.")

# ---------------- Offers slide: Baseline button + pop-up ----------------
btn = box(s6, 0.45, 4.6, 3.2, 0.4, TEAL, adj=0.2, name='baselineButton')
write(btn, [('▶  Baseline for every PCQ', True)], 10, WHITE, font='Play')
hint = tbox(s6, 3.8, 4.6, 5.0, 0.4, 'baselineHint')
write(hint, [('Click once to show the list, click again to close.', False)], 9, SOFT, align=PP_ALIGN.LEFT, margin=0)

items = [
    ('Acknowledge', "Use the Empathy Statements macro, matched to the customer's specific situation."),
    ("Probe, don't withhold", 'We already know the root cause. Give the resolution while letting the customer vent.'),
    ('Refund eligibility', "72 hours from the job. Customer self-submits, CP has 7 days to accept or decline. Skip if they don't want one."),
    ('Premium fee', 'Refund it if applicable. It fires on the confirmed service issue, with or without refund intent.'),
    ('Pairing and CP', 'Block C/CP pairing unless they want the same CP. Coach or penalize with the CP-facing PCQ macro.'),
    ('Flag and RC', 'Flag customer_unsatisfied if no review yet. Cancel RC and upcoming jobs unless they want to keep them.'),
]
panel = box(s6, 0.38, 1.25, 9.24, 3.95, WHITE, line=BORDER, name='pop1_2Panel', adj=0.03)
t = tbox(s6, 0.65, 1.38, 8.7, 0.32, 'pop1_2Title')
write(t, [('Baseline for every PCQ', True)], 14, INK, align=PP_ALIGN.LEFT, font='Play', margin=0)
for k, (head, body) in enumerate(items):
    col, row = k % 2, k // 2
    x, y = 0.65 + col * 4.5, 1.85 + row * 1.05
    tile = box(s6, x, y, 4.3, 0.92, TEALSOFT, name=f'pop1_2Tile{k}', adj=0.08)
    write(tile, [], 9, paras=[[(head, True, INK, 10.5, 'Play')], [(body, False, SOFT, 9)]],
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.15)

s6.notes_slide.notes_text_frame.text = (
    s6.notes_slide.notes_text_frame.text +
    "\nBaseline button (click to show, click again to close): " + s5.notes_slide.notes_text_frame.text)

# ---------------- Remove the Baseline slide ----------------
sldIdLst = prs.slides._sldIdLst
sid = sldIdLst[4]; rId = sid.get(qn('r:id'))
prs.part.drop_rel(rId); sldIdLst.remove(sid)
offers_part = s6.part.partname
buf = io.BytesIO(); prs.save(buf)

# ---------------- Click animation on the Offers slide only ----------------
ns = {}
exec(open(__file__.replace('pcq-decision-tree.py', 'add_story_anim.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
target = str(offers_part).lstrip('/')
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == target:
        x = data.decode('utf8'); assert '<p:timing' not in x
        ids = re.findall(r'<p:cNvPr id="(\d+)" name="pop1_2[^"]*"', x)
        cid, ins, outs, bld = 3, '', '', []
        for i, sid_ in enumerate(ids):
            ins += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid_}"/></p:tgtEl>', 'clickEffect' if i == 0 else 'withEffect'); cid += 3
        c1 = ns['click'](cid, ins); cid += 2
        for i, sid_ in enumerate(ids):
            outs += ns['leave'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid_}"/></p:tgtEl>', 'clickEffect' if i == 0 else 'withEffect'); cid += 3
        c2 = ns['click'](cid, outs)
        for sid_ in ids:
            bld += [f'<p:bldP spid="{sid_}" grpId="0" animBg="1"/>', f'<p:bldP spid="{sid_}" grpId="1" animBg="1"/>']
        x = x.replace('</p:sld>', ns['timing']([c1, c2], bld) + '</p:sld>')
        data = x.encode('utf8'); print('offers pop-up shapes', len(ids))
    zout.writestr(item, data)
zout.close()
