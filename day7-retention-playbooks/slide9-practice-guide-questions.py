"""Day 7 deck slide 9 (Practice ticket): replace the ticket card and steps with a
"trainer provides a live ticket" note plus four guide questions.

Usage: python3 slide9-practice-guide-questions.py in.pptx out.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, WHITE, BORDER, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', '9A6410', 'FFF3D6'
prs = Presentation(sys.argv[1])
s = prs.slides[8]
sl = list(s.shapes)
assert sl[1].text_frame.text.startswith('Practice ticket')
keep = {id(sh._element) for sh in sl[:3]} | {id(sh._element) for sh in sl if sh.is_placeholder}
for sh in sl:
    if id(sh._element) not in keep:
        sh._element.getparent().remove(sh._element)

sub = sl[2].text_frame.paragraphs[0].runs
sub[0].text = 'Your trainer will provide a live ticket. Use these guide questions as you investigate.'
for r in sub[1:]:
    r._r.getparent().remove(r._r)


def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)


def box(x, y, w, h, fill, line=None, adj=0.08):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
    b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
    else: b.line.fill.background()
    return b


def text(shape, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, after=2):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
        for t, bold, col, sz in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'Play' if bold else 'Arial'
    return shape


tb = lambda x, y, w, h: s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))

# trainer note
box(0.45, 1.4, 9.1, 0.5, GOLDSOFT)
text(tb(0.65, 1.4, 8.8, 0.5), [[('▶  Live ticket:  ', True, GOLD, 11),
                               ('your trainer provides a live PCQ cancel-intent ticket to investigate together.', False, INK, 10)]])

qs = [('What is the customer\'s situation?', 'e.g. FC: PCQ job not the first (within MCT, not consecutive)'),
      ('What actions are you going to take?', None),
      ('What will you offer?', None),
      ('Draft your reply.', None)]
y0, gap = 2.05, 0.1
h = (4.9 - y0 - gap * (len(qs) - 1)) / len(qs)
for i, (q, ex) in enumerate(qs):
    y = y0 + i * (h + gap)
    last = i == len(qs) - 1
    box(0.45, y, 9.1, h, TEALSOFT if not last else WHITE, None if not last else TEAL)
    text(box(0.62, y + (h - 0.34) / 2, 0.34, 0.34, TEAL, adj=0.5), [[(str(i + 1), True, WHITE, 11)]], align=PP_ALIGN.CENTER)
    paras = [[(q, True, TEAL, 12)]]
    if ex:
        paras.append([(ex, False, SOFT, 9.5)])
    text(tb(1.15, y, 8.2, h), paras)

s.notes_slide.notes_text_frame.text = (
    'Practice ticket activity. Provide a live PCQ cancel-intent ticket and investigate it together, using the guide questions. '
    '1) What is the customer\'s situation? Have trainees name the row of the Quick Reference that applies, for example FC: PCQ job not the first (within MCT, not consecutive), and say whether they are within or beyond MCT. '
    '2) What actions are you going to take? The baseline for every PCQ (empathy, refund eligibility, Premium fee, pairing ban, CP coaching, flag) plus the scenario steps. '
    '3) What will you offer? Voucher, free month and the ETF spiel from the Quick Reference, and what happens if C rejects it (Exit Strategy within MCT; cancel or set a TR beyond MCT). '
    '4) Draft your reply. Have each trainee write the response to C, then compare as a group.')
prs.save(sys.argv[2]); print('saved', sys.argv[2])
