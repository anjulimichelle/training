"""Day 7 deck: after slide 7 (Beyond MCT) add
  8. a PCQ Retention Playbook knowledge check (answers reveal one per click), and
  9. a practice ticket activity (trainer shares a live ticket; trainees investigate).
Content comes from slides 4-7 of the deck. Reuses the click helpers in
day3-platform-mechanics/add_story_anim.py.

Usage: python3 add-slides8-9-pcq-kc-and-practice.py in.pptx out.pptx
"""
import copy, io, os, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, WHITE, BORDER, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', '9A6410', 'FFF3D6'
prs = Presentation(src)
S = list(prs.slides)
ref = S[6]
assert ref.shapes[1].text_frame.text.strip() == 'Beyond MCT'
assert S[7].shapes[1].text_frame.text.strip() == 'Questions?'
num = [x for x in ref.shapes if x.is_placeholder][0]


def new_slide(eyebrow, title, sub, after):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
        s.shapes._spTree.append(copy.deepcopy(el))
    for shape, t in ((s.shapes[0], eyebrow), (s.shapes[1], title), (s.shapes[2], sub)):
        r = shape.text_frame.paragraphs[0].runs; r[0].text = t
        for x in r[1:]: x._r.getparent().remove(x._r)
        for p in shape.text_frame.paragraphs[1:]: p._p.getparent().remove(p._p)
    lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); list(lst)[after].addnext(el)
    return s


def kit(s):
    def strip(x):
        st = x._element.find(P_NS + 'style')
        if st is not None: x._element.remove(st)
    def box(x, y, w, h, fill, line=None, adj=0.08, name=None, dash=False):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
        b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line:
            b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(1 if dash else 0.75)
            if dash:
                from pptx.enum.dml import MSO_LINE
                b.line.dash_style = MSO_LINE.DASH
        else: b.line.fill.background()
        if name: b.name = name
        return b
    def text(shape, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, after=3):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
            for t, bold, col, sz in runs:
                r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'Play' if bold else 'Arial'
        return shape
    def tb(x, y, w, h, name=None):
        t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name: t.name = name
        return t
    return box, text, tb


# ---------- 8. Knowledge check ----------
kc = new_slide('PCQ RETENTION PLAYBOOK  ·  QUICK KNOWLEDGE CHECK', 'Quick Knowledge Check',
               'Ask each question, then click to reveal the answer.', 6)
box, text, tb = kit(kc)
qa = [('A non-FC member wants to cancel after a PCQ job. What can you offer?',
       'Up to a 1-hour voucher or $20 credits. No free month and no ETF: there is no membership.'),
      ('FC, within MCT, first job (3 hours) was PCQ. Which voucher?',
       '1-hour voucher (2 hours if the job was 4 hours or more). If the system already auto-issued a 3-hour voucher, that replaces it.'),
      ('Within MCT, C rejects the resolution and insists on cancelling free. What do you check first?',
       'Were ALL necessary resolutions for the root cause provided? No: provide the missing one. Yes: go to Exit Strategy.'),
      ('What are the two Exit Strategy levers?',
       'ETF reduction (up to 50% off) or MCT reduction (up to 2 months, start with 1). Pick whichever costs C less.'),
      ('Beyond MCT, C is retainable. When do you set the TR?',
       '3 days before the next MF date. If the next MF is within 5 days, issue a free month first.'),
      ('The TR triggers: no reply, but C has a new booking. What do you do?',
       "Don't cancel the FC. Contact C to confirm, add 1–2 weeks to the due date and set a new TR.")]
n = len(qa); h = (3.75 - 0.08 * (n - 1)) / n
for i, (q, a) in enumerate(qa):
    y = 1.4 + i * (h + 0.08)
    box(0.45, y, 5.05, h, WHITE, BORDER)
    text(box(0.6, y + (h - 0.3) / 2, 0.3, 0.3, TEALSOFT, adj=0.5), [[(str(i + 1), True, TEAL, 9)]], align=PP_ALIGN.CENTER)
    text(tb(1.05, y, 4.35, h), [[(q, False, INK, 10)]])
    box(5.6, y, 3.95, h, TEALSOFT, name=f'step{i + 1}Box')
    text(tb(5.75, y, 3.7, h, name=f'step{i + 1}Ans'), [[(a, False, INK, 9)]])
kc.notes_slide.notes_text_frame.text = (
    'Quick knowledge check: PCQ Retention Playbook. Ask each question and let trainees answer before revealing. Each click reveals the next answer. '
    '1) Non-FC members get up to a 1-hour voucher or $20 credits, with no free month or ETF, because there is no membership. '
    '2) For FC within MCT the voucher is 1 hour for jobs under 4 hours and 2 hours for 4 hours or more, unless the system already issued a 3-hour voucher, which replaces it. '
    '3) Before moving on when C insists on cancelling at no charge, check whether every resolution the root cause calls for was provided. If not, provide the missing one and keep retaining. If yes, move to Exit Strategy. '
    '4) Exit Strategy compares ETF reduction (maximum 50% off) with MCT reduction (up to 2 months, ideally starting with 1, capped at 1 if 2 would finish the MCT). Choose whichever costs the customer less; ETF reduction is usually cheaper. Max out the chosen lever before Graceful Closure. '
    '5) Beyond MCT, set the TR 3 days before the next MF date. If the next MF is within 5 days, issue a free month first, then set the TR 3 days before the new MF date. '
    '6) When the TR triggers with no reply but the customer is using the service, such as a new booking, do not cancel. Contact C to confirm, push the due date out 1 to 2 weeks, and set a new TR. With no reply and no new bookings, cancel the FC.')

# ---------- 9. Practice ticket ----------
pt = new_slide('PCQ RETENTION PLAYBOOK  ·  PRACTICE', "Practice ticket: let's investigate together!",
               'Your trainer will share a live ticket. Work it through with the playbook.', 7)
box, text, tb = kit(pt)
# left: the ticket card
box(0.45, 1.4, 3.6, 3.3, WHITE, TEAL, dash=True)
text(tb(0.65, 1.55, 3.2, 0.3), [[('▶  Live ticket', True, TEAL, 12)]])
text(tb(0.65, 1.95, 3.2, 0.5), [[('Your trainer shares the ticket on screen. Note these as you investigate:', False, SOFT, 9)]], anchor=MSO_ANCHOR.TOP)
for i, lab in enumerate(['Ticket ID', 'Customer ID', 'Job ID (PCQ job)', 'Model: FCF or DHJ', 'Within or beyond MCT']):
    y = 2.5 + i * 0.4
    text(tb(0.65, y, 1.7, 0.3), [[(lab, True, INK, 9)]])
    box(2.45, y + 0.03, 1.4, 0.26, 'F5F6F8', BORDER, adj=0.2)
# right: the steps
steps = [('Confirm the trigger', 'Is PCQ really why C wants to cancel? Check the job and the messages.'),
         ('Check the account', 'FC or non-FC? FCF or DHJ? Within or beyond MCT? 1st job or back-to-back issue jobs?'),
         ('Start with the baseline', 'Empathy, refund eligibility, Premium fee, pairing ban, CP coaching, flag.'),
         ('Pick the offer', 'Use the Quick Reference: voucher, free month, ETF spiel.'),
         ('Plan the next step', 'Within MCT: Exit Strategy. Beyond MCT: cancel, or set the TR.'),
         ('Draft your reply', 'Write the response to C, then share your path with the group.')]
sh = (3.3 - 0.07 * (len(steps) - 1)) / len(steps)
for i, (head, body) in enumerate(steps):
    y = 1.4 + i * (sh + 0.07)
    box(4.25, y, 5.3, sh, TEALSOFT if i < 5 else GOLDSOFT, None)
    text(box(4.38, y + (sh - 0.28) / 2, 0.28, 0.28, TEAL if i < 5 else GOLD, adj=0.5),
         [[(str(i + 1), True, WHITE, 9)]], align=PP_ALIGN.CENTER)
    text(tb(4.8, y, 4.65, sh), [[(head + ':  ', True, TEAL if i < 5 else GOLD, 9.5), (body, False, INK, 9)]])
pt.notes_slide.notes_text_frame.text = (
    'Practice ticket activity. Share a live PCQ cancel-intent ticket on screen and investigate it together. '
    'Have trainees note the ticket, customer and job IDs, the model (FCF or DHJ), and whether the customer is within or beyond MCT. '
    'Walk the steps in order: confirm PCQ is the real reason for the cancel intent; check the account (FC or not, model, MCT, first job or back-to-back issue jobs) and pick the scenario from the decision tree; apply the baseline responses; choose the offer from the Quick Reference; plan what happens if C rejects (Exit Strategy within MCT, cancel or set a TR beyond MCT); then have each trainee draft the reply and compare paths as a group.')

buf = io.BytesIO(); prs.save(buf)
ns = {}
anim = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'day3-platform-mechanics', 'add_story_anim.py')
exec(open(anim).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == kc.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        groups = {}
        for sid, k in re.findall(r'<p:cNvPr id="(\d+)" name="step(\d+)[^"]*"', x): groups.setdefault(int(k), []).append(sid)
        cid, clicks, bld = 3, [], []
        for k in sorted(groups):
            inner = ''
            for j, sid in enumerate(groups[k]):
                inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if j == 0 else 'withEffect'); cid += 3
                bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
            clicks.append(ns['click'](cid, inner)); cid += 2
        x = x.replace('</p:sld>', ns['timing'](clicks, bld) + '</p:sld>')
        data = x.encode('utf8'); print('clicks:', len(clicks))
    zout.writestr(item, data)
zout.close()
