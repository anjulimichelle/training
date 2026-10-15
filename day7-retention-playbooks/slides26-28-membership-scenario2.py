"""Day 7 deck, Membership Retention Scenario 2 (cannot afford, no indication of extreme hardship).

Slide 26: overview to tell extreme financial hardship (skip retention: 50%-off ETF exit, then
          full waiver) from an affordability issue (retention process), with sample distinctions.
New 27:   flow chart for ETF applies (Legacy DHJ within MCT, FCF within MCT with a completed job).
New 28:   flow chart for ETF does not apply, with the "When the TR is triggered" pop-up.
Helpers come from slides23-25-membership-scenario1.py.

Usage: python3 slides26-28-membership-scenario2.py in.pptx out.pptx
"""
import copy, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
S = list(prs.slides)
assert txt(S[25].shapes[1]) == "Scenario 2: can't afford, no extreme hardship", txt(S[25].shapes[1])
s = None
s8_notes = "The same Timed Reminder rules apply in every beyond-MCT branch. If the next MF is more than 5 days away, set the reminder at least 3 days before it. If it is 5 days or less, give 1 free month first and set the reminder 3 days before the new MF date. When it triggers with no reply and no engagement, cancel FC; the auto comms are enough unless an unresolved pain point needs a manual email. If there is no reply but the customer is using the service, such as a new booking, don't cancel. Confirm with them, push the due date out 1 to 2 weeks, and set a new reminder 3 days before it."


def clear(slide):
    sl = list(slide.shapes)
    keep = {id(e._element) for e in sl[:3]} | {id(e._element) for e in sl if e.is_placeholder}
    for sh in sl:
        if id(sh._element) not in keep:
            sh._element.getparent().remove(sh._element)


def set_text(sh, t, size=None):
    runs = sh.text_frame.paragraphs[0].runs
    runs[0].text = t
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in sh.text_frame.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    if size:
        runs[0].font.size = Pt(size)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill_paras(tf, paras, align):
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, font, size, bold, color in runs:
            r = p.add_run(); r.text = t
            f = r.font; f.name = font; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tf


def hang(tf):
    for p in tf.paragraphs:
        lead = p.runs[0].text if p.runs else ''
        if lead.startswith('•'):
            p._p.get_or_add_pPr().set('marL', str(11 * 12700)); p._p.get_or_add_pPr().set('indent', str(-11 * 12700))
        elif lead.startswith('–'):
            p._p.get_or_add_pPr().set('marL', str(24 * 12700)); p._p.get_or_add_pPr().set('indent', str(-10 * 12700))


def box(x, y, wd, h, fill, line, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, lw=1):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(wd), Pt(h))
    strip(sh)
    sh.adjustments[0] = min(0.3, 8 / h)
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(7); tf.margin_top = tf.margin_bottom = Pt(3)
    fill_paras(tf, paras, align)
    hang(tf)
    return sh


def text(x, y, wd, h, paras, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(wd), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_paras(tf, paras, align)
    return tb


def arrow(x1, y1, x2, y2, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    strip(c)
    c.line.color.rgb = RGBColor.from_string(MUTED); c.line.width = Pt(1)
    if head:
        ln = c._element.find('.//' + qn('a:ln'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
    return c


P = lambda t, size=9, bold=True, color=WHITE: (t, 'Play', size, bold, color)
A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)
B = lambda t, size=8.5: [('•  ', 'Arial', size, True, TEAL), A(t, size)]
B2 = lambda t, size=8.5: [('–  ', 'Arial', size, False, MUTED), A(t, size)]


A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)
B = lambda t, size=8.5: [('•  ', 'Arial', size, True, TEAL), A(t, size)]
B2 = lambda t, size=8.5: [('–  ', 'Arial', size, False, MUTED), A(t, size)]

def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def leave(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="1" presetClass="exit" presetSubtype="0" fill="hold" grpId="1" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>'
            f'</p:childTnLst></p:cTn></p:par>')


def click(cid, inner):
    return (f'<p:par><p:cTn id="{cid}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{cid+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


def apply_timing(pop):
    ids = [sh.shape_id for sh in pop]
    cid = 5
    c1 = ''
    for i, spid in enumerate(ids):
        c1 += effect(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
    start2 = cid; cid += 2
    c2 = ''
    for i, spid in enumerate(ids):
        c2 += leave(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 2
    bld = ''.join(f'<p:bldP spid="{i}" grpId="0"/><p:bldP spid="{i}" grpId="1"/>' for i in ids
                  if s.shapes._spTree.xpath(f'.//p:sp[p:nvSpPr/p:cNvPr/@id="{i}"]'))
    NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           + click(3, c1) + click(start2, c2) +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
    old = s._element.find(qn('p:timing'))
    if old is not None:
        s._element.remove(old)
    t_el = etree.fromstring(tim)
    ext = s._element.find(qn('p:extLst'))
    (ext.addprevious(t_el) if ext is not None else s._element.append(t_el))



def add_note(slide, text_):
    nt = slide.notes_slide.notes_text_frame
    if text_[:40] in nt.text:
        return
    p = copy.deepcopy(nt.paragraphs[-1]._p)
    for r in p.findall(qn('a:r'))[1:]:
        p.remove(r)
    p.findall(qn('a:r'))[0].find(qn('a:t')).text = text_
    nt.paragraphs[-1]._p.addnext(p)


# ======================= slide 26: overview =======================
s = S[25]
clear(s)
set_text(s.shapes[1], 'Scenario 2: Cannot Afford, No Indication of Extreme Hardship', 18)
set_text(s.shapes[2], 'First ask: can they genuinely not afford it at all, or does the price just feel too high?')
# mini flow on the left
LC, RC = 111, 281
box(32, 102, 328, 34, AMBER_PALE, AMBER, [[P('Re-read C\'s message: job loss, bankruptcy, homelessness,', 8.5, True, AMBER)],
                                         [P('or anything comparably severe?', 8.5, True, AMBER)]])
arrow(196, 136, 196, 143, head=False)
arrow(LC, 143, RC, 143, head=False)
arrow(LC, 143, LC, 151); arrow(RC, 143, RC, 151)
text(LC + 5, 140, 30, 9, [[P('YES', 7.5, True, INK)]]); text(RC + 5, 140, 30, 9, [[P('NO', 7.5, True, INK)]])
box(32, 151, 158, 22, AMBER, None, [[P('Extreme financial hardship', 8.5)]])
arrow(LC, 173, LC, 180)
box(32, 180, 158, 36, WHITE, AMBER, [[P('Skip retention.', 8.5, True, AMBER)], [A('Exit offer: 50% off the ETF', 8)]])
arrow(LC, 216, LC, 223)
box(LC - 45, 223, 90, 18, AMBER_PALE, AMBER, [[P('Declined?', 8.5, True, AMBER)]])
arrow(LC, 241, LC, 248)
text(LC + 5, 241, 30, 8, [[P('YES', 7, True, INK)]])
box(32, 248, 158, 22, AMBER_PALE, AMBER, [[P('Full ETF waiver', 9, True, AMBER)]])
text(32, 274, 158, 22, [[A('ETF Waivers > Bankruptcy / extreme hardship. No proof required.', 7.5, False, MUTED)]], PP_ALIGN.CENTER)
box(202, 151, 158, 22, TEAL, None, [[P('Affordability issue', 8.5)]])
arrow(RC, 173, RC, 180)
box(202, 180, 158, 36, TEAL_PALE, TEAL, [[P('Proceed to retention', 8.5, True, TEAL)], [A('the process is on the next slides', 8)]])
text(202, 222, 158, 34, [[A('Affordability concerns alone are not a reason to stop attempting retention.', 7.5, False, MUTED)]], PP_ALIGN.CENTER)
# sample distinctions on the right
text(375, 102, 300, 12, [[P('SAMPLE DISTINCTIONS', 8.5, True, AMBER)]])
samples = [
    ('"I just don\'t think $59 a month is worth it for me anymore"', 'Scenario 2', 'General cost sensitivity, no hardship signal.', TEAL, TEAL_PALE),
    ('"I lost my job last month and can\'t keep up with this"', 'Extreme financial hardship', 'Skip to the 50%-off exit offer. Declined: full ETF waiver.', AMBER, AMBER_PALE),
    ('"Money\'s tight right now, it\'s been a rough few months"', 'Still Scenario 2', 'Vague strain without a named event isn\'t confirmed hardship. Don\'t upgrade it on your own.', TEAL, TEAL_PALE),
]
for i, (quote, verdict, why, col, pale) in enumerate(samples):
    y = 118 + i * 64
    box(375, y, 312, 58, pale, col, [
        [('', 'Arial', 8.5, False, INK), (quote, 'Arial', 8.5, False, INK)],
        [P('→  ' + verdict + ':  ', 8.5, True, col), A(why, 8)]], align=PP_ALIGN.LEFT)
for r in s.shapes[-1].text_frame.paragraphs[0].runs: pass
box(32, 312, 655, 26, WHITE, LINE, [
    [P('Why the split:  ', 8.5, True, TEAL), A('hardship gets a deliberately generous exit. Running ordinary "I can\'t afford this" through it would hand out ETF waivers on request.', 8)]],
    align=PP_ALIGN.LEFT)
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith('"'):
        for r in sh.text_frame.paragraphs[0].runs:
            r.font.italic = True
s.notes_slide.notes_text_frame.text = (
    'Before you touch any offer, answer one question honestly: is this customer in a financial situation where they genuinely can\'t afford the membership at all, or are they simply saying the price feels too high? '
    'Those are different problems, and the policy treats them differently on purpose. Financial hardship (job loss, bankruptcy, homelessness) skips retention entirely and goes straight to a 50%-off-then-full-waiver exit, with no proof required, because we take people at their word when the stakes are that high. '
    'If ordinary "I can\'t afford this" complaints went through that path, we would effectively be handing out ETF waivers on request, which would undercut the reason the ETF and MF exist: they are what makes offering a free or discounted first cleaning sustainable (see ForeverClean Membership). '
    'This scenario is for the much more common case: "this is too expensive" or "I can\'t afford this" with no major life event attached. Per the General Retention Playbook, affordability concerns alone are not a reason to stop attempting retention. '
    'A customer who can\'t afford $59/month might stay for $44/month, so attempt retention with the tools built for a cost complaint (MF reduction) rather than waiving anything. '
    'Before anything else, re-read the customer\'s message for job loss, bankruptcy, homelessness or anything comparably severe. If it\'s there, stop: go to ETF Waivers > Bankruptcy / extreme hardship instead. '
    'Examples: "I just don\'t think $59 a month is worth it for me anymore" is Scenario 2 (general cost sensitivity, no hardship signal). "I lost my job last month and can\'t keep up with this" is extreme financial hardship: skip straight to the 50%-off exit offer, and if declined, a full ETF waiver. '
    '"Money\'s tight right now, it\'s been a rough few months" is still Scenario 2: vague financial strain without a named event isn\'t confirmed hardship, so don\'t upgrade it on your own.')

# ======================= new slide 27: ETF applies =======================
ref = S[25]
s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
num = [x for x in ref.shapes if x.is_placeholder][0]
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); list(lst)[25].addnext(el)
set_text(s.shapes[1], 'Scenario 2: ETF applies', 28)
set_text(s.shapes[2], 'The process for Legacy DHJ within MCT  |  FCF within MCT with a completed job')
CX, LX, RX = 360, 181, 538
box(CX - 200, 99, 400, 22, TEAL, None,
    [[P('FC retention – Cannot afford, no indication of extreme hardship', 8.5)], [P('(ETF applies)', 7.5, False, TEAL_PALE)]])
arrow(CX, 121, CX, 127)
box(CX - 90, 127, 180, 20, AMBER_PALE, AMBER, [[P('Is the C still retainable?', 9, True, AMBER)]])
arrow(CX, 147, CX, 153, head=False)
arrow(LX, 153, RX, 153, head=False)
arrow(LX, 153, LX, 161); arrow(RX, 153, RX, 161)
text(LX + 5, 151, 30, 9, [[P('YES', 7.5, True, INK)]]); text(RX + 5, 151, 30, 9, [[P('NO', 7.5, True, INK)]])
box(LX - 149, 161, 298, 44, TEAL_PALE, TEAL, [
    B('Explain FC details (MF + ETF)', 8), B('Offer an MF reduction + free months (if applicable)', 8),
    B('Apply free weeks/months to move the billing date as a one-time courtesy', 8)], align=PP_ALIGN.LEFT)
arrow(LX, 205, LX, 212)
box(LX - 140, 212, 280, 20, INK, None, [[P('C rejects the resolution and insists on cancelling the FC', 8, False)]])
arrow(LX, 232, LX, 239)
box(LX - 75, 239, 150, 18, AMBER_PALE, AMBER, [[P('Can we reduce further?', 8.5, True, AMBER)]])
arrow(LX + 75, 248, 365, 248, head=False)
text(LX + 80, 238, 30, 9, [[P('NO', 7.5, True, INK)]])
arrow(LX, 257, LX, 264)
text(LX + 5, 256, 30, 8, [[P('YES', 7.5, True, INK)]])
box(LX - 140, 264, 280, 30, AMBER_PALE, AMBER, [[A('Offer it directly if it is within your authority ($10-15 off the MF), or raise it to CSQ for a further reduction', 8)]])
arrow(LX, 294, LX, 301)
box(LX - 140, 301, 280, 20, INK, None, [[P('C rejects the resolution and insists on cancelling the FC', 8, False)]])
arrow(LX + 140, 311, 365, 311, head=False)
arrow(365, 311, 365, 180, head=False)
arrow(365, 180, RX - 148, 180)
box(RX - 148, 161, 296, 46, TEAL, None, [
    [P('EXIT Strategy', 9.5)], [P('MCT or ETF reduction', 8, False, TEAL_PALE)],
    [P('Apply free weeks/months to move the billing date as a one-time courtesy', 7.5, False, TEAL_PALE)]])
arrow(RX, 207, RX, 214)
box(RX - 140, 214, 280, 20, INK, None, [[P('C rejects the resolution and insists on cancelling the FC', 8, False)]])
arrow(RX, 234, RX, 241)
box(RX - 75, 241, 150, 22, AMBER_PALE, AMBER, [[P('GRACEFUL CLOSURE', 9, True, AMBER)]])
box(RX - 148, 272, 296, 36, WHITE, LINE, [[P('Exit at the lowest cost:  ', 8, True, TEAL),
    A('MCT reduction (up to 2 months) or ETF reduction (up to 50%), whichever is cheaper. Max it out before Graceful Closure.', 7.5)]],
    align=PP_ALIGN.LEFT)
box(RX - 148, 314, 296, 24, AMBER_PALE, AMBER, [[P('Your authority:  ', 8, True, AMBER),
    A('$10-15 off the MF. More needs CSQ approval.', 7.5)]], align=PP_ALIGN.LEFT)
s.notes_slide.notes_text_frame.text = (
    'Flow (ETF applies: Legacy DHJ within MCT, or FCF within MCT with a completed job). Assess retainability first. '
    'Retainable: explain the FC details (MF and ETF), offer an MF reduction plus free months if applicable, and move the billing date with free weeks/months as a one-time courtesy. '
    'If C rejects and insists on cancelling, ask whether we can reduce further: if yes, offer it directly when it is within your authority ($10 to 15 off), or raise it to CSQ for more. If they still reject, or we cannot reduce further, go to Exit Strategy. '
    'Not retainable: go straight to Exit Strategy, using whichever of MCT reduction (up to 2 months) or ETF reduction (up to 50%) costs the customer less, plus free weeks/months to move the billing date as a one-time courtesy. Max out the lever, then Graceful Closure if they still reject.')

# ======================= new slide 28: ETF does not apply =======================
ref = S[25]
s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
num = [x for x in ref.shapes if x.is_placeholder][0]
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); list(lst)[26].addnext(el)
set_text(s.shapes[1], 'Scenario 2: ETF does not apply')
set_text(s.shapes[2], 'Legacy DHJ beyond MCT  |  FCF with no completed job (within or beyond MCT)  |  FCF beyond MCT with a completed job', 11)
s.notes_slide.notes_text_frame.text = (
    'Flow (ETF does not apply: Legacy DHJ beyond MCT, FCF with no completed job, or FCF beyond MCT with a completed job). Is the customer still retainable? If not, address any pain points and cancel the FC. '
    'If retainable, attempt to retain and offer an MF reduction plus free months if applicable, explicitly confirm no further MF will be charged unless they agree to continue, and set a TR. '
    '$10 to 15 off the MF is within your authority; more needs CSQ approval. If C rejects the resolution and insists on cancelling, cancel the FC.')
L, LW = 32, 400
CX = L + LW / 2
R, RW = 456, 231
RX = R + RW / 2

box(L, 98, LW, 22, TEAL, None,
    [[P('FC retention – Cannot afford, no indication of extreme hardship', 8.5)], [P('(ETF does not apply)', 8, False, TEAL_PALE)]])
arrow(CX, 120, CX, 127)
box(CX - 110, 127, 220, 26, AMBER_PALE, AMBER,
    [[P('Is the C still retainable?', 9.5, True, AMBER)],
     [P('Use the Five Key Factors', 7.5, False, AMBER)]])
arrow(CX + 110, 140, R, 140)
text(CX + 120, 129, 30, 10, [[P('NO', 8, True, INK)]])
box(R, 125, RW, 30, TEAL_PALE, TEAL,
    [[P('Address any pain points and cancel the FC', 9, False, TEAL)]])
box(R, 164, RW, 50, WHITE, LINE,
    [[P('Assessing retainability:  ', 8.5, True, TEAL),
      A('use the General Retention Playbook Five Key Factors. Weigh what is actually present, not a factor count.', 8)]],
    align=PP_ALIGN.LEFT)

box(R, 222, RW, 40, AMBER_PALE, AMBER,
    [[P('Your authority:  ', 8.5, True, AMBER), A('$10-15 off the MF is within your authority. More needs CSQ approval.', 8)]],
    align=PP_ALIGN.LEFT)
arrow(CX, 153, CX, 163)
text(CX + 5, 153, 30, 10, [[P('YES', 8, True, INK)]])
box(L, 163, LW, 88, TEAL_PALE, TEAL, [
    B('Attempt to retain C and offer an MF reduction + free months (if applicable: unused paid MFs, not more than asked, not more than 3)'),
    B('Explicitly confirm no further MF will be charged unless they agree to continue'),
    B('Set a TR 3 days before the next MF date. If the next MF is within 5 days, issue a free month first'),
], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT)

# split into two branches
LB, RB = L + 100, L + LW - 100
arrow(CX, 251, CX, 260, head=False)
arrow(LB, 260, RB, 260, head=False)
arrow(LB, 260, LB, 269)
arrow(RB, 260, RB, 269)
box(LB - 92, 269, 184, 26, INK, None, [[P('C rejects the resolution', 9, False)]])
arrow(LB, 295, LB, 304)
box(LB - 70, 304, 140, 24, AMBER_PALE, AMBER, [[P('Cancel the FC', 9.5, True, AMBER)]])
box(RB - 92, 269, 184, 26, INK, None, [[P('No reply, but the TR was triggered', 9, False)]])
arrow(RB, 295, RB, 304)
box(RB - 70, 304, 140, 24, TEAL, None, [[P('What to do next  ▸', 9, True)]])
text(RB - 92, 331, 184, 10, [[A('click to see both outcomes', 7, False, MUTED)]], PP_ALIGN.CENTER)

# ---- pop-up: when the TR is triggered (from slide 8) ----
pop = []
card = box(110, 102, 500, 218, WHITE, TEAL, [[A('')]], lw=1.5)
sp = card._element.find(qn('p:spPr'))
eff = etree.SubElement(sp, qn('a:effectLst'))
sh_ = etree.SubElement(eff, qn('a:outerShdw'), blurRad='76200', dist='25400', dir='5400000', algn='t', rotWithShape='0')
clr = etree.SubElement(sh_, qn('a:srgbClr'), val='000000'); etree.SubElement(clr, qn('a:alpha'), val='25000')
pop.append(card)
pop.append(box(250, 114, 220, 30, AMBER_PALE, AMBER, [[P('When the TR is triggered', 11, True, AMBER)]]))
PL, PR = 245, 475
pop.append(arrow(360, 144, 360, 154, head=False))
pop.append(arrow(PL, 154, PR, 154, head=False))
pop.append(arrow(PL, 154, PL, 164))
pop.append(arrow(PR, 154, PR, 164))
pop.append(box(PL - 110, 164, 220, 30, AMBER, None, [[P('No reply and no new bookings', 9.5)]]))
pop.append(box(PL - 110, 200, 220, 84, AMBER_PALE, AMBER,
               [[A('Cancel the FC.', 9, True)],
                [A("Send comms only if there's an unresolved pain point the auto comms wouldn't cover.", 9)]]))
pop.append(box(PR - 110, 164, 220, 30, TEAL, None, [[P('No reply, but C is using the service', 9.5)]]))
pop.append(box(PR - 110, 200, 220, 84, TEAL_PALE, TEAL, [
    B("Don't cancel the FC", 9), B('Contact C to confirm the cancellation', 9),
    B('Add 1–2 weeks to move the due date, and set a new TR', 9)], align=PP_ALIGN.LEFT))
pop.append(text(130, 296, 460, 12, [[A('Click to close', 7.5, False, MUTED)]], PP_ALIGN.RIGHT))



apply_timing(pop)
add_note(s, 'TRAINER: click once to open the "When the TR is triggered" pop-up, click again to close it. ' + s8_notes)

prs.save(sys.argv[2]); print('saved', sys.argv[2])
