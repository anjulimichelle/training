"""Day 7 deck: add the ETF column from slide 6 into slide 5's offers table, then delete slide 6.

Usage: python3 slide5-6-etf-merge.py in.pptx out.pptx
"""
import copy, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.oxml.ns import qn

prs = Presentation(sys.argv[1])
s5, s6 = prs.slides[4], prs.slides[5]
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
assert txt(s5.shapes[1]).startswith('Offers by situation: voucher and free month')
assert txt(s6.shapes[1]) == 'Offers by situation: ETF'

t5 = next(sh for sh in s5.shapes if sh.has_table)
t6 = next(sh for sh in s6.shapes if sh.has_table)
rows5, rows6 = list(t5.table.rows), list(t6.table.rows)
assert len(rows5) == len(rows6) == 7
etf = [r.cells[1].text for r in rows6]
assert etf[0] == 'ETF'

tbl = t5._element.find('.//' + qn('a:tbl'))
grid = tbl.find(qn('a:tblGrid'))
cols = grid.findall(qn('a:gridCol'))
total = sum(int(c.get('w')) for c in cols)
newcol = copy.deepcopy(cols[-1]); cols[-1].addnext(newcol)
widths = [0.21, 0.255, 0.27, 0.265]
acc = 0
for i, c in enumerate(grid.findall(qn('a:gridCol'))):
    w = int(total * widths[i]) if i < 3 else total - acc
    c.set('w', str(w)); acc += w


def set_cell_text(tc, text):
    body = tc.find(qn('a:txBody'))
    ps = body.findall(qn('a:p'))
    for p in ps[1:]:
        body.remove(p)
    p = ps[0]
    runs = p.findall(qn('a:r'))
    for r in runs[1:]:
        p.remove(r)
    runs[0].find(qn('a:t')).text = text


for i, tr in enumerate(tbl.findall(qn('a:tr'))):
    last = tr.findall(qn('a:tc'))[-1]
    new = copy.deepcopy(last)
    set_cell_text(new, etf[i])
    last.addnext(new)
    tr.set('h', str(int((22 if i == 0 else 30) * 12700)))
    # body text 8pt, header 9pt; cell margins a bit tighter
    for tc in tr.findall(qn('a:tc')):
        for r in tc.iter(qn('a:r')):
            if r.find(qn('a:rPr')) is None:
                r.insert(0, r.makeelement(qn('a:rPr'), {'lang': 'en'}))
        for tag in ('a:rPr', 'a:endParaRPr'):
            for r in tc.iter(qn(tag)):
                r.set('sz', '900' if i == 0 else '800')
        pr = tc.find(qn('a:tcPr'))
        pr.set('marT', '19050'); pr.set('marB', '19050')

# title and Cap bar
title = s5.shapes[1]
title.text_frame._txBody.findall('.//' + qn('a:t'))[0].text = 'Offers by situation'
sub = s5.shapes[2]
sub.text_frame._txBody.findall('.//' + qn('a:t'))[0].text = 'Voucher, free month and ETF for PCQ with cancel intent.'
for t in sub.text_frame._txBody.findall('.//' + qn('a:t'))[1:]:
    t.text = ''
for t in title.text_frame._txBody.findall('.//' + qn('a:t'))[1:]:
    t.text = ''
sl = list(s5.shapes)
k = next(i for i, sh in enumerate(sl) if txt(sh).startswith('Cap:'))
dy = Pt(344) - sl[k - 2].top
for sh in sl[k - 2:k + 1]:
    sh.top = sh.top + dy

# notes: append slide 6's ETF notes
n5, n6 = s5.notes_slide.notes_text_frame, s6.notes_slide.notes_text_frame.text
if 'The ETF column' not in n5.text:
    p = copy.deepcopy(n5.paragraphs[-1]._p)
    for r in p.findall(qn('a:r'))[1:]:
        p.remove(r)
    p.findall(qn('a:r'))[0].find(qn('a:t')).text = n6
    n5.paragraphs[-1]._p.addnext(p)

# delete slide 6
lst = prs.slides._sldIdLst
sld = lst[5]
rid = sld.get(qn('r:id'))
lst.remove(sld)
prs.part.drop_rel(rid)

prs.save(sys.argv[2]); print('saved', sys.argv[2])
