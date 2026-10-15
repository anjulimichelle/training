"""Day 7 deck slide 30 (Membership Scenario 4): customer has issues with membership terms.
Brief overview (MF too expensive or MCT too long; a reduction, not a refund) plus a flow chart:
retainable -> MF or MCT reduction -> reduce further? -> Exit Strategy -> Graceful Closure.
Helpers come from slide29-membership-scenario3-pause.py.

Usage: python3 slide30-membership-scenario4-terms.py in.pptx out.pptx
"""
import copy, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
S = list(prs.slides)
assert txt(S[29].shapes[1]) == 'Scenario 4: issues with the terms', txt(S[29].shapes[1])
s = None
s8_notes = "The same Timed Reminder rules apply in every beyond-MCT branch. If the next MF is more than 5 days away, set the reminder at least 3 days before it. If it is 5 days or less, give 1 free month first and set the reminder 3 days before the new MF date. When it triggers with no reply and no engagement, cancel FC; the auto comms are enough unless an unresolved pain point needs a manual email. If there is no reply but the customer is using the service, such as a new booking, don't cancel. Confirm with them, push the due date out 1 to 2 weeks, and set a new reminder 3 days before it."


def clear(slide):
    sl = list(slide.shapes)
    keep = {id(e._element) for e in sl[:3]} | {id(e._element) for e in sl if e.is_placeholder}
    for sh in sl:
        if id(sh._element) not in keep:
            sh._element.getparent().remove(sh._element)


def set_text(sh, t, size=None):
    runs = sh.text_frame.paragraphs[0].runs
    runs[0].text = t
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in sh.text_frame.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    if size:
        runs[0].font.size = Pt(size)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill_paras(tf, paras, align):
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, font, size, bold, color in runs:
            r = p.add_run(); r.text = t
            f = r.font; f.name = font; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tf


def hang(tf):
    for p in tf.paragraphs:
        lead = p.runs[0].text if p.runs else ''
        if lead.startswith('•'):
            p._p.get_or_add_pPr().set('marL', str(11 * 12700)); p._p.get_or_add_pPr().set('indent', str(-11 * 12700))
        elif lead.startswith('–'):
            p._p.get_or_add_pPr().set('marL', str(24 * 12700)); p._p.get_or_add_pPr().set('indent', str(-10 * 12700))


def box(x, y, wd, h, fill, line, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, lw=1):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(wd), Pt(h))
    strip(sh)
    sh.adjustments[0] = min(0.3, 8 / h)
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(7); tf.margin_top = tf.margin_bottom = Pt(3)
    fill_paras(tf, paras, align)
    hang(tf)
    return sh


def text(x, y, wd, h, paras, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(wd), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_paras(tf, paras, align)
    return tb


def arrow(x1, y1, x2, y2, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    strip(c)
    c.line.color.rgb = RGBColor.from_string(MUTED); c.line.width = Pt(1)
    if head:
        ln = c._element.find('.//' + qn('a:ln'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
    return c


P = lambda t, size=9, bold=True, color=WHITE: (t, 'Play', size, bold, color)
A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)
B = lambda t, size=8.5: [('•  ', 'Arial', size, True, TEAL), A(t, size)]
B2 = lambda t, size=8.5: [('–  ', 'Arial', size, False, MUTED), A(t, size)]


A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)
B = lambda t, size=8.5: [('•  ', 'Arial', size, True, TEAL), A(t, size)]
B2 = lambda t, size=8.5: [('–  ', 'Arial', size, False, MUTED), A(t, size)]


def add_note(slide, text_):
    nt = slide.notes_slide.notes_text_frame
    if text_[:40] in nt.text:
        return
    p = copy.deepcopy(nt.paragraphs[-1]._p)
    for r in p.findall(qn('a:r'))[1:]:
        p.remove(r)
    p.findall(qn('a:r'))[0].find(qn('a:t')).text = text_
    nt.paragraphs[-1]._p.addnext(p)


s = S[29]
clear(s)
set_text(s.shapes[1], 'Scenario 4: Customer Has Issues With Membership Terms', 19)
set_text(s.shapes[2], 'The monthly fee feels too expensive, or the minimum commitment term feels too long.')
box(32, 100, 655, 22, WHITE, LINE, [[P('A price objection gets a reduction, not a refund.   ', 8.5, True, TEAL),
    A('Anchor on the dollar figure ("$15 off", not "25%"). Check escalation and payment history.', 8)]],
    align=PP_ALIGN.LEFT)

CX, LX, RX = 360, 181, 538
box(CX - 120, 128, 240, 18, TEAL, None, [[P('FC retention – Issues with FC terms', 9)]])
arrow(CX, 146, CX, 152)
box(CX - 90, 152, 180, 18, AMBER_PALE, AMBER, [[P('Is the C still retainable?', 9, True, AMBER)]])
arrow(CX, 170, CX, 176, head=False)
arrow(LX, 176, RX, 176, head=False)
arrow(LX, 176, LX, 184); arrow(RX, 176, RX, 184)
text(LX + 5, 174, 30, 9, [[P('YES', 7.5, True, INK)]]); text(RX + 5, 174, 30, 9, [[P('NO', 7.5, True, INK)]])
# two objections under YES
ML, MR = 32, 186
box(ML, 184, 144, 16, AMBER, None, [[P('MF too expensive', 8.5)]])
box(MR, 184, 144, 16, TEAL, None, [[P('MCT too long', 8.5)]])
common = [B('Explain FC terms (MF + ETF)', 7.5), B('Offer free months (max 3)', 7.5),
          B('Apply free weeks/months to move the billing date as a one-time courtesy', 7.5)]
box(ML, 202, 144, 62, AMBER_PALE, AMBER, [B('MF reduction', 7.5)] + common, align=PP_ALIGN.LEFT)
box(MR, 202, 144, 62, TEAL_PALE, TEAL, [B('MCT reduction', 7.5)] + common, align=PP_ALIGN.LEFT)
arrow(ML + 72, 264, ML + 72, 268, head=False)
arrow(MR + 72, 264, MR + 72, 268, head=False)
arrow(ML + 72, 268, MR + 72, 268, head=False)
arrow(LX, 268, LX, 274)
box(LX - 140, 274, 280, 18, INK, None, [[P('C rejects the resolution and insists on cancelling the FC', 8, False)]])
arrow(LX, 292, LX, 298)
box(LX - 75, 298, 150, 16, AMBER_PALE, AMBER, [[P('Can we reduce further?', 8.5, True, AMBER)]])
arrow(LX + 75, 306, 365, 306, head=False)
text(LX + 80, 296, 30, 9, [[P('NO', 7.5, True, INK)]])
arrow(LX, 314, LX, 320)
text(LX + 5, 313, 30, 8, [[P('YES', 7.5, True, INK)]])
box(LX - 140, 320, 280, 24, AMBER_PALE, AMBER, [[A('Offer it directly if within your authority ($10-15), or raise it to CSQ for a further reduction', 7.5)]])
arrow(LX, 344, LX, 349)
box(LX - 140, 349, 280, 18, INK, None, [[P('C rejects the resolution and insists on cancelling the FC', 8, False)]])
arrow(LX + 140, 358, 365, 358, head=False)
arrow(365, 358, 365, 203, head=False)
arrow(365, 203, RX - 148, 203)
# NO / exit path
box(RX - 148, 184, 296, 44, TEAL, None, [
    [P('EXIT Strategy', 9.5)], [P('MCT or ETF reduction', 8, False, TEAL_PALE)],
    [P('Apply free weeks/months to move the billing date as a one-time courtesy', 7.5, False, TEAL_PALE)]])
arrow(RX, 228, RX, 235)
box(RX - 140, 235, 280, 18, INK, None, [[P('C rejects the resolution and insists on cancelling the FC', 8, False)]])
arrow(RX, 253, RX, 260)
box(RX - 75, 260, 150, 20, AMBER_PALE, AMBER, [[P('GRACEFUL CLOSURE', 9, True, AMBER)]])
box(RX - 148, 292, 296, 30, WHITE, LINE, [[P('Exit at the lowest cost:  ', 8, True, TEAL),
    A('pick the cheaper of MCT reduction or ETF reduction.', 7.5)]], align=PP_ALIGN.LEFT)
box(RX - 148, 328, 296, 30, AMBER_PALE, AMBER, [[P('Before Graceful Closure:  ', 8, True, AMBER),
    A('the ETF reduction offer must be maxed out to the full 50%.', 7.5)]], align=PP_ALIGN.LEFT)

add_note(s, 'Scenario 4 covers two term objections: the monthly fee feels too expensive, or the minimum commitment term feels too long. '
            'Flow: assess retainability. Retainable, MF too expensive: offer an MF reduction; MCT too long: offer an MCT reduction. Either way, explain the FC terms (MF and ETF), offer free months (max 3), and move the billing date with free weeks/months as a one-time courtesy. '
            'If C rejects and insists on cancelling, can we reduce further? Yes: offer it directly if within your authority ($10 to 15 off), or raise it to CSQ. If they still reject, or no further reduction is possible, go to Exit Strategy. '
            'Not retainable: Exit Strategy with the cheaper of MCT or ETF reduction, plus free weeks/months to move the billing date as a one-time courtesy. If they reject the exit offer, Graceful Closure, but only after the ETF reduction has been maxed out at 50%.')

prs.save(sys.argv[2]); print('saved', sys.argv[2])
