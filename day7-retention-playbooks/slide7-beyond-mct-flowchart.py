"""Day 7 deck slide 7 (Beyond MCT): replace the tiles with a flow chart in the deck
theme. FC retention PCQ (no ETF) -> still retainable? (NO: address + cancel FC)
-> offer / confirm no further MF / set TR -> two branches: C rejects -> cancel FC,
and No reply but TR triggered, which opens a pop-up with slide 8's two TR outcomes
(shows on click 1, closes on click 2). Slide 8 is then deleted.

Usage: python3 slide7-beyond-mct-flowchart.py in.pptx out.pptx
"""
import copy, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
s, s8 = prs.slides[6], prs.slides[7]
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
assert txt(s.shapes[1]) == 'Beyond MCT', txt(s.shapes[1])
assert txt(s8.shapes[1]) == 'When the Timed Reminder triggers'
s8_notes = s8.notes_slide.notes_text_frame.text

sl = list(s.shapes)
keep = {id(e._element) for e in sl[:3] + [sl[-1]]}
for sh in sl:
    if id(sh._element) not in keep:
        sh._element.getparent().remove(sh._element)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill_paras(tf, paras, align):
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, font, size, bold, color in runs:
            r = p.add_run(); r.text = t
            f = r.font; f.name = font; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tf


def hang(tf):
    for p in tf.paragraphs:
        lead = p.runs[0].text if p.runs else ''
        if lead.startswith('•'):
            p._p.get_or_add_pPr().set('marL', str(11 * 12700)); p._p.get_or_add_pPr().set('indent', str(-11 * 12700))
        elif lead.startswith('–'):
            p._p.get_or_add_pPr().set('marL', str(24 * 12700)); p._p.get_or_add_pPr().set('indent', str(-10 * 12700))


def box(x, y, wd, h, fill, line, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, lw=1):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(wd), Pt(h))
    strip(sh)
    sh.adjustments[0] = min(0.3, 8 / h)
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(7); tf.margin_top = tf.margin_bottom = Pt(3)
    fill_paras(tf, paras, align)
    hang(tf)
    return sh


def text(x, y, wd, h, paras, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(wd), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_paras(tf, paras, align)
    return tb


def arrow(x1, y1, x2, y2, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    strip(c)
    c.line.color.rgb = RGBColor.from_string(MUTED); c.line.width = Pt(1)
    if head:
        ln = c._element.find('.//' + qn('a:ln'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
    return c


P = lambda t, size=9, bold=True, color=WHITE: (t, 'Play', size, bold, color)
A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)
B = lambda t, size=8.5: [('•  ', 'Arial', size, True, TEAL), A(t, size)]
B2 = lambda t, size=8.5: [('–  ', 'Arial', size, False, MUTED), A(t, size)]

L, LW = 32, 400
CX = L + LW / 2
R, RW = 456, 231
RX = R + RW / 2

box(CX - 125, 98, 250, 22, TEAL, None,
    [[P('FC retention – PCQ', 9.5)], [P('(ETF does not apply)', 8, False, TEAL_PALE)]])
arrow(CX, 120, CX, 127)
box(CX - 110, 127, 220, 26, AMBER_PALE, AMBER,
    [[P('Is the C still retainable?', 9.5, True, AMBER)],
     [P('Use the Five Key Factors', 7.5, False, AMBER)]])
arrow(CX + 110, 140, R, 140)
text(CX + 120, 129, 30, 10, [[P('NO', 8, True, INK)]])
box(R, 125, RW, 30, TEAL_PALE, TEAL,
    [[P('Address the PCQ issue and cancel the FC', 9, False, TEAL)]])
box(R, 164, RW, 50, WHITE, LINE,
    [[P('Assessing retainability:  ', 8.5, True, TEAL),
      A('use the General Retention Playbook Five Key Factors. Weigh what is actually present, not a factor count.', 8)]],
    align=PP_ALIGN.LEFT)

arrow(CX, 153, CX, 163)
text(CX + 5, 153, 30, 10, [[P('YES', 8, True, INK)]])
box(L, 163, LW, 88, TEAL_PALE, TEAL, [
    B('Address the PCQ issue and offer a resolution per Quick Reference:'),
    B2('Voucher'),
    B2('Free month per model (unused paid MFs become free months, not more than asked, not more than 3)'),
    B('Explicitly confirm no further MF will be charged unless they agree to continue'),
    B('Set a TR 3 days before the next MF date. If the next MF is within 5 days, issue a free month first'),
], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT)

# split into two branches
LB, RB = L + 100, L + LW - 100
arrow(CX, 251, CX, 260, head=False)
arrow(LB, 260, RB, 260, head=False)
arrow(LB, 260, LB, 269)
arrow(RB, 260, RB, 269)
box(LB - 92, 269, 184, 26, INK, None, [[P('C rejects the resolution', 9, False)]])
arrow(LB, 295, LB, 304)
box(LB - 70, 304, 140, 24, AMBER_PALE, AMBER, [[P('Cancel the FC', 9.5, True, AMBER)]])
box(RB - 92, 269, 184, 26, INK, None, [[P('No reply, but the TR was triggered', 9, False)]])
arrow(RB, 295, RB, 304)
box(RB - 70, 304, 140, 24, TEAL, None, [[P('What to do next  ▸', 9, True)]])
text(RB - 92, 331, 184, 10, [[A('click to see both outcomes', 7, False, MUTED)]], PP_ALIGN.CENTER)

# ---- pop-up: when the TR is triggered (from slide 8) ----
pop = []
card = box(110, 102, 500, 218, WHITE, TEAL, [[A('')]], lw=1.5)
sp = card._element.find(qn('p:spPr'))
eff = etree.SubElement(sp, qn('a:effectLst'))
sh_ = etree.SubElement(eff, qn('a:outerShdw'), blurRad='76200', dist='25400', dir='5400000', algn='t', rotWithShape='0')
clr = etree.SubElement(sh_, qn('a:srgbClr'), val='000000'); etree.SubElement(clr, qn('a:alpha'), val='25000')
pop.append(card)
pop.append(box(250, 114, 220, 30, AMBER_PALE, AMBER, [[P('When the TR is triggered', 11, True, AMBER)]]))
PL, PR = 245, 475
pop.append(arrow(360, 144, 360, 154, head=False))
pop.append(arrow(PL, 154, PR, 154, head=False))
pop.append(arrow(PL, 154, PL, 164))
pop.append(arrow(PR, 154, PR, 164))
pop.append(box(PL - 110, 164, 220, 30, AMBER, None, [[P('No reply and no new bookings', 9.5)]]))
pop.append(box(PL - 110, 200, 220, 84, AMBER_PALE, AMBER,
               [[A('Cancel the FC.', 9, True)],
                [A("Send comms only if there's an unresolved pain point the auto comms wouldn't cover.", 9)]]))
pop.append(box(PR - 110, 164, 220, 30, TEAL, None, [[P('No reply, but C is using the service', 9.5)]]))
pop.append(box(PR - 110, 200, 220, 84, TEAL_PALE, TEAL, [
    B("Don't cancel the FC", 9), B('Contact C to confirm the cancellation', 9),
    B('Add 1–2 weeks to move the due date, and set a new TR', 9)], align=PP_ALIGN.LEFT))
pop.append(text(130, 296, 460, 12, [[A('Click to close', 7.5, False, MUTED)]], PP_ALIGN.RIGHT))


def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def leave(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="1" presetClass="exit" presetSubtype="0" fill="hold" grpId="1" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>'
            f'</p:childTnLst></p:cTn></p:par>')


def click(cid, inner):
    return (f'<p:par><p:cTn id="{cid}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{cid+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


ids = [sh.shape_id for sh in pop]
cid = 5
c1 = ''
for i, spid in enumerate(ids):
    c1 += effect(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
start2 = cid; cid += 2
c2 = ''
for i, spid in enumerate(ids):
    c2 += leave(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 2
bld = ''.join(f'<p:bldP spid="{i}" grpId="0"/><p:bldP spid="{i}" grpId="1"/>' for i in ids
              if s.shapes._spTree.xpath(f'.//p:sp[p:nvSpPr/p:cNvPr/@id="{i}"]'))
NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
       '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
       + click(3, c1) + click(start2, c2) +
       '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
       '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
       '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
old = s._element.find(qn('p:timing'))
if old is not None:
    s._element.remove(old)
t_el = etree.fromstring(tim)
ext = s._element.find(qn('p:extLst'))
(ext.addprevious(t_el) if ext is not None else s._element.append(t_el))

# notes: TR rules move here from slide 8
nt = s.notes_slide.notes_text_frame
if 'same Timed Reminder rules' not in nt.text:
    p = copy.deepcopy(nt.paragraphs[-1]._p)
    for r in p.findall(qn('a:r'))[1:]:
        p.remove(r)
    p.findall(qn('a:r'))[0].find(qn('a:t')).text = (
        'TRAINER: click once to open the "When the TR is triggered" pop-up, click again to close it. ' + s8_notes)
    nt.paragraphs[-1]._p.addnext(p)

# delete slide 8
lst = prs.slides._sldIdLst
sld = lst[7]
rid = sld.get(qn('r:id'))
lst.remove(sld)
prs.part.drop_rel(rid)

prs.save(sys.argv[2]); print('saved', sys.argv[2])
