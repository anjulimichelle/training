"""Day 7 deck slide 29 (Membership Scenario 3): customer wants to pause the membership.
Brief overview (three ways it starts, free months simulate a pause) plus a four-branch
flow chart: cancel but temporary, pause with timeframe, pause without timeframe, CS
notices the signal. Keeps the 6-month cap bar. Helpers come from slides26-28-membership-scenario2.py.

Usage: python3 slide29-membership-scenario3-pause.py in.pptx out.pptx
"""
import copy, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
S = list(prs.slides)
assert txt(S[28].shapes[1]) == 'Scenario 3: pause the membership', txt(S[28].shapes[1])
s = None
s8_notes = "The same Timed Reminder rules apply in every beyond-MCT branch. If the next MF is more than 5 days away, set the reminder at least 3 days before it. If it is 5 days or less, give 1 free month first and set the reminder 3 days before the new MF date. When it triggers with no reply and no engagement, cancel FC; the auto comms are enough unless an unresolved pain point needs a manual email. If there is no reply but the customer is using the service, such as a new booking, don't cancel. Confirm with them, push the due date out 1 to 2 weeks, and set a new reminder 3 days before it."


def clear(slide):
    sl = list(slide.shapes)
    keep = {id(e._element) for e in sl[:3]} | {id(e._element) for e in sl if e.is_placeholder}
    for sh in sl:
        if id(sh._element) not in keep:
            sh._element.getparent().remove(sh._element)


def set_text(sh, t, size=None):
    runs = sh.text_frame.paragraphs[0].runs
    runs[0].text = t
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in sh.text_frame.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    if size:
        runs[0].font.size = Pt(size)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill_paras(tf, paras, align):
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, font, size, bold, color in runs:
            r = p.add_run(); r.text = t
            f = r.font; f.name = font; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tf


def hang(tf):
    for p in tf.paragraphs:
        lead = p.runs[0].text if p.runs else ''
        if lead.startswith('•'):
            p._p.get_or_add_pPr().set('marL', str(11 * 12700)); p._p.get_or_add_pPr().set('indent', str(-11 * 12700))
        elif lead.startswith('–'):
            p._p.get_or_add_pPr().set('marL', str(24 * 12700)); p._p.get_or_add_pPr().set('indent', str(-10 * 12700))


def box(x, y, wd, h, fill, line, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, lw=1):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(wd), Pt(h))
    strip(sh)
    sh.adjustments[0] = min(0.3, 8 / h)
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(7); tf.margin_top = tf.margin_bottom = Pt(3)
    fill_paras(tf, paras, align)
    hang(tf)
    return sh


def text(x, y, wd, h, paras, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(wd), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_paras(tf, paras, align)
    return tb


def arrow(x1, y1, x2, y2, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    strip(c)
    c.line.color.rgb = RGBColor.from_string(MUTED); c.line.width = Pt(1)
    if head:
        ln = c._element.find('.//' + qn('a:ln'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
    return c


P = lambda t, size=9, bold=True, color=WHITE: (t, 'Play', size, bold, color)
A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)
B = lambda t, size=8.5: [('•  ', 'Arial', size, True, TEAL), A(t, size)]
B2 = lambda t, size=8.5: [('–  ', 'Arial', size, False, MUTED), A(t, size)]


A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)
B = lambda t, size=8.5: [('•  ', 'Arial', size, True, TEAL), A(t, size)]
B2 = lambda t, size=8.5: [('–  ', 'Arial', size, False, MUTED), A(t, size)]


def add_note(slide, text_):
    nt = slide.notes_slide.notes_text_frame
    if text_[:40] in nt.text:
        return
    p = copy.deepcopy(nt.paragraphs[-1]._p)
    for r in p.findall(qn('a:r'))[1:]:
        p.remove(r)
    p.findall(qn('a:r'))[0].find(qn('a:t')).text = text_
    nt.paragraphs[-1]._p.addnext(p)


s = S[28]
sl = list(s.shapes)
w = next(i for i, sh in enumerate(sl) if txt(sh).startswith('Cap:'))
cap = sl[w - 2:w + 1]
keep = {id(e._element) for e in sl[:3] + cap} | {id(e._element) for e in sl if e.is_placeholder}
for sh in sl:
    if id(sh._element) not in keep:
        sh._element.getparent().remove(sh._element)
set_text(s.shapes[1], 'Scenario 3: Customer Wants to Pause the Membership', 21)
set_text(s.shapes[2], "There is no literal pause button. We simulate one by issuing free months so the MF doesn't charge.")

# overview strip
box(32, 100, 655, 24, WHITE, LINE, [[P('It starts one of three ways:   ', 8.5, True, TEAL),
    A('C asks to pause directly   •   C asks to cancel, but the real reason is temporary   •   you spot the signal before C has said anything', 8)]],
    align=PP_ALIGN.LEFT)

# flow chart
CX = 360
box(CX - 80, 130, 160, 20, AMBER_PALE, AMBER, [[P('Pausing of Membership', 9.5, True, AMBER)]])
W, G = 154, 12
xs = [32 + i * (W + G) for i in range(4)]
cx = [x + W / 2 for x in xs]
arrow(CX, 150, CX, 156, head=False)
arrow(cx[0], 156, cx[3], 156, head=False)
for c in cx:
    arrow(c, 156, c, 162)
heads = [('C asks to cancel, but the reason may be TEMPORARY', AMBER, AMBER_PALE),
         ('C asks to pause with a timeframe', TEAL, TEAL_PALE),
         ('C asks to pause without a timeframe', TEAL, TEAL_PALE),
         ('CS notices the signal first', INK, WHITE)]
for (h, col, _), x in zip(heads, xs):
    box(x, 162, W, 26, col, None, [[P(h, 8, True)]])
# column 1
box(xs[0], 192, W, 34, AMBER_PALE, AMBER, [[A('Probe for the reason and offer to pause the FC', 8)]])
c1 = cx[0]; a1, b1 = xs[0] + 36, xs[0] + W - 36
arrow(c1, 226, c1, 234, head=False)
arrow(a1, 234, b1, 234, head=False)
arrow(a1, 234, a1, 241); arrow(b1, 234, b1, 241)
box(xs[0], 241, 74, 18, TEAL, None, [[P('C accepts', 8)]])
box(xs[0] + 80, 241, 74, 18, AMBER, None, [[P('C rejects', 8)]])
box(xs[0], 262, 74, 36, TEAL_PALE, TEAL, [[A('Move into the Pause Flow', 7.5)]])
box(xs[0] + 80, 262, 74, 36, AMBER_PALE, AMBER, [[A('Move into the Cancellation Flow', 7.5)]])
# column 2
box(xs[1], 192, W, 52, TEAL_PALE, TEAL, [B('Issue free months (capped at 6)', 8), B('Confirm the resume date', 8),
                                         B('Disable RC / upcoming jobs', 8)], align=PP_ALIGN.LEFT)
# column 3
box(xs[2], 192, W, 62, TEAL_PALE, TEAL, [B('Issue 1 free month', 8), B('Confirm the resume date', 8),
                                         B('Ask for the duration', 8), B('Disable RC / upcoming jobs', 8)], align=PP_ALIGN.LEFT)
arrow(cx[2], 254, cx[2], 262)
box(xs[2], 262, W, 20, TEAL, None, [[P('C provided the duration', 8)]])
box(xs[2], 285, W, 40, TEAL_PALE, TEAL, [B('Top up the remaining free months (capped at 6)', 8), B('Confirm the resume date', 8)],
    align=PP_ALIGN.LEFT)
# column 4
box(xs[3], 192, W, 34, WHITE, LINE, [[A("Don't assume: confirm intent first", 8)]])
arrow(cx[3], 226, cx[3], 234)
box(xs[3], 234, W, 20, INK, None, [[P('C confirmed', 8)]])
box(xs[3], 257, W, 30, TEAL_PALE, TEAL, [[A('Move into the Pause Flow', 8)]])

# cap bar at the bottom
bg, icon, ct = cap
dy = Pt(338) - bg.top
for sh in cap:
    sh.top = sh.top + dy
bg.height = Pt(26); icon.top = bg.top + Pt(6.5); ct.top = bg.top + Pt(5)

add_note(s, 'Scenario 3 starts one of three ways: the customer asks to pause directly; the customer asks to cancel but the real reason is temporary; or you spot the signal before they have said anything. '
            'There is no literal pause button, so we simulate one by issuing free months so the MF does not charge. '
            'Flow: asks to cancel but the reason may be temporary, probe and offer to pause (accepts: Pause Flow; rejects: Cancellation Flow). '
            'Asks to pause with a timeframe: issue free months capped at 6, confirm the resume date, disable RC and upcoming jobs. '
            'Asks to pause without a timeframe: issue 1 free month, confirm the resume date, ask for the duration, disable RC and upcoming jobs; once they give the duration, top up the remaining free months (capped at 6) and confirm the resume date. '
            'You notice the signal first: don\'t assume, confirm intent first; once confirmed, move into the Pause Flow.')

prs.save(sys.argv[2]); print('saved', sys.argv[2])
