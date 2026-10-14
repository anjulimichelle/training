"""Day 7 deck slide 6 (Within MCT): replace the three tiles with a flow chart
(FC retention PCQ -> offer -> C rejects -> all resolutions given? -> Exit Strategy
-> C rejects -> Graceful Closure), in the deck theme. Keeps the Graceful Closure
warning bar. The Exit Strategy box opens a pop-up with slide 7's lever table
(shows on click 1, closes on click 2), then slide 7 is deleted.

Usage: python3 slide6-within-mct-flowchart.py in.pptx out.pptx
"""
import copy, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, GREY = '1D1D1F', '6E6E73', 'FFFFFF', '6E6E73'

prs = Presentation(sys.argv[1])
s = prs.slides[5]
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
assert txt(s.shapes[1]) == 'Within MCT', txt(s.shapes[1])
sl = list(s.shapes)
s7 = prs.slides[6]
assert txt(s7.shapes[1]) == 'Exit strategy levers'
lever_rows = [[c.text for c in r.cells] for r in next(sh for sh in s7.shapes if sh.has_table).table.rows][1:]
lever_sub = txt(s7.shapes[2])
lever_notes = s7.notes_slide.notes_text_frame.text
w = next(i for i, sh in enumerate(sl) if txt(sh).startswith('Before Graceful Closure'))
keep = {id(e._element) for e in sl[:3] + sl[w - 2:w + 1] + [sl[-1]]}
for sh in sl:
    if id(sh._element) not in keep:
        sh._element.getparent().remove(sh._element)
warn = sl[w - 2:w + 1]


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def box(x, y, wd, h, fill, line, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(wd), Pt(h))
    strip(sh)
    sh.adjustments[0] = min(0.3, 8 / h)
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(6); tf.margin_top = tf.margin_bottom = Pt(2)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, font, size, bold, color in runs:
            r = p.add_run(); r.text = text
            f = r.font; f.name = font; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return sh


def text(x, y, wd, h, paras, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(wd), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, font, size, bold, color in runs:
            r = p.add_run(); r.text = t
            f = r.font; f.name = font; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tb


def arrow(x1, y1, x2, y2):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    strip(c)
    c.line.color.rgb = RGBColor.from_string(GREY); c.line.width = Pt(1)
    ln = c._element.find('.//' + qn('a:ln'))
    ln.append(etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med'))
    return c


P = lambda t, size=9, bold=True, color=WHITE: (t, 'Play', size, bold, color)
A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)

L, LW = 32, 360          # main column
CX = L + LW / 2
R, RW = 420, 267         # side column
RX = R + RW / 2

box(L + 55, 98, LW - 110, 22, TEAL, None,
    [[P('FC retention – PCQ', 9.5)], [P('(ETF applies)', 8, False, 'E3F4F1')]])
arrow(CX, 120, CX, 127)
box(L, 127, LW, 58, TEAL_PALE, TEAL,
    [[P('Address the service issue and offer a resolution per Quick Reference', 8.5, True, TEAL)]],
    anchor=MSO_ANCHOR.TOP)
bul = lambda t: [('•  ', 'Arial', 8.5, True, TEAL), A(t)]
text(L + 14, 146, 150, 36, [bul('Voucher'), bul('Free month per model')])
text(L + 165, 146, 188, 36, [bul('ETF-waiver spiel, where applicable'),
                             bul('Free weeks/months to move the billing date as needed')])
arrow(CX, 185, CX, 193)
box(L + 25, 193, LW - 50, 22, INK, None,
    [[P('C rejects the resolution and insists on cancelling at no charge', 8.5, False)]])
arrow(CX, 215, CX, 222)
box(L + 15, 222, LW - 30, 32, AMBER_PALE, AMBER,
    [[P('Have ', 9, False, AMBER), P('ALL', 9, True, AMBER),
      P(' necessary resolutions based on the root cause been provided?', 9, False, AMBER)]])
arrow(CX, 254, CX, 264)
text(CX + 5, 254, 30, 10, [[P('YES', 8, True, INK)]])
exit_node = box(L + 70, 264, LW - 140, 26, TEAL, None,
    [[P('EXIT Strategy', 9.5)], [P('ETF or MCT reduction  ·  click for levers ▸', 7.5, False, 'E3F4F1')]])
arrow(CX, 290, CX, 297)
box(L + 70, 297, LW - 140, 18, INK, None, [[P('C rejects the EXIT strategy?', 8.5, False)]])
arrow(CX, 315, CX, 322)
box(L + 70, 322, LW - 140, 22, AMBER_PALE, AMBER, [[P('GRACEFUL CLOSURE', 9.5, True, AMBER)]])

# NO branch
arrow(L + LW - 15, 238, R, 238)
text(L + LW - 8, 226, 40, 10, [[P('NO', 8, True, INK)]], PP_ALIGN.CENTER)
box(R, 222, RW, 32, TEAL_PALE, TEAL,
    [[P('Provide the missing resolution and attempt to retain', 9, False, TEAL)]])
arrow(RX, 254, RX, 264)
box(R, 264, RW, 26, INK, None,
    [[P('C rejects the resolution and insists on cancelling at no charge', 8.5, False)]])
arrow(R, 277, L + LW - 70, 277)


# warning bar at the bottom
bg, icon, wt = warn
dy = Pt(351) - bg.top
for sh in warn:
    sh.top = sh.top + dy
bg.height = Pt(24)
icon.top = bg.top + Pt(5.5)
wt.top = bg.top + Pt(5)
for r in wt._element.iter(qn('a:rPr')):
    r.set('sz', '850')

# ---- pop-up: exit strategy levers (from slide 7) ----
pop = []
card = box(140, 104, 440, 236, WHITE, TEAL, [[A('')]]); card.line.width = Pt(1.5)
sp = card._element.find(qn('p:spPr'))
eff = etree.SubElement(sp, qn('a:effectLst'))
sh_ = etree.SubElement(eff, qn('a:outerShdw'), blurRad='76200', dist='25400', dir='5400000', algn='t', rotWithShape='0')
clr = etree.SubElement(sh_, qn('a:srgbClr'), val='000000'); etree.SubElement(clr, qn('a:alpha'), val='25000')
pop.append(card)
pop.append(box(140, 104, 440, 30, TEAL, None,
               [[P('Exit strategy levers', 11), P('   ' + lever_sub, 8.5, False, 'E3F4F1')]], align=PP_ALIGN.LEFT))
y = 142
heights = [70, 62, 48]
for (name, how), h in zip(lever_rows, heights):
    fill = AMBER_PALE if name.lower().startswith('which') else TEAL_PALE
    line = AMBER if fill == AMBER_PALE else TEAL
    pop.append(box(152, y, 416, h - 6, fill, line,
                   [[P(name, 9.5, True, line)], [A(how, 8.5)]], align=PP_ALIGN.LEFT))
    y += h
pop.append(text(152, 324, 416, 12, [[A('Click to close', 7.5, False, MUTED)]], PP_ALIGN.RIGHT))

# timing: click 1 shows the pop-up, click 2 hides it
def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')

def leave(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="1" presetClass="exit" presetSubtype="0" fill="hold" grpId="1" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>'
            f'</p:childTnLst></p:cTn></p:par>')

def click(cid, inner):
    return (f'<p:par><p:cTn id="{cid}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{cid+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')

ids = [sh.shape_id for sh in pop]
cid = 3
c1 = ''
for i, spid in enumerate(ids):
    c1 += effect(cid + 2, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
click1 = click(3, c1); cid += 2
c2 = ''
start2 = cid
cid += 2
for i, spid in enumerate(ids):
    c2 += leave(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 2
click2 = click(start2, c2)
bld = ''.join(f'<p:bldP spid="{i}" grpId="0"/><p:bldP spid="{i}" grpId="1"/>' for i in ids)
P_NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
tim = ('<p:timing ' + P_NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
       '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
       + click1 + click2 +
       '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
       '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
       '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
old = s._element.find(qn('p:timing'))
if old is not None:
    s._element.remove(old)
t_el = etree.fromstring(tim)
ext = s._element.find(qn('p:extLst'))
(ext.addprevious(t_el) if ext is not None else s._element.append(t_el))

# notes: lever detail moves here from slide 7
nt = s.notes_slide.notes_text_frame
if 'MCT reduction is up to 2 months' not in nt.text:
    p = copy.deepcopy(nt.paragraphs[-1]._p)
    for r in p.findall(qn('a:r'))[1:]:
        p.remove(r)
    p.findall(qn('a:r'))[0].find(qn('a:t')).text = ('TRAINER: click once to open the Exit strategy levers pop-up, '
                                                    'click again to close it. ' + lever_notes)
    nt.paragraphs[-1]._p.addnext(p)

# delete slide 7
lst = prs.slides._sldIdLst
sld = lst[6]
rid = sld.get(qn('r:id'))
lst.remove(sld)
prs.part.drop_rel(rid)

prs.save(sys.argv[2]); print('saved', sys.argv[2])
