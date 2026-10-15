"""Day 7 deck: after slide 36 (Scenario 6, beyond MCT) add
  37. a Membership Retention Playbook knowledge check (answers reveal one per click), and
  38. a practice ticket slide (trainer provides a live ticket; four guide questions).
Content comes from slides 22-36 of the deck. Reuses the click helpers in
day3-platform-mechanics/add_story_anim.py.

Usage: python3 add-slides37-38-membership-kc-and-practice.py in.pptx out.pptx
"""
import copy, io, os, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, WHITE, BORDER, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', '9A6410', 'FFF3D6'
prs = Presentation(src)
S = list(prs.slides)
ref = S[35]
assert ref.shapes[1].text_frame.text.strip().endswith('(Beyond MCT)')
assert S[36].shapes[1].text_frame.text.strip() == 'Questions?'
num = [x for x in ref.shapes if x.is_placeholder][0]
TITLE_SIZE = S[16].shapes[1].text_frame.paragraphs[0].runs[0].font.size  # slide 17 knowledge check title


def new_slide(eyebrow, title, sub, after):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
        s.shapes._spTree.append(copy.deepcopy(el))
    for shape, t in ((s.shapes[0], eyebrow), (s.shapes[1], title), (s.shapes[2], sub)):
        r = shape.text_frame.paragraphs[0].runs; r[0].text = t
        for x in r[1:]: x._r.getparent().remove(x._r)
        for p in shape.text_frame.paragraphs[1:]: p._p.getparent().remove(p._p)
    s.shapes[1].text_frame.paragraphs[0].runs[0].font.size = TITLE_SIZE
    lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); list(lst)[after].addnext(el)
    return s


def kit(s):
    def strip(x):
        st = x._element.find(P_NS + 'style')
        if st is not None: x._element.remove(st)
    def box(x, y, w, h, fill, line=None, adj=0.08, name=None, dash=False):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
        b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line:
            b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(1 if dash else 0.75)
            if dash:
                from pptx.enum.dml import MSO_LINE
                b.line.dash_style = MSO_LINE.DASH
        else: b.line.fill.background()
        if name: b.name = name
        return b
    def text(shape, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, after=3):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
            for t, bold, col, sz in runs:
                r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'Play' if bold else 'Arial'
        return shape
    def tb(x, y, w, h, name=None):
        t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name: t.name = name
        return t
    return box, text, tb


# ---------- 37. Knowledge check ----------
kc = new_slide('MEMBERSHIP RETENTION  ·  QUICK KNOWLEDGE CHECK', 'Quick Knowledge Check',
               'Ask each question, then click to reveal the answer.', 35)
box, text, tb = kit(kc)
qa = [('C says "money\'s tight, it\'s been a rough few months." Which path?',
       'Still Scenario 2: vague strain with no named event. Attempt retention with an MF reduction; don\'t upgrade it to hardship.'),
      ('C lost their job and can\'t keep up with payments. What do you offer?',
       'Extreme hardship: skip retention. Exit offer of 50% off the ETF; if declined, a full ETF waiver. No proof required.'),
      ('A gift-voucher (dhj_gift) customer says they never knew about the FC. What do you avoid saying?',
       'Don\'t say they "purchased" the voucher. Say they redeemed it, and show where the terms were shared.'),
      ('C asks to pause but gives no timeframe. What do you do?',
       'Issue 1 free month, confirm the resume date, ask for the duration, disable RC/upcoming jobs. Top up once they answer (cap 6).'),
      ('Inactive C: 12 paid months, 6-month MCT, asks for a refund. How much?',
       'Only MFs from the last 120 days that are beyond MCT: about 4 months, not 6.'),
      ('C is moving outside the service area, within MCT. What do you do?',
       'Cancel the FC and waive the ETF (ETF Waiver: Relocation, Out of Coverage). No free month.')]
n = len(qa); h = (3.75 - 0.08 * (n - 1)) / n
for i, (q, a) in enumerate(qa):
    y = 1.4 + i * (h + 0.08)
    box(0.45, y, 5.05, h, WHITE, BORDER)
    text(box(0.6, y + (h - 0.3) / 2, 0.3, 0.3, TEALSOFT, adj=0.5), [[(str(i + 1), True, TEAL, 9)]], align=PP_ALIGN.CENTER)
    text(tb(1.05, y, 4.35, h), [[(q, False, INK, 10)]])
    box(5.6, y, 3.95, h, TEALSOFT, name=f'step{i + 1}Box')
    text(tb(5.75, y, 3.7, h, name=f'step{i + 1}Ans'), [[(a, False, INK, 9)]])
kc.notes_slide.notes_text_frame.text = (
    'Quick knowledge check: Membership Retention Playbook. Ask each question and let trainees answer before revealing. Each click reveals the next answer. '
    '1) "Money\'s tight right now, it\'s been a rough few months" is still Scenario 2: vague financial strain without a named event is not confirmed hardship, so attempt retention with an MF reduction ($10 to 15 is within your authority; more needs CSQ). '
    '2) Job loss, bankruptcy, homelessness or anything comparably severe is extreme financial hardship: skip retention and go to ETF Waivers > Bankruptcy / extreme hardship, a 50%-off exit offer and, if declined, a full ETF waiver. No proof required. '
    '3) For gift voucher (dhj_gift) customers, never say they "purchased" the voucher; say they redeemed it. Show where the terms were shared: the guest welcome email, the Host\'s gift email (cropped), and the first booking for the new booking-flow cohort. '
    '4) A pause with no timeframe: issue 1 free month right away, confirm the resume date, ask for the duration, and disable RC and upcoming jobs; once they give the duration, top up the remaining free months, capped at 6. '
    '5) Refunds after long inactivity cover only MFs charged in the last 120 days that are beyond MCT: with 12 paid months and a 6-month MCT, six months are beyond MCT but only about four fall in the last 120 days, so the refund is 4 months, not 6. '
    '6) Moving outside the service area within MCT: cancel the FC and waive the ETF under ETF Waiver: Relocation (Out of Coverage). No free month, since there is nothing to apply it to.')

# ---------- 38. Practice ticket ----------
pt = new_slide('MEMBERSHIP RETENTION  ·  PRACTICE', "Practice ticket: let's investigate together!",
               'Your trainer will provide a live ticket. Use these guide questions as you investigate.', 36)
box, text, tb = kit(pt)
box(0.45, 1.4, 9.1, 0.5, GOLDSOFT)
text(tb(0.65, 1.4, 8.8, 0.5), [[('\u25b6  Live ticket:  ', True, GOLD, 11),
                               ('your trainer provides a live membership cancel-intent ticket to investigate together.', False, INK, 10)]])
qs = [("What is the customer's situation?", 'e.g. Scenario 2: can\'t afford, no extreme hardship (FCF within MCT with a completed job, ETF applies)'),
      ('What actions are you going to take?', None),
      ('What will you offer?', None),
      ('Draft your reply.', None)]
y0, gap = 2.05, 0.1
h = (4.9 - y0 - gap * (len(qs) - 1)) / len(qs)
for i, (q, ex) in enumerate(qs):
    y = y0 + i * (h + gap)
    last = i == len(qs) - 1
    box(0.45, y, 9.1, h, TEALSOFT if not last else WHITE, None if not last else TEAL)
    text(box(0.62, y + (h - 0.34) / 2, 0.34, 0.34, TEAL, adj=0.5), [[(str(i + 1), True, WHITE, 11)]], align=PP_ALIGN.CENTER)
    paras = [[(q, True, TEAL, 12)]]
    if ex:
        paras.append([(ex, False, SOFT, 9.5)])
    text(tb(1.15, y, 8.2, h), paras)
pt.notes_slide.notes_text_frame.text = (
    'Practice ticket activity. Provide a live membership cancel-intent ticket and investigate it together, using the guide questions. '
    "1) What is the customer's situation? Have trainees name the scenario (unaware, can't afford, pause, issues with terms, inactivity, moving), and whether the ETF applies (within or beyond MCT, FCF with or without a completed job). Check for FC Sales signups and gift vouchers first. "
    '2) What actions are you going to take? Follow that scenario\'s flow chart: probe the root cause, confirm retainability, and note any TR or free months to set. '
    '3) What will you offer? MF or MCT reduction, credits, free months, a pause, or the exit path (ETF or MCT reduction), within your authority or raised to CSQ. '
    '4) Draft your reply. Have each trainee write the response to C, then compare as a group.')

buf = io.BytesIO(); prs.save(buf)
ns = {}
anim = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'day3-platform-mechanics', 'add_story_anim.py')
exec(open(anim).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == kc.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        groups = {}
        for sid, k in re.findall(r'<p:cNvPr id="(\d+)" name="step(\d+)[^"]*"', x): groups.setdefault(int(k), []).append(sid)
        cid, clicks, bld = 3, [], []
        for k in sorted(groups):
            inner = ''
            for j, sid in enumerate(groups[k]):
                inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if j == 0 else 'withEffect'); cid += 3
                bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
            clicks.append(ns['click'](cid, inner)); cid += 2
        x = x.replace('</p:sld>', ns['timing'](clicks, bld) + '</p:sld>')
        data = x.encode('utf8'); print('clicks:', len(clicks))
    zout.writestr(item, data)
zout.close()
