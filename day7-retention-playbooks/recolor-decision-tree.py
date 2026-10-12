"""Day 7 copy, slide 4: recolor the PCQ decision tree to the deck's teal/gold theme.
Questions = teal, Yes = teal header, No = grey header, scenario outcomes = gold. Patches the live export."""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
src, dst = sys.argv[1], sys.argv[2]
C = RGBColor.from_string
TEAL, TEALSOFT, INK, SOFT, WHITE, GOLD, GOLDSOFT, LINE = map(C, ('0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'FFFFFF', '9A6410', 'FFF3D6', '9AA5A3'))
prs = Presentation(src); s = prs.slides[3]
assert 'The trigger' in ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)

def fill(sh):
    try: return str(sh.fill.fore_color.rgb) if sh.fill.type == 1 else None
    except Exception: return None
def style(sh, bg, text, line=None, bold=None):
    sh.fill.solid(); sh.fill.fore_color.rgb = bg
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(1)
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = text
            if bold is not None: r.font.bold = bold
n = 0
for sh in s.shapes:
    f = fill(sh); t = sh.text_frame.text if sh.has_text_frame else ''
    if f == 'E69138': style(sh, TEAL, WHITE)                                   # start
    elif f == '3D85C6': style(sh, TEALSOFT, INK, TEAL, True)                    # first question
    elif f == 'E06666': style(sh, TEAL, WHITE)                                  # Yes header
    elif f == '666666': style(sh, SOFT, WHITE)                                  # No header
    elif f in ('F4CCCC', 'D9D9D9'):
        if t.startswith('Use'): style(sh, GOLDSOFT, GOLD, GOLD, True)           # scenario outcome
        else: style(sh, TEALSOFT, INK, TEAL, False)                             # follow-up question
    elif sh.shape_type == 9 and str(sh.line.color.rgb) == '595959': sh.line.color.rgb = LINE
    else: continue
    n += 1
print('restyled', n)
prs.save(dst)
