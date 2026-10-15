"""Day 7 deck: after slide 16 (Service Reliability, ETF does not apply) add
  17. a Service Reliability Retention Playbook knowledge check (answers reveal one per click), and
  18. a practice ticket slide (trainer provides a live ticket; four guide questions).
Content comes from slides 12-16 of the deck. Reuses the click helpers in
day3-platform-mechanics/add_story_anim.py.

Usage: python3 add-slides17-18-sr-kc-and-practice.py in.pptx out.pptx
"""
import copy, io, os, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, WHITE, BORDER, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', '9A6410', 'FFF3D6'
prs = Presentation(src)
S = list(prs.slides)
ref = S[15]
assert ref.shapes[1].text_frame.text.strip() == 'ETF does not apply'
assert S[16].shapes[1].text_frame.text.strip() == 'Questions?'
num = [x for x in ref.shapes if x.is_placeholder][0]


def new_slide(eyebrow, title, sub, after):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
        s.shapes._spTree.append(copy.deepcopy(el))
    for shape, t in ((s.shapes[0], eyebrow), (s.shapes[1], title), (s.shapes[2], sub)):
        r = shape.text_frame.paragraphs[0].runs; r[0].text = t
        for x in r[1:]: x._r.getparent().remove(x._r)
        for p in shape.text_frame.paragraphs[1:]: p._p.getparent().remove(p._p)
    lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); list(lst)[after].addnext(el)
    return s


def kit(s):
    def strip(x):
        st = x._element.find(P_NS + 'style')
        if st is not None: x._element.remove(st)
    def box(x, y, w, h, fill, line=None, adj=0.08, name=None, dash=False):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
        b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line:
            b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(1 if dash else 0.75)
            if dash:
                from pptx.enum.dml import MSO_LINE
                b.line.dash_style = MSO_LINE.DASH
        else: b.line.fill.background()
        if name: b.name = name
        return b
    def text(shape, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, after=3):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
            for t, bold, col, sz in runs:
                r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'Play' if bold else 'Arial'
        return shape
    def tb(x, y, w, h, name=None):
        t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name: t.name = name
        return t
    return box, text, tb


# ---------- 17. Knowledge check ----------
kc = new_slide('SERVICE RELIABILITY  ·  QUICK KNOWLEDGE CHECK', 'Quick Knowledge Check',
               'Ask each question, then click to reveal the answer.', 15)
box, text, tb = kit(kc)
qa = [('Which four triggers does the Service Reliability playbook cover?',
       'No CP Claim, CP No-Show, CP Cancellation/LMC, or a CP being late, on a single issue job.'),
      ('C had two consecutive Service Reliability issue jobs. What do you use?',
       'The Two Consecutive Issue Jobs article, not this playbook.'),
      ('FCF member with a completed job, CP no-show. Which credits and free month?',
       '$10-20 credits. Free months: unused MFs first, else 1 month if the next MF is within 2 weeks.'),
      ('The CP was late. What is the credit range?',
       '$10-25, a wider band than the other triggers.'),
      ('Should you lead with a free month?',
       'No. It is optional: offer it with cancel intent or if C brings up unused months. Always block the pairing and coach or penalize the CP.'),
      ('ETF applies and C rejects the resolution, insisting on cancelling free. Next step?',
       'Check every resolution was given, then Exit Strategy: ETF reduction (up to 50%) or MCT reduction (up to 2 months), whichever is cheaper.')]
n = len(qa); h = (3.75 - 0.08 * (n - 1)) / n
for i, (q, a) in enumerate(qa):
    y = 1.4 + i * (h + 0.08)
    box(0.45, y, 5.05, h, WHITE, BORDER)
    text(box(0.6, y + (h - 0.3) / 2, 0.3, 0.3, TEALSOFT, adj=0.5), [[(str(i + 1), True, TEAL, 9)]], align=PP_ALIGN.CENTER)
    text(tb(1.05, y, 4.35, h), [[(q, False, INK, 10)]])
    box(5.6, y, 3.95, h, TEALSOFT, name=f'step{i + 1}Box')
    text(tb(5.75, y, 3.7, h, name=f'step{i + 1}Ans'), [[(a, False, INK, 9)]])
kc.notes_slide.notes_text_frame.text = (
    'Quick knowledge check: Service Reliability Retention Playbook. Ask each question and let trainees answer before revealing. Each click reveals the next answer. '
    '1) The playbook covers cancel intent after a single reliability issue job: No CP Claim, CP No-Show, CP Cancellation/LMC, or a late CP. '
    '2) Two consecutive Service Reliability issue jobs go to the Two Consecutive Issue Jobs article instead. '
    '3) Per the Quick Reference, FCF with a completed job gets $10-20 credits; free months go to unused MFs first, else 1 month if the next MF is within 2 weeks. Free months are never more than asked and never more than 3. '
    '4) A late CP has the wider $10-25 credit band, in all cases. '
    '5) Free months are optional and will not fail QA either way. Hold off by default; offer them only with cancellation intent or if the customer brings up the unused months. Always block the C/CP pairing and coach or penalize the CP. '
    '6) When the ETF applies and the customer rejects the resolution, check that every resolution for the root cause was given, then go to Exit Strategy: ETF reduction (up to 50%) or MCT reduction (up to 2 months, capped at 1 if 2 would waive the ETF), whichever is cheaper. Max out the lever before Graceful Closure.')

# ---------- 18. Practice ticket ----------
pt = new_slide('SERVICE RELIABILITY  ·  PRACTICE', "Practice ticket: let's investigate together!",
               'Your trainer will provide a live ticket. Use these guide questions as you investigate.', 16)
box, text, tb = kit(pt)
box(0.45, 1.4, 9.1, 0.5, GOLDSOFT)
text(tb(0.65, 1.4, 8.8, 0.5), [[('\u25b6  Live ticket:  ', True, GOLD, 11),
                               ('your trainer provides a live Service Reliability cancel-intent ticket to investigate together.', False, INK, 10)]])
qs = [("What is the customer's situation?", 'e.g. FCF, with completed job: CP No-Show (ETF applies)'),
      ('What actions are you going to take?', None),
      ('What will you offer?', None),
      ('Draft your reply.', None)]
y0, gap = 2.05, 0.1
h = (4.9 - y0 - gap * (len(qs) - 1)) / len(qs)
for i, (q, ex) in enumerate(qs):
    y = y0 + i * (h + gap)
    last = i == len(qs) - 1
    box(0.45, y, 9.1, h, TEALSOFT if not last else WHITE, None if not last else TEAL)
    text(box(0.62, y + (h - 0.34) / 2, 0.34, 0.34, TEAL, adj=0.5), [[(str(i + 1), True, WHITE, 11)]], align=PP_ALIGN.CENTER)
    paras = [[(q, True, TEAL, 12)]]
    if ex:
        paras.append([(ex, False, SOFT, 9.5)])
    text(tb(1.15, y, 8.2, h), paras)
pt.notes_slide.notes_text_frame.text = (
    'Practice ticket activity. Provide a live Service Reliability cancel-intent ticket and investigate it together, using the guide questions. '
    "1) What is the customer's situation? Have trainees name the trigger (No CP Claim, No-Show, Cancellation/LMC or Late), the Quick Reference row, and whether the ETF applies. "
    '2) What actions are you going to take? Block the C/CP pairing, coach or penalize the CP, and follow the ETF applies or ETF does not apply flow. '
    '3) What will you offer? Credits and free months from the Quick Reference, the ETF-waiver spiel where it applies, and what happens if C rejects it (Exit Strategy if the ETF applies; cancel or set a TR if it does not). '
    '4) Draft your reply. Have each trainee write the response to C, then compare as a group.')

buf = io.BytesIO(); prs.save(buf)
ns = {}
anim = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'day3-platform-mechanics', 'add_story_anim.py')
exec(open(anim).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == kc.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        groups = {}
        for sid, k in re.findall(r'<p:cNvPr id="(\d+)" name="step(\d+)[^"]*"', x): groups.setdefault(int(k), []).append(sid)
        cid, clicks, bld = 3, [], []
        for k in sorted(groups):
            inner = ''
            for j, sid in enumerate(groups[k]):
                inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if j == 0 else 'withEffect'); cid += 3
                bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
            clicks.append(ns['click'](cid, inner)); cid += 2
        x = x.replace('</p:sld>', ns['timing'](clicks, bld) + '</p:sld>')
        data = x.encode('utf8'); print('clicks:', len(clicks))
    zout.writestr(item, data)
zout.close()
