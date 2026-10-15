"""Day 7 deck, Membership Retention Scenario 6 (customer wants to cancel because they're moving).

Slide 35: flow chart within MCT (ETF applies): new address shared? -> in service area? ->
          pause with free weeks/months -> exit strategy -> graceful closure. Keeps the "Always start" bar.
New 36:   flow chart beyond MCT (ETF does not apply): same questions, ends in cancel the FC.
Helpers come from slide30-membership-scenario4-terms.py.

Usage: python3 slides35-36-membership-scenario6-moving.py in.pptx out.pptx
"""
import copy, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
S = list(prs.slides)
assert txt(S[34].shapes[1]) == 'Scenario 6: moving', txt(S[34].shapes[1])
s = None
s8_notes = "The same Timed Reminder rules apply in every beyond-MCT branch. If the next MF is more than 5 days away, set the reminder at least 3 days before it. If it is 5 days or less, give 1 free month first and set the reminder 3 days before the new MF date. When it triggers with no reply and no engagement, cancel FC; the auto comms are enough unless an unresolved pain point needs a manual email. If there is no reply but the customer is using the service, such as a new booking, don't cancel. Confirm with them, push the due date out 1 to 2 weeks, and set a new reminder 3 days before it."


def clear(slide):
    sl = list(slide.shapes)
    keep = {id(e._element) for e in sl[:3]} | {id(e._element) for e in sl if e.is_placeholder}
    for sh in sl:
        if id(sh._element) not in keep:
            sh._element.getparent().remove(sh._element)


def set_text(sh, t, size=None):
    runs = sh.text_frame.paragraphs[0].runs
    runs[0].text = t
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in sh.text_frame.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    if size:
        runs[0].font.size = Pt(size)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill_paras(tf, paras, align):
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, font, size, bold, color in runs:
            r = p.add_run(); r.text = t
            f = r.font; f.name = font; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tf


def hang(tf):
    for p in tf.paragraphs:
        lead = p.runs[0].text if p.runs else ''
        if lead.startswith('•'):
            p._p.get_or_add_pPr().set('marL', str(11 * 12700)); p._p.get_or_add_pPr().set('indent', str(-11 * 12700))
        elif lead.startswith('–'):
            p._p.get_or_add_pPr().set('marL', str(24 * 12700)); p._p.get_or_add_pPr().set('indent', str(-10 * 12700))


def box(x, y, wd, h, fill, line, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, lw=1):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(wd), Pt(h))
    strip(sh)
    sh.adjustments[0] = min(0.3, 8 / h)
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(7); tf.margin_top = tf.margin_bottom = Pt(3)
    fill_paras(tf, paras, align)
    hang(tf)
    return sh


def text(x, y, wd, h, paras, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(wd), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_paras(tf, paras, align)
    return tb


def arrow(x1, y1, x2, y2, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    strip(c)
    c.line.color.rgb = RGBColor.from_string(MUTED); c.line.width = Pt(1)
    if head:
        ln = c._element.find('.//' + qn('a:ln'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
    return c


P = lambda t, size=9, bold=True, color=WHITE: (t, 'Play', size, bold, color)
A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)
B = lambda t, size=8.5: [('•  ', 'Arial', size, True, TEAL), A(t, size)]
B2 = lambda t, size=8.5: [('–  ', 'Arial', size, False, MUTED), A(t, size)]


A = lambda t, size=8.5, bold=False, color=INK: (t, 'Arial', size, bold, color)
B = lambda t, size=8.5: [('•  ', 'Arial', size, True, TEAL), A(t, size)]
B2 = lambda t, size=8.5: [('–  ', 'Arial', size, False, MUTED), A(t, size)]



def moving_flow(etf):
    """Draw the Scenario 6 flow on the global slide s. etf=True: within MCT; False: beyond MCT."""
    CX = 400
    box(CX - 130, 99, 260, 22, TEAL, None,
        [[P('FC retention – Moving', 9)], [P('(ETF applies)' if etf else '(ETF does not apply)', 7.5, False, TEAL_PALE)]])
    arrow(CX, 121, CX, 127)
    box(CX - 110, 127, 220, 18, AMBER_PALE, AMBER, [[P('Has C shared their new address?', 8.5, True, AMBER)]])
    arrow(CX - 110, 136, 236, 136)
    text(CX - 136, 125, 24, 9, [[P('NO', 7.5, True, INK)]])
    left = [B('Ask for the new address', 7.5), B('Offer to pause their FC', 7.5)]
    if etf:
        left += [B('Next MF in under 7 days: issue 1 free month', 7.5),
                 B('Next MF more than 7 days away: set a TR to check back before the next billing date', 7.5)]
    else:
        left += [B('Explicitly confirm no further MF will be charged unless they agree to continue', 7.5),
                 B('Set a TR', 7.5)]
    box(32, 112, 204, 66, WHITE, LINE, left, align=PP_ALIGN.LEFT)
    arrow(CX, 145, CX, 153)
    text(CX + 5, 144, 30, 9, [[P('YES', 7.5, True, INK)]])
    box(CX - 130, 153, 260, 18, AMBER_PALE, AMBER, [[P('Is the new address within our service area?', 8.5, True, AMBER)]])
    arrow(CX + 130, 162, 553, 162)
    text(CX + 133, 151, 18, 9, [[P('NO', 7.5, True, INK)]])
    if etf:
        box(553, 146, 134, 40, AMBER_PALE, AMBER, [[P('Cancel FC + waive the ETF', 8.5, True, AMBER)],
                                                    [A('ETF Waiver: Relocation (Out of Coverage)', 6.5, False, AMBER)]])
    else:
        box(553, 148, 134, 30, AMBER_PALE, AMBER, [[P('Cancel FC and address other pain points', 8, True, AMBER)]])
    arrow(CX, 171, CX, 178)
    text(CX + 5, 170, 30, 9, [[P('YES', 7.5, True, INK)]])
    mid = [B('Issue free weeks/months to cover it', 8), B('Cancel RC / upcoming jobs', 8)]
    if not etf:
        mid += [B('Explicitly confirm no further MF will be charged unless they agree to continue', 8), B('Set a TR', 8)]
    mh = 30 if etf else 60
    box(CX - 140, 178, 280, mh, TEAL_PALE, TEAL, mid, align=PP_ALIGN.LEFT)
    y = 178 + mh
    arrow(CX, y, CX, y + 7)
    r1 = y + 7
    box(CX - 130, r1, 260, 18, INK, None, [[P('C rejects the resolution and insists on cancelling', 8, False)]])
    arrow(134, 178, 134, r1 + 9, head=False)
    arrow(134, r1 + 9, CX - 130, r1 + 9)
    arrow(CX, r1 + 18, CX, r1 + 25)
    y = r1 + 25
    if etf:
        box(CX - 100, y, 200, 24, TEAL, None, [[P('EXIT Strategy', 9)], [P('MCT or ETF reduction', 7.5, False, TEAL_PALE)]])
        arrow(CX, y + 24, CX, y + 31)
        box(CX - 130, y + 31, 260, 18, INK, None, [[P('C rejects the resolution and insists on cancelling', 8, False)]])
        arrow(CX, y + 49, CX, y + 56)
        box(CX - 75, y + 56, 150, 20, AMBER_PALE, AMBER, [[P('GRACEFUL CLOSURE', 9, True, AMBER)]])
        return y + 76
    box(CX - 70, y, 140, 20, AMBER_PALE, AMBER, [[P('Cancel the FC', 9, True, AMBER)]])
    return y + 20


# ======================= slide 35: within MCT =======================
s = S[34]
sl = list(s.shapes)
w = next(i for i, sh in enumerate(sl) if txt(sh).startswith('Always start'))
bar = sl[w - 2:w + 1]
keep = {id(e._element) for e in sl[:3] + bar} | {id(e._element) for e in sl if e.is_placeholder}
for sh in sl:
    if id(sh._element) not in keep:
        sh._element.getparent().remove(sh._element)
set_text(s.shapes[1], "Scenario 6: Customer Wants to Cancel Because They're Moving (Within MCT)", 14)
set_text(s.shapes[2], 'Where they are going decides pause or exit. The ETF applies.')
moving_flow(True)
bg, icon, bt = bar
dy = Pt(346) - bg.top
for sh in bar:
    sh.top = sh.top + dy
bg.height = Pt(24); icon.top = bg.top + Pt(5.5); bt.top = bg.top + Pt(4)
s.notes_slide.notes_text_frame.text = (
    'Scenario 6, within MCT (ETF applies). Always start with: has the customer shared their new address? '
    'If not, ask for it and offer to pause their FC. If the next MF is less than 7 days away, issue 1 free month; if it is more than 7 days away, set a TR to check back before the next billing date. '
    'If they have shared it and the new address is outside our service area, cancel the FC and waive the ETF (ETF Waiver: Relocation, Out of Coverage); no free month, since there is nothing to apply it to. '
    'If it is within our service area, treat it as a pause: issue free weeks or months to cover the move and cancel RC and upcoming jobs. '
    'If C rejects the resolution and insists on cancelling, go to Exit Strategy (MCT or ETF reduction); if they reject that too, Graceful Closure.')

# ======================= new slide 36: beyond MCT =======================
ref = S[34]
s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
num = [x for x in ref.shapes if x.is_placeholder][0]
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); list(lst)[34].addnext(el)
set_text(s.shapes[1], "Scenario 6: Customer Wants to Cancel Because They're Moving (Beyond MCT)", 14)
set_text(s.shapes[2], 'Where they are going decides pause or cancel. The ETF does not apply.')
moving_flow(False)
s.notes_slide.notes_text_frame.text = (
    'Scenario 6, beyond MCT (ETF does not apply). Always start with: has the customer shared their new address? '
    'If not, ask for it, offer to pause their FC, explicitly confirm no further MF will be charged unless they agree to continue, and set a TR. '
    'If the new address is outside our service area, cancel the FC and address any other pain points. '
    'If it is within our service area: issue free weeks or months to cover it, cancel RC and upcoming jobs, explicitly confirm no further MF will be charged unless they agree to continue, and set a TR. '
    'If C rejects the resolution and insists on cancelling, cancel the FC.')

prs.save(sys.argv[2]); print('saved', sys.argv[2])
