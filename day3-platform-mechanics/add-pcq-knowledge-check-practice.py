"""Day 3 copy (KB - Poor Cleaning Quality): after slide 88 add a quick knowledge check (answers reveal per click)
and a "Let's investigate together!" live-ticket slide. Patches the live Google export."""
import copy, io, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, WHITE, BORDER, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', '9A6410', 'FFF3D6'
prs = Presentation(src)
S = list(prs.slides)
ref = S[87]
assert ref.shapes[1].text_frame.text.strip() == 'Handling a cleaning supplies issue'
num = [x for x in ref.shapes if x.is_placeholder][0]

def new_slide(eyebrow, title, sub, after):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
        s.shapes._spTree.append(copy.deepcopy(el))
    for shape, t in ((s.shapes[0], eyebrow), (s.shapes[1], title), (s.shapes[2], sub)):
        r = shape.text_frame.paragraphs[0].runs; r[0].text = t
        for x in r[1:]: x._r.getparent().remove(x._r)
    lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); list(lst)[after].addnext(el)
    return s
def kit(s):
    def strip(x):
        st = x._element.find(P_NS + 'style')
        if st is not None: x._element.remove(st)
    def box(x, y, w, h, fill, line=None, adj=0.08, name=None):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
        b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
        else: b.line.fill.background()
        if name: b.name = name
        return b
    def text(shape, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, after=3):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
            for t, bold, col, sz in runs:
                r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'Play' if bold else 'Arial'
        return shape
    def tb(x, y, w, h, name=None):
        t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name: t.name = name
        return t
    return box, text, tb

# ---------- Knowledge check ----------
kc = new_slide('POOR CLEANING QUALITY  ·  QUICK KNOWLEDGE CHECK', 'Quick Knowledge Check', 'Ask each question, then click to reveal the answer.', 87)
box, text, tb = kit(kc)
qa = [('Missed areas, but C hasn\'t said they want to cancel. Which guide applies?', 'KB - PCQ. Cancel intent at any point? Move to the PCQ Retention Framework.'),
      ('Can you tell C they have 96 hours to request a refund?', 'No. Say 72 hours. The 96-hour window is internal: disclosing it is ZTP Level 2.'),
      ('DHJ customer, confirmed PCQ on a 3-hr job, clearly upset. Offer?', '1 hr voucher (2 hr voucher if the job was 4 hrs or more).'),
      ('Refund request vs admin refund?', 'Refund request: the cleaner\'s call, never promise it. Admin refund: ours, CP still paid.'),
      ('PCQ for a high-loyalty C (13 invoiced jobs, 9 paid months). Offer?', 'Full admin refund.'),
      ('A CA CP arrived without complete supplies. Penalize?', 'No. Supplies aren\'t required in CA/CA-WA. Encourage it; check Cleaning Notes.')]
n = len(qa); h = (3.75 - 0.08 * (n - 1)) / n
for i, (q, a) in enumerate(qa):
    y = 1.4 + i * (h + 0.08)
    box(0.45, y, 5.05, h, WHITE, BORDER)
    text(box(0.6, y + (h - 0.3) / 2, 0.3, 0.3, TEALSOFT, adj=0.5), [[(str(i + 1), True, TEAL, 9)]], align=PP_ALIGN.CENTER)
    text(tb(1.05, y, 4.35, h), [[(q, False, INK, 10)]])
    box(5.6, y, 3.95, h, TEALSOFT, name=f'step{i + 1}Box')
    text(tb(5.75, y, 3.7, h, name=f'step{i + 1}Ans'), [[(a, False, INK, 9)]])
kc.notes_slide.notes_text_frame.text = (
    'Quick knowledge check: KB - Poor Cleaning Quality. Ask each question and let trainees answer before revealing. Each click reveals the next answer. '
    '1) This KB covers a customer unhappy about a specific cleaning with no cancel intent. The moment they want to cancel, for any reason, move to the PCQ Retention Framework; ETF and waivers live there only. '
    '2) Customers can self-submit a refund request within 72 hours of invoice. The system allows up to 96 hours, but that is internal only: disclosing it is a ZTP Level 2 violation. '
    '3) Recovery matrix: PCQ on a job under 4 hours: 1 hr voucher (DHJ and FCF; non-FC: 1 hr voucher or $20 credits); 4 hours or more: 2 hr voucher. Offer only with a confirmed service issue and a customer who is clearly upset, hesitant to continue, or has lost confidence. '
    '4) A refund request is the cleaner\'s call: the customer submits within 72 hours, the CP may accept or decline and has 7 days to respond; never promise it. An admin refund comes from Homeaglow and the CP is still paid in full. '
    '5) High-loyalty customers (e.g. 13 invoiced jobs, 9 paid months): full admin refund. '
    '6) CA / CA-WA CPs aren\'t required to bring complete supplies: don\'t penalize; encourage it and check Cleaning Notes for the customer\'s preference.')

# ---------- Let's investigate together ----------
pr = new_slide('POOR CLEANING QUALITY  ·  PRACTICE', 'Let\'s investigate together!', 'Review the live ticket and investigate.', 88)
box, text, tb = kit(pr)
top, bottom, lw = 1.42, 5.0, 3.7
box(0.45, top, lw, bottom - top, WHITE, BORDER)
text(tb(0.65, top + 0.3, lw - 0.4, 0.4), [[('Live ticket', True, INK, 14)]])
text(tb(0.65, top + 0.8, lw - 0.4, 1.2), [[('Your trainer will provide a live PCQ ticket. Review it and investigate before you decide anything.', False, SOFT, 11)]], anchor=MSO_ANCHOR.TOP)
text(box(0.65, bottom - 0.6, 1.6, 0.38, GOLDSOFT, adj=0.2), [[('Trainer provides', True, GOLD, 10)]], align=PP_ALIGN.CENTER)
prompts = [('Is it PCQ? Any cancel intent?', 'KB - PCQ, or the PCQ Retention Framework?'),
           ('What do you check first?', 'Job admin, messages, refund status, voucher, pattern.'),
           ('What actions and offer?', 'Recovery matrix, refund rules, account and CP steps.')]
rx = 0.45 + lw + 0.15; rw = 9.55 - rx; rh = (bottom - top - 2 * 0.15) / 3
for k, (a, b) in enumerate(prompts):
    y = top + k * (rh + 0.15)
    box(rx, y, rw, rh, WHITE, BORDER)
    text(box(rx + 0.2, y + (rh - 0.32) / 2, 0.32, 0.32, TEALSOFT, adj=0.5), [[(str(k + 1), True, TEAL, 9.5)]], align=PP_ALIGN.CENTER)
    text(tb(rx + 0.7, y + 0.18, rw - 0.9, 0.35), [[(a, True, INK, 13)]])
    text(tb(rx + 0.7, y + 0.58, rw - 0.9, 0.3), [[(b, False, SOFT, 10.5)]])
pr.notes_slide.notes_text_frame.text = (
    'TRAINER: share a live PCQ ticket. Give trainees a few minutes to review it on their own, then work through it together. '
    '1) Is it PCQ (missed areas, incomplete tasks, visibly dirty areas, incomplete supplies that hurt the job)? Any cancel intent? If so, it moves to the PCQ Retention Framework. '
    '2) What do you check first? Job admin and history, customer-CP and customer-Care messages, supplies overlap, refund request status and the 72-hour window, voucher status (admin courtesy voucher), and whether it\'s isolated or a pattern. '
    '3) What actions and offer? Use the recovery matrix (only for a confirmed issue and a clearly upset customer), check refund request rules and blockers, refund the Premium Upsell Fee if charged, flag customer_unsatisfied, block the pairing, and coach or penalize the CP. Then move to Questions.')

buf = io.BytesIO(); prs.save(buf)
ns = {}
exec(open(__file__.replace('day3_pcq_kc.py', 'add_story_anim.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == kc.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        groups = {}
        for sid, n in re.findall(r'<p:cNvPr id="(\d+)" name="step(\d+)[^"]*"', x): groups.setdefault(int(n), []).append(sid)
        cid, clicks, bld = 3, [], []
        for n in sorted(groups):
            inner = ''
            for j, sid in enumerate(groups[n]):
                inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if j == 0 else 'withEffect'); cid += 3
                bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
            clicks.append(ns['click'](cid, inner)); cid += 2
        x = x.replace('</p:sld>', ns['timing'](clicks, bld) + '</p:sld>')
        data = x.encode('utf8'); print('clicks:', len(clicks))
    zout.writestr(item, data)
zout.close()
