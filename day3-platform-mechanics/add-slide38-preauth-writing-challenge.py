"""Day 3 original deck: insert a 5-minute writing challenge after slide 37 ('Before you explain a pre-auth').
Context + customer message on the left; the sample CS response stays hidden until the click. Patches the live export."""
import copy, io, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, BG, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', 'F6F7F9', '9A6410', 'FFF3D6'
prs = Presentation(src)
S = list(prs.slides)
ref = S[36]
assert ref.shapes[1].text_frame.text.strip() == 'Before you explain a pre-auth'
s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
num = [x for x in ref.shapes if x.is_placeholder][0]
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
for shape, t in ((s.shapes[0], 'PRE-AUTHORIZATION HOLDS  ·  PRACTICE'), (s.shapes[1], '5-minute writing challenge'), (s.shapes[2], 'Read the customer\'s message. How would you respond?')):
    r = shape.text_frame.paragraphs[0].runs; r[0].text = t
    for x in r[1:]: x._r.getparent().remove(x._r)
lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); list(lst)[36].addnext(el)

def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
def box(x, y, w, h, fill, line=None, adj=0.06, name=None, dash=False):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
    b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
        if dash:
            ln = b._element.spPr.find(qn('a:ln')); ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))
    else: b.line.fill.background()
    if name: b.name = name
    return b
def text(shape, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=4, m=0.0):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(m); tf.margin_top = tf.margin_bottom = Inches(m / 2)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
        for t, bold, col, sz, *it in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold; r.font.italic = bool(it)
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'Play' if bold else 'Arial'
    return shape
def tb(x, y, w, h, name=None):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: t.name = name
    return t

# Left: context, customer message, timer.
text(box(0.45, 1.42, 3.85, 1.25, GOLDSOFT, GOLD), [[('CONTEXT', True, GOLD, 8)],
    [('Groupon customer booked a 5-hour cleaning with a 3-hour voucher. The system pre-authorized the extra 2 hours for $84.', False, INK, 10.5)]], m=0.16, anchor=MSO_ANCHOR.MIDDLE)
text(box(0.45, 2.8, 3.85, 1.15, BG, BORDER, adj=0.1), [[('CUSTOMER', True, SOFT, 8)],
    [('“Why did you charge me? I did not have any cleaning yet and I also have a voucher.”', False, INK, 11.5, 1)]], m=0.16, anchor=MSO_ANCHOR.MIDDLE)
text(box(0.45, 4.08, 3.85, 0.95, TEALSOFT, adj=0.1), [[('5 minutes', True, TEAL, 11)],
    [('Write your reply first. Then compare it with the sample.', False, INK, 9.5)]], m=0.16, anchor=MSO_ANCHOR.MIDDLE)
# Right: placeholder, cue, hidden response.
box(4.45, 1.42, 5.1, 3.61, WHITE, BORDER, dash=True)
text(tb(4.45, 2.75, 5.1, 0.25), [[('Write your reply first.', False, SOFT, 10)]], align=PP_ALIGN.CENTER)
cue = box(4.45 + (5.1 - 2.4) / 2, 3.08, 2.4, 0.4, TEAL, adj=0.3)
text(cue, [[('▶  Show CS response', True, WHITE, 10.5)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
box(4.45, 1.42, 5.1, 3.61, TEALSOFT, TEAL, name='step1Box')
resp = ["I'm happy to clarify.",
        "The $84 you're seeing is a pre-authorization, not a final charge. Your voucher covers 3 hours, so the additional 2 hours of your 5-hour booking were pre-authorized.",
        "This is a temporary hold placed by your bank to confirm the funds for your upcoming [date] appointment. The hold will be automatically released, usually within 7 days.",
        "Let me know if you have questions or need anything else."]
text(tb(4.63, 1.52, 4.74, 3.42, name='step1Text'), [[('SAMPLE CS RESPONSE', True, TEAL, 8)]] + [[(p, False, INK, 10)] for p in resp], after=7)

s.notes_slide.notes_text_frame.text = (
    '5-minute writing challenge. Context: a Groupon customer booked a 5-hour cleaning with a 3-hour voucher. The system pre-authorized the extra 2 hours for $84. '
    'Customer\'s message: "Why did you charge me? I did not have any cleaning yet and I also have a voucher." '
    'Give trainees 5 minutes to write their reply, read a few aloud, then click to reveal the sample CS response. What a good reply does: names it as a pre-authorization, not a final charge; '
    'explains why it exists (the voucher covers 3 hours, the extra 2 hours of the 5-hour booking were pre-authorized); says it is a temporary hold to confirm funds for the upcoming appointment; '
    'and sets the release expectation (usually within 7 days, bank-dependent). '
    'Sample response: ' + ' '.join(resp))
buf = io.BytesIO(); prs.save(buf)
ns = {}
exec(open(__file__.replace('add-slide38-preauth-writing-challenge.py', 'add_story_anim.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == s.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        groups = {}
        for sid, n in re.findall(r'<p:cNvPr id="(\d+)" name="step(\d+)[^"]*"', x): groups.setdefault(int(n), []).append(sid)
        cid, clicks, bld = 3, [], []
        for n in sorted(groups):
            inner = ''
            for j, sid in enumerate(groups[n]):
                inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if j == 0 else 'withEffect'); cid += 3
                bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
            clicks.append(ns['click'](cid, inner)); cid += 2
        x = x.replace('</p:sld>', ns['timing'](clicks, bld) + '</p:sld>')
        data = x.encode('utf8'); print('clicks:', len(clicks))
    zout.writestr(item, data)
zout.close()
