"""Day 3 original deck (Vouchers in CRM): after slide 50 add a quick knowledge check (answers reveal per click).
Questions follow the KB article 'Vouchers in the CRM'. Patches the live Google export."""
import copy, io, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, WHITE, BORDER, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', '9A6410', 'FFF3D6'
prs = Presentation(src)
S = list(prs.slides)
ref = S[48]  # header source: a voucher slide with eyebrow, title and subtitle
assert ref.shapes[1].text_frame.text.strip() == 'When a customer wants a refund'
assert S[49].shapes[1].text_frame.text.strip() == 'How the system handles combos'
num = [x for x in ref.shapes if x.is_placeholder][0]

def new_slide(eyebrow, title, sub, after):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
        s.shapes._spTree.append(copy.deepcopy(el))
    for shape, t in ((s.shapes[0], eyebrow), (s.shapes[1], title), (s.shapes[2], sub)):
        r = shape.text_frame.paragraphs[0].runs; r[0].text = t
        for x in r[1:]: x._r.getparent().remove(x._r)
    lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); list(lst)[after].addnext(el)
    return s
def kit(s):
    def strip(x):
        st = x._element.find(P_NS + 'style')
        if st is not None: x._element.remove(st)
    def box(x, y, w, h, fill, line=None, adj=0.08, name=None):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
        b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
        else: b.line.fill.background()
        if name: b.name = name
        return b
    def text(shape, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, after=3):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
            for t, bold, col, sz in runs:
                r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'Play' if bold else 'Arial'
        return shape
    def tb(x, y, w, h, name=None):
        t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name: t.name = name
        return t
    return box, text, tb

# ---------- Knowledge check ----------
kc = new_slide('VOUCHERS IN CRM  ·  QUICK KNOWLEDGE CHECK', 'Quick Knowledge Check', 'Ask each question, then click to reveal the answer.', 49)
box, text, tb = kit(kc)
qa = [('An FCF customer says "I never bought a voucher." Do you correct them?', "No. They're right. They signed up for a membership; the system issued the free voucher behind the scenes."),
      ('C books a 5-hr cleaning with a 3-hr voucher. What about the extra 2 hours?', "They're charged at Homeaglow's standard rate. Voucher hours come off whatever the CP's rate."),
      ('A voucher in the new CRM has no suffix. How do you check if it was used?', 'Click View all in legacy CRM. Job details in the JOB column = used. A dash = unused.'),
      ('A Groupon voucher expires unused. What happens to it?', 'It converts into credits equal to what the customer paid for it.'),
      ('C wants their Groupon voucher refunded. What must you do?', 'Address the reason, send the self-refund steps (Groupon refund macro), and invalidate the voucher.'),
      ('Groupon voucher used on a job shorter than the voucher. Courtesy voucher for the rest?', 'No. Groupon vouchers can\'t be split. The system only does this for DHJ vouchers.')]
n = len(qa); h = (3.75 - 0.08 * (n - 1)) / n
for i, (q, a) in enumerate(qa):
    y = 1.4 + i * (h + 0.08)
    box(0.45, y, 5.05, h, WHITE, BORDER)
    text(box(0.6, y + (h - 0.3) / 2, 0.3, 0.3, TEALSOFT, adj=0.5), [[(str(i + 1), True, TEAL, 9)]], align=PP_ALIGN.CENTER)
    text(tb(1.05, y, 4.35, h), [[(q, False, INK, 10)]])
    box(5.6, y, 3.95, h, TEALSOFT, name=f'step{i + 1}Box')
    text(tb(5.75, y, 3.7, h, name=f'step{i + 1}Ans'), [[(a, False, INK, 9)]])
kc.notes_slide.notes_text_frame.text = (
    'Quick knowledge check: Vouchers in the CRM. Ask each question and let trainees answer before revealing. Each click reveals the next answer. '
    '1) FCF customers sign up for a membership; the system issues a free DHJ voucher behind the scenes. An FCF customer who says "I never bought any voucher" is correct from their point of view, so don\'t contradict them. '
    '2) Voucher hours are deducted from the cleaning\'s total hours regardless of the CP\'s hourly rate. If the booking is longer than the voucher, the excess hours are charged at Homeaglow\'s standard rate. '
    '3) In the new CRM, read the code: _expired means expired, _customer_refunded means refunded. No suffix means active: click View all in legacy CRM and check the JOB column. Job details mean the voucher was used; a dash means it is unused. '
    '4) Third-party merchant vouchers (e.g. Groupon) expire 90 days after redemption and are then converted into credits equal to what the customer paid for the voucher. '
    '5) Groupon refund: find and address the reason first, use the Refund voucher (groupon) macro for the self-refund steps, and invalidate the voucher, otherwise Groupon may be unable to refund. '
    '6) Groupon and other industry-standard vouchers cannot be split across appointments, so the system does not issue an admin courtesy voucher for unused hours. It does this only for DHJ vouchers.')

buf = io.BytesIO(); prs.save(buf)
ns = {}
exec(open(__file__.replace('add-slide51-vouchers-knowledge-check.py', 'add_story_anim.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == kc.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        groups = {}
        for sid, n in re.findall(r'<p:cNvPr id="(\d+)" name="step(\d+)[^"]*"', x): groups.setdefault(int(n), []).append(sid)
        cid, clicks, bld = 3, [], []
        for n in sorted(groups):
            inner = ''
            for j, sid in enumerate(groups[n]):
                inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if j == 0 else 'withEffect'); cid += 3
                bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
            clicks.append(ns['click'](cid, inner)); cid += 2
        x = x.replace('</p:sld>', ns['timing'](clicks, bld) + '</p:sld>')
        data = x.encode('utf8'); print('clicks:', len(clicks))
    zout.writestr(item, data)
zout.close()
