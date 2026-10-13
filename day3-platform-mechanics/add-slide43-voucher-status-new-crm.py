"""Day 3 original deck: insert a new slide after 42 showing how to check a voucher's status in the
new C CRM (expired / refunded from the code; active -> View all in legacy CRM -> used vs unused).
Patches the live export so every other slide stays exactly as it is."""
import copy, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from PIL import Image
from lxml import etree

src, dst = sys.argv[1:3]
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'voucher-status')
C = RGBColor.from_string
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE = C('0B7A6F'), C('E3F4F1'), C('1D1D1F'), C('6E6E73'), C('DCE1E7'), C('FFFFFF')
RED, PINK, GREEN, GREENSOFT, GOLD, GOLDSOFT, GREY = C('B3261E'), C('FCEBEA'), C('2E7D32'), C('E4F2E1'), C('9A6410'), C('FFF3D6'), C('EEF0F2')

prs = Presentation(src)
ref = prs.slides[41]
assert ref.shapes[1].text_frame.text.strip() == 'Reading a voucher in the new CRM'
s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
for sh in list(ref.shapes)[:3] + [x for x in ref.shapes if x.is_placeholder]:
    s.shapes._spTree.append(copy.deepcopy(sh._element))
hdr = list(s.shapes)
hdr[1].text_frame.paragraphs[0].runs[0].text = 'Checking voucher status in the C CRM'
for r in hdr[1].text_frame.paragraphs[0].runs[1:]: r.text = ''
hdr[2].text_frame.paragraphs[0].runs[0].text = 'Read the code first. If the voucher is active, check the legacy CRM.'
for r in hdr[2].text_frame.paragraphs[0].runs[1:]: r.text = ''

def nostyle(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None: sh._element.remove(st)
    return sh
def box(x, y, w, h, fill, line=None, kind=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06):
    sh = nostyle(s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h)))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = adj
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    return sh
def text(x, y, w, h, runs, size=9.5, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, shape=None):
    tb = shape or s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = align
    for t, bold, color, *font in runs:
        r = p.add_run(); r.text = t; f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = font[0] if font else 'Arial'
    return tb
def pill(x, y, label, fg, bg):
    w = 0.3 + 0.092 * len(label)
    p = box(x, y, w, 0.24, bg, adj=0.5)
    text(0, 0, 0, 0, [(label, True, fg, 'Play')], 9.5, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, shape=p)
    p.text_frame.word_wrap = False
def shot(name, x, y, maxw, maxh, mark=None):
    path = os.path.join(IMG, name); W, H = Image.open(path).size
    scale = min(maxw / W, maxh / H); w, h = W * scale, H * scale
    pic = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    pic.line.color.rgb = BORDER; pic.line.width = Pt(0.75)
    if mark:  # highlight a region given in source pixels
        mx, my, mw, mh = mark
        box(x + mx * scale, y + my * scale, mw * scale, mh * scale, None, TEAL, MSO_SHAPE.RECTANGLE).line.width = Pt(1.75)
    return h

# Row 1: what the code tells you.
cards = [('Expired', SOFT, GREY, [('The code ends in ', False, SOFT), ('_expired', True, INK), ('.', False, SOFT)], 'expired.png', None),
         ('Refunded', RED, PINK, [('The code ends in ', False, SOFT), ('_customer_refunded', True, INK), ('.', False, SOFT)], 'refunded.png', None),
         ('Active', TEAL, TEALSOFT, [('No suffix. Click ', False, SOFT), ('View all in legacy CRM', True, INK), (' to see if it was used.', False, SOFT)],
          'active-new-crm-crop.png', (108, 63, 136, 26))]
for i, (label, fg, bg, desc, img, mark) in enumerate(cards):
    x = 0.45 + i * 3.08
    box(x, 1.42, 2.95, 1.52, WHITE, BORDER)
    pill(x + 0.15, 1.54, label, fg, bg)
    text(x + 0.15, 1.86, 2.65, 0.36, desc, 9)
    shot(img, x + 0.15, 2.24, 2.65, 0.6, mark)

# Row 2: active vouchers in the legacy CRM.
ar = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(8.08), Inches(2.94), Inches(8.08), Inches(3.14))
ar.line.color.rgb = TEAL; ar.line.width = Pt(1.25)
etree.SubElement(ar.line._get_or_add_ln(), qn('a:tailEnd'), type='triangle', w='med', len='med')
text(0.45, 3.06, 7.4, 0.24, [('Active voucher? In the legacy CRM, look at the ', False, INK), ('JOB', True, TEAL), (' column.', False, INK)], 10.5)
rows = [('Used', GREEN, GREENSOFT, [('It has job details: the job ID and the date it was applied to.', False, SOFT)], 'legacy-used.png', (500, 4, 203, 98)),
        ('Unused', GOLD, GOLDSOFT, [('No job details yet. The JOB column shows a dash.', False, SOFT)], 'legacy-unused.png', (552, 4, 40, 70))]
for i, (label, fg, bg, desc, img, mark) in enumerate(rows):
    x = 0.45 + i * 4.62
    box(x, 3.36, 4.48, 1.62, WHITE, BORDER)
    pill(x + 0.15, 3.48, label, fg, bg)
    text(x + 0.15, 3.8, 4.18, 0.22, desc, 9)
    shot(img, x + 0.15, 4.1, 4.18, 0.78, mark)

s.notes_slide.notes_text_frame.text = (
    "How to check a voucher's status in the C CRM. Start with the code in the Vouchers section. "
    "If it ends in _expired, the voucher has expired. If it ends in _customer_refunded, it was refunded. "
    "If there is no suffix, the voucher is active: click View all in legacy CRM to see whether it has been used. "
    "In the legacy CRM, look at the JOB column. A used voucher has job details, the job ID and the date it was applied to. "
    "An unused voucher has no job details yet; the JOB column shows a dash.")

lst = prs.slides._sldIdLst; new = lst[-1]; lst.remove(new); lst.insert(42, new)
prs.save(dst)
