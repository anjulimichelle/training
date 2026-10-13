"""Day 3 original deck, slide 90 (PCQ: Two kinds of refund): remove the "You control" bar and add a brief
'Refund request: system behavior' panel from the KB. Patches the live export."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
C = RGBColor.from_string
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE = C('0B7A6F'), C('E3F4F1'), C('1D1D1F'), C('6E6E73'), C('DCE1E7'), C('FFFFFF')
prs = Presentation(src)
s = prs.slides[89]
assert s.shapes[1].text_frame.text.strip() == 'Two kinds of refund', s.shapes[1].text_frame.text
bar = [sh for sh in s.shapes if Inches(3.4) < sh.top < Inches(4.0)]
assert any(sh.has_text_frame and sh.text_frame.text.startswith('You control') for sh in bar)
for sh in bar: sh._element.getparent().remove(sh._element)

def box(x, y, w, h, fill, line=None, adj=0.05):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    st = b._element.find(qn('p:style'))
    if st is not None: b._element.remove(st)
    b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None: b.line.fill.background()
    else: b.line.color.rgb = line; b.line.width = Pt(0.75)
    return b
def text(x, y, w, h, paras):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame
    tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.LEFT; p.space_after = Pt(2)
        for t_, bold, color, size, font in runs:
            r = p.add_run(); r.text = t_; f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = font

Y, H = 3.45, 1.72
box(0.45, Y, 9.1, H, TEALSOFT)
text(0.65, Y + 0.12, 8.7, 0.25, [[('Refund request: system behavior', True, TEAL, 11.5, 'Play')]])
items = [('Voucher on the job blocks it', 'Invalidate the voucher first if C should be able to request one. The 72-hour window still applies.'),
         ('Failed job payment', "C can't submit a refund request if the job payment failed."),
         ('CP may accept or decline', "CPs are independent contractors and aren't obligated to respond."),
         ('7 days to respond', "If the CP doesn't respond in time, the system may issue C a courtesy voucher.")]
for k, (head, body) in enumerate(items):
    x = 0.65 + (k % 2) * 4.45; y = Y + 0.46 + (k // 2) * 0.62
    box(x, y + 0.05, 0.1, 0.1, TEAL, adj=0.5)
    text(x + 0.22, y, 4.0, 0.45, [[(head, True, INK, 9.5, 'Arial')], [(body, False, SOFT, 9, 'Arial')]])

n = s.notes_slide.notes_text_frame
t = n.text
cut = t.find('What the agent controls')
if cut != -1: t = t[:cut].rstrip()
n.text = (t + ' Refund request system behavior: a customer can\'t submit a refund request if the system already issued a voucher for the job; '
          'invalidate the voucher first if the customer should be allowed to request one (the 72-hour window still applies, both conditions must be met). '
          'A refund request is also unavailable if the job payment failed. CPs are independent contractors and may accept or decline; they aren\'t obligated to respond. '
          'CPs have 7 days to respond; if that window expires without a response, the system may issue the customer a courtesy voucher automatically.')
prs.save(dst)
