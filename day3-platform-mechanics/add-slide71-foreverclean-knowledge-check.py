"""Day 3 original deck (ForeverClean Membership): after slide 70 add a quick knowledge check (answers reveal per click).
Questions follow the KB article 'ForeverClean Membership'. Patches the live Google export."""
import copy, io, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, WHITE, BORDER, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', '9A6410', 'FFF3D6'
prs = Presentation(src)
S = list(prs.slides)
ref = S[69]
assert ref.shapes[1].text_frame.text.strip() == 'Early Termination Fee (ETF)'
num = [x for x in ref.shapes if x.is_placeholder][0]

def new_slide(eyebrow, title, sub, after):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
        s.shapes._spTree.append(copy.deepcopy(el))
    for shape, t in ((s.shapes[0], eyebrow), (s.shapes[1], title), (s.shapes[2], sub)):
        r = shape.text_frame.paragraphs[0].runs; r[0].text = t
        for x in r[1:]: x._r.getparent().remove(x._r)
    lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); list(lst)[after].addnext(el)
    return s
def kit(s):
    def strip(x):
        st = x._element.find(P_NS + 'style')
        if st is not None: x._element.remove(st)
    def box(x, y, w, h, fill, line=None, adj=0.08, name=None):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
        b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
        else: b.line.fill.background()
        if name: b.name = name
        return b
    def text(shape, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, after=3):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
            for t, bold, col, sz in runs:
                r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'Play' if bold else 'Arial'
        return shape
    def tb(x, y, w, h, name=None):
        t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name: t.name = name
        return t
    return box, text, tb

# ---------- Knowledge check ----------
kc = new_slide('FOREVERCLEAN MEMBERSHIP  ·  QUICK KNOWLEDGE CHECK', 'Quick Knowledge Check', 'Ask each question, then click to reveal the answer.', 69)
box, text, tb = kit(kc)
qa = [('How do you tell if a customer is on FCF or DHJ?', 'Look for the FCF banner in the New or Legacy CRM. No banner: treat them as DHJ.'),
      ('When does the membership start for DHJ? For FCF?', 'DHJ: after the first completed cleaning. FCF: immediately at sign-up.'),
      ('An FCF customer picked a 4-hour first cleaning. What did they pay at sign-up?', '$38: the $19 first MF plus $19 for the longer cleaning (plus tax if applicable). Then $59/month.'),
      ('FCF customer, 3 months into MCT, no cleaning yet, wants to cancel. ETF?', 'No. The ETF only applies after a completed cleaning, even though MCT started at sign-up.'),
      ('How is the ETF calculated?', '$35/hr x the duration of the first cleaning, for both models. Older $49-MF DHJ customers: check the FC table.'),
      ("A DHJ customer's first job is a lockout. What happens to the membership?", 'It is paused until a job is invoiced. An FCF membership keeps renewing.')]
n = len(qa); h = (3.75 - 0.08 * (n - 1)) / n
for i, (q, a) in enumerate(qa):
    y = 1.4 + i * (h + 0.08)
    box(0.45, y, 5.05, h, WHITE, BORDER)
    text(box(0.6, y + (h - 0.3) / 2, 0.3, 0.3, TEALSOFT, adj=0.5), [[(str(i + 1), True, TEAL, 9)]], align=PP_ALIGN.CENTER)
    text(tb(1.05, y, 4.35, h), [[(q, False, INK, 10)]])
    box(5.6, y, 3.95, h, TEALSOFT, name=f'step{i + 1}Box')
    text(tb(5.75, y, 3.7, h, name=f'step{i + 1}Ans'), [[(a, False, INK, 9)]])
kc.notes_slide.notes_text_frame.text = (
    'Quick knowledge check: ForeverClean Membership. Ask each question and let trainees answer before revealing. Each click reveals the next answer. '
    '1) The quickest way to identify an FCF customer is the FCF banner, visible on both the New CRM and Legacy CRM. If the banner is not there, treat the customer as DHJ. Since August 3, 2026, all new enrollments use FCF. '
    '2) DHJ starts with a voucher purchase, so the membership only begins after the first completed cleaning. FCF customers become members immediately at sign-up. Monthly billing, MCT and self-service all follow from when membership starts. '
    '3) FCF sign-up: $19 discounted first MF, plus $19 for a 4-hour first cleaning ($38 total) or $59 for a 6-hour one ($78 total), taxes may apply. After that, $59/month. Never describe FCF as purchasing a voucher. '
    '4) For FCF, the MCT clock and the ETF clock are different. MCT starts at sign-up, but ETF only applies once a cleaning has been completed and before MCT ends. No completed cleaning, no ETF. '
    '5) ETF = $35/hour x the duration of the first cleaning, for both models. Some older Legacy DHJ customers with a $49 monthly fee may have a different formula; the FC table is the source of truth. Explain the ETF before charging it. '
    '6) If a DHJ customer\'s first job is a lockout, the membership is paused; it un-pauses once their next job is invoiced. An FCF customer is already an active member, so billing keeps going. Washington customers are exempt from MCT and ETF under either model.')
buf = io.BytesIO(); prs.save(buf)
ns = {}
exec(open(__file__.replace('add-slide71-foreverclean-knowledge-check.py', 'add_story_anim.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == kc.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        groups = {}
        for sid, n in re.findall(r'<p:cNvPr id="(\d+)" name="step(\d+)[^"]*"', x): groups.setdefault(int(n), []).append(sid)
        cid, clicks, bld = 3, [], []
        for n in sorted(groups):
            inner = ''
            for j, sid in enumerate(groups[n]):
                inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if j == 0 else 'withEffect'); cid += 3
                bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
            clicks.append(ns['click'](cid, inner)); cid += 2
        x = x.replace('</p:sld>', ns['timing'](clicks, bld) + '</p:sld>')
        data = x.encode('utf8'); print('clicks:', len(clicks))
    zout.writestr(item, data)
zout.close()
