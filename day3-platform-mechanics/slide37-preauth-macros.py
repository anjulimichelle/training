"""Day 3 original deck, slide 37 (Before you explain a pre-auth): add a 'Suggested macros' strip
in the same style as the strips on slides 71, 74 and 75. Patches the live export."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK = '0B7A6F', 'E3F4F1', '1D1D1F'
macros = ['Pre-authorization hold: explain how it works', 'Pre-Auth – DHJ First MF (Job Completed)', 'Pre-Auth – DHJ First MF (No job)']
prs = Presentation(src)
s = prs.slides[36]
assert s.shapes[1].text_frame.text.strip() == 'Before you explain a pre-auth', s.shapes[1].text_frame.text
assert not any(x.has_text_frame and x.text_frame.text.startswith('Suggested macro') for x in s.shapes)
b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(4.2), Inches(9.1), Inches(0.82))
st = b._element.find(P_NS + 'style')
if st is not None: b._element.remove(st)
b.adjustments[0] = 0.15; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(TEALSOFT)
b.line.color.rgb = RGBColor.from_string(TEAL); b.line.width = Pt(0.75); b.name = 'suggestedMacro'
tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.18); tf.margin_right = Inches(0.18); tf.margin_top = tf.margin_bottom = Inches(0.06)
def para(first):
    q = tf.paragraphs[0] if first else tf.add_paragraph(); q.alignment = PP_ALIGN.LEFT; return q
def run(q, t, bold, col):
    r = q.add_run(); r.text = t; r.font.name = 'Arial'; r.font.size = Pt(9.5); r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(col)
run(para(True), 'Suggested macros:', True, TEAL)
for m in macros:
    q = para(False); run(q, '•  ', True, TEAL); run(q, m, True, INK)
notes = s.notes_slide.notes_text_frame
notes.text = (notes.text + ' ' if notes.text else '') + 'Suggested macros: ' + '; '.join(macros) + '.'
prs.save(dst)
