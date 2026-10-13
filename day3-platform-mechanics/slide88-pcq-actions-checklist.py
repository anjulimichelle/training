"""Day 3 original deck, slide 88 (PCQ: Required agent actions): replace the Area/Action table with a
four-card checklist (Customer, Resolution, Account, CP) that follows the KB's Required Agent Actions."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
C = RGBColor.from_string
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE = C('0B7A6F'), C('E3F4F1'), C('1D1D1F'), C('6E6E73'), C('DCE1E7'), C('FFFFFF')
prs = Presentation(src)
s = prs.slides[87]
assert s.shapes[1].text_frame.text.strip() == 'Required agent actions', s.shapes[1].text_frame.text
for sh in [x for x in s.shapes if x.has_table]: sh._element.getparent().remove(sh._element)

def box(x, y, w, h, fill, line=None, adj=0.05):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    st = b._element.find(qn('p:style'))
    if st is not None: b._element.remove(st)
    b.adjustments[0] = adj
    if fill is None: b.fill.background()
    else: b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None: b.line.fill.background()
    else: b.line.color.rgb = line; b.line.width = Pt(1)
    return b
def text(x, y, w, h, runs, size):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame
    tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    for t_, bold, color, *font in runs:
        r = p.add_run(); r.text = t_; f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = font[0] if font else 'Arial'
    return t

N = lambda t: (t, False, INK)
B = lambda t: (t, True, INK)
G = lambda t: (t, False, SOFT)
cards = {
    'Customer': (0.45, 1.4, 4.48, 1.08, [
        [N('Acknowledge and empathize with the '), B('specific'), N(' experience. '), G('No generic "I understand your frustration."')],
        [N('Probe if the explanation is vague, so C can share what actually went wrong.')]]),
    'Resolution': (0.45, 2.58, 4.48, 2.5, [
        [N('Check refund request eligibility ('), B('72 hours'), N(') and blockers.')],
        [N('Refund the '), B('Premium Upsell Fee'), N(' if charged and the service issue is confirmed, even without refund intent.')],
        [N('Refund the '), B('Priority Markup Fee'), N(' only if C asks for it.')],
        [B('Leave the Job ID blank when you issue a PCQ voucher. '), G('A voucher tied to the job blocks C from requesting a refund inside the window.')]]),
    'Account': (5.07, 1.4, 4.48, 2.42, [
        [N('Flag the job in job admin > Do as '), B('customer_unsatisfied'), N(' if C hasn\'t left a review. '),
         G('Job over 30 days old: block comms first, check the auto-message is blocked, then remove the block.')],
        [N('Block the '), B('C/CP pairing'), N(', unless C wants to retry the same CP.')],
        [N('Cancel upcoming appointments/RC only if C no longer wants the service. '), G('Keep them if C shows signs of wanting to.')]]),
    'CP': (5.07, 3.92, 4.48, 1.16, [
        [N('Coach or penalize the CP as necessary.')],
        [N('Handle a cleaning supplies issue if mentioned.')]]),
}
LINE_H = 0.155  # 9 pt Arial, approx inches per wrapped line
CPL = 68        # approx characters per line in a card
def lines(runs): return max(1, -(-sum(len(r[0]) for r in runs) // CPL))
def need(items): return 0.42 + sum(lines(r) * LINE_H + 0.1 for r in items) + 0.06
TOP, BOTTOM, GAP = 1.4, 5.08, 0.12
for col in (('Customer', 'Resolution'), ('Account', 'CP')):
    first, second = col
    x, _, w, _, it1 = cards[first]; _, _, _, _, it2 = cards[second]
    h1 = need(it1); h2 = BOTTOM - TOP - GAP - h1
    assert h2 >= need(it2), (second, h2, need(it2))
    cards[first] = (x, TOP, w, h1, it1); cards[second] = (x, TOP + h1 + GAP, w, h2, it2)
for title, (x, y, w, h, items) in cards.items():
    box(x, y, w, h, WHITE, BORDER)
    box(x, y, 0.07, h, TEAL, adj=0.5)
    text(x + 0.2, y + 0.1, w - 0.4, 0.25, [(title, True, TEAL, 'Play'), (f'   {len(items)} steps', False, SOFT)], 11)
    cy = y + 0.42
    for runs in items:
        box(x + 0.22, cy + 0.02, 0.13, 0.13, WHITE, TEAL, adj=0.15)
        text(x + 0.45, cy, w - 0.62, lines(runs) * LINE_H, runs, 9)
        cy += lines(runs) * LINE_H + 0.1

n = s.notes_slide.notes_text_frame
n.text = ("Use this as the checklist; trainees tick each box. Customer: empathize with the specific experience (no generic \"I understand your frustration\") and probe if the explanation is vague. "
          "Resolution: check refund request eligibility (72 hours) and blockers; refund the Premium Upsell Fee if charged and a service issue is confirmed, even without refund intent; refund the Priority Markup Fee only if the customer asks; "
          "leave the Job ID blank when issuing a PCQ voucher, because a voucher tied to the job blocks the customer from requesting a refund inside the window. "
          "Account: flag the job in job admin > Do as customer_unsatisfied if there is no review yet (for jobs over 30 days old, block comms first, ensure the auto-message is blocked, then remove the block); block the C/CP pairing unless the customer wants to retry the same CP; "
          "cancel upcoming appointments or RC only if the customer no longer wants the service, unless there are signs they want to keep them. CP: coach or penalize as necessary, and handle a cleaning supplies issue if mentioned.")
prs.save(dst)
