"""Day 3 Platform & Operating Mechanics deck (Value of ForeverClean): add 'Suggested macro' strips to
slide 71 (ForeverClean overview / FCF membership explanation), slide 74 (free month offer) and slide 75 (ETF terms)."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK = '0B7A6F', 'E3F4F1', '1D1D1F'
prs = Presentation(src)
S = list(prs.slides)
plan = {71: ('Lead with value', 'Suggested macros:', ['Membership details: ForeverClean overview', 'FCF - Membership Explanation']),
        74: ('Customer never used the membership', 'Suggested macro:', ['[Fragment] Free month offer']),
        75: ('Explaining the ETF', 'Suggested macro:', ['[Fragment] ETF terms explanation'])}
for n, (title, label, macros) in plan.items():
    s = S[n - 1]
    assert s.shapes[1].text_frame.text.strip() == title, (n, s.shapes[1].text_frame.text)
    assert not any(x.has_text_frame and x.text_frame.text.startswith('Suggested macro') for x in s.shapes)
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(4.15), Inches(9.1), Inches(0.42))
    st = b._element.find(P_NS + 'style')
    if st is not None: b._element.remove(st)
    b.adjustments[0] = 0.15; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(TEALSOFT)
    b.line.color.rgb = RGBColor.from_string(TEAL); b.line.width = Pt(0.75); b.name = 'suggestedMacro'
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.18); tf.margin_right = Inches(0.18); tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    runs = [(label + '  ', True, TEAL)]
    for i, m in enumerate(macros):
        if i: runs.append(('   ·   ', False, TEAL))
        runs.append((m, True, INK))
    for t, bold, col in runs:
        r = p.add_run(); r.text = t; r.font.name = 'Arial'; r.font.size = Pt(10); r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(col)
    notes = s.notes_slide.notes_text_frame
    notes.text = (notes.text + ' ' if notes.text else '') + f'{label} ' + '; '.join(macros) + '.'
prs.save(dst)
