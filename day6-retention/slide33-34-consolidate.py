"""Day 6 deck: fold slide 33 (Fix the real problem) into slide 34 (The 5 key factors).

Drops the old-vs-new comparison, keeps the "root cause first" tile as a strip
above the 5-factor table, merges the speaker notes, then removes slide 33.

Usage: python3 slide33-34-consolidate.py in.pptx out.pptx
"""
import copy, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn

prs = Presentation(sys.argv[1])
s33, s34 = prs.slides[32], prs.slides[33]
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
assert txt(s33.shapes[1]) == 'Fix the real problem'
assert txt(s34.shapes[1]) == 'The 5 key factors'


def geom(sh, x, y, w, h):
    sh.left, sh.top, sh.width, sh.height = Pt(x), Pt(y), Pt(w), Pt(h)


hdr = next(sh for sh in s33.shapes if txt(sh) == 'Now: root cause first')
body = next(sh for sh in s33.shapes if txt(sh).startswith('Find why they want to leave'))
bg = list(s33.shapes)[list(s33.shapes).index(hdr) - 1]
assert bg.shape_type == 1 and not txt(bg)

tree = s34.shapes._spTree
title_el = s34.shapes[2]._element  # insert after the subtitle
new = []
for src in (bg, hdr, body):
    el = copy.deepcopy(src._element)
    title_el.addnext(el) if not new else new[-1].addnext(el)
    new.append(el)
shapes = {sh._element: sh for sh in s34.shapes}
nbg, nhdr, nbody = (shapes[e] for e in new)
for sh in (nbg, nhdr, nbody):
    c = sh._element.find('.//' + qn('p:cNvPr'))
    c.set('id', str(s34.shapes._next_shape_id)); c.set('name', 'root cause first')

geom(nbg, 32, 102, 655, 34)
geom(nhdr, 32, 102, 150, 34)
geom(nbody, 196, 102, 484, 34)
nhdr.text_frame._txBody.findall('.//' + qn('a:t'))[0].text = 'Root cause first'
for t in nhdr.text_frame._txBody.findall('.//' + qn('a:t'))[1:]:
    t.text = ''

# three bullets -> one line of steps
tb = nbody.text_frame._txBody
ps = tb.findall(qn('a:p'))
steps = [''.join(t.text or '' for t in p.iter(qn('a:t'))).strip() for p in ps]
assert steps == ['Find why they want to leave', 'Match the offer to that reason', 'Show its value clearly'], steps
p0 = ps[0]
for p in ps[1:]:
    tb.remove(p)
ppr = p0.find(qn('a:pPr'))
ppr.set('marL', '0'); ppr.set('indent', '0')
for tag in ('a:buChar', 'a:buAutoNum', 'a:buFont', 'a:buSzPts', 'a:buClr'):
    for e in ppr.findall(qn(tag)):
        ppr.remove(e)
if ppr.find(qn('a:buNone')) is None:
    ppr.append(ppr.makeelement(qn('a:buNone'), {}))
runs = p0.findall(qn('a:r'))
runs[0].find(qn('a:t')).text = '  →  '.join(steps)
for r in runs[1:]:
    p0.remove(r)
nbody.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

# table and History bar move down
table = next(sh for sh in s34.shapes if sh.has_table)
table.top = Pt(144)
for tr in table._element.iter(qn('a:tr')):
    tr.set('h', str(int(27.5 * 12700)))
hist_txt = next(sh for sh in s34.shapes if txt(sh).startswith('History:'))
k = list(s34.shapes).index(hist_txt)
for sh in list(s34.shapes)[k - 2:k + 1]:
    sh.top = sh.top + Pt(322 - 299)

# notes: slide 33's root-cause framing goes first (without the old-ladder comparison)
n34 = s34.notes_slide.notes_text_frame
add = ("Every resolution should go straight to the root cause of why the customer is thinking about "
       "cancelling, not a generic offer thrown at the problem. Find why they want to leave, match the "
       "offer to what the customer actually needs, and show its value clearly. The goal is for the "
       "customer to feel understood and like their specific problem is being solved.")
if not n34.text.startswith('Every resolution'):
    first = n34.paragraphs[0]._p
    newp = copy.deepcopy(first)
    for r in newp.findall(qn('a:r'))[1:]:
        newp.remove(r)
    newp.findall(qn('a:r'))[0].find(qn('a:t')).text = add
    first.addprevious(newp)

# remove slide 33
lst = prs.slides._sldIdLst
sld = lst[32]
rid = sld.get(qn('r:id'))
lst.remove(sld)
prs.part.drop_rel(rid)

prs.save(sys.argv[2])
print('saved', sys.argv[2])
