"""Day 6 deck, merged slide 33 (The 5 key factors): add back the goal bar from the
old "Fix the real problem" slide under the History bar.

Usage: python3 slide33-add-goal.py current.pptx original-with-old-slide33.pptx out.pptx
"""
import copy, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.oxml.ns import qn

prs = Presentation(sys.argv[1]); old = Presentation(sys.argv[2])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
s = prs.slides[32]; o = old.slides[32]
assert txt(s.shapes[1]) == 'The 5 key factors' and txt(o.shapes[1]) == 'Fix the real problem'
assert not any(txt(sh).startswith('The goal') for sh in s.shapes)

ol = list(o.shapes)
k = next(i for i, sh in enumerate(ol) if txt(sh).startswith('The goal'))
gbg, gicon, gtxt = ol[k - 2], ol[k - 1], ol[k]
assert gicon.shape_type == 13

# tighten table rows, move History bar up, goal bar below it
table = next(sh for sh in s.shapes if sh.has_table)
for tr in table._element.iter(qn('a:tr')):
    tr.set('h', str(int(26 * 12700)))
sl = list(s.shapes)
h = next(i for i, sh in enumerate(sl) if txt(sh).startswith('History:'))
for sh in sl[h - 2:h + 1]:
    sh.top = sh.top - Pt(322 - 306)

blob = gicon._element.find('.//' + qn('a:blip'))
img = o.part.related_part(blob.get(qn('r:embed')))
# copy the image part bytes into this deck via a picture add, then swap geometry
tmp = s.shapes.add_picture(__import__('io').BytesIO(img.blob), 0, 0)
rid = tmp._element.find('.//' + qn('a:blip')).get(qn('r:embed'))
tmp._element.getparent().remove(tmp._element)

anchor = sl[h]._element
els = [copy.deepcopy(e._element) for e in (gbg, gicon, gtxt)]
els[1].find('.//' + qn('a:blip')).set(qn('r:embed'), rid)
for el in els:
    anchor.addnext(el); anchor = el
dy = Pt(340) - gbg.top
for el in els:
    sh = next(x for x in s.shapes if x._element is el)
    c = el.find('.//' + qn('p:cNvPr')); c.set('id', str(s.shapes._next_shape_id)); c.set('name', 'goal bar')
    sh.top = sh.top + dy

prs.save(sys.argv[3]); print('saved', sys.argv[3])
