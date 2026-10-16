"""Build the "Projects" training deck in the New Hire Care Training theme.
Template: an export of the Day 8 deck (cover, topic divider, content header and Questions slides
are reused for the theme). Content: the Care "Projects" process guide (Notion), as pasted by the
trainer: FC Starts, Redemption Assist, FC Reactivations, dhj_bookmark_sms, CC Fails and
Third-Party Poor Reviews, in that order.

Usage: python3 build-deck.py day8-template.pptx out.pptx
"""
import copy, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

C = {'TEAL': '0B7A6F', 'TEAL_PALE': 'E3F4F1', 'AMBER': '9A6410', 'AMBER_PALE': 'FFF3D6',
     'INK': '1D1D1F', 'MUTED': '6E6E73', 'WHITE': 'FFFFFF', 'LINE': 'DCE1E7', 'GREY': '6E6E73', 'ALT': 'F5F6F8'}
X, W = 32.4, 655.2

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
S = list(prs.slides)
lst = prs.slides._sldIdLst
el_of = {id(s): e for s, e in zip(S, list(lst))}
cover = S[0]
divider = next(s for s in S if txt(s.shapes[0]).startswith('TOPIC') and len(s.shapes) > 20)
questions = next(s for s in S if txt(s.shapes[1]) == 'Questions?')
content_ref = next(s for s in S if txt(s.shapes[0]) == 'ASSESSING DISPUTE LIKELIHOOD' and txt(s.shapes[1]) == 'What is this?')
keep = {id(cover), id(divider), id(questions)}


def set_runs(shape, t):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    lines = t.split('\n')
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    p0.runs[0].text = lines[0]
    for ln in lines[1:]:
        np_ = copy.deepcopy(p0._p); p0._p.getparent().append(np_)
        rr = np_.findall(qn('a:r'))
        rr[0].find(qn('a:t')).text = ln
    return shape



def new_slide(eyebrow, title, sub):
    s = prs.slides.add_slide(content_ref.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    num = [x for x in content_ref.shapes if x.is_placeholder][0]
    for e in (content_ref.shapes[0]._element, content_ref.shapes[1]._element, content_ref.shapes[2]._element, num._element):
        s.shapes._spTree.append(copy.deepcopy(e))
    set_runs(s.shapes[0], eyebrow); set_runs(s.shapes[1], title); set_runs(s.shapes[2], sub)
    order.append(s)
    return s


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill_tf(tf, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, after=2):
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, item in enumerate(paras):
        num, runs = item if isinstance(item, tuple) else ('', item)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after); p.alignment = align
        if num:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int(12 * 12700))); pPr.set('indent', str(-10 * 12700))
            runs = [(num + ' ', runs[0][1], True, 'TEAL' if num != '•' else 'MUTED')] + runs
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(C[color])


def box(s, x, y, w, h, bg, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=9,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=6, lw=0.75):
    sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(C[bg])
    if line:
        sh.line.color.rgb = RGBColor.from_string(C[line]); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(3)
    fill_tf(tf, paras, anchor, align)
    return sh


def text(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, after=2, name=None):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, paras, anchor, align, after)
    if name:
        tb.name = name
    return tb


def arrow(s, x1, y1, x2, y2, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)); strip(c)
    c.line.color.rgb = RGBColor.from_string(C['GREY']); c.line.width = Pt(1)
    if head:
        ln = c._element.find('.//' + qn('a:ln'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')


def dot(s, x, y, label, bg='TEAL', d=20, sz=9):
    return box(s, x, y, d, d, bg, paras=[[(label, sz, True, 'WHITE')]], align=PP_ALIGN.CENTER, margin=0, shape=MSO_SHAPE.OVAL)


B = lambda t, c='INK', sz=8.5: (t, sz, True, c)
A = lambda t, c='INK', sz=8.5: (t, sz, False, c)


def tile(s, x, y, w, h, n, head, body, col='TEAL', pale='TEAL_PALE', sz=8.5):
    box(s, x, y, w, h, 'WHITE', 'LINE')
    dot(s, x + 10, y + 10, n, col, 22)
    text(s, x + 40, y + 12, w - 50, 18, [[B(head, col, 10.5)]])
    text(s, x + 12, y + 40, w - 24, h - 46, body, after=3)



def reveal_timing(s, groups):
    """One click per group: fade every shape in the group in together."""
    def effect(cid, spid, node):
        tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
        return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
                f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
                f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
                f'</p:childTnLst></p:cTn></p:par>')
    cid, clicks = 3, ''
    for ids in groups:
        start, inner = cid, ''
        cid += 2
        for i, spid in enumerate(ids):
            inner += effect(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
        clicks += (f'<p:par><p:cTn id="{start}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                   f'<p:par><p:cTn id="{start+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                   f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    bld = ''.join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for g in groups for i in g)
    NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + clicks +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
    ext = s._element.find(qn('p:extLst'))
    t_el = etree.fromstring(tim)
    (ext.addprevious(t_el) if ext is not None else s._element.append(t_el))


# ---- cover and divider ----
set_runs(cover.shapes[0], 'HOMEAGLOW  ·  NEW HIRE CARE TRAINING')
set_runs(cover.shapes[1], 'Projects')
set_runs(cover.shapes[2], 'When a ticket belongs to a project team, and how to reassign it.')
cover.notes_slide.notes_text_frame.text = ('Projects. Source: the Care "Projects" process guide in Notion. Six active projects are covered, '
                                           'in this order: FC Starts, Redemption Assist, FC Reactivations, dhj_bookmark_sms, CC Fails and Third-Party Poor Reviews.')
set_runs(divider.shapes[0], 'PROJECTS')
set_runs(divider.shapes[1], 'Projects')
set_runs(divider.shapes[2], 'Some tickets belong to a project team. Check first, then reassign or handle BAU.')
for k, t in zip((6, 9, 12, 15), ('The six active projects', 'What to check before you reassign',
                                   'The reassignment steps', 'Special cases: STOP, TRs, Project Leads')):
    set_runs(divider.shapes[k], t)
divider.notes_slide.notes_text_frame.text = (
    'Projects are teams of agents working specific groups of customers: getting voucher holders to book, reactivating FC, following up on '
    'failed charges and poor public reviews. When one of those customers writes in, the ticket may belong to the project instead of to us. '
    'By the end, trainees can tell when a ticket belongs to a project, reassign it the right way, and handle the special cases.')
order = [cover, divider]

EB = 'PROJECTS'
GAP = 8


def checks_box(s, x, y, w, h, items, title='What to check'):
    """Teal checklist panel. items: list of run-lists."""
    box(s, x, y, w, h, 'TEAL_PALE', anchor=MSO_ANCHOR.TOP)
    text(s, x + 12, y + 10, w - 24, 16, [[B(title, 'TEAL', 10)]])
    text(s, x + 12, y + 32, w - 24, h - 38, [('✓', r) for r in items], after=5)


def reassign_box(s, x, y, w, h, dm=False, title='Reassign', col='TEAL', tail='Resolve the ticket.'):
    """The standard reassignment steps as a numbered panel."""
    box(s, x, y, w, h, 'WHITE', col, anchor=MSO_ANCHOR.TOP)
    text(s, x + 12, y + 10, w - 24, 16, [[B(title, col, 10)]])
    steps = [[A('Log the ticket on the '), B('Project Reassignment Form')],
             [A('Add the internal note '), B('(format below)')]]
    if dm:
        steps.append([A('Send a '), B('Slack DM'), A(' to the Project agent who handled it')])
    steps.append([A(tail)])
    text(s, x + 12, y + 32, w - 24, h - 38, [(str(i + 1), r) for i, r in enumerate(steps)], after=4)


def note_format(s, x, y, w, h):
    box(s, x, y, w, h, 'ALT', 'LINE', anchor=MSO_ANCHOR.MIDDLE, margin=12, paras=[
        [B('Internal note:  ', 'MUTED', 8), A('Reassigned to Project: ', 'INK', 8.5), A('<project>', 'MUTED', 8.5),
         A('    Link to the customer\'s pending message: ', 'INK', 8.5), A('<link>', 'MUTED', 8.5)]])


def scenarios(s, y, h, items):
    """Amber 'Specific scenarios' row. items: (when, action runs)."""
    n = len(items); w = (W - GAP * (n - 1)) / n
    for i, (when, act) in enumerate(items):
        x = X + i * (w + GAP)
        box(s, x, y, w, h, 'AMBER_PALE', anchor=MSO_ANCHOR.TOP)
        text(s, x + 11, y + 8, w - 22, h - 12, [[B(when, 'AMBER', 8.5)], act], after=3)


# ---- Overview: the six projects ----
s = new_slide(EB, 'What is this?', 'Six active projects. Click a project to jump to it.')
projects = [('FC Starts', [A('Unused DHJ voucher, and an '), B('FC Starts'), A(' agent reached out to get C to book')]),
            ('Redemption Assist', [A('Unused DHJ voucher, and a '), B('Redemption Assist'), A(' agent reached out to get C to book')]),
            ('FC Reactivations', [A('FC is '), B('deactivated'), A(', and a project agent reached out with offers to reactivate')]),
            ('dhj_bookmark_sms', [A('C replied to the '), B('dhj_bookmark_sms'), A(' marketing SMS and has '), B('no DHJ voucher')]),
            ('CC Fails', [A('Failed '), B('job'), A(' charge, job 72 hours to 30 days ago, C ID not ending in 2 or 3')]),
            ('Third-Party Poor Reviews', [B('public_poor_review'), A(' flag, and an internal note from a Project Agent')])]
tw, th = (W - GAP) / 2, 58
overview_tiles = []
for i, (name, body) in enumerate(projects):
    c, r = divmod(i, 3)
    x, y = X + c * (tw + GAP), 100 + r * (th + 7)
    t = box(s, x, y, tw, th, 'WHITE', 'LINE')
    dot(s, x + 12, y + (th - 24) / 2, f'{i + 1:02d}', 'TEAL', 24, 8.5)
    text(s, x + 46, y + 9, tw - 72, th - 14, [[B(name, 'INK', 10.5)], body], after=2)
    chev = text(s, x + tw - 22, y, 12, th, [[B('›', 'TEAL', 14)]], anchor=MSO_ANCHOR.MIDDLE)
    overview_tiles.append((t, chev))
box(s, X, 100 + 3 * (th + 7) + 3, W, 34, 'AMBER_PALE', margin=14, paras=[
    [B('None of the checks match?  ', 'AMBER', 9.5), A('Handle the ticket BAU (business as usual), like any other ticket.', 'INK', 9.5)]])
s.notes_slide.notes_text_frame.text = (
    'Six active projects, in the order we will cover them. FC Starts and Redemption Assist both work customers with an unused DHJ voucher; '
    'the difference is which project the agent who reached out belongs to. FC Reactivations works deactivated FC customers. dhj_bookmark_sms '
    'is a marketing SMS: a reply from a customer with no DHJ voucher goes to the project. CC Fails covers failed job charges from 72 hours to '
    '30 days ago, for customers whose C ID does not end in 2 or 3. Third-Party Poor Reviews covers accounts with a public_poor_review flag and '
    'an internal note from a Project Agent. If none of the checks match, handle the ticket BAU.')

# ---- The reassignment steps ----
s = new_slide(EB, 'How to reassign a ticket', 'The same four steps for every project. The Slack DM is the only one that varies.')
labels = [('1', 'Log it', [A('Log the ticket on the '), B('Project Reassignment Form'), A('.')]),
          ('2', 'Add an internal note', [A('On the account, in the format below.')]),
          ('3', 'Slack DM', [A('Message the Project agent who handled it. '), B('FC Starts, Redemption Assist and FC Reactivations only.')]),
          ('4', 'Resolve', [A('Resolve the ticket. (CC Fails: clear it.)')])]
sw = (W - 3 * 18) / 4
for i, (n, h, body) in enumerate(labels):
    x = X + i * (sw + 18)
    col = 'AMBER' if n == '3' else 'TEAL'
    box(s, x, 100, sw, 100, 'TEAL_PALE' if col == 'TEAL' else 'AMBER_PALE', anchor=MSO_ANCHOR.TOP)
    dot(s, x + 10, 110, n, col, 22)
    text(s, x + 40, 113, sw - 48, 18, [[B(h, col, 10)]])
    text(s, x + 12, 140, sw - 24, 56, [body], after=2)
    if i < 3:
        arrow(s, x + sw + 2, 150, x + sw + 16, 150)
note_format(s, X, 210, W, 30)
box(s, X, 250, W, 92, 'AMBER_PALE', 'AMBER', anchor=MSO_ANCHOR.TOP, margin=14, paras=[
    [B('C says "STOP", "Unsubscribe" or anything like it', 'AMBER', 10)],
    [A('Unsubscribe C, then reassign to '), B('DNC'), A(' (do not contact) on the Project Reassignment Form instead of the project. '
       'Applies to FC Starts, Redemption Assist, FC Reactivations and dhj_bookmark_sms.', 'INK', 9)],
    [B('Not unsubscribing a customer who clearly asked not to be contacted is a ZTP (zero tolerance policy) violation.', 'INK', 9)]])
s.notes_slide.notes_text_frame.text = (
    'Every project uses the same reassignment steps. One: log the ticket on the Project Reassignment Form. Two: create an internal note on the account '
    'in this format: "Reassigned to Project:" followed by the project, and "Link to the customer\'s pending message:" followed by the link. '
    'Three, for FC Starts, Redemption Assist and FC Reactivations only: send a Slack DM to the Project agent who previously handled the ticket. '
    'Four: resolve the ticket. For CC Fails, the guide says you may clear the ticket once done.\n'
    'The one rule that overrides everything: if the customer replies "STOP", "Unsubscribe" or anything similar that shows they do not want to be contacted, '
    'unsubscribe them and reassign to DNC on the Project Reassignment Form instead. Failing to unsubscribe a customer after they clearly asked is a ZTP violation.')

DHJ_CHECKS = [[A('The account has an '), B('unused DHJ voucher')],
              [A('A project agent reached out to C to '), B('encourage booking'), A(' with it')],
              [A('Look up the agent\'s profile on '), B('Slack'), A(' to see which project they\'re in')]]

# ---- 01 FC Starts ----
s = new_slide(EB, '01  FC Starts', 'Unused DHJ voucher, and the agent who reached out is under FC Starts.')
project_slides = [s]
checks_box(s, X, 100, 210, 150, DHJ_CHECKS)
box(s, X + 218, 100, 200, 150, 'WHITE', 'LINE', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('Handle BAU if', 'AMBER', 10)],
    [A('The FC Starts agent '), B('already gave offers'), A(', and C still insists on a refund or cancellation.')]])
reassign_box(s, X + 426, 100, W - 426, 112, dm=True, title='Otherwise: reassign')
note_format(s, X + 426, 218, W - 426, 32)
scenarios(s, 260, 82, [
    ('C says "STOP" or "Unsubscribe"', [A('Unsubscribe C and reassign to '), B('DNC'), A(' on the Project Reassignment Form. Missing it is a ZTP violation.')]),
    ('The action item is a TR from the project agent', [A('A timed reminder (TR) to encourage C to book: '), B('follow the reassignment process'), A('.')])])
s.notes_slide.notes_text_frame.text = (
    'FC Starts. What to check: the account has an unused DHJ voucher, and a project agent has reached out to the customer encouraging them to book with it. '
    'Look for the Project Agent\'s profile on Slack to see which project they belong to. This slide is for agents under FC Starts.\n'
    'If the FC Starts agent has already provided offers and the customer still insists on a refund or cancellation, handle BAU. Otherwise: log the ticket on the '
    'Project Reassignment Form, add the internal note (Reassigned to Project, Link to the customer\'s pending message), send a Slack DM to the Project agent who '
    'previously handled the ticket, then resolve the ticket.\n'
    'Specific scenarios: if the customer says STOP or Unsubscribe, unsubscribe them and reassign to DNC on the form; failing to do so is a ZTP violation. '
    'If the action item is a timed reminder from the project agent to encourage the customer to book, follow the reassignment process.')

# ---- 02 Redemption Assist ----
s = new_slide(EB, '02  Redemption Assist', 'Unused DHJ voucher, and the agent who reached out is under Redemption Assist.')
project_slides.append(s)
checks_box(s, X, 100, 210, 150, DHJ_CHECKS)
box(s, X + 218, 100, 200, 150, 'WHITE', 'LINE', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('Handle BAU if C', 'AMBER', 10)],
    ('•', [A('Wants to '), B('cancel or refund'), A(' the voucher')]),
    ('•', [A('Has a '), B('general inquiry')]),
    ('•', [A('Is messaging about '), B('something else')])])
reassign_box(s, X + 426, 100, W - 426, 112, dm=True, title='C wants to book: reassign')
note_format(s, X + 426, 218, W - 426, 32)
scenarios(s, 260, 82, [
    ('C says "STOP" or "Unsubscribe"', [A('Unsubscribe C and reassign to '), B('DNC'), A(' on the Project Reassignment Form. Missing it is a ZTP violation.')]),
    ('The action item is a TR from the project agent', [A('A timed reminder (TR) to encourage C to book: '), B('follow the reassignment process'), A('.')])])
s.notes_slide.notes_text_frame.text = (
    'Redemption Assist. Same checks as FC Starts: an unused DHJ voucher, and a project agent reached out encouraging the customer to book. Check the agent\'s Slack '
    'profile; this slide is for agents under Redemption Assist.\n'
    'Here the rule is narrower. If the customer wants to cancel or refund the voucher, has a general inquiry, or is messaging us about something else, handle BAU. '
    'Only if the customer wants to book an appointment: log the ticket on the Project Reassignment Form, add the internal note, send a Slack DM to the Project agent '
    'who previously handled the ticket, then resolve the ticket.\n'
    'Specific scenarios: STOP or Unsubscribe means unsubscribe and reassign to DNC (a ZTP violation if missed). A timed reminder from the project agent to encourage '
    'booking: follow the reassignment process.')

# ---- 03 FC Reactivations ----
s = new_slide(EB, '03  FC Reactivations', 'Deactivated FC, and a project agent reached out with offers to reactivate.')
project_slides.append(s)
checks_box(s, X, 100, 210, 150, [[A('FC status is '), B('deactivated')],
                                 [A('A project agent reached out with '), B('offers'), A(' encouraging C to reactivate FC')]])
box(s, X + 218, 100, 200, 150, 'WHITE', 'LINE', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('Handle BAU if', 'AMBER', 10)],
    [A('C '), B('rejects the offer'), A(' to reactivate.')]])
reassign_box(s, X + 426, 100, W - 426, 112, dm=True, title='Interested? Reassign')
note_format(s, X + 426, 218, W - 426, 32)
scenarios(s, 260, 82, [
    ('C says "STOP" or "Unsubscribe"', [A('Unsubscribe C and reassign to '), B('DNC'), A(' on the Project Reassignment Form. Missing it is a ZTP violation.')]),
    ('The action item is that agent\'s TR', [A('A timed reminder (TR) from the project agent, assigned to them: '), B('follow the reassignment process'), A('.')]),
    ('FC already reactivated, C has concerns', [A('See the next slide: some go to the '), B('Project Lead'), A('.')])])
s.notes_slide.notes_text_frame.text = (
    'FC Reactivations. What to check: FC status is deactivated, and a project agent reached out with offers encouraging the customer to reactivate.\n'
    'If the customer rejects the offer to reactivate, handle BAU. If they are interested in reactivating, show hints that they may want to keep using the service, '
    'or can be negotiated with: log the ticket on the Project Reassignment Form, add the internal note, send a Slack DM to the Project agent who previously handled '
    'the ticket, then resolve.\n'
    'Specific scenarios: STOP or Unsubscribe means unsubscribe and reassign to DNC (a ZTP violation if missed). If the action item is a timed reminder from the project '
    'agent and is assigned specifically to that agent, follow the reassignment process. If the FC is already reactivated and the customer has concerns, see the next slide.')

# ---- 03 FC Reactivations: already reactivated ----
s = new_slide(EB, '03  FC Reactivations: already reactivated', 'C has concerns about a reactivation. These go to the Project Lead.')
LW = 400
box(s, X, 100, LW, 200, 'WHITE', 'LINE', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('Reactivated 48 hours ago or less, and C', 'TEAL', 10)],
    ('•', [A('Says they '), B('did not agree'), A(' to reactivate')]),
    ('•', [A('Disputes the '), B('changes'), A(' the Project agent made')]),
    ('•', [A('Has questions '), B('not covered in the post-call email'), A(' (e.g. asks if the credits are recurring, and the email didn\'t say)')]),
    ('•', [A('Wants their '), B('ETF refunded'), A(' after the reactivation')]),
    ('•', [A('Is looking for the Project agent they spoke with, for a concern not listed above')])])
box(s, X, 308, LW, 34, 'ALT', margin=12, paras=[
    [B('Reactivated more than 48 hours ago:  ', 'TEAL', 9), A('only when C disputes the membership terms the agent set up.', 'INK', 9)]])
RX, RW = X + LW + 12, W - LW - 12
box(s, RX, 100, RW, 242, 'TEAL_PALE', 'TEAL', anchor=MSO_ANCHOR.TOP, margin=14, paras=[
    [B('Reassign to the Project Lead', 'TEAL', 10.5)],
    [A('The fc_reactivation team lead in CRM: '), B('John Adrian Suiza'), A('.')],
    [A(' ', 'INK', 4)],
    ('1', [A('Log the ticket on the '), B('Project Reassignment Form')]),
    ('2', [A('Add the internal note')]),
    ('3', [A('Reassign to the Project Lead')])])
s.notes_slide.notes_text_frame.text = (
    'When the FC is already reactivated and the customer has concerns, some cases go to the Project Lead instead of the agent. If the FC was reactivated 48 hours ago or '
    'less and the customer: claims they did not agree to reactivate; is disputing the changes made by the Project agent they spoke with; has questions or concerns not '
    'covered in the post-call email (for example, asking whether the credits issued are recurring when the email did not say); wants their ETF refunded after the FC was '
    'reactivated; or is looking for the Project agent they spoke with, for a concern not covered above. Also if the FC was reactivated more than 48 hours ago and the '
    'customer is disputing the membership terms the FC Reactivations agent implemented.\n'
    'For all of these, reassign to the Project Lead, the fc_reactivation team lead in CRM, John Adrian Suiza, and follow the reassignment process using the Project '
    'Reassignment Form.')

# ---- 04 dhj_bookmark_sms ----
s = new_slide(EB, '04  dhj_bookmark_sms', 'C replied to our marketing SMS and has no DHJ voucher.')
project_slides.append(s)
checks_box(s, X, 100, 210, 150, [[A('C responded to our marketing SMS '), B('dhj_bookmark_sms'), A(' (whatever the message)')],
                                 [A('C does '), B('not'), A(' have a DHJ voucher on their account')]])
box(s, X + 218, 100, 200, 150, 'WHITE', 'LINE', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('Handle BAU if', 'AMBER', 10)],
    [A('Neither condition is true.')],
    [A('C replied to the SMS but '), B('already has a DHJ voucher'), A(' on the account.')]])
reassign_box(s, X + 426, 100, W - 426, 112, dm=False, title='Both are true: reassign')
note_format(s, X + 426, 218, W - 426, 32)
scenarios(s, 260, 82, [
    ('C says "STOP" or "Unsubscribe"', [A('Unsubscribe C and reassign to '), B('DNC'), A(' on the Project Reassignment Form. Missing it is a ZTP violation.')]),
    ('C replied but already has a DHJ voucher', [B('Handle BAU.'), A(' The project is for customers without one.')])])
s.notes_slide.notes_text_frame.text = (
    'dhj_bookmark_sms. What to check: the customer responded to our marketing SMS dhj_bookmark_sms, regardless of the message, and the customer does not have a DHJ '
    'voucher on their account. If neither is true, handle BAU.\n'
    'If both are true: log the ticket on the Project Reassignment Form, create the internal note (Reassigned to Project, Link to the customer\'s pending message), '
    'then resolve the ticket. No Slack DM for this one.\n'
    'Specific scenarios: if the customer replies STOP, Unsubscribe or similar, unsubscribe them and reassign to DNC on the form instead; failing to unsubscribe is a '
    'ZTP violation. If the customer replies to the SMS but already has a DHJ voucher on the account, handle BAU.')

# ---- 05 CC Fails ----
s = new_slide(EB, '05  CC Fails', 'A failed job charge. Resolve it BAU, then reassign only if C wants a call.')
project_slides.append(s)
checks_box(s, X, 100, 210, 150, [[A('C has a '), B('failed job charge'), A(' (not the transaction fee)')],
                                 [A('The job took place '), B('more than 72 hours'), A(' and '), B('less than 30 days'), A(' ago')],
                                 [A('The C ID does '), B('not end in 2 or 3')]])
box(s, X + 218, 100, 200, 150, 'TEAL_PALE', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('First: resolve it BAU', 'TEAL', 10)],
    [A('Handle the ticket as usual, whatever the project checks say.')]])
reassign_box(s, X + 426, 100, W - 426, 112, dm=False, title='C also wants a call: reassign', tail='Once done, you may clear the ticket.')
note_format(s, X + 426, 218, W - 426, 32)
box(s, X, 260, W, 50, 'AMBER_PALE', margin=14, paras=[
    [B('Reassign only when both are true:  ', 'AMBER', 9.5), A('C wants a phone call, and C qualifies as a Project sample (all three checks match).', 'INK', 9.5)]])
s.notes_slide.notes_text_frame.text = (
    'CC Fails. What to check: the customer has a failed job charge, not the transaction fee; the job took place more than 72 hours ago and less than 30 days ago; '
    'and the C ID does not end in 2 or 3.\n'
    'What to do: resolve the ticket BAU. If the customer also wants a phone call and qualifies as a Project sample, log the ticket on the Project Reassignment Form, '
    'create the internal note (Reassigned to Project, Link to the customer\'s pending message), and once done you may clear the ticket.')

# ---- 06 Third-Party Poor Reviews ----
s = new_slide(EB, '06  Third-Party Poor Reviews', 'A poor public review, already being handled by a Project Agent.')
project_slides.append(s)
checks_box(s, X, 100, 210, 150, [[A('The account has a '), B('public_poor_review'), A(' flag')],
                                 [A('The account has an '), B('internal note'), A(' created by a Project Agent')]])
box(s, X + 218, 100, 200, 150, 'WHITE', 'LINE', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('Handle BAU if', 'AMBER', 10)],
    [A('The account also has a '), B('public_poor_review_removed'), A(' or '), B('public_poor_review_updated'), A(' flag.')]])
reassign_box(s, X + 426, 100, W - 426, 112, dm=False, title='Both are true: reassign')
note_format(s, X + 426, 218, W - 426, 32)
scenarios(s, 260, 82, [
    ('The only action item is a TR to issue a voucher', [A('Handle the ticket as usual and '), B('issue the voucher or credits'), A(' the TR says. Set another TR for next month\'s, if needed.')]),
    ('Subject line: "Resolutions Team - Perks at Work"', [B('Don\'t clear the action item.'), A(' Log it in the Perks At Work - C/CP concerns sheet with Status '), B('Pending'),
                                                          A(', then reassign to '), B('janafajardo'), A(' or '), B('jemhylaguzman'), A('.')])])
s.notes_slide.notes_text_frame.text = (
    'Third-Party Poor Reviews. What to check: the account has a public_poor_review flag, and an internal note created by a Project Agent. If the account also has a '
    'public_poor_review_removed or public_poor_review_updated flag, handle BAU.\n'
    'What to do: log the ticket on the Project Reassignment Form, create the internal note (Reassigned to Project, Link to the customer\'s pending message), then resolve.\n'
    'Specific scenarios: if the only action item is a timed reminder about issuing a voucher to the customer, handle the ticket as usual, issue the voucher or credits as '
    'the TR instructs, and create another TR for the following month\'s voucher or credit if necessary. If the customer reached out with "Resolutions Team - Perks at Work" '
    'as the subject line: do not clear the action item, log the ticket in the Perks At Work - C/CP concerns sheet with the Status column set to Pending, and reassign the '
    'ticket to janafajardo or jemhylaguzman.')

# overview tiles jump to each project slide
for (t, chev), target in zip(overview_tiles, project_slides):
    t.click_action.target_slide = target
    chev.click_action.target_slide = target

# ---- Knowledge check ----
s = new_slide('PROJECTS  ·  KNOWLEDGE CHECK', 'Let’s See What You’ve Got!', 'Ask each question, then click to reveal the answer.')
qa = [('C has an unused DHJ voucher. A Redemption Assist agent reached out, and C asks to refund the voucher. Reassign?',
       'No. Redemption Assist only takes C who want to book. Handle BAU.'),
      ('An FC Starts agent reached out about C\'s voucher. C replies "STOP". What do you do?',
       'Unsubscribe C and reassign to DNC on the Project Reassignment Form. Missing it is a ZTP violation.'),
      ('C\'s FC was reactivated yesterday. C says they never agreed to it. Who gets the ticket?',
       'The Project Lead: John Adrian Suiza (fc_reactivation team lead), via the Project Reassignment Form.'),
      ('C replied to dhj_bookmark_sms and already has a DHJ voucher. Reassign?',
       'No. Handle BAU. The project is for C without a DHJ voucher.'),
      ('Failed job charge, job 5 days ago, C ID ends in 7, and C wants a call. What do you do?',
       'Resolve it BAU, then reassign (form + internal note) since C is a Project sample. Then you may clear it.'),
      ('public_poor_review flag, but also public_poor_review_removed. Reassign?',
       'No. A removed or updated flag means handle BAU.')]
n = len(qa); gap = 5; h = (272 - gap * (n - 1)) / n
groups = []
for i, (q, a) in enumerate(qa):
    y = 100 + i * (h + gap)
    box(s, X, y, 364, h, 'WHITE', 'LINE')
    box(s, X + 10, y + (h - 20) / 2, 20, 20, 'TEAL_PALE', paras=[[(str(i + 1), 9, True, 'TEAL')]], align=PP_ALIGN.CENTER, margin=0, shape=MSO_SHAPE.OVAL)
    text(s, X + 40, y, 316, h, [[A(q, 'INK', 8.5)]], anchor=MSO_ANCHOR.MIDDLE)
    b = box(s, X + 371, y, W - 371, h, 'TEAL_PALE')
    t = text(s, X + 381, y, W - 391, h, [[A(a, 'INK', 8.5)]], anchor=MSO_ANCHOR.MIDDLE)
    groups.append([b.shape_id, t.shape_id])
reveal_timing(s, groups)
s.notes_slide.notes_text_frame.text = (
    'Quick knowledge check: Projects. Ask each question and let trainees answer before revealing. Each click reveals the next answer.\n'
    '1) Redemption Assist only takes customers who want to book; a voucher refund is BAU. 2) STOP always means unsubscribe and reassign to DNC. '
    '3) Reactivated 48 hours ago or less and C says they did not agree: Project Lead, John Adrian Suiza. 4) dhj_bookmark_sms with an existing DHJ voucher is BAU. '
    '5) CC Fails: job between 72 hours and 30 days ago, C ID not ending in 2 or 3, so resolve BAU and reassign because C wants a call. '
    '6) A public_poor_review_removed or _updated flag means BAU.')

order.append(questions)
questions.notes_slide.notes_text_frame.text = 'Open the floor for questions. When in doubt, check the Projects guide or ask your Trainer or Team Lead.'

# keep only the slides we built, in order
want = [id(s) for s in order]
all_slides = list(prs.slides)
els = list(lst)
pairs = dict((id(sl), e) for sl, e in zip(all_slides, els))
for sl, e in zip(all_slides, els):
    if id(sl) not in want:
        rid = e.get(qn('r:id')); lst.remove(e); prs.part.drop_rel(rid)
for e in list(lst):
    lst.remove(e)
for sl in order:
    lst.append(pairs[id(sl)])
prs.save(sys.argv[2]); print('saved', sys.argv[2], len(prs.slides), 'slides')
