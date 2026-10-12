"""Day 2 Cleaner Foundations deck: insert two quick knowledge-check slides (after slide 35: CP Profile Status +
CP-Facing Deactivation; after slide 60: Job Statuses). Answers reveal one per click. Patches the live export."""
import copy, io, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, WHITE, BORDER = '0B7A6F', 'E3F4F1', '1D1D1F', 'FFFFFF', 'DCE1E7'
prs = Presentation(src)
S = list(prs.slides)
assert S[34].shapes[1].text_frame.text.strip() == 'Coach, then document'
assert S[59].shapes[1].text_frame.text.strip() == 'Spotting a Pending Invoice job'
ref = S[15]
num = [x for x in ref.shapes if x.is_placeholder][0]

KC1 = ('CP PROFILE STATUS & CP-FACING DEACTIVATION  ·  QUICK KNOWLEDGE CHECK', [
    ('A CP is on norequests. Can they keep jobs they already claimed?', 'Yes. They just can\'t claim new jobs.'),
    ('norequests vs suspended: what\'s the difference?', 'Suspended also loses claimed jobs, and the dashboard is limited to Completed Jobs.'),
    ('What do you check before changing a CP\'s status?', 'The Flags table history. Ask: "Would I want this cleaner for my cleaning?"'),
    ('A new CP on probation has a no-show. What does the system do?', 'Moves them to suspended (probation: any no-show).'),
    ('Where do you enter the CP-facing reason, and how?', 'Change CP step window (Legacy CP CRM). Exact reason code format, never freeform.'),
    ('A customer reports a CP no-show. Which reason code?', 'cpq_noshow: "Your customer <C\'s name> reported you didn\'t show up to their <day, date> job…"'),
], 'Quick knowledge check: CP Profile Status and CP-Facing Deactivation Reasons. Ask each question and let trainees answer before revealing. Each click reveals the next answer. '
   'Answers: 1) Yes: norequests can\'t claim new jobs but keeps claimed ones, with full dashboard access. 2) Suspended can\'t claim new jobs, claimed jobs are removed, and the dashboard is limited to Completed Jobs; reactivation for both is a $15–25 fee (not available with a do_not_reactivate flag or in CA/CA-WA). '
   '3) Check the reason or flag history in the Flags table first, and ask "If I were the customer, would I want this cleaner for my cleaning?" 4) New cleaners (0–2 invoiced jobs since activation) are on probation: any no-show moves them to suspended. '
   '5) In the Change CP step window in the Legacy CP CRM; always use the exact reason code format, never a freeform reason, because the cleaner sees it on their dashboard and it\'s stored as the flag value. '
   '6) cpq_noshow Your customer <C\'S NAME> reported you didn\'t show up to their <DAY, DATE> job. Submit reactivation appeal through dashboard.')
KC2 = ('JOB STATUSES  ·  QUICK KNOWLEDGE CHECK', [
    ('A customer wants a cleaning in 8 hours. Can they book it?', 'Not themselves (12-hr minimum). CS can book it in the CRM, with no priority markup.'),
    ('A Submitted job is still unclaimed 24 hrs before start. What happens?', 'The system adds alternate times (unless the do_not_auto_add_alternate_start_times flag is on).'),
    ('Claimed job: until when can the customer reschedule or rematch?', 'Up to 6 hours before the start time.'),
    ('How does an Invoiced job become Cancelled?', 'Only through a full refund.'),
    ('The booked time passed and the job isn\'t invoiced. Status? Was it done?', 'Pending Invoice. It doesn\'t mean done: it needs action.'),
    ('C got confirmations from 2 different cleaners. First step?', 'Check the job history: did the first cleaner cancel, or were they swapped? Never mention swapping.'),
], 'Quick knowledge check: Job Statuses. Ask each question and let trainees answer before revealing. Each click reveals the next answer. '
   'Answers: 1) Customers can book at least 12 hours ahead (under 48 hours adds a priority fee); CS can book less than 12 hours ahead through the CRM, and priority markups don\'t apply to CRM bookings. '
   '2) If a Submitted job is unclaimed 24 hours before start, the system adds alternate times; if the customer can\'t do other times, add the do_not_auto_add_alternate_start_times flag (New CRM > Flags > + Add Flags). Within 24–48 hours of start it\'s also boosted as a Priority Request. '
   '3) Once claimed, the customer can reschedule or rematch (Manage Appointment > Change My Cleaner) up to 6 hours before the start. 4) An Invoiced job only becomes Cancelled through a full refund (by the cleaner or CS). '
   '5) Pending Invoice: the booked time passed and it isn\'t invoiced yet; the status changes on a timer, so it doesn\'t mean the cleaning happened. Act on it even if the customer asked about something else. '
   '6) Check the job history: did the first cleaner cancel, or were they swapped out? Swapping is internal (ZTP): never mention it.')

def build(after_idx, eyebrow, qa, notes):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
        s.shapes._spTree.append(copy.deepcopy(el))
    for shape, t in ((s.shapes[0], eyebrow), (s.shapes[1], 'Quick Knowledge Check'), (s.shapes[2], 'Ask each question, then click to reveal the answer.')):
        r = shape.text_frame.paragraphs[0].runs; r[0].text = t
        for x in r[1:]: x._r.getparent().remove(x._r)
    def strip(x):
        st = x._element.find(P_NS + 'style')
        if st is not None: x._element.remove(st)
    def box(x, y, w, h, fill, line=None, adj=0.08, name=None):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
        b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
        else: b.line.fill.background()
        if name: b.name = name
        return b
    def text(shape, t, size, bold, col, font, align=PP_ALIGN.LEFT, ml=0.0):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(ml); tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = align
        r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold; r.font.name = font; r.font.color.rgb = RGBColor.from_string(col)
    n = len(qa); h = (3.75 - 0.08 * (n - 1)) / n
    for i, (q, a) in enumerate(qa):
        y = 1.4 + i * (h + 0.08)
        box(0.45, y, 5.05, h, WHITE, BORDER)
        b = box(0.6, y + (h - 0.3) / 2, 0.3, 0.3, TEALSOFT, adj=0.5); text(b, str(i + 1), 9, True, TEAL, 'Play', PP_ALIGN.CENTER)
        t = s.shapes.add_textbox(Inches(1.05), Inches(y), Inches(4.35), Inches(h)); text(t, q, 10, False, INK, 'Arial')
        box(5.6, y, 3.95, h, TEALSOFT, name=f'step{i + 1}Box')
        t = s.shapes.add_textbox(Inches(5.75), Inches(y), Inches(3.7), Inches(h)); t.name = f'step{i + 1}Ans'; text(t, a, 9, False, INK, 'Arial')
    s.notes_slide.notes_text_frame.text = notes
    lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el)
    list(lst)[after_idx].addnext(el)
    return s
# Insert the later one first so the earlier index stays valid.
k2 = build(59, *KC2)
k1 = build(34, *KC1)

buf = io.BytesIO(); prs.save(buf)
ns = {}
exec(open(__file__.replace('day2_quick_kc.py', 'add_story_anim.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
targets = {k1.part.partname.lstrip('/'), k2.part.partname.lstrip('/')}
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename in targets:
        x = data.decode('utf8'); assert '<p:timing' not in x
        groups = {}
        for sid, n in re.findall(r'<p:cNvPr id="(\d+)" name="step(\d+)[^"]*"', x): groups.setdefault(int(n), []).append(sid)
        cid, clicks, bld = 3, [], []
        for n in sorted(groups):
            inner = ''
            for j, sid in enumerate(groups[n]):
                inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if j == 0 else 'withEffect'); cid += 3
                bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
            clicks.append(ns['click'](cid, inner)); cid += 2
        x = x.replace('</p:sld>', ns['timing'](clicks, bld) + '</p:sld>')
        data = x.encode('utf8'); print(item.filename, 'clicks:', len(clicks))
    zout.writestr(item, data)
zout.close()
