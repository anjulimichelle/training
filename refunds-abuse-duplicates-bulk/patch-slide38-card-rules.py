"""Refunds deck, Failed Charge / Overdue MF 'What is this?': replace the three tiles with a short description that
answers 'should I remind C to update their card?', then Rule 1 and Rule 2. Also drops the duplicate slide-number
placeholders left on the merged topic 5–7 slides. Patches the live Google export."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', '9A6410', 'FFF3D6'
RED, PINK, GREEN, GREENSOFT = 'B3261E', 'FCEBEA', '2E7D32', 'E4F2E1'
prs = Presentation(src)

# Duplicate slide-number placeholders (same spot) on any slide: keep the first.
for s in prs.slides:
    phs = [x for x in s.shapes if x.is_placeholder and x.has_text_frame and x.text_frame.text.strip() in ('‹#›', '<#>')]
    for x in phs[1:]:
        if x.left == phs[0].left and x.top == phs[0].top: x._element.getparent().remove(x._element)

S = list(prs.slides)
idx = [i for i, s in enumerate(S) if s.shapes[0].has_text_frame and s.shapes[0].text_frame.text.strip().upper() == 'FAILED CHARGE / OVERDUE MF' and s.shapes[1].text_frame.text.strip() == 'What is this?']
assert len(idx) == 1; s = S[idx[0]]
body = [x for x in s.shapes if not x.is_placeholder and x.top >= Inches(1.4)]
assert len(body) == 15, len(body)
for x in body: x._element.getparent().remove(x._element)
r = s.shapes[2].text_frame.paragraphs[0].runs; r[0].text = 'Should I remind C to update their card?'
for x in r[1:]: x._r.getparent().remove(x._r)

def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
def box(x, y, w, h, fill, line=None, adj=0.06):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
    b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(1)
    else: b.line.fill.background()
    return b
def text(shape, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=4, ml=0.18, mt=0.1):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(ml); tf.margin_top = tf.margin_bottom = Inches(mt)
    for k, runs in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
        for t, bold, col, sz in runs:
            rr = p.add_run(); rr.text = t; rr.font.size = Pt(sz); rr.font.bold = bold
            rr.font.color.rgb = RGBColor.from_string(col); rr.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
    return shape

text(box(0.45, 1.42, 9.1, 1.05, WHITE, BORDER), [
    [('You spot a failed job charge or an overdue MF ', True, INK, 10.5), ('(C wrote in about it, or you noticed it while working on something else).', False, SOFT, 10)],
    [('It depends on timing. ', True, TEAL, 10.5), ('Asking at the wrong moment feels like chasing payment instead of taking care of C, especially with an open complaint. Check the two rules below before you mention the card.', False, INK, 10)],
], anchor=MSO_ANCHOR.MIDDLE)

def rule(x, title, items, col, soft):
    box(x, 2.62, 4.45, 1.95, soft, col)
    text(box(x, 2.62, 4.45, 0.38, col, adj=0.12), [[(title, True, WHITE, 10.5)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, mt=0)
    t = s.shapes.add_textbox(Inches(x), Inches(3.05), Inches(4.45), Inches(1.5))
    text(t, [[('•  ', True, col, 9.5), (a, True, INK, 9.5), (b, False, INK, 9.5)] for a, b in items], after=5, ml=0.2, mt=0.02)
rule(0.45, 'Rule 1: Don\'t chase payment while an issue is open', [
    ('Unresolved issue? ', 'Don\'t ask for a card update or send the update link yet.'),
    ('Resolve it first: ', 'we did what C asked, or C explicitly accepted our resolution.'),
    ('Example: ', 'PCQ complaint + failed job charge. Handle the PCQ, then the card.')], RED, PINK)
rule(5.1, 'Rule 2: Upcoming appointment + invalid card', [
    ('The one exception: ', 'advise C to update their card right away.'),
    ('Even if an issue is still open: ', 'protecting the upcoming job comes first.'),
    ('Why: ', 'otherwise that appointment risks being cancelled.')], GREEN, GREENSOFT)
text(box(0.45, 4.7, 9.1, 0.4, GOLDSOFT, adj=0.15), [[('Next:  ', True, GOLD, 9.5), ('how many overdue MFs to bring current, and free months: Understanding Offer Logic › Free Months (Job 2).', False, INK, 9.5)]], anchor=MSO_ANCHOR.MIDDLE, mt=0)

s.notes_slide.notes_text_frame.text = (
    'Should I remind the customer to update their card? This covers what to do the moment you notice a failed job charge or an overdue membership fee, whether that\'s why the customer contacted you or something you spotted while working on a different issue. '
    'The instinct when you spot money owed is to fix it: ask C to update their card and get the charge current. But doing that at the wrong moment can make the interaction feel like we\'re chasing payment instead of taking care of the customer, especially if there\'s an unresolved complaint on the account. '
    'Rule 1: don\'t chase payment while an issue is still open. If C has an unresolved issue and a failed job charge or overdue MF, don\'t ask them to update their card, and don\'t send the link to update it, until the issue is resolved. "Resolved" means either we did what the customer asked, or the customer explicitly accepted our resolution. Example: a customer raises a PCQ complaint and you notice a failed job charge while investigating: handle the PCQ first, then bring up the card. '
    'Rule 2: an upcoming appointment with an invalid card is the one exception. Proactively advise C to update their card, otherwise they risk that appointment getting cancelled. This applies even if Rule 1\'s unresolved issue is in play: protecting the upcoming job takes priority. '
    'For how many overdue MFs to bring current, when to issue free months, and cancellation-intent branching, see Understanding Offer Logic > Free Months (Job 2).')
prs.save(dst)
print('slide', idx[0] + 1)
