"""Refunds deck: insert 'MF Refunds — sorted by root cause' (FC active vs cancelled diagram) right after
'The MCT test', linking to the root-cause table that follows. Patches the live Google export."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, INK, SOFT, WHITE = '0B7A6F', '1D1D1F', '6E6E73', 'FFFFFF'
RED, PINK, PURPLE, PURPLESOFT, PEACH = 'B3261E', 'FCEBEA', '5E4A9C', 'EEEAF8', 'C0612B'
prs = Presentation(src)
titles = [x.shapes[1].text_frame.text.strip() if x.shapes[1].has_text_frame else '' for x in prs.slides]
i_ref = titles.index('The MCT test'); ref = prs.slides[i_ref]
table = prs.slides[titles.index('MF refunds by root cause')]
assert titles.index('MF refunds by root cause') == i_ref + 1

s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
num = [x for x in ref.shapes if x.is_placeholder][0]
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
for shape, t in ((s.shapes[1], 'MF Refunds — sorted by root cause'), (s.shapes[2], 'Start with the FC status.')):
    r = shape.text_frame.paragraphs[0].runs; r[0].text = t
    for x in r[1:]: x._r.getparent().remove(x._r)
lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); lst.insert(i_ref + 1, el)

def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
def box(x, y, w, h, fill, line=None, adj=0.08):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
    b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(1)
    else: b.line.fill.background()
    return b
def text(shape, paras, size=10, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, after=4, margin=0):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
        for t, bold, col, *sz in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(sz[0] if sz else size); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
    return shape
def tb(x, y, w, h): return s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
def line(x1, y1, x2, y2, arrow=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); strip(c)
    c.line.color.rgb = RGBColor.from_string(SOFT); c.line.width = Pt(1)
    if arrow:
        ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))

text(box(3.25, 1.42, 3.5, 0.5, PEACH), [[('MF refund request', True, WHITE, 12)]], align=PP_ALIGN.CENTER)
lx, rx = 2.675, 7.325
line(5.0, 1.92, 5.0, 2.12); line(lx, 2.12, rx, 2.12)
line(lx, 2.12, lx, 2.32, True); line(rx, 2.12, rx, 2.32, True)
text(box(0.45, 2.32, 4.45, 0.5, RED), [[('FC is ', False, WHITE, 11.5), ('active', True, WHITE, 11.5), ('  (triggered)', False, WHITE, 10)]], align=PP_ALIGN.CENTER)
text(box(5.1, 2.32, 4.45, 0.5, PURPLE), [[('FC is ', False, WHITE, 11.5), ('cancelled', True, WHITE, 11.5)]], align=PP_ALIGN.CENTER)

box(0.45, 2.95, 4.45, 1.95, PINK, RED)
text(tb(0.65, 3.03, 4.05, 1.8), [
    [('Active and retainable?', True, RED, 10.5)],
    [('Lead with free months equal to the eligible amount (max 3).', False, INK, 10)],
    [('Exception:', True, RED, 10.5)],
    [('Extenuating circumstances that warrant a refund (billing errors, etc.)? Proceed with the refund.', False, INK, 10)],
], anchor=MSO_ANCHOR.TOP, after=3)

box(5.1, 2.95, 4.45, 1.95, PURPLESOFT, PURPLE)
text(tb(5.3, 3.03, 4.05, 1.15), [
    [('Follow the root cause table', True, PURPLE, 10.5)],
    [('Also use it when C declines the free months, or isn\'t retainable.', False, INK, 10)],
], anchor=MSO_ANCHOR.TOP, after=3)
btn = box(5.3 + (4.05 - 2.4) / 2, 4.22, 2.4, 0.42, TEAL, adj=0.25)
text(btn, [[('▶  Root cause table', True, WHITE, 10.5)]], align=PP_ALIGN.CENTER)
btn.click_action.target_slide = table

s.notes_slide.notes_text_frame.text = (
    'MF refund requests start with the FC status. FC is active (triggered): if the customer is retainable, lead with a free-months offer equal to the eligible amount, capped at 3. '
    '"Eligible" means the customer has a legitimate claim to the value of those unused MFs; it doesn\'t mean the default resolution is a cash refund. '
    'Exception: if there are extenuating circumstances that warrant a refund (billing errors, etc.), proceed with the refund. Agent errors, billing errors and high-risk/legal situations are refunded directly regardless of FC status. '
    'FC is cancelled, the customer declines the free months, or the customer isn\'t retainable: follow the root cause table on the next slide to decide whether the MF is refundable. '
    'TRAINER: in slideshow mode, the "Root cause table" button jumps to the next slide.')
prs.save(dst)
print('inserted at', i_ref + 2)
