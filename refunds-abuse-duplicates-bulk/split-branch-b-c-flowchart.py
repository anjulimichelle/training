"""Refunds deck, Bulk Charges: slide 27 becomes Branch B only; a new Branch C flowchart slide is inserted right after it
(the ETF-credit-vs-MCT example moves down one). Patches the live Google export."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', '9A6410', 'FFF3D6'
prs = Presentation(src)
titles = [x.shapes[1].text_frame.text.strip() if x.shapes[1].has_text_frame else '' for x in prs.slides]
i = titles.index('Branches B & C'); b = prs.slides[i]
rules = prs.slides[titles.index('MF refunds by root cause')]
example = prs.slides[i + 1]; assert 'ETF credit or MCT reduction' in titles[i + 1]

def retext(shape, t):
    r = shape.text_frame.paragraphs[0].runs; r[0].text = t
    for y in r[1:]: y._r.getparent().remove(y._r)
def kit(s):
    def strip(x):
        st = x._element.find(P_NS + 'style')
        if st is not None: x._element.remove(st)
    def box(x, y, w, h, fill, line=None, adj=0.08):
        bx = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(bx)
        bx.adjustments[0] = adj; bx.fill.solid(); bx.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line: bx.line.color.rgb = RGBColor.from_string(line); bx.line.width = Pt(1)
        else: bx.line.fill.background()
        return bx
    def text(shape, paras, size=9.5, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=2):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.08); tf.margin_top = tf.margin_bottom = 0
        for k, runs in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
            for t, bold, col, *sz in runs:
                r = p.add_run(); r.text = t; r.font.size = Pt(sz[0] if sz else size); r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
        return shape
    def tb(x, y, w, h): return s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    def line(x1, y1, x2, y2, arrow=False):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); strip(c)
        c.line.color.rgb = RGBColor.from_string(SOFT); c.line.width = Pt(1)
        if arrow:
            ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    def label(x, y, w, t): text(tb(x, y, w, 0.2), [[(t, True, SOFT, 8.5)]], align=PP_ALIGN.LEFT)
    return box, text, tb, line, label

# ---------- Slide 27: Branch B only ----------
retext(b.shapes[1], 'Branch B: FC already cancelled'); retext(b.shapes[2], 'C wants a refund for the bulk charge. Nothing bulk-specific here.')
for x in list(b.shapes):
    if x.has_text_frame and x.text_frame.text.strip().startswith('C: FC ACTIVE') or (not x.is_placeholder and x.left >= Inches(3.4) and x.top >= Inches(1.4)):
        x._element.getparent().remove(x._element)
for x in b.shapes:   # widen the B panel (box, header, text) to the left half
    if not x.is_placeholder and x.top >= Inches(1.4):
        x.width = Inches(4.45 - (0.36 if x.left > Inches(0.5) else 0)) if x.has_text_frame and x.left > Inches(0.5) else Inches(4.45)
        if x.height > Inches(2.5): x.height = Inches(2.6)
box, text, tb, line, label = kit(b)
box(5.1, 1.45, 4.45, 2.6, WHITE, BORDER)
text(tb(5.3, 1.6, 4.05, 1.5), [[('Standard refund policy', True, TEAL, 11)],
     [('Unused MF charges beyond MCT, made within the past 120 days, are refundable.', False, INK)],
     [('System-error exceptions, high-dispute-risk overrides and the rest: see MF & ETF Refund Requests.', False, SOFT)]], align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=5)
btn = box(5.3 + (4.05 - 2.6) / 2, 3.35, 2.6, 0.42, TEAL, adj=0.25)
text(btn, [[('▶  MF refunds by root cause', True, WHITE, 10)]]); btn.click_action.target_slide = rules
b.notes_slide.notes_text_frame.text = (
    'Branch B: ForeverClean is already cancelled, and the customer wants a refund for the bulk charge. This is just the standard refund policy, nothing bulk-charge-specific: '
    'unused MF charges beyond MCT, made within the past 120 days, are refundable. See MF & ETF Refund Requests for the full nuance (system-error exceptions, high-dispute-risk overrides, etc.). '
    'TRAINER: the button jumps to the "MF refunds by root cause" table.')

# ---------- New slide 28: Branch C flowchart ----------
s = prs.slides.add_slide(b.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
num = [x for x in b.shapes if x.is_placeholder][0]
for el in (b.shapes[0]._element, b.shapes[1]._element, b.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
retext(s.shapes[1], 'Branch C: FC active, cancelling over it'); retext(s.shapes[2], 'Check if it\'s still overdue, then the MCT: save mode or exit mode.')
lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); lst.insert(i + 1, el)
box, text, tb, line, label = kit(s)
# Row 1
text(box(0.45, 1.42, 2.75, 0.55, WHITE, BORDER), [[('FC active. C wants to cancel over the bulk charge', True, INK, 9)]])
line(3.2, 1.695, 3.5, 1.695, True)
text(box(3.5, 1.42, 2.75, 0.55, GOLDSOFT, GOLD), [[('Still overdue after the bulk charge?', True, GOLD, 9)], [('Check first', False, SOFT, 8)]])
line(6.25, 1.695, 6.55, 1.695, True)
text(box(6.55, 1.42, 3.0, 0.55, GOLDSOFT, GOLD), [[('Within or beyond the MCT?', True, GOLD, 9.5)], [('Decides save mode or exit mode', False, SOFT, 8)]])
lx, rx = 2.675, 7.325
line(8.05, 1.97, 8.05, 2.12); line(lx, 2.12, 8.05, 2.12)
line(lx, 2.12, lx, 2.3, True); line(rx, 2.12, rx, 2.3, True)
# Left: beyond MCT
text(box(0.45, 2.3, 4.45, 0.45, TEAL), [[('Beyond MCT: no ETF. C is free to leave', True, WHITE, 10)]])
line(lx, 2.75, lx, 2.95, True)
text(box(0.45, 2.95, 4.45, 0.95, TEALSOFT, TEAL), [[('Offer free months', True, TEAL, 10)],
     [('Value back for months already paid and can\'t use. Not a discount.', False, INK, 9)],
     [('+ $20 credit to get them back in for a cleaning', True, INK, 9)]])
line(lx, 3.9, lx, 4.1, True)
text(box(0.45, 4.1, 4.45, 0.5, WHITE, BORDER), [[('Still wants to leave? Cancel FC. No ETF.', True, INK, 9.5)]])
# Right: within MCT
text(box(5.1, 2.3, 4.45, 0.45, GOLD), [[('Within MCT: ETF applies', True, WHITE, 10)]])
line(rx, 2.75, rx, 2.9, True)
text(box(5.1, 2.9, 4.45, 0.4, WHITE, BORDER), [[('Retention attempt first', True, INK, 9.5)]])
line(rx, 3.3, rx, 3.48, True); label(rx + 0.08, 3.28, 1.2, "Doesn't land")
text(box(5.85, 3.48, 2.95, 0.4, GOLDSOFT, GOLD), [[('FCF, no first cleaning yet?', True, GOLD, 9.5)]])
l2, r2 = 6.2, 8.45
line(rx, 3.88, rx, 3.98); line(l2, 3.98, r2, 3.98)
line(l2, 3.98, l2, 4.12, True); line(r2, 3.98, r2, 4.12, True)
label(l2 + 0.06, 3.96, 0.5, 'Yes'); label(r2 + 0.06, 3.96, 0.5, 'No')
text(box(5.1, 4.12, 2.2, 0.85, WHITE, BORDER), [[('No ETF', True, INK, 9.5)], [('FCF\'s ETF starts after the first completed cleaning', False, SOFT, 8.5)]])
text(box(7.35, 4.12, 2.2, 0.85, TEALSOFT, TEAL), [[('The cheaper option', True, TEAL, 9.5)], [('Credit unused bulk MFs to the ETF, or reduce the MCT (max 2 mo)', False, INK, 8.5)]])
s.notes_slide.notes_text_frame.text = (
    'Branch C: FC is still active, and the bulk charge is why the customer wants to cancel. First, check whether the account is still overdue even after the bulk charge landed, then check MCT status, because that determines whether you\'re in save mode or exit mode. '
    'Beyond MCT: the ETF does not apply; the customer is free to leave. The only question is whether it\'s worth trying to keep them and what they\'re owed either way. The free-month offer isn\'t a discount: it gives back value for months they already paid for and can\'t use anymore. '
    'The $20 credit is a separate, bulk-charge-specific incentive to get them back in for a cleaning. '
    'Within MCT: the ETF applies, and the retention attempt happens first. If it doesn\'t land, the exit path branches on the membership model: FCF customers who haven\'t had a first cleaning have no ETF at all, since FCF\'s ETF only kicks in after the first completed cleaning. '
    'For everyone else, apply whichever option costs the customer less: crediting the unused bulk MFs against the ETF, or reducing the MCT. MCT reduction is capped at 2 months and only usable when it\'s cheaper than a reduced ETF (see the worked example on the next slide).')
prs.save(dst)
print('slide', i + 1, 'Branch B; new slide', i + 2, 'Branch C')
