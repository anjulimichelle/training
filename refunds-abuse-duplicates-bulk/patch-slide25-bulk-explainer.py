"""Refunds deck, Bulk Charges 'What is this?': replace the three cards with a short how-it-happens explanation
plus a sample Payment history screenshot (four same-day $59 MFs). Keeps the 'three reasons' row. Patches the live export."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst, img = sys.argv[1:4]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', '9A6410', 'FFF3D6'
prs = Presentation(src)
idx = [i for i, x in enumerate(prs.slides) if x.shapes[0].has_text_frame and x.shapes[0].text_frame.text.strip() == 'BULK CHARGES' and x.shapes[1].text_frame.text.strip() == 'What is this?']
assert len(idx) == 1; s = prs.slides[idx[0]]
# Remove the three top cards (everything between the subtitle and the 'three reasons' eyebrow).
top = [x for x in s.shapes if not x.is_placeholder and Inches(1.4) <= x.top < Inches(3.3)]
assert len(top) == 15, len(top)
for x in top: x._element.getparent().remove(x._element)

def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
def box(x, y, w, h, fill, line=None, adj=0.08):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
    b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
    else: b.line.fill.background()
    return b
def text(shape, runs, size=9.5, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = align
    for t, bold, col in runs:
        r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
    return shape
def tb(x, y, w, h): return s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))

steps = [('MF overdue', ' for several billing cycles'), ('A payment finally goes through', ''), ('Every missed month is charged at once', ': same day or same month')]
for k, (lead, rest) in enumerate(steps):
    y = 1.45 + k * 0.6
    box(0.45, y, 3.85, 0.52, WHITE, BORDER)
    b = box(0.57, y + 0.11, 0.3, 0.3, TEALSOFT, adj=0.5); text(b, [(str(k + 1), True, TEAL)], 9, PP_ALIGN.CENTER)
    text(tb(1.0, y, 3.2, 0.52), [(lead, True, INK), (rest, False, SOFT)])

w = 5.1; h = w * 212 / 763
box(4.45, 1.45, w + 0.1, h + 0.1, WHITE, BORDER, adj=0.03)
s.shapes.add_picture(img, Inches(4.5), Inches(1.5), Inches(w), Inches(h)).name = 'bulkPaymentHistory'
text(tb(4.45, 1.45 + h + 0.15, w + 0.1, 0.3), [('Four $59 MFs, one day: ', True, GOLD), ('Jan, Feb, Mar and Apr paid-through. Different months, not duplicates.', False, INK)], 9)

s.notes_slide.notes_text_frame.text = (
    'Bulk charges happen when a customer\'s ForeverClean (FC) monthly fee has been overdue for multiple billing cycles, and a payment attempt finally goes through. '
    'Instead of billing one overdue month at a time, the system charges all the past-due months in a single transaction. The customer ends up seeing several MF charges posted on the same day or within the same month: each one for a different month they missed, not a duplicate of the same charge. '
    'Screenshot (sample Payment history): four $59 MF charges on Oct 1, each with a different foreverclean_paid_through date (1/27, 2/27, 3/27, 4/27): four missed months, not duplicates. '
    'This applies to both DHJ and FCF memberships. Since September 12, 2024, the system no longer auto-issues free months to true up overdue accounts, so the proactive true-up (Branch A) falls to agents. '
    'A wall of same-day charges reads as a billing error to the customer, even when every dollar was owed. Your job is to figure out whether the customer got real value for those months (using the service, or didn\'t know they were billed) and let that drive how much you fix. '
    'Three reasons it reaches you: A) FC overdue, but that\'s not why C wrote in; B) FC already cancelled and C wants a refund; C) FC active and C wants to cancel because of the bulk charge.')
prs.save(dst)
