"""Refunds deck: turn 'Steps 0–1: Thumbtack, then cancel' and 'Steps 2–3: refund or free months' (Duplicate Accounts)
into one flowchart slide. Slide 21 is rebuilt in place (header kept), slide 22 is removed. Patches the live Google export."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', '9A6410', 'FFF3D6'
prs = Presentation(src)
titles = [x.shapes[1].text_frame.text.strip() if x.shapes[1].has_text_frame else '' for x in prs.slides]
i = titles.index('Steps 0–1: Thumbtack, then cancel'); s = prs.slides[i]; old = prs.slides[i + 1]
assert titles[i + 1] == 'Steps 2–3: refund or free months'
fm_title = [x for x in old.shapes if x.has_text_frame and 'Free months' in x.text_frame.text][0].text_frame.text.strip()
for x in prs.slides:
    for sh in x.shapes:
        assert not (sh.click_action.target_slide is not None and sh.click_action.target_slide == old), 'something links to slide 22'

# Drop slide 22.
lst = prs.slides._sldIdLst
for sid in list(lst):
    if prs.part.related_part(sid.get(qn('r:id'))) is old.part:
        prs.part.drop_rel(sid.get(qn('r:id'))); lst.remove(sid)

# Rebuild slide 21's body.
num = [x for x in s.shapes if x.is_placeholder][0]
keep = {s.shapes[0].shape_id, s.shapes[1].shape_id, s.shapes[2].shape_id, num.shape_id}
for x in list(s.shapes):
    if x.shape_id not in keep: x._element.getparent().remove(x._element)
def retext(shape, t):
    r = shape.text_frame.paragraphs[0].runs; r[0].text = t
    for y in r[1:]: y._r.getparent().remove(y._r)
retext(s.shapes[1], 'The process'); retext(s.shapes[2], 'Check Thumbtack, stop the next charge, then let C choose.')

def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
def box(x, y, w, h, fill, line=None, adj=0.08):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
    b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(1)
    else: b.line.fill.background()
    return b
def text(shape, paras, size=10, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=2):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.08); tf.margin_top = tf.margin_bottom = 0
    for k, runs in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
        for t, bold, col, *sz in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(sz[0] if sz else size); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
    return shape
def label(x, y, w, t):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.2)); text(tb, [[(t, True, SOFT, 8.5)]], align=PP_ALIGN.LEFT)
def line(x1, y1, x2, y2, arrow=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); strip(c)
    c.line.color.rgb = RGBColor.from_string(SOFT); c.line.width = Pt(1)
    if arrow:
        ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
def panel(x, y, w, h, title, items, col, soft):
    box(x, y, w, h, soft, col)
    text(box(x, y, w, 0.34, col, adj=0.15), [[(title, True, WHITE, 10.5)]])
    t = s.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.42), Inches(w - 0.3), Inches(h - 0.5))
    text(t, [[('•  ' + it, False, INK, 9.5)] for it in items], align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=3)

# Row 1: Thumbtack check.
text(box(0.45, 1.42, 2.4, 0.5, WHITE, BORDER), [[('Duplicate account found', True, INK)]])
line(2.85, 1.67, 3.25, 1.67, True)
text(box(3.25, 1.42, 3.0, 0.5, GOLDSOFT, GOLD), [[('project_thumbtack flag?', True, GOLD)], [('Legacy CRM › View › Flag', False, SOFT, 8.5)]])
line(6.25, 1.67, 6.65, 1.67, True); label(6.3, 1.45, 0.4, 'Yes')
text(box(6.65, 1.42, 2.9, 0.5, INK), [[('Reassign to the Thumbtack queue. Stop.', True, WHITE)]])
# Row 2: Step 1 on the duplicate.
line(4.75, 1.92, 4.75, 2.12, True); label(4.83, 1.92, 0.5, 'No')
text(box(0.45, 2.12, 9.1, 0.62, TEAL), [
    [('Step 1, on the duplicate:  ', True, WHITE, 10.5), ('cancel the FC right away', True, WHITE, 10.5)],
    [('Cancel upcoming jobs and the RC (unless C wants them)  ·  internal note on both accounts with the other ticket link / C ID', False, WHITE, 9)]])
# Row 3: the choice.
line(4.75, 2.74, 4.75, 2.92, True)
text(box(3.0, 2.92, 3.5, 0.4, GOLDSOFT, GOLD), [[('Let C choose', True, GOLD)]])
lx, rx = 2.675, 7.325
line(4.75, 3.32, 4.75, 3.45); line(lx, 3.45, rx, 3.45)
line(lx, 3.45, lx, 3.6, True); line(rx, 3.45, rx, 3.6, True)
panel(0.45, 3.6, 4.45, 1.55, fm_title, ['Up to 3 free months on the main account; adjust the MCT', 'Contesting? +1 bonus month', 'Paid months reach 6? MCT effectively fulfilled'], TEAL, TEALSOFT)
panel(5.1, 3.6, 4.45, 1.55, 'Refund', ['Duplicate charges within 120 days', 'Beyond 120 days: only for a system or agent error', 'No bonus month'], GOLD, GOLDSOFT)

s.notes_slide.notes_text_frame.text = (
    'IMPORTANT: before you touch anything on the ticket, check whether the duplicate account is associated with Thumbtack. Thumbtack customers have their own handling process. '
    'How to check: Legacy CRM > View > Flag > look for project_thumbtack. If the flag is present: reassign the ticket to the Thumbtack queue, and stop. If not: continue. '
    'Step 1, confirm the duplicate and stop it from recurring: cancel the FC membership on the duplicate account immediately, regardless of which branch applies; it stops a third charge while you sort out the rest. '
    'Cancel any upcoming appointments or recurring cleaning plan on the duplicate account too, unless the customer has given you a reason to keep them. Leave an internal note on both accounts with the ticket link or C ID of the duplicate/main account. '
    'Steps 2–3: whether the customer is paying two MFs, unaware and not contesting, or actively contesting, present the choice: refund the duplicate charges, or apply them as free months on the main account (max 3), adjusting the MCT. '
    'If they contest and choose free months, add 1 bonus month on top as a goodwill gesture; if total paid months reach 6, explain they\'ve effectively fulfilled the MCT. If they choose a refund: no bonus month. '
    'Customer insists on a refund and rejects free months: refund the duplicate within 120 days; beyond 120 days only if the duplicate was caused by a system/agent error. '
    'Why the asymmetry: a free month costs Homeaglow far less than a refund, so the goodwill gesture comes only with the cheaper option. '
    'Tiering: count the customer\'s history across both accounts wherever it matters (e.g. lockout customer category): 22 paid months and 8 jobs split across two accounts = Very High-Value.')
prs.save(dst)
print('rebuilt slide', i + 1, '; removed old slide', i + 2, '; free-months title:', fm_title)
