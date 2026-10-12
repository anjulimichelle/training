"""Refunds deck: Branch B (slide 27) down to one tile; Branch C (slide 28) beyond-MCT path as one box
(retainable vs not retainable). Patches the live Google export."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73'
prs = Presentation(src)
titles = [x.shapes[1].text_frame.text.strip() if x.shapes[1].has_text_frame else '' for x in prs.slides]
b = prs.slides[titles.index('Branch B: FC already cancelled')]
c = prs.slides[titles.index('Branch C: FC active, cancelling over it')]
def by_text(s, start):
    m = [x for x in s.shapes if x.has_text_frame and x.text_frame.text.strip().startswith(start)]
    assert len(m) == 1, start; return m[0]
def rm(x): x._element.getparent().remove(x._element)

# ---------- Branch B: one tile ----------
right_text = by_text(b, 'Standard refund policy\nUnused MF charges')
right_box = [x for x in b.shapes if x.shape_type == 1 and abs(x.left - Inches(5.1)) < Inches(0.05) and abs(x.top - Inches(1.45)) < Inches(0.05)]
assert len(right_box) == 1
rm(right_text); rm(right_box[0])
hdr = by_text(b, 'B: FC already cancelled'); body = by_text(b, 'Standard refund policy')
panel = [x for x in b.shapes if x.shape_type == 1 and x.left == hdr.left and x.top == hdr.top and x.height > hdr.height][0]
for x in (panel, hdr): x.width = Inches(9.1)
panel.height = Inches(1.85)
body.width = Inches(8.74); body.height = Inches(1.0)
btn = [x for x in b.shapes if x.click_action.target_slide is not None][0]
btn.left = Inches(0.45 + (9.1 - 2.6) / 2); btn.top = Inches(2.68)

# ---------- Branch C: one box under 'Beyond MCT' ----------
for start in ('Offer free months', 'Still wants to leave?'): rm(by_text(c, start))
for x in list(c.shapes):   # the two arrows inside the left column below the header
    if x.shape_type == 9 and abs(x.left - Inches(2.675)) < Inches(0.02) and x.top > Inches(2.7): rm(x)
def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
ln = c.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2.675), Inches(2.76), Inches(2.675), Inches(2.95)); strip(ln)
ln.line.color.rgb = RGBColor.from_string(SOFT); ln.line.width = Pt(1)
l = ln.line._get_or_add_ln(); l.append(l.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
bx = c.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(2.95), Inches(4.45), Inches(2.02)); strip(bx)
bx.adjustments[0] = 0.06; bx.fill.solid(); bx.fill.fore_color.rgb = RGBColor.from_string(TEALSOFT)
bx.line.color.rgb = RGBColor.from_string(TEAL); bx.line.width = Pt(1)
tf = bx.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = tf.margin_right = Inches(0.2); tf.margin_top = tf.margin_bottom = Inches(0.08)
paras = [[('Retainable?', True, TEAL, 11)],
         [('Offer free months + $20 credit to get them back in for a cleaning', True, INK, 10)],
         [('Value back for months already paid and can\'t use. Not a discount.', False, SOFT, 9.5)],
         [('Not retainable?  ', True, TEAL, 11), ('Cancel FC', True, INK, 10.5)]]
for k, runs in enumerate(paras):
    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(10 if k == 2 else 3)
    for t, bold, col, sz in runs:
        r = p.add_run(); r.text = t; r.font.size = Pt(sz); r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
prs.save(dst)
