"""Refunds deck, Bulk Charges: Branch C (slide 28) gets a refund note under 'Beyond MCT' and a retention-steps tile
under 'Within MCT' with a button to slide 29; slide 29 becomes the 'Did C accept the offer?' flowchart (with the FCF
no-ETF branch). Patches the live Google export."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', '9A6410', 'FFF3D6'
RED, PINK, GREEN, GREENSOFT, PEACH, PEACHSOFT = 'B3261E', 'FCEBEA', '2E7D32', 'E4F2E1', 'C0612B', 'FCEADF'
prs = Presentation(src)
titles = [x.shapes[1].text_frame.text.strip() if x.shapes[1].has_text_frame else '' for x in prs.slides]
c = prs.slides[titles.index('Branch C: FC active, cancelling over it')]
e = prs.slides[titles.index('Within MCT: ETF credit or MCT reduction?')]
assert titles.index('Within MCT: ETF credit or MCT reduction?') == titles.index('Branch C: FC active, cancelling over it') + 1

def kit(s):
    def strip(x):
        st = x._element.find(P_NS + 'style')
        if st is not None: x._element.remove(st)
    def box(x, y, w, h, fill, line=None, adj=0.08):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
        b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(1)
        else: b.line.fill.background()
        return b
    def text(shape, paras, size=9.5, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=2, margin=0.08):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = 0
        for k, runs in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
            for t, bold, col, *sz in runs:
                r = p.add_run(); r.text = t; r.font.size = Pt(sz[0] if sz else size); r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
        return shape
    def tb(x, y, w, h): return s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    def line(x1, y1, x2, y2, arrow=False):
        cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); strip(cn)
        cn.line.color.rgb = RGBColor.from_string(SOFT); cn.line.width = Pt(1)
        if arrow:
            ln = cn.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    def label(x, y, w, t): text(tb(x, y, w, 0.2), [[(t, True, SOFT, 8.5)]], align=PP_ALIGN.LEFT, margin=0)
    return box, text, tb, line, label

# ---------- Slide 28 ----------
# Left box: add the refund note.
left = [x for x in c.shapes if x.has_text_frame and x.text_frame.text.startswith('Retainable?')][0]
p = left.text_frame.add_paragraph(); p.alignment = PP_ALIGN.LEFT
for t, bold, col in (('Refund request?  ', True, TEAL), ('Follow the standard refund policy.', False, INK)):
    r = p.add_run(); r.text = t; r.font.size = Pt(10 if bold else 9.5); r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
left.text_frame.paragraphs[-2].space_after = Pt(8)
# Right side: drop everything below 'Retention attempt first'.
ret = [x for x in c.shapes if x.has_text_frame and x.text_frame.text.strip() == 'Retention attempt first'][0]
for x in list(c.shapes):
    if not x.is_placeholder and x.left >= Inches(5.0) and x.top > ret.top + Inches(0.1): x._element.getparent().remove(x._element)
box, text, tb, line, label = kit(c)
line(7.325, 3.3, 7.325, 3.45, True)
tile = box(5.1, 3.45, 4.45, 1.0, WHITE, BORDER, adj=0.06)
text(tile, [[('Offer free months = the unused bulk MF charges', True, INK, 9.5)],
            [('Explain the FC / ETF terms', True, INK, 9.5)],
            [('Tailor further if other pain points exist ', True, INK, 9.5), ('(cost complaint → reduce MF + free months; PCQ → add a voucher + resolution)', False, SOFT, 8.5)]],
     align=PP_ALIGN.LEFT, after=3, margin=0.15)
line(7.325, 4.45, 7.325, 4.57, True)
btn = box(7.325 - 1.35, 4.57, 2.7, 0.4, TEAL, adj=0.25)
text(btn, [[('▶  C responds to the offer', True, WHITE, 10)]]); btn.click_action.target_slide = e
c.notes_slide.notes_text_frame.text = c.notes_slide.notes_text_frame.text.split(' Within MCT:')[0] + (
    ' Not retainable: cancel FC. If the customer has a refund request, follow the standard refund policy (MF & ETF Refund Requests). '
    'Within MCT: the ETF applies, and the retention attempt happens first: offer free months equal to the unused bulk MF charges, explain the FC/ETF terms, and tailor further if other pain points exist '
    '(e.g. a cost complaint: reduce the MF + free months; PCQ: add a voucher + resolution). TRAINER: click "C responds to the offer" to go to the next slide (did C accept or not).')

# ---------- Slide 29: flowchart ----------
for x in list(e.shapes):
    if not x.is_placeholder and x.top >= Inches(1.38): x._element.getparent().remove(x._element)
def retext(shape, t):
    r = shape.text_frame.paragraphs[0].runs; r[0].text = t
    for y in r[1:]: y._r.getparent().remove(y._r)
retext(e.shapes[1], 'Within MCT: did C accept the offer?'); retext(e.shapes[2], 'Accepted? Apply it. If not, exit at the lowest cost to C.')
box, text, tb, line, label = kit(e)
text(box(3.4, 1.4, 3.6, 0.4, INK), [[('Did C accept the offer?', True, WHITE, 10.5)]])
line(5.2, 1.8, 5.2, 1.9); line(1.75, 1.9, 5.2, 1.9); line(1.75, 1.9, 1.75, 2.02, True); line(5.2, 1.9, 5.2, 2.02, True)
label(1.83, 1.88, 0.6, 'Yes'); label(5.28, 1.88, 0.6, 'No')
text(box(0.45, 2.02, 2.6, 0.65, GREENSOFT, GREEN), [[('Acknowledge and apply the offers', True, GREEN, 9.5)], [('if not yet issued', False, INK, 8.5)]])
text(box(3.6, 2.02, 3.2, 0.45, GOLDSOFT, GOLD), [[('FCF, no completed cleaning yet?', True, GOLD, 9.5)]])
line(6.8, 2.245, 7.3, 2.245, True); label(6.88, 2.03, 0.4, 'Yes')
text(box(7.3, 2.02, 2.25, 0.45, INK), [[('No ETF. Cancel FC.', True, WHITE, 9.5)]])
line(5.2, 2.47, 5.2, 2.68, True); label(5.28, 2.47, 0.4, 'No')
text(box(3.45, 2.68, 3.5, 0.55, PINK, RED), [[('Reduce the MCT or the ETF,', True, RED, 9.5)], [('whichever is most cost-effective for C', False, INK, 8.5)]])
line(5.2, 3.23, 5.2, 3.32); line(3.5, 3.32, 6.9, 3.32); line(3.5, 3.32, 3.5, 3.42, True); line(6.9, 3.32, 6.9, 3.42, True)
text(box(2.0, 3.42, 3.0, 0.55, PEACHSOFT, PEACH), [[('MCT reduction', True, PEACH, 9.5)], [('Up to 2 months off the original MCT', False, INK, 8.5)]])
text(box(5.4, 3.42, 3.0, 0.55, GOLDSOFT, GOLD), [[('ETF reduction', True, GOLD, 9.5)], [('Apply the unused bulk MFs to the ETF', False, INK, 8.5)]])
line(6.9, 3.97, 6.9, 4.05); line(5.8, 4.05, 8.35, 4.05); line(5.8, 4.05, 5.8, 4.14, True); line(8.35, 4.05, 8.35, 4.14, True)
text(box(4.65, 4.14, 2.35, 0.75, GREENSOFT, GREEN), [[('Paid MFs < ETF', True, GREEN, 9.5)], [('Reduce the ETF by the paid, unused MF amount', False, INK, 8.5)]])
text(box(7.15, 4.14, 2.4, 0.75, PINK, RED), [[('Paid MFs ≥ ETF', True, RED, 9.5)], [('Waive the ETF, cancel FC', False, INK, 8.5)]])
text(box(0.45, 4.14, 4.0, 0.75, GOLDSOFT), [[('Near the end of the term? Compare both.', True, GOLD, 9)],
     [('1 month left, $150 ETF, $100 unused MFs: ETF credit leaves $50; reducing the last month = $0', False, INK, 8.5)]], align=PP_ALIGN.LEFT, margin=0.15)
e.notes_slide.notes_text_frame.text = (
    'Within MCT, after the retention attempt. Did the customer accept the offer? Yes: acknowledge and apply the offers if they haven\'t been issued yet. '
    'No: FCF customers who haven\'t had a first completed cleaning have no ETF at all (FCF\'s ETF only kicks in after the first completed cleaning): cancel FC. '
    'For everyone else, reduce the MCT or the ETF, whichever is most cost-effective for the customer. MCT reduction: up to 2 months off the original MCT, and only when it\'s cheaper than a reduced ETF. '
    'ETF reduction: apply the unused bulk MFs to the ETF. If the paid MFs are less than the ETF, reduce the ETF by the paid-unused-MF amount; if the paid MFs are equal to or more than the ETF, waive the ETF and cancel FC. '
    'Worked example: the customer has 1 month left on MCT, owes a $150 ETF, and has $100 in unused bulk MFs. ETF credit path: $100 credit, $50 ETF still owed. MCT reduction path: reduce the last month, MCT complete, ETF waived, $0 owed. '
    'MCT reduction wins here. Run this comparison whenever the customer is close to the end of their term; don\'t default to the ETF credit just because it\'s the more familiar move.')
prs.save(dst)
