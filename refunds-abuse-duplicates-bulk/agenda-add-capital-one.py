"""Day 9 deck: add "Capital One Card on File" to the Agenda (slide 2) as topic 07.

Capital One sits between Failed Charge (06) and Graceful Closure in the deck, so it takes
the 07 card and Graceful Closure moves to a new 08 card in the right column's 4th row.
Every "TOPIC 0N OF 07" divider kicker becomes "OF 08", and the Capital One divider kicker
becomes "TOPIC 07 OF 08". Cards keep their click-to-jump links.

Usage: python3 agenda-add-capital-one.py in.pptx out.pptx
"""
import copy
import sys

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

prs = Presentation(sys.argv[1])
slides = list(prs.slides)


def texts(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame]


def find_slide(kicker):
    for s in slides:
        for sh in texts(s):
            if sh.text_frame.text.strip() == kicker:
                return s, sh
    raise SystemExit(f"slide with kicker {kicker!r} not found")


def set_text(shape, *lines):
    for para, line in zip(shape.text_frame.paragraphs, lines):
        para.runs[0].text = line
        for r in para.runs[1:]:
            r.text = ""


def relink(shape, slide, target):
    rid = slide.part.relate_to(target.part, RT.SLIDE)
    for el in shape._element.iter(qn("a:hlinkClick")):
        el.set(qn("r:id"), rid)


agenda = next(s for s in slides if any(sh.text_frame.text.strip() == "Agenda" for sh in texts(s)))
capone, capone_kicker = find_slide("RETENTION AND MEMBERSHIP")

shapes = list(agenda.shapes)
sub = next(sh for sh in texts(agenda) if sh.text_frame.text.startswith("Seven topics"))
card = next(sh for sh in texts(agenda) if sh.text_frame.text.startswith("Graceful Closure"))
num = next(sh for sh in texts(agenda) if sh.text_frame.text.strip() == "07")
chev = next(sh for sh in texts(agenda)
            if sh.text_frame.text.strip() == "›" and sh.top == card.top and sh.left > card.left)
row4 = next(sh for sh in texts(agenda) if sh.text_frame.text.startswith("Bulk Charges")).top
dy = row4 - card.top

# 08: copy of the Graceful Closure card (keeps its link), placed in row 4 of the right column.
for sh in (card, num, chev):
    agenda.shapes._spTree.append(copy.deepcopy(sh._element))
new = list(agenda.shapes)[len(shapes):]
for sh in new:
    sh.top = sh.top + dy
    if sh.text_frame.text.strip() == "07":
        set_text(sh, "08")

# 07: the original Graceful Closure card becomes Capital One and jumps to its divider.
set_text(card, "Capital One Card on File",
         "Every card is Capital One: fix the payment.")
for sh in (card, num, chev):
    relink(sh, agenda, capone)

set_text(sub, "Eight topics. Click a topic to jump to it.")

# Divider kickers.
for s in slides:
    for sh in texts(s):
        t = sh.text_frame.text.strip()
        if t.startswith("TOPIC ") and t.endswith(" OF 07"):
            n = "08" if t == "TOPIC 07 OF 07" else t[6:8]
            set_text(sh, f"TOPIC {n} OF 08")
set_text(capone_kicker, "TOPIC 07 OF 08")

prs.save(sys.argv[2])
