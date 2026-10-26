"""Day 10 deck: label the "Referencing Help Center searches" slide "TOPIC 06 OF 06".

The Training Agenda Library starts a topic only at a slide whose top label reads "TOPIC NN OF NN".
Without it, slides 34-35 were counted as part of Projects. This changes only that one label.

Usage: python3 relabel-help-center-topic.py in.pptx out.pptx
"""
import sys

from pptx import Presentation

prs = Presentation(sys.argv[1])
hits = 0
for s in prs.slides:
    tx = [sh for sh in s.shapes if sh.has_text_frame]
    if len(tx) > 1 and tx[1].text_frame.text.strip() == "Referencing Help Center searches" \
            and tx[0].text_frame.text.strip() == "HELP CENTER SEARCHES":
        runs = tx[0].text_frame.paragraphs[0].runs
        runs[0].text = "TOPIC 06 OF 06"
        for r in runs[1:]:
            r.text = ""
        hits += 1
if hits != 1:
    raise SystemExit(f"expected 1 slide to relabel, found {hits}")
prs.save(sys.argv[2])
