"""Day 10 deck: add "Referencing Help Center searches" to the Agenda (slide 2) as topic 06.

The new card goes in the right column's third row (a copy of the 05 Projects card) and jumps to the
"Referencing Help Center searches" slide. The subtitle becomes "Six topics", and every
"TOPIC 0N OF 05" divider kicker becomes "OF 06".

Usage: python3 agenda-add-help-center-searches.py in.pptx out.pptx
"""
import copy
import sys

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

prs = Presentation(sys.argv[1])
slides = list(prs.slides)


def texts(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame]


def set_text(shape, *lines):
    for para, line in zip(shape.text_frame.paragraphs, lines):
        para.runs[0].text = line
        for r in para.runs[1:]:
            r.text = ""


agenda = next(s for s in slides if any(sh.text_frame.text.strip() == "Agenda" for sh in texts(s)))
target = next(s for s in slides if len(texts(s)) > 1 and texts(s)[1].text_frame.text.strip() == "Referencing Help Center searches")

card05 = next(sh for sh in texts(agenda) if sh.text_frame.text.startswith("Projects"))
group = [sh for sh in agenda.shapes
         if not sh.is_placeholder and sh.top == card05.top and card05.left <= sh.left < card05.left + card05.width]
row_h = card05.top - next(sh for sh in texts(agenda) if sh.text_frame.text.startswith("Personal Info Not Sold")).top

rid = agenda.part.relate_to(target.part, RT.SLIDE)
for sh in group:
    el = copy.deepcopy(sh._element)
    agenda.shapes._spTree.append(el)
    new = list(agenda.shapes)[-1]
    new.top = new.top + row_h
    for h in el.iter(qn("a:hlinkClick")):
        h.set(qn("r:id"), rid)
    t = new.text_frame.text.strip()
    if t == "05":
        set_text(new, "06")
    elif t.startswith("Projects"):
        set_text(new, "Referencing Help Center searches", "Respond to what C sent, not what C searched.")

sub = next(sh for sh in texts(agenda) if sh.text_frame.text.startswith("Five topics"))
set_text(sub, "Six topics. Click a topic to jump to it.")

for s in slides:
    for sh in texts(s):
        t = sh.text_frame.text.strip()
        if t.startswith("TOPIC ") and t.endswith(" OF 05"):
            set_text(sh, t[:-2] + "06")

prs.save(sys.argv[2])
