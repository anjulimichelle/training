"""Day 10 deck (Dispute History & Privacy): slides 5-7.

Slide 5 "Has a dispute reached us?": add two sample internal notes under the table,
  "Received, auto-lost (RDR)" and "Received and awaiting a response".
Slide 6 "Identifying a PayPal dispute": add the Payment history "Disputed" flag screenshot
  inside the Bank dispute panel.
Slide 7 "Refunds and cancellations": replace the two panels with three scenario tiles from
  the Knowledge Library ("Cancellation Intent & Dispute Scenarios"), keep the bottom callout
  for the 120-day exception.

Usage: python3 slides5-7-screenshots-and-scenarios.py in.pptx out.pptx IMG_DIR
IMG_DIR holds rdr-auto-lost.png, awaiting-response.png, payment-history-disputed.png.
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt
from PIL import Image

C = dict(TEAL="0B7A6F", TEAL_PALE="E3F4F1", AMBER="9A6410", AMBER_PALE="FFF3D6",
         INK="1D1D1F", MUTED="6E6E73", LINE="DCE1E7", WHITE="FFFFFF")
EMU = 12700  # 1 pt

src, dst, img_dir = sys.argv[1], sys.argv[2], sys.argv[3]
prs = Presentation(src)


def slide_titled(title):
    for s in prs.slides:
        for sh in s.shapes:
            # Title position only: the topic divider's "You'll learn" list repeats some titles.
            if (sh.has_text_frame and sh.text_frame.text.strip() == title
                    and abs(sh.top / EMU - 36) < 4):
                return s
    raise SystemExit(f"slide {title!r} not found")


def rgb(k):
    return RGBColor.from_string(C[k])


def strip(sh):
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05):
    sh = slide.shapes.add_shape(shape, x * EMU, y * EMU, w * EMU, h * EMU)
    strip(sh)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = adj
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    return sh


def frame(sh, anchor=MSO_ANCHOR.TOP, pad=0):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = pad * EMU
    return tf


def run(p, text, size, color="INK", bold=False, font="Arial"):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = rgb(color)
    return r


def bullet(p, color="INK"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(9 * EMU))
    pPr.set("indent", str(-9 * EMU))
    for tag, attrs in (("a:buClr", None), ("a:buSzPct", {"val": "100000"}),
                       ("a:buFont", {"typeface": "Arial"}), ("a:buChar", {"char": "•"})):
        el = pPr.makeelement(qn(tag), attrs or {})
        if tag == "a:buClr":
            el.append(el.makeelement(qn("a:srgbClr"), {"val": C[color]}))
        pPr.append(el)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    """paras: list of (runs, opts); runs = [(text, size, color, bold, font)], opts: bullet, space."""
    sh = slide.shapes.add_textbox(x * EMU, y * EMU, w * EMU, h * EMU)
    tf = frame(sh, anchor)
    for i, (runs, opts) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if opts.get("align"):
            p.alignment = opts["align"]
        if opts.get("space"):
            p.space_before = Pt(opts["space"])
        if opts.get("bullet"):
            bullet(p)
        for r in runs:
            run(p, *r)
    return sh


def picture(slide, path, x, y, w):
    iw, ih = Image.open(path).size
    h = w * ih / iw
    pic = slide.shapes.add_picture(path, int(x * EMU), int(y * EMU), int(w * EMU), int(h * EMU))
    pic.line.color.rgb = rgb("LINE")
    pic.line.width = Pt(0.75)
    return h


def prepend_notes(slide, line):
    tf = slide.notes_slide.notes_text_frame
    tf.text = line + "\n" + tf.text


def label(slide, x, y, w, s, color="TEAL"):
    text(slide, x, y, w, 12, [([(s, 9, color, True, "Play")], {})])


# ---------------- Slide 5: two sample notes ----------------
s5 = slide_titled("Has a dispute reached us?")
y = 274
label(s5, 32, y, 320, "Received, auto-lost (RDR)")
picture(s5, os.path.join(img_dir, "rdr-auto-lost.png"), 32, y + 14, 320)
label(s5, 367, y, 320, "Received and awaiting a response")
picture(s5, os.path.join(img_dir, "awaiting-response.png"), 367, y + 14, 320)
prepend_notes(s5, "TRAINER: Point to the two sample notes. Left: RDR auto-lost the dispute in Stripe "
              "(\"Auto-actions taken: RDR auto-lost dispute\"), so it's settled and Care has nothing to do. "
              "Right: \"Auto-actions taken: None\" and the instructions say to prepare a dispute response, "
              "so the DRT owns it.")

# ---------------- Slide 6: Disputed flag in the Bank dispute panel ----------------
s6 = slide_titled("Identifying a PayPal dispute")
label(s6, 45, 196, 294, "The flag: Payment history in the CRM", "MUTED")
picture(s6, os.path.join(img_dir, "payment-history-disputed.png"), 45, 210, 294)
prepend_notes(s6, "TRAINER: Show the flag in the screenshot: Payment history shows "
              "\"1 Disputed/Inquired Stripe Charge\" and a Disputed tag on the charge.")

# ---------------- Slide 7: three scenario tiles ----------------
s7 = slide_titled("Refunds and cancellations")
keep_texts = {"DISPUTE HISTORY", "Refunds and cancellations"}
shapes = list(s7.shapes)
callout = [sh for sh in shapes if round(sh.top / EMU) >= 300 and not sh.is_placeholder]
for sh in shapes:
    if sh.is_placeholder or sh in callout:
        continue
    if sh.has_text_frame and (sh.text_frame.text.strip() in keep_texts
                              or sh.text_frame.text.startswith("Only the charge")):
        continue
    sh._element.getparent().remove(sh._element)

sub = next(sh for sh in s7.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Only the charge"))
sub.text_frame.paragraphs[0].runs[0].text = "Three scenarios. Find the one C is in, then follow it."
for r in sub.text_frame.paragraphs[0].runs[1:]:
    r.text = ""

# Bottom callout: the 120-day exception (cancellation and ETF guidance moved into the tiles).
ct = next(sh for sh in callout if sh.has_text_frame and sh.text_frame.text.startswith("Also"))
runs = ct.text_frame.paragraphs[0].runs
runs[0].text = "Exception: "
runs[1].text = "a successfully disputed charge older than 120 days still counts as disputable."
for r in runs[2:]:
    r.text = ""

TOP, H, W, GAP = 104, 194, 210, 12.5
tiles = [
    ("TEAL", "TEAL_PALE", "C asks to cancel, before or after filing a dispute", [
        [("Proactively cancel the FC membership", 9, "INK", False, "Arial")],
        [("Waive the ETF, if applicable", 9, "INK", False, "Arial")],
        [("Same if cancellation intent was pending before the dispute: an open cancel ticket, "
          "a dashboard cancel attempt, or rejected offers", 9, "INK", False, "Arial")],
    ]),
    ("TEAL", "TEAL_PALE", "C asks for an ETF refund after the disputed charges were refunded", [
        [("Refund the ETF", 9, "INK", False, "Arial")],
        [("Do it proactively: ", 9, "INK", True, "Arial"),
         ("don't wait for C to ask separately", 9, "INK", False, "Arial")],
        [("The disputed charges were settled in C's favor, so the ETF follows the same logic",
          9, "INK", False, "Arial")],
    ]),
    ("AMBER", "AMBER_PALE", "C asks to refund other charges (not the ETF)", None),
]
for i, (dark, pale, title, items) in enumerate(tiles):
    x = 32 + i * (W + GAP)
    box(s7, x, TOP, W, H, pale, dark, adj=0.03)
    hdr = box(s7, x, TOP, W, 24, dark, dark, adj=0.17)
    tf = frame(hdr, MSO_ANCHOR.MIDDLE)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run(tf.paragraphs[0], f"SCENARIO {i + 1}", 9.5, "WHITE", True, "Play")
    paras = [([(title, 10, dark, True, "Play")], {})]
    if items:
        paras += [(r, {"bullet": True, "space": 4}) for r in items]
    else:
        paras += [
            ([("Charge is disputable", 9, "INK", True, "Arial"),
              (" (within 120 days, or the exception below)", 9, "MUTED", False, "Arial")], {"space": 5}),
            ([("Cancel FC, waive the ETF if applicable", 9, "INK", False, "Arial")], {"bullet": True, "space": 2}),
            ([("Refund the disputable charges or what C asks for, whichever is lower", 9, "INK", False, "Arial")],
             {"bullet": True, "space": 2}),
            ([("No longer disputable", 9, "INK", True, "Arial")], {"space": 6}),
            ([("No assessment or offers needed", 9, "INK", False, "Arial")], {"bullet": True, "space": 2}),
            ([("Politely tell C no further refund can be issued", 9, "INK", False, "Arial")],
             {"bullet": True, "space": 2}),
        ]
    text(s7, x + 11, TOP + 32, W - 22, H - 38, paras)

prepend_notes(s7, "TRAINER: Three scenarios, one tile each. 1: cancellation before or after a dispute, "
              "cancel FC and waive the ETF. 2: disputed charges already refunded, refund the ETF without "
              "waiting to be asked. 3: other refunds depend on whether the charge is still disputable.")

prs.save(dst)
