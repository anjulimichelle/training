"""Day 10 deck: split "Referencing Help Center searches" into two slides, before the Thank you slide.

Slide A "Referencing Help Center searches": what this is, why it matters, and the sample
  Help Center queries worth addressing (assets/help-center-query-examples.png).
Slide B "Search or message?": the Generic topic vs. Specific narrative table and the 3-step process.
Reuses the single slide added by add-help-center-search-slide.py as slide A (cleared and rebuilt).

Source: Care Knowledge Library tab "Referencing Help Center Search Activity" (Care Philosophy and
Communications), plus the trainer's screenshots of its example, table and process sections.

Usage: python3 help-center-search-two-slides.py in.pptx out.pptx ASSETS_DIR
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

C = {'TEAL': '0B7A6F', 'TEAL_PALE': 'E3F4F1', 'AMBER': '9A6410', 'AMBER_PALE': 'FFF3D6',
     'INK': '1D1D1F', 'MUTED': '6E6E73', 'WHITE': 'FFFFFF', 'LINE': 'DCE1E7', 'ALT': 'F5F6F8'}
X, W = 32.4, 655.2

prs = Presentation(sys.argv[1])
slides = list(prs.slides)
txt = lambda sh: sh.text_frame.text.strip() if sh.has_text_frame else ''
ref = next(s for s in slides if txt(s.shapes[0]) == 'DO NOT SELL' and txt(s.shapes[1]) == 'What is this?')
thanks = next(s for s in slides if txt(s.shapes[1]) == 'Thank you!')


def set_runs(shape, t):
    p0 = shape.text_frame.paragraphs[0]
    p0.runs[0].text = t
    for r in p0.runs[1:]:
        r.text = ''


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill_tf(tf, paras, anchor=MSO_ANCHOR.TOP, after=3):
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, item in enumerate(paras):
        num, runs = item if isinstance(item, tuple) else ('', item)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after); p.alignment = PP_ALIGN.LEFT
        if num:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(12 * 12700)); pPr.set('indent', str(-10 * 12700))
            runs = [(num + ' ', runs[0][1], True, 'TEAL' if num != '•' else 'MUTED')] + runs
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(C[color])


def box(s, x, y, w, h, bg, line=None, paras=(), anchor=MSO_ANCHOR.TOP, margin=11, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, 6 / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(C[bg])
    if line:
        sh.line.color.rgb = RGBColor.from_string(C[line]); sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(8)
    fill_tf(tf, paras, anchor)
    return sh


def dot(s, x, y, label, bg='TEAL', d=20):
    sh = box(s, x, y, d, d, bg, paras=[[(label, 9, True, 'WHITE')]], anchor=MSO_ANCHOR.MIDDLE, margin=0, shape=MSO_SHAPE.OVAL)
    sh.text_frame.margin_top = sh.text_frame.margin_bottom = 0
    sh.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return sh


def text(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, after=3):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, paras, anchor, after)
    return tb


B = lambda t, c='INK', sz=8.5: (t, sz, True, c)
A = lambda t, c='INK', sz=8.5: (t, sz, False, c)

assets = sys.argv[3]
lst = prs.slides._sldIdLst
ids = list(lst)

def place_before_thanks(el):
    lst.remove(el)
    thanks_el = next(x for x in lst if prs.part.related_part(x.get(qn('r:id'))) is thanks.part)
    thanks_el.addprevious(el)


def header(s, eyebrow, title, sub):
    set_runs(s.shapes[0], eyebrow); set_runs(s.shapes[1], title); set_runs(s.shapes[2], sub)
    return s


def new_slide(eyebrow, title, sub):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    num = next(x for x in ref.shapes if x.is_placeholder)
    for e in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
        s.shapes._spTree.append(copy.deepcopy(e))
    place_before_thanks(list(lst)[-1])
    return header(s, eyebrow, title, sub)


def reuse_slide(s, eyebrow, title, sub):
    """Keep the earlier single slide (and its place) but clear everything except its header."""
    for sh in list(s.shapes)[3:]:
        if not sh.is_placeholder:
            sh._element.getparent().remove(sh._element)
    return header(s, eyebrow, title, sub)


# the earlier single slide, if present, becomes slide A
old = next((sl for sl in slides if txt(sl.shapes[1]) == 'Referencing Help Center searches'), None)


EB = 'HELP CENTER SEARCHES'

# ---------- Slide A: what this is and why it matters ----------
args = ('TOPIC 06 OF 06', 'Referencing Help Center searches', "Respond to what C sent us. Don't bring up what C searched for.")
s = reuse_slide(old, *args) if old is not None else new_slide(*args)
LW = 268
box(s, X, 100, LW, 74, 'TEAL_PALE', paras=[
    [B('What is this?', 'TEAL', 10)],
    [A('Tickets sometimes show C\'s '), B('Help Center search queries'), A(' next to their message. This covers when to reference them, and when to leave them out.')]])
box(s, X, 182, LW, 162, 'WHITE', 'LINE', paras=[
    [B('Why this matters', 'TEAL', 10)],
    [A('Most of the time, C is just looking around on their own. '), B('They didn\'t tell us they searched'), A(', so we shouldn\'t act like we saw it.')],
    [A('Bringing it up can make C feel '), B('watched'), A(' and break their trust, even if the rest of the conversation went well.')],
    [B('The exception: ', 'AMBER'), A('some Cs don\'t know the search bar isn\'t a way to reach a person. If they\'re clearly trying to contact us, '),
     B('ignoring it'), A(' means ignoring someone who needs help.')]])
RX, RW = X + LW + 12, W - LW - 12
text(s, RX, 100, RW, 14, [[B('Help Center queries worth addressing', 'INK', 10)]])
from PIL import Image
img = f'{assets}/help-center-query-examples.png'
iw, ih = Image.open(img).size
ph = RW * ih / iw
pic = s.shapes.add_picture(img, Pt(RX), Pt(118), Pt(RW), Pt(ph))
pic.line.color.rgb = RGBColor.from_string(C['LINE']); pic.line.width = Pt(0.75)
box(s, RX, 118 + ph + 8, RW, 344 - (118 + ph + 8), 'AMBER_PALE', anchor=MSO_ANCHOR.MIDDLE, paras=[
    [B('Why these count: ', 'AMBER', 8.5), A('each one asks Care to act ("Call me", "I need to speak to a person") or tells us C\'s own plan. '
                                            'None of them is a topic someone would search for.', 'INK', 8.5)]])
s.notes_slide.notes_text_frame.text = (
    "Referencing Help Center search activity. Tickets sometimes show a customer's Help Center search queries next to their message. "
    "This covers when it's okay to reference what a customer searched for, and when to leave it out entirely.\n"
    "Why this matters: most of the time, a customer searching the Help Center is just looking around on their own. They didn't tell us they "
    "searched, so we shouldn't act like we saw it. Imagine typing into a company's search bar and getting a call hours later: \"Hey, I saw you "
    "searched for X.\" Even if they meant well, it would feel weird, even creepy. That can ruin the customer's trust, even if the rest of the "
    "conversation went well.\n"
    "The exception: some customers don't understand that the search bar isn't a way to talk to a real person. Something like \"please help me, "
    "my cleaner never showed up and I've called three times\" isn't a search; it's a customer trying to reach a human and not knowing how. If we "
    "ignore that because it showed up in a search field, we're ignoring someone who genuinely needs help.\n"
    "Walk through the examples: Legacy CRM \"Call me no show again\" and \"No I need speak to person who booked my original apt please.\" New CRM "
    "\"Don't want to renew membership after 6months is up.\" Each is addressed to Care or states the customer's own situation, so we respond to it.")

# ---------- Slide B: the test and the process ----------
s = new_slide(EB, 'Search or message?', "Don't judge by tone alone. Is it a short, general phrase, or a specific story about their situation?")
rows = [('Type', 'Example', 'Treat as'),
        ('Generic topic phrase', '"How do I cancel?" / "cancel my membership" / "cancel the service"', "Search: don't bring it up"),
        ('Generic topic phrase', '"refund my money" / "how do i get a refund"', "Search: don't bring it up"),
        ('Generic topic phrase', '"no show" / "cleaner cancelled"', "Search: don't bring it up"),
        ('Specific, personal, conversational', "Describes their exact situation with detail, reads like it's addressed to a person, not just a topic",
         'Message: acknowledge and respond to it')]
widths = [150, 330, W - 480]
heights = [18, 20, 20, 20, 28]
gt = s.shapes.add_table(len(rows), 3, Pt(X), Pt(100), Pt(W), Pt(sum(heights)))
tbl = gt.table
for j, w_ in enumerate(widths):
    tbl.columns[j].width = Pt(w_)
tblPr = tbl._tbl.tblPr
for a in ('bandRow', 'firstRow'):
    tblPr.set(a, '0')
sid = tblPr.find(qn('a:tableStyleId'))
if sid is not None:
    tblPr.remove(sid)
for i, row in enumerate(rows):
    tbl.rows[i].height = Pt(heights[i])
    msg = i == len(rows) - 1
    for j, v in enumerate(row):
        c = tbl.cell(i, j)
        c.margin_left = c.margin_right = Pt(9); c.margin_top = c.margin_bottom = Pt(3)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        bg = 'TEAL' if i == 0 else ('AMBER_PALE' if msg else 'WHITE')
        c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string(C[bg])
        tcPr = c._tc.get_or_add_tcPr()
        for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
            ln = tcPr.makeelement(qn(tag), {'w': '9525'})
            sf = ln.makeelement(qn('a:solidFill'), {}); sf.append(sf.makeelement(qn('a:srgbClr'), {'val': C['LINE']})); ln.append(sf)
            tcPr.append(ln)
        fe = tcPr.find(qn('a:solidFill')); tcPr.remove(fe); tcPr.insert(0, fe)
        p = c.text_frame.paragraphs[0]; c.text_frame.word_wrap = True
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        r = p.add_run(); r.text = v
        bold = i == 0 or j == 2
        r.font.size = Pt(9 if i == 0 else 8.5); r.font.bold = bold; r.font.name = 'Play' if bold else 'Arial'
        r.font.color.rgb = RGBColor.from_string(C['WHITE' if i == 0 else (('AMBER' if msg else 'TEAL') if j == 2 else 'INK')])

y0 = 100 + sum(heights) + 10
ws = [330, (W - 330 - 16) / 2, (W - 330 - 16) / 2]
steps = [('Look at what was typed, not just where', [
            [A('Does it read like a search query, or like a message written '), B('to'), A(' Support?')],
            ('•', [A('Reads like a search ("How do I cancel?", "Membership fee", "refund policy"): self-service browsing. '), B('Don\'t bring it up.'),
                   A(' Reply only to what C sent elsewhere.')]),
            ('•', [B('Obviously using the search box to contact Care?'), A(' Treat it as if C contacted us directly. Usually C is continuing a '
                   'conversation or asking Care to act, not knowing how to reach us.')])]),
         ('When in doubt, assume it\'s a search', [
            [A('The exception is only for cases where it\'s '), B('obvious'), A(' C meant it as outreach, not a judgment call you have to stretch to make.')]]),
         ('Respond, but don\'t say where from', [
            [A('Don\'t call out that it came from a Help Center search. '), B('Respond to the substance'), A(', the same as any message from a normal channel.')]])]
x = X
for i, (h, body) in enumerate(steps):
    w_ = ws[i]
    box(s, x, y0, w_, 344 - y0, 'TEAL_PALE' if i == 0 else 'WHITE', None if i == 0 else 'LINE')
    dot(s, x + 10, y0 + 9, str(i + 1))
    text(s, x + 36, y0 + 11, w_ - 44, 16, [[B(h, 'TEAL', 9)]])
    text(s, x + 12, y0 + 36, w_ - 22, 344 - y0 - 40, body, after=3)
    x += w_ + 8
s.notes_slide.notes_text_frame.text = (
    "The test: generic topic versus specific narrative. Don't judge by tone alone, because a lot of real searches sound like commands. Ask: is this "
    "a short, general phrase someone would type to find an article, or a specific story about their situation?\n"
    "Generic topic phrases like \"How do I cancel?\", \"cancel my membership\", \"cancel the service\", \"refund my money\", \"how do i get a refund\", "
    "\"no show\" and \"cleaner cancelled\" are searches: don't bring them up. They're keyword-shaped, not narrative-shaped. What flips something into "
    "\"this was meant for us\" is specificity: it names details unique to their situation, reads like it's speaking to a person rather than labeling a "
    "topic, or wouldn't make sense as a search-engine-style query at all. Then: acknowledge and respond to it.\n"
    "The process. 1) Look at what was typed, not just where it was typed. Reads like a search, such as \"How do I cancel?\", \"Membership fee\" or "
    "\"refund policy\": genuine self-service browsing, so don't bring it up and respond only to what the customer explicitly sent elsewhere. If it's "
    "obvious the customer is using the search box to contact Customer Care, treat it as though they contacted us directly; this typically happens when "
    "they appear to be continuing a conversation or asking Care to take action because they don't know the proper way to reach us. 2) When in doubt, "
    "assume it's a search; the exception is only for cases where it's obvious the customer meant it as outreach. 3) If you do respond, don't call out "
    "that it came from a Help Center search. Just respond to the substance, the same way you would through a normal channel.")

prs.save(sys.argv[2])
