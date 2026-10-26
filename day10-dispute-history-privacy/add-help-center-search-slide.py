"""Day 10 deck: add one slide, "Referencing Help Center Search Activity", before the Thank you slide.

Source: Care Knowledge Library tab "Referencing Help Center Search Activity" (Care Philosophy and
Communications). The header (eyebrow, title, subtitle, slide number) is cloned from the
Do Not Sell "What is this?" slide so the theme matches.

Usage: python3 add-help-center-search-slide.py in.pptx out.pptx
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

C = {'TEAL': '0B7A6F', 'TEAL_PALE': 'E3F4F1', 'AMBER': '9A6410', 'AMBER_PALE': 'FFF3D6',
     'INK': '1D1D1F', 'MUTED': '6E6E73', 'WHITE': 'FFFFFF', 'LINE': 'DCE1E7', 'ALT': 'F5F6F8'}
X, W = 32.4, 655.2

prs = Presentation(sys.argv[1])
slides = list(prs.slides)
txt = lambda sh: sh.text_frame.text.strip() if sh.has_text_frame else ''
ref = next(s for s in slides if txt(s.shapes[0]) == 'DO NOT SELL' and txt(s.shapes[1]) == 'What is this?')
thanks = next(s for s in slides if txt(s.shapes[1]) == 'Thank you!')


def set_runs(shape, t):
    p0 = shape.text_frame.paragraphs[0]
    p0.runs[0].text = t
    for r in p0.runs[1:]:
        r.text = ''


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill_tf(tf, paras, anchor=MSO_ANCHOR.TOP, after=3):
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, item in enumerate(paras):
        num, runs = item if isinstance(item, tuple) else ('', item)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after); p.alignment = PP_ALIGN.LEFT
        if num:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(12 * 12700)); pPr.set('indent', str(-10 * 12700))
            runs = [(num + ' ', runs[0][1], True, 'TEAL' if num != '•' else 'MUTED')] + runs
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(C[color])


def box(s, x, y, w, h, bg, line=None, paras=(), anchor=MSO_ANCHOR.TOP, margin=11, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, 6 / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(C[bg])
    if line:
        sh.line.color.rgb = RGBColor.from_string(C[line]); sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(8)
    fill_tf(tf, paras, anchor)
    return sh


def dot(s, x, y, label, bg='TEAL', d=20):
    sh = box(s, x, y, d, d, bg, paras=[[(label, 9, True, 'WHITE')]], anchor=MSO_ANCHOR.MIDDLE, margin=0, shape=MSO_SHAPE.OVAL)
    sh.text_frame.margin_top = sh.text_frame.margin_bottom = 0
    sh.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return sh


def text(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, after=3):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, paras, anchor, after)
    return tb


B = lambda t, c='INK', sz=8.5: (t, sz, True, c)
A = lambda t, c='INK', sz=8.5: (t, sz, False, c)

# new slide with the cloned header, placed before Thank you
s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
num = next(x for x in ref.shapes if x.is_placeholder)
for e in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(e))
set_runs(s.shapes[0], 'CARE COMMUNICATIONS')
set_runs(s.shapes[1], 'Referencing Help Center searches')
set_runs(s.shapes[2], "Respond to what C sent us. Don't bring up what C searched for.")
lst = prs.slides._sldIdLst
ids = list(lst)
new_el = ids[-1]
lst.remove(new_el)
ids[slides.index(thanks)].addprevious(new_el)

# left: why it matters
LW = 196
box(s, X, 100, LW, 134, 'TEAL_PALE', paras=[
    [B('Why it matters', 'TEAL', 10)],
    [A('C didn\'t tell us they searched. If we bring it up, it feels like being '), B('watched'),
     A(', and that can break their trust even if the rest of the reply is great.')],
    [A('Think of a company calling you to say "I saw you searched for X."', 'MUTED', 8)]])

# right: the test
RX, RW = X + LW + 10, W - LW - 10
text(s, RX, 100, RW, 14, [[B('The test: a generic topic, or a specific story?', 'INK', 10)]])
cw = (RW - 8) / 2
box(s, RX, 120, cw, 114, 'WHITE', 'LINE', paras=[
    [B('Generic topic phrase', 'TEAL', 9.5)],
    [A('"How do I cancel?"  ·  "refund my money"  ·  "no show"  ·  "cleaner cancelled"  ·  "Membership fee"', 'MUTED', 8)],
    [B('Search: don\'t bring it up.'), A(' Reply only to what C sent elsewhere.')],
    [B('Don\'t judge by tone: ', 'AMBER', 8), A('many real searches sound like commands.', 'INK', 8)]])
box(s, RX + cw + 8, 120, cw, 114, 'AMBER_PALE', 'AMBER', paras=[
    [B('Specific, personal story', 'AMBER', 9.5)],
    [A('"please help me, my cleaner never showed up and I\'ve called three times"', 'MUTED', 8)],
    [B('Message: respond to it'), A(' as if C contacted us directly. C is trying to reach a person and doesn\'t know how.')]])

# bottom: the process
steps = [('Look at what was typed', [A('Not where it was typed. Does it read like a search, or like a message written '), B('to'), A(' Support?')]),
         ('When in doubt, it\'s a search', [A('Treat it as outreach only when that\'s '), B('obvious'), A(', not a stretch.')]),
         ('Never say where it came from', [A('Don\'t mention the Help Center search. Answer the substance like any other message.')])]
sw = (W - 16) / 3
for i, (h, body) in enumerate(steps):
    x = X + i * (sw + 8)
    box(s, x, 246, sw, 84, 'WHITE', 'LINE')
    dot(s, x + 10, 256, str(i + 1))
    text(s, x + 38, 259, sw - 46, 16, [[B(h, 'TEAL', 9)]])
    text(s, x + 12, 284, sw - 22, 42, [body])

s.notes_slide.notes_text_frame.text = (
    "Referencing Help Center search activity. Tickets sometimes show a customer's Help Center search queries next to their message. "
    "Most of the time, a customer searching the Help Center is just looking around on their own. They didn't tell us they searched, so we "
    "shouldn't act like we saw it. Imagine typing something into a company's search bar and getting a call hours later saying \"I saw you "
    "searched for X.\" Even if it was meant well, it feels creepy, and it can ruin the customer's trust even if the rest of the conversation went well.\n"
    "The exception: some customers don't realize the search bar isn't a way to reach a person. Something like \"please help me, my cleaner never "
    "showed up and I've called three times\" isn't a search; it's a customer trying to reach us. Ignoring it because it landed in a search field "
    "means ignoring someone who needs help.\n"
    "The test is generic topic versus specific story, not tone, because many real searches sound like commands. Short keyword phrases like "
    "\"How do I cancel?\", \"refund my money\", \"no show\" or \"cleaner cancelled\" are searches: don't bring them up. A message that names details "
    "of their own situation and reads like it's addressed to a person is outreach: respond to it.\n"
    "The process: 1) look at what was typed, not where it was typed. 2) When in doubt, assume it's a search; treat it as outreach only when it's "
    "obvious, for example when C seems to be continuing a conversation or asking Care to take action. 3) If you respond, don't say it came from "
    "a Help Center search; respond to the substance the same way you would through any normal channel.\n"
    "Source: Care Knowledge Library, Referencing Help Center Search Activity (Care Philosophy and Communications).")

prs.save(sys.argv[2])
