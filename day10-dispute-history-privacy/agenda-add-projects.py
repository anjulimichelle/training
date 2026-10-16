"""Day 10 deck: add "Projects" to the Agenda (slide 2) as topic 05.

The agenda goes from a 2x2 grid to two columns of three (read down each column):
01-03 on the left, 04-05 on the right. Projects (05) jumps to its divider slide.
Divider kickers become "TOPIC 0N OF 05"; the Projects divider kicker "PROJECTS"
becomes "TOPIC 05 OF 05".

Usage: python3 agenda-add-projects.py in.pptx out.pptx
"""
import copy
import sys

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

EMU = 12700
prs = Presentation(sys.argv[1])
slides = list(prs.slides)


def texts(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame]


def set_text(shape, *lines):
    for para, line in zip(shape.text_frame.paragraphs, lines):
        para.runs[0].text = line
        for r in para.runs[1:]:
            r.text = ""


agenda = next(s for s in slides if any(sh.text_frame.text.strip() == "Agenda" for sh in texts(s)))
divider = next(s for s in slides
               if any(sh.text_frame.text.strip() == "PROJECTS" for sh in texts(s))
               and any(sh.text_frame.text.strip() == "YOU'LL LEARN" for sh in texts(s)))


def group(title_start):
    """The card, its number and its chevron: the shapes sharing the card's row and column."""
    card = next(sh for sh in texts(agenda) if sh.text_frame.text.startswith(title_start))
    return [sh for sh in agenda.shapes
            if not sh.is_placeholder and sh.top == card.top and card.left <= sh.left < card.left + card.width]


def move(shapes, dx, dy):
    for sh in shapes:
        sh.left, sh.top = sh.left + dx, sh.top + dy


cards = [group(t) for t in ("Customer Has Dispute History", "Account or Personal Data Deleted",
                            "Arbitration Clause", "Personal Info Not Sold")]
left_x = cards[0][0].left
right_x = cards[3][0].left
row_h = cards[1][0].top - cards[0][0].top
top = 124 * EMU  # three rows, centred where the two rows sat

# 01, 02 stay in the left column; 03 drops to the left column's third row; 04 moves to the right column's first row.
for i, g in enumerate(cards[:2]):
    move(g, 0, top + i * row_h - g[0].top)
move(cards[2], left_x - cards[2][0].left, top + 2 * row_h - cards[2][0].top)
move(cards[3], right_x - cards[3][0].left, top - cards[3][0].top)

# 05 Projects: a copy of the 04 card in the right column's second row.
new = []
for sh in cards[3]:
    el = copy.deepcopy(sh._element)
    agenda.shapes._spTree.append(el)
    new.append(list(agenda.shapes)[-1])
move(new, 0, row_h)
rid = agenda.part.relate_to(divider.part, RT.SLIDE)
for sh in new:
    for el in sh._element.iter(qn("a:hlinkClick")):
        el.set(qn("r:id"), rid)
    t = sh.text_frame.text.strip()
    if t == "04":
        set_text(sh, "05")
    elif t.startswith("Personal Info Not Sold"):
        set_text(sh, "Projects", "Check first, then reassign or handle BAU.")

sub = next(sh for sh in texts(agenda) if sh.text_frame.text.startswith("Four topics"))
set_text(sub, "Five topics. Click a topic to jump to it.")

for s in slides:
    for sh in texts(s):
        t = sh.text_frame.text.strip()
        if t.startswith("TOPIC ") and t.endswith(" OF 04"):
            set_text(sh, t[:-2] + "05")
set_text(next(sh for sh in texts(divider) if sh.text_frame.text.strip() == "PROJECTS"), "TOPIC 05 OF 05")

prs.save(sys.argv[2])
