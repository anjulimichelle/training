"""Day 8 deck slide 31: replace "Early Fraud Warning example" with "HDR Ticket 6: Job charged
more than 15 days ago" as a response table (1st Response steps 1-2, 2nd Response step 1).
The old slide text moves to the speaker notes.

Usage: python3 slide31-hdr-ticket6.py in.pptx out.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE, RED = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', 'A33A2B'

prs = Presentation(sys.argv[1])
s = prs.slides[30]
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
assert txt(s.shapes[1]) == 'Early Fraud Warning example', txt(s.shapes[1])
old_text = ' '.join(txt(sh) for sh in list(s.shapes)[2:] if txt(sh) and not sh.is_placeholder)
old_notes = s.notes_slide.notes_text_frame.text
sl = list(s.shapes)
keep = {id(x._element) for x in sl[:3]} | {id(x._element) for x in sl if x.is_placeholder}
for sh in sl:
    if id(sh._element) not in keep:
        sh._element.getparent().remove(sh._element)


def set_text(shape, t):
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = t
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


set_text(s.shapes[1], 'HDR Ticket 6')
set_text(s.shapes[2], 'Job charged more than 15 days ago')
r = s.shapes[2].text_frame.paragraphs[0].runs[0]
r.font.bold = True; r.font.name = 'Play'; r.font.color.rgb = RGBColor.from_string(TEAL)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def rect(x, y, w, h, fill):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    sh.line.color.rgb = RGBColor.from_string(LINE); sh.line.width = Pt(0.75)
    return sh


def write(x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, after=1.5):
    """paras: list of (level, runs); level 0 plain, 1 bullet, 2 sub-bullet, 3 sub-sub-bullet."""
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (lvl, runs) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(after)
        if lvl:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int((8 + 12 * (lvl - 1)) * 12700 + 8 * 12700)))
            pPr.set('indent', str(-8 * 12700))
            pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
            pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '•◦▪'[lvl - 1]}))
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tb


SZ = 8
B = lambda t, c=INK, sz=SZ: (t, sz, True, c)
A = lambda t, c=INK, sz=SZ: (t, sz, False, c)

X, W, LW, SW = 32.4, 655.2, 88, 54
SX, CX = X + LW, X + LW + SW
CW = W - LW - SW
y0, h1, h2, h3 = 100, 62, 128, 82

# 1st response
rect(X, y0, LW, h1 + h2, AMBER)
write(X, y0, LW, h1 + h2, [(0, [B('1st Response', WHITE, 10)])], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
for y, h, label in ((y0, h1, 'Step 1'), (y0 + h1, h2, 'Step 2')):
    rect(SX, y, SW, h, AMBER_PALE)
    write(SX, y, SW, h, [(0, [A(label, AMBER, 8.5)])], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    rect(CX, y, CW, h, AMBER_PALE)
write(CX + 12, y0 + 7, CW - 24, h1 - 10, [
    (0, [B('First, sanity-check the charge itself:')]),
    (1, [A('Looks like a '), B('false invoice'), A(' → usual False Invoice process.')]),
    (1, [A('Looks '), B('overcharged'), A(' → usual OCH process.')]),
    (1, [A('Charge is legit and the job happened → keep going.')])])
write(CX + 12, y0 + h1 + 7, CW - 24, h2 - 10, [
    (0, [B("Then check whether C has actually said anything about it:")]),
    (1, [B('No', TEAL), A(", C didn't reach out about the job charge:")]),
    (2, [A('Create an internal note of your findings and simply clear the ticket.')]),
    (1, [B('Yes', RED), A(', C reached out '), B('without refund intent but raised issues', RED), A(' about the job:')]),
    (2, [A('Handle the ticket as usual.')]),
    (1, [B('Yes', RED), A(', C reached out '), B('with refund intent', RED), A(':')]),
    (2, [A("Handle the refund request as normal, based on the issue C raised. If C didn't mention the reason for the refund request:")]),
    (3, [A('Probe to identify the root cause of the refund request')]),
    (3, [A('Explain the validity of the job charge')]),
    (3, [A('Address all other pain points as usual')])])

# 2nd response
y2 = y0 + h1 + h2
rect(X, y2, LW, h3, TEAL)
write(X, y2, LW, h3, [(0, [B('2nd Response', WHITE, 10)])], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
rect(SX, y2, SW, h3, TEAL_PALE)
write(SX, y2, SW, h3, [(0, [A('Step 1', TEAL, 8.5)])], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
rect(CX, y2, CW, h3, WHITE)
write(CX + 12, y2 + 8, CW - 24, h3 - 12, [
    (0, [B('If the job charge is valid and C is still insisting on a refund:')]),
    (1, [A('Assess against the card-based refund threshold:')]),
    (2, [B('Less than the card-based threshold: ', TEAL), A('process the refund.')]),
    (2, [B('Greater than the card-based threshold: ', AMBER), A('no refund, unless other Transformation Guidelines apply (e.g. an Attorney General threat).')])])

s.notes_slide.notes_text_frame.text = (
    'HDR Ticket 6: job charged more than 15 days ago.\n'
    '1st response, step 1: sanity-check the charge itself. If it looks like a false invoice, use the usual False Invoice process; if it looks overcharged, use the usual OCH process; '
    'if the charge is legit and the job happened, keep going.\n'
    "Step 2: check whether C has actually said anything about it. No, C didn't reach out about the job charge: create an internal note of your findings and simply clear the ticket. "
    'Yes, C reached out without refund intent but raised issues about the job: handle the ticket as usual. Yes, C reached out with refund intent: handle the refund request as normal based on the issue raised; '
    "if C didn't mention the reason, probe for the root cause, explain the validity of the job charge, and address all other pain points as usual.\n"
    '2nd response: if the job charge is valid and C is still insisting on a refund, assess against the card-based refund threshold. Less than the threshold: process the refund. '
    'Greater than the threshold: no refund, unless other Transformation Guidelines apply, for example an Attorney General threat.\n'
    'From the previous version of this slide (Early Fraud Warning example): ' + old_text + '\n' + old_notes)
prs.save(sys.argv[2]); print('saved', sys.argv[2])
