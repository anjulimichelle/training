"""Day 8 deck slide 25 (HDR Ticket Handling, "What is this?"): rebuild as an overview of
system-flagged HDR tickets: how to spot the flag (internal note screenshot), the card-based
dispute threshold table (Discover uses the percentage-based assessment), and a brief note on
guideline exceptions. Slide 26 (Card-based dispute threshold) is then deleted.

Usage: python3 slide25-hdr-overview.py in.pptx out.pptx
"""
import os, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets', 'hdr-flag-internal-note.png')

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
s, s26 = prs.slides[24], prs.slides[25]
assert txt(s.shapes[0]) == 'HDR TICKET HANDLING' and txt(s.shapes[1]) == 'What is this?'
assert txt(s26.shapes[1]) == 'Card-based dispute threshold'
old_notes = s.notes_slide.notes_text_frame.text
s26_notes = s26.notes_slide.notes_text_frame.text

sl = list(s.shapes)
keep = {id(x._element) for x in sl[:3]} | {id(x._element) for x in sl if x.is_placeholder}
for sh in sl:
    if id(sh._element) not in keep:
        sh._element.getparent().remove(sh._element)

sub = s.shapes[2].text_frame.paragraphs[0]
sub.runs[0].text = 'Pertains to tickets the system flagged as HDR (High Dispute Risk).'
for r in sub.runs[1:]:
    r._r.getparent().remove(r._r)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill(tf, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=2):
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(after)
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)


def box(x, y, w, h, bg, line=None, paras=(), anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT, margin=10, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, 6 / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(bg)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(2)
    fill(tf, paras, align, anchor)
    return sh


def text(x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill(tf, paras, align, anchor)
    return tb


def dot(x, y, label):
    return box(x, y, 18, 18, TEAL, paras=[[(label, 9, True, WHITE)]], align=PP_ALIGN.CENTER, margin=0, shape=MSO_SHAPE.OVAL)


L, TOP, BOT = 32.4, 108, 300
LW = 322
RX = L + LW + 11
RW = 687.6 - RX

# ---- left: how to spot it ----
box(L, TOP, LW, BOT - TOP, TEAL_PALE)
dot(L + 10, TOP + 10, '1')
text(L + 34, TOP + 10, LW - 44, 18, [[('How to spot it', 11, True, TEAL)]], anchor=MSO_ANCHOR.MIDDLE)
text(L + 12, TOP + 36, LW - 24, 44, [[('The system adds an ', 9, False, INK), ('internal note', 9, True, INK),
                                       (' with the action item ', 9, False, INK), ('Customer.flag_as_dispute_risk', 9, True, TEAL),
                                       (', saying why C was flagged as HDR.', 9, False, INK)]])
iw = LW - 24
ih = iw * 152 / 702
frame = box(L + 8, TOP + 86, iw + 8, ih + 8, WHITE, LINE, margin=0)
s.shapes.add_picture(IMG, Pt(L + 12), Pt(TOP + 90), Pt(iw), Pt(ih))
text(L + 12, TOP + 98 + ih, LW - 24, 30, [[('Example: ', 8.5, True, MUTED),
                                         ('flagged because C called the phone number on their credit card statement.', 8.5, False, MUTED)]])

# ---- right: card-based dispute threshold ----
box(RX, TOP, RW, BOT - TOP, WHITE, LINE)
dot(RX + 10, TOP + 10, '2')
text(RX + 34, TOP + 10, RW - 44, 18, [[('Card-based dispute threshold', 11, True, TEAL)]], anchor=MSO_ANCHOR.MIDDLE)
text(RX + 12, TOP + 34, RW - 24, 30, [[('A flat dollar amount per card network. If the disputable amount is below it, refund.', 9, False, INK)]])
rows = [('Card', 'Threshold'), ('Amex', '< $50'), ('MasterCard', '< $58'), ('Visa', '< $70'),
        ('Discover', 'No flat threshold: use the Percentage-based Dispute Threshold Assessment')]
tx, ty, tw = RX + 12, TOP + 66, RW - 24
heights = [20, 20, 20, 20, 32]
gt = s.shapes.add_table(len(rows), 2, Pt(tx), Pt(ty), Pt(tw), Pt(sum(heights)))
tbl = gt.table
tbl.columns[0].width = Pt(82); tbl.columns[1].width = Pt(tw - 82)
tblPr = tbl._tbl.tblPr
for a in ('bandRow', 'firstRow'):
    tblPr.set(a, '0')
sid = tblPr.find(qn('a:tableStyleId'))
if sid is not None:
    tblPr.remove(sid)
for i, (a, b) in enumerate(rows):
    tbl.rows[i].height = Pt(heights[i])
    for j, v in enumerate((a, b)):
        c = tbl.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor.from_string(TEAL if i == 0 else (AMBER_PALE if a == 'Discover' else (WHITE if i % 2 else 'F5F6F8')))
        c.margin_left = c.margin_right = Pt(8); c.margin_top = c.margin_bottom = Pt(2)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        r = p.add_run(); r.text = v; f = r.font
        head = i == 0 or j == 0
        f.name = 'Play' if head else 'Arial'; f.bold = head
        f.size = Pt(9 if i == 0 else 9); f.color.rgb = RGBColor.from_string(WHITE if i == 0 else (AMBER if a == 'Discover' and j == 1 else INK))

# ---- bottom: guideline exceptions ----
box(L, BOT + 10, 655.2, 42, AMBER_PALE, AMBER,
    paras=[[('Guideline exceptions:  ', 9, True, AMBER),
            ('a charge above the threshold can still be refunded if another Transformation Guideline applies, most often a material legal threat or extreme financial hardship. '
             'The table is a default, not an absolute rule.', 9, False, INK)]])

s.notes_slide.notes_text_frame.text = (
    'These are tickets the system flagged as HDR (High Dispute Risk). To spot one, look for the internal note with the action item '
    'Customer.flag_as_dispute_risk. It says why the customer was flagged, for example that they just called the phone number listed on their credit card statement.\n'
    'Card-based dispute threshold: a flat dollar amount per card network. Amex under $50, MasterCard under $58, Visa under $70: if the disputable amount is under it, refund. '
    'Discover has no flat threshold: use the Percentage-based Dispute Threshold Assessment (if you are not trained on it, reassign to the Risk Reduction queue).\n'
    'Guideline exceptions: a charge above the threshold can still be refunded if another Transformation Guideline applies, most commonly a material legal threat or an extreme financial hardship case. '
    'The threshold table is a default, not an absolute rule.\n'
    'Background from the previous version of this slide: ' + old_notes + '\n'
    'From the removed card-threshold slide: ' + s26_notes)

lst = prs.slides._sldIdLst
sld = lst[25]
rid = sld.get(qn('r:id'))
lst.remove(sld)
prs.part.drop_rel(rid)
prs.save(sys.argv[2]); print('saved', sys.argv[2])
