"""Day 8 deck slide 49 (Internal Note Action Items: What is this?): remove the "Where to look" tile,
stack the "Three kinds" and "Two jobs" tiles on the left, and add a "How to spot it in CRM" panel
(the Sender field is a system code) with a sample internal note screenshot.

Usage: python3 slide49-action-items-how-to-spot.py in.pptx out.pptx
"""
import os, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, INK, MUTED, WHITE, LINE = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets', 'action-item-sender-sample.png')

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
s = next(sl for sl in prs.slides if txt(sl.shapes[0]) == 'INTERNAL NOTE ACTION ITEMS' and txt(sl.shapes[1]) == 'What is this?')
sl = list(s.shapes)


def tile_group(title):
    """The card, icon circle, icon picture, heading and body that make up one tile."""
    k = next(i for i, sh in enumerate(sl) if txt(sh) == title)
    return sl[k - 3:k + 2]


where, kinds, jobs = tile_group('Where to look'), tile_group('Three kinds'), tile_group('Two jobs')
for sh in where:
    sh._element.getparent().remove(sh._element)

X = 32.4
LW, TH, G = 262, 95, 6
for n, (card, circ, icon, head, body) in enumerate((kinds, jobs)):
    y = 104 + n * (TH + G)
    dx_icon = icon.left - circ.left; dy_icon = icon.top - circ.top
    card.left, card.top, card.width, card.height = Pt(X), Pt(y), Pt(LW), Pt(TH)
    circ.left, circ.top = Pt(X + 12), Pt(y + 12)
    icon.left, icon.top = circ.left + dx_icon, circ.top + dy_icon
    head.left, head.top, head.width, head.height = Pt(X + 48), Pt(y + 12), Pt(LW - 58), Pt(20)
    body.left, body.top, body.width, body.height = Pt(X + 48), Pt(y + 34), Pt(LW - 58), Pt(TH - 40)

bar = next(sh for sh in sl if txt(sh).startswith('Unlisted trigger'))
k = sl.index(bar)
dy = Pt(306) - sl[k - 2].top
for sh in sl[k - 2:k + 1]:
    sh.top = sh.top + dy


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


RX = X + LW + 12; RW = 655.2 - LW - 12
panel = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(RX), Pt(104), Pt(RW), Pt(196)); strip(panel)
panel.adjustments[0] = 0.04
panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor.from_string(TEAL_PALE)
panel.line.color.rgb = RGBColor.from_string(TEAL); panel.line.width = Pt(0.75)
tb = s.shapes.add_textbox(Pt(RX + 12), Pt(112), Pt(RW - 24), Pt(40))
tf = tb.text_frame; tf.word_wrap = True
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
for i, runs in enumerate([[('How to spot it in CRM', 10, True, TEAL)],
                          [('The ', 8.5, False, INK), ('Sender', 8.5, True, INK),
                           (' field of the internal note action item (top right) is a code from the system, e.g. ', 8.5, False, INK),
                           ('utils.report_double_booked_clients', 8.5, True, TEAL), ('. Look it up, then follow its steps.', 8.5, False, INK)]]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(2)
    for t, size, bold, color in runs:
        r = p.add_run(); r.text = t; f = r.font
        f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold; f.color.rgb = RGBColor.from_string(color)
iw = RW - 24; ih = iw * 267 / 735
s.shapes.add_picture(IMG, Pt(RX + 12), Pt(104 + 196 - ih - 10), Pt(iw), Pt(ih))

nt = s.notes_slide.notes_text_frame
nt.text = nt.text + ('\nHow to spot it in CRM: an action item is an internal note whose Sender field (top right of the note) is a system code. '
                     'In the sample, the Sender is utils.report_double_booked_clients: the customer has double-booked appointments, and the Agent Do steps say to '
                     'call or confirm whether C wants both; if they obviously do, silence the flag for 30 days (C CRM > Do > Add Flag > admin_approve_double_book); '
                     'otherwise cancel the unwanted one(s) to avoid LMCs and lockouts.')
prs.save(sys.argv[2]); print('saved', sys.argv[2])
