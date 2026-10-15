"""Day 8 deck (helpers copied from add-practice-and-hdr-kc.py):
  - After "Two consecutive issue jobs: the process": add a knowledge check covering ETF waivers
    and two consecutive serious issue jobs (one answer revealed per click).

Usage: python3 add-etf-waivers-kc.py in.pptx out.pptx
"""
import copy, os, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets', 'cancel-after-mct-customer-message.png')

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''


def find(title, eyebrow_start=None):
    for i, sl in enumerate(prs.slides):
        if txt(sl.shapes[1]) == title and (eyebrow_start is None or txt(sl.shapes[0]).startswith(eyebrow_start)):
            return i, sl
    raise SystemExit('not found: ' + title)


def new_slide(ref, n_head, after_idx, texts):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    num = [x for x in ref.shapes if x.is_placeholder][0]
    for el in [ref.shapes[k]._element for k in range(n_head)] + [num._element]:
        s.shapes._spTree.append(copy.deepcopy(el))
    for k, t in enumerate(texts):
        r = s.shapes[k].text_frame.paragraphs[0].runs
        r[0].text = t
        for x in r[1:]:
            x._r.getparent().remove(x._r)
    lst = prs.slides._sldIdLst
    el = list(lst)[-1]; lst.remove(el); list(lst)[after_idx].addnext(el)
    return s


def kit(s):
    def strip(sh):
        st = sh._element.find(qn('p:style'))
        if st is not None:
            sh._element.remove(st)

    def fill(tf, paras, align, anchor, after=3):
        tf.word_wrap = True; tf.vertical_anchor = anchor
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align; p.space_after = Pt(after)
            for t, size, bold, color in runs:
                r = p.add_run(); r.text = t; f = r.font
                f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
                f.color.rgb = RGBColor.from_string(color)

    def box(x, y, w, h, bg, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=8,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=False, name=None):
        sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = min(0.5, 6 / min(w, h))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(bg)
        if line:
            sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(1 if dash else 0.75)
            if dash:
                sh.line.dash_style = MSO_LINE.DASH
        else:
            sh.line.fill.background()
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(3)
        fill(tf, paras, align, anchor)
        if name:
            sh.name = name
        return sh

    def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
        tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        fill(tf, paras, align, anchor)
        if name:
            tb.name = name
        return tb
    return box, text


def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def click(cid, inner):
    return (f'<p:par><p:cTn id="{cid}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{cid+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


def reveal_timing(s, groups):
    """Each group of shapes appears on its own click, in order."""
    cid, clicks = 3, ''
    for ids in groups:
        start, inner = cid, ''
        cid += 2
        for i, spid in enumerate(ids):
            inner += effect(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
        clicks += click(start, inner)
    bld = ''.join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for g in groups for i in g)
    NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + clicks +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
    old = s._element.find(qn('p:timing'))
    if old is not None:
        s._element.remove(old)
    t_el = etree.fromstring(tim)
    ext = s._element.find(qn('p:extLst'))
    (ext.addprevious(t_el) if ext is not None else s._element.append(t_el))



# ---------------- ETF waivers knowledge check (after "Two consecutive issue jobs: the process") ----------------
proc_i = next(i for i, sl in enumerate(prs.slides) if txt(sl.shapes[0]) == 'ETF WAIVERS' and txt(sl.shapes[1]).startswith('Two consecutive') and 'process' in txt(sl.shapes[1]))
ref = next(sl for sl in prs.slides if txt(sl.shapes[0]).endswith('KNOWLEDGE CHECK'))
kc = new_slide(ref, 2, proc_i, ['ETF WAIVERS  ·  KNOWLEDGE CHECK', 'Let’s See What You’ve Got!'])
box, text = kit(kc)
X, W = 32.4, 655.2
qa = [("C is upset, but nothing unfair happened. Do you waive the ETF to smooth things over?",
       "No. Waive only when charging would be genuinely unfair or hurt us more (legal risk, reputation, a likely dispute)."),
      ("C is moving abroad (a valid reason) but is fine paying the ETF. Do you offer the waiver?",
       "No. C is willing to pay, so there's nothing unfair to fix."),
      ("C threatens to complain to the Attorney General and wants to cancel. What do you do?",
       "Material legal threat: cancel the FC and waive the ETF right away. No retention attempt first."),
      ("The same CP moves the same job twice, each time with more than 24 hours' notice, and it isn't fulfilled. How many serious issues?",
       "One issue. With under 24 hours' notice each time, it would count as 2."),
      ("C mentions one bad job, and you spot a second consecutive issue job they didn't mention. Do you offer the waiver?",
       "Not yet. Probe openly and remove the self-cancel link. Never explain or offer the waiver until eligibility is confirmed."),
      ("Two consecutive serious issue jobs confirmed, and C is retainable. What next?",
       "Offer voucher/credits and free months, with the self-cancel link removed. If C rejects: waive the ETF and cancel the FC via CRM.")]
n = len(qa); gap = 5; h = (272 - gap * (n - 1)) / n
groups = []
for i, (q, a) in enumerate(qa):
    y = 100 + i * (h + gap)
    box(X, y, 364, h, WHITE, LINE)
    box(X + 10, y + (h - 20) / 2, 20, 20, TEAL_PALE, paras=[[(str(i + 1), 9, True, TEAL)]], align=PP_ALIGN.CENTER, margin=0,
        shape=MSO_SHAPE.OVAL)
    text(X + 40, y, 316, h, [[(q, 8.5, False, INK)]], anchor=MSO_ANCHOR.MIDDLE)
    b = box(X + 371, y, W - 371, h, TEAL_PALE, name=f'step{i + 1}Box')
    t = text(X + 381, y, W - 391, h, [[(a, 8.5, False, INK)]], anchor=MSO_ANCHOR.MIDDLE, name=f'step{i + 1}Ans')
    groups.append([b.shape_id, t.shape_id])
reveal_timing(kc, groups)
kc.notes_slide.notes_text_frame.text = (
    'Quick knowledge check: ETF waivers and two consecutive serious issue jobs. Ask each question and let trainees answer before revealing. Each click reveals the next answer.\n'
    "1) Don't waive just because someone is upset. The ETF is waived only when charging it would be genuinely unfair or would hurt us more than it helps (legal risk, reputation, a likely dispute).\n"
    "2) Assess intent: if C is willing to pay the ETF or finish the MCT, there's nothing unfair to fix, so don't offer the waiver even for a valid reason such as moving abroad.\n"
    '3) Material legal threats (Attorney General complaints, legal documentation, a lawyer CC\'d, class action, law firm contact, small claims) mean cancel the FC and waive the ETF immediately if that is what C asks, with no retention attempt first.\n'
    "4) Repeat cancellations from the same CP and job ID: more than 24 hours' notice, even repeatedly, counts as 1 issue; less than 24 hours' notice counts as 2. Different CPs failing the same job always count separately.\n"
    '5) Never explain or offer the waiver until you know C is eligible. Address what they raised, probe openly about the second issue, and remove the self-cancel link right away. If it was not a factor, restore the link and explain the ETF.\n'
    '6) Retainable: offer voucher or credits and free months with the self-cancel link removed. If they accept, apply the offers and book a new job; if they reject, waive the ETF and cancel the FC via CRM. Not retainable: waive and cancel.')
prs.save(sys.argv[2]); print('saved', sys.argv[2])
