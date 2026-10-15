"""Day 8 deck:
  - After slide 8 (Keep MCT knowledge check): add a "Let's practice!" slide: sample customer
    message (screenshot), context (4 paid months, last MCT charge Dec 15), 5-minute writing
    challenge, and a sample CS answer that stays hidden until click 1.
  - After the "HDR tickets 8 to 12" slide: add a knowledge check covering Assessing Dispute
    Likelihood and the HDR Ticket Handling Process (one answer revealed per click).

Usage: python3 add-practice-and-hdr-kc.py in.pptx out.pptx
"""
import copy, os, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets', 'cancel-after-mct-customer-message.png')

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''


def find(title, eyebrow_start=None):
    for i, sl in enumerate(prs.slides):
        if txt(sl.shapes[1]) == title and (eyebrow_start is None or txt(sl.shapes[0]).startswith(eyebrow_start)):
            return i, sl
    raise SystemExit('not found: ' + title)


def new_slide(ref, n_head, after_idx, texts):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    num = [x for x in ref.shapes if x.is_placeholder][0]
    for el in [ref.shapes[k]._element for k in range(n_head)] + [num._element]:
        s.shapes._spTree.append(copy.deepcopy(el))
    for k, t in enumerate(texts):
        r = s.shapes[k].text_frame.paragraphs[0].runs
        r[0].text = t
        for x in r[1:]:
            x._r.getparent().remove(x._r)
    lst = prs.slides._sldIdLst
    el = list(lst)[-1]; lst.remove(el); list(lst)[after_idx].addnext(el)
    return s


def kit(s):
    def strip(sh):
        st = sh._element.find(qn('p:style'))
        if st is not None:
            sh._element.remove(st)

    def fill(tf, paras, align, anchor, after=3):
        tf.word_wrap = True; tf.vertical_anchor = anchor
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align; p.space_after = Pt(after)
            for t, size, bold, color in runs:
                r = p.add_run(); r.text = t; f = r.font
                f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
                f.color.rgb = RGBColor.from_string(color)

    def box(x, y, w, h, bg, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=8,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=False, name=None):
        sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = min(0.5, 6 / min(w, h))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(bg)
        if line:
            sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(1 if dash else 0.75)
            if dash:
                sh.line.dash_style = MSO_LINE.DASH
        else:
            sh.line.fill.background()
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(3)
        fill(tf, paras, align, anchor)
        if name:
            sh.name = name
        return sh

    def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
        tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        fill(tf, paras, align, anchor)
        if name:
            tb.name = name
        return tb
    return box, text


def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def click(cid, inner):
    return (f'<p:par><p:cTn id="{cid}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{cid+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


def reveal_timing(s, groups):
    """Each group of shapes appears on its own click, in order."""
    cid, clicks = 3, ''
    for ids in groups:
        start, inner = cid, ''
        cid += 2
        for i, spid in enumerate(ids):
            inner += effect(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
        clicks += click(start, inner)
    bld = ''.join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for g in groups for i in g)
    NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + clicks +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
    old = s._element.find(qn('p:timing'))
    if old is not None:
        s._element.remove(old)
    t_el = etree.fromstring(tim)
    ext = s._element.find(qn('p:extLst'))
    (ext.addprevious(t_el) if ext is not None else s._element.append(t_el))


# ---------------- Let's practice (after Keep MCT knowledge check) ----------------
kc_i, kc8 = find('Let’s See What You’ve Got!', 'KEEP MCT')
_, ref7 = find('When the TR triggers')
pr = new_slide(ref7, 3, kc_i, ['KEEP MCT  ·  LET\'S PRACTICE', "Let's practice!",
                                '5-minute writing challenge: read C\'s message and draft your reply.'])
box, text = kit(pr)
X, W = 32.4, 655.2
LW = 380
text(X, 102, LW, 12, [[('CUSTOMER MESSAGE', 8, True, TEAL)]])
iw = LW - 16; ih = iw * 86 / 604
box(X, 117, LW, ih + 14, WHITE, LINE, margin=0)
pr.shapes.add_picture(IMG, Pt(X + 8), Pt(124), Pt(iw), Pt(ih))
cy = 117 + ih + 24
text(X, cy, LW, 12, [[('CONTEXT', 8, True, TEAL)]])
box(X, cy + 15, LW, 52, TEAL_PALE, paras=[
    [('•  C has ', 9.5, False, INK), ('4 paid months', 9.5, True, INK), ('.', 9.5, False, INK)],
    [('•  The last charge to complete the MCT is on ', 9.5, False, INK), ('Dec 15', 9.5, True, INK), ('.', 9.5, False, INK)]], margin=12)
ty = cy + 77
box(X, ty, LW, 60, AMBER_PALE, AMBER, paras=[
    [('⏱  5-minute writing challenge', 10.5, True, AMBER)],
    [('Write your response to C. When time is up, click to reveal a sample answer and compare.', 9, False, INK)]], margin=12)
RX, RW = X + LW + 14, W - LW - 14
box(RX, 102, RW, 270, WHITE, TEAL, dash=True, align=PP_ALIGN.CENTER,
    paras=[[('Sample CS answer', 11, True, TEAL)], [('Click to reveal after the 5 minutes are up.', 8.5, False, MUTED)]])
ans_bg = box(RX, 102, RW, 270, TEAL_PALE, TEAL, name='step1AnsBox')
ans = text(RX + 14, 112, RW - 28, 252, [
    [('SAMPLE CS ANSWER', 8, True, TEAL)],
    [('Great question! Happy to walk you through it.', 9, False, INK)],
    [("You've completed 4 membership payments so far, so you only have 2 more to go to fulfill the 6-month minimum commitment. "
      'Your 6th payment is scheduled for December 15. After that, you can cancel your membership anytime without having to pay an early termination fee.', 9, False, INK)],
    [("Since you're just planning ahead, may I ask what's making you consider canceling after the 6 months? Also, how has your experience with the "
      "cleanings been so far? I'd love to hear your feedback and see if there's anything I can do to help make your experience even better!", 9, False, INK)]],
    name='step1AnsText')
reveal_timing(pr, [[ans_bg.shape_id, ans.shape_id]])
pr.notes_slide.notes_text_frame.text = (
    "Let's practice: 5-minute writing challenge. Show C's message: \"When and how do I cancel membership without a charge. I want to cancel at end of current membership.\" "
    'Context: C has 4 paid months, and the last charge to complete the MCT is on Dec 15. Give trainees 5 minutes to write their reply, then click once to reveal the sample answer.\n'
    'What a good answer does: confirms how many payments are left (2 more to fulfill the 6-month minimum commitment), gives the date of the final payment (Dec 15), '
    "explains that after that they can cancel anytime with no ETF, and then probes kindly: why they're considering cancelling, and how their cleanings have been so far. "
    'It does not remind them about the ETF as a threat, does not offer to waive anything, and keeps the tone helpful.\n'
    "Sample answer: Great question, happy to walk you through it. You've completed 4 membership payments so far, so you only have 2 more to go to fulfill the 6-month "
    'minimum commitment. Your 6th payment is scheduled for December 15. After that, you can cancel your membership anytime without having to pay an early termination fee. '
    "Since you're just planning ahead, may I ask what's making you consider canceling after the 6 months? Also, how has your experience with the cleanings been so far? "
    "I'd love to hear your feedback and see if there's anything I can do to help make your experience even better!\n"
    'Remind trainees: after replying, set a TR for the day after the MF that fulfills the MCT is charged.')

# ---------------- Knowledge check (after HDR tickets 8 to 12) ----------------
hdr_i, _ = find('HDR tickets 8 to 12')
_, kc8 = find('Let’s See What You’ve Got!', 'KEEP MCT')
kc = new_slide(kc8, 2, hdr_i, ['DISPUTE LIKELIHOOD & HDR TICKETS  ·  KNOWLEDGE CHECK', 'Let’s See What You’ve Got!'])
box, text = kit(kc)
qa = [('C wants money back. Do you calculate a dispute likelihood score?',
       'No. It is a judgement call, not a formula: weigh the signs together.'),
      ('C has 3+ paid FC months and asks us to explain a charge. Does that raise or lower dispute likelihood?',
       'Lowers it. Both are signs that decrease dispute likelihood.'),
      ('HDR flag, $49 Mastercard charge, refund intent only. C insists again within 48 hours of the flag. What do you do?',
       '$49 is below the $58 Mastercard threshold: leave an internal note of your assessment and refund.'),
      ('Same case, but C comes back 3 days after the HDR flag. What changes?',
       'Past 48 hours: handle as usual under the Transformation Guidelines.'),
      ('Amex: $105 charge, C only asks for $20 back. Do you refund the $20?',
       'Assess the full $105. It is above the $50 Amex threshold, so no refund unless another guideline applies.'),
      ('Stripe Dispute Received, and the dispute response has NOT been submitted yet. What do you do?',
       "Don't resolve the ticket. Change the queue to Dispute Response Team.")]
n = len(qa); gap = 5; h = (272 - gap * (n - 1)) / n
for i, (q, a) in enumerate(qa):
    y = 100 + i * (h + gap)
    box(X, y, 364, h, WHITE, LINE)
    box(X + 10, y + (h - 20) / 2, 20, 20, TEAL_PALE, paras=[[(str(i + 1), 9, True, TEAL)]], align=PP_ALIGN.CENTER, margin=0,
        shape=MSO_SHAPE.OVAL)
    text(X + 40, y, 316, h, [[(q, 8.5, False, INK)]], anchor=MSO_ANCHOR.MIDDLE)
    b = box(X + 371, y, W - 371, h, TEAL_PALE, name=f'step{i + 1}Box')
    t = text(X + 381, y, W - 391, h, [[(a, 8.5, False, INK)]], anchor=MSO_ANCHOR.MIDDLE, name=f'step{i + 1}Ans')
    qa[i] = (q, a, [b.shape_id, t.shape_id])
reveal_timing(kc, [ids for _, _, ids in qa])
kc.notes_slide.notes_text_frame.text = (
    'Quick knowledge check: Assessing Dispute Likelihood and the HDR Ticket Handling Process. Ask each question and let trainees answer before revealing. '
    'Each click reveals the next answer.\n'
    '1) Dispute likelihood is a judgement call, not a formula: read the signs together rather than scoring them.\n'
    '2) 3+ paid FC months and wanting a charge explained are both signs that decrease dispute likelihood.\n'
    '3) Card-based threshold: Mastercard is under $58, so a $49 charge is below it. On the 2nd response within 48 hours of the HDR flag, leave an internal note of the card-based assessment and refund.\n'
    '4) After 48 hours the urgency window has closed: handle as usual under the Transformation Guidelines.\n'
    '5) For a partial refund, assess against the full original Stripe amount. $105 is above the $50 Amex threshold, so no refund unless another Transformation Guideline applies, '
    'such as a material legal threat or extreme financial hardship.\n'
    '6) HDR Ticket 7: check the CRM and Stripe. If the dispute response has not been submitted, do not resolve the ticket; change the queue to Dispute Response Team. If it has been submitted, handle BAU.')
prs.save(sys.argv[2]); print('saved', sys.argv[2])
