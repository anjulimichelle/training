"""Day 8 deck slide 28: replace the "HDR tickets 1-4" handling table with "HDR tickets at a
glance": which HDR ticket triggers follow the New HDR Ticket Handling Process (card-based
threshold) and which follow the Percentage-based Dispute Likelihood Assessment.
The old per-ticket handling text moves to the speaker notes.

Usage: python3 slide28-hdr-tickets-at-a-glance.py in.pptx out.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, AMBER, AMBER_PALE, INK, WHITE, ALT = '0B7A6F', '9A6410', 'FFF3D6', '1D1D1F', 'FFFFFF', 'F5F6F8'

prs = Presentation(sys.argv[1])
s = prs.slides[27]
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
assert txt(s.shapes[1]) == 'HDR tickets 1-4', txt(s.shapes[1])
old = next(sh for sh in s.shapes if sh.has_table)
old_rows = [[c.text for c in r.cells] for r in old.table.rows][1:]
old_notes = s.notes_slide.notes_text_frame.text
old._element.getparent().remove(old._element)


def set_text(shape, t):
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = t
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


set_text(s.shapes[1], 'HDR tickets at a glance')
set_text(s.shapes[2], 'Which process each HDR ticket follows.')

left = ["C's RefundRequest expired & received a courtesy voucher",
        'C just attempted, and we prevented them from updating the card on file to a prepaid card',
        'C just called our phone number listed on their Credit Card statement',
        'C just viewed their credit card statement, visited our cleaningbill.com website, and searched for the charge',
        '[Dispute Risk] Unrecognized phone number called our billing phone number',
        'Job charged more than 15 days ago',
        'C messaged HelpBot containing dispute risk',
        'Stripe Early Fraud Warning Received']
right = ['BBB Complaint Case received via Support Gmail',
         'System just prevented C from redeeming <DHJ voucher> as a 2nd voucher & C has a submitted or claimed job']
rows = [('New HDR Ticket Handling Process (card-based threshold)', 'Percentage-based Dispute Likelihood Assessment')]
rows += [(l, right[i] if i < len(right) else '') for i, l in enumerate(left)]

x, y, w = 32.4, 100, 655.2
heights = [28] + [30] * len(left)
gt = s.shapes.add_table(len(rows), 2, Pt(x), Pt(y), Pt(w), Pt(sum(heights)))
tbl = gt.table
tbl.columns[0].width = Pt(w / 2); tbl.columns[1].width = Pt(w / 2)
tblPr = tbl._tbl.tblPr
for a in ('bandRow', 'firstRow'):
    tblPr.set(a, '0')
sid = tblPr.find(qn('a:tableStyleId'))
if sid is not None:
    tblPr.remove(sid)
for i, row in enumerate(rows):
    tbl.rows[i].height = Pt(heights[i])
    for j, v in enumerate(row):
        c = tbl.cell(i, j)
        c.fill.solid()
        if i == 0:
            bg = TEAL if j == 0 else AMBER_PALE
        else:
            bg = WHITE if i % 2 else ALT
        c.fill.fore_color.rgb = RGBColor.from_string(bg)
        c.margin_left = c.margin_right = Pt(10); c.margin_top = c.margin_bottom = Pt(2)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        tcPr = c._tc.get_or_add_tcPr()
        for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
            ln = tcPr.find(qn(tag))
            if ln is not None:
                tcPr.remove(ln)
            ln = tcPr.makeelement(qn(tag), {'w': '9525' if tag in ('a:lnT', 'a:lnB') or (tag == 'a:lnR' and j == 0) else '0'})
            if ln.get('w') == '0':
                ln.append(ln.makeelement(qn('a:noFill'), {}))
            else:
                sf = ln.makeelement(qn('a:solidFill'), {}); sf.append(sf.makeelement(qn('a:srgbClr'), {'val': 'DCE1E7'})); ln.append(sf)
            tcPr.insert(0, ln) if False else tcPr.append(ln)
        fill_el = tcPr.find(qn('a:solidFill'))
        if fill_el is not None:
            tcPr.remove(fill_el); tcPr.append(fill_el)
        r = p.add_run(); r.text = v; f = r.font
        f.name = 'Play' if i == 0 else 'Arial'; f.bold = i == 0; f.size = Pt(9.5 if i == 0 else 8.5)
        f.color.rgb = RGBColor.from_string(WHITE if (i == 0 and j == 0) else (AMBER if i == 0 else INK))

s.notes_slide.notes_text_frame.text = (
    'HDR tickets at a glance: which process each HDR ticket trigger follows. '
    'New HDR Ticket Handling Process (card-based threshold): ' + '; '.join(left) + '. '
    'Percentage-based Dispute Likelihood Assessment: ' + '; '.join(right) + '.\n'
    'Handling detail from the previous version of this slide: '
    + ' '.join(f'{a}: {b}.' for a, b in old_rows) + '\n' + old_notes)
prs.save(sys.argv[2]); print('saved', sys.argv[2])
