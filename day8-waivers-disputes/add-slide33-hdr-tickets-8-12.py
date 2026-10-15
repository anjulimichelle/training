"""Day 8 deck: add slide 33 "HDR tickets 8 to 12" after HDR Ticket 7, one tile per ticket
(8 HelpBot dispute risk, 9 FC search with "refund", 10 Stripe Early Fraud Warning,
11 blocked 2nd DHJ voucher, 12 FC-cancel keywords). Same tile style as slide 30.

Usage: python3 add-slide33-hdr-tickets-8-12.py in.pptx out.pptx
"""
import copy, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
S = list(prs.slides)
ref = S[29]
assert txt(ref.shapes[1]) == 'HDR tickets 2 to 5', txt(ref.shapes[1])
assert txt(S[31].shapes[1]) == 'HDR Ticket 7'
num = [x for x in ref.shapes if x.is_placeholder][0]

s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
r = s.shapes[1].text_frame.paragraphs[0].runs
r[0].text = 'HDR tickets 8 to 12'
for x in r[1:]:
    x._r.getparent().remove(x._r)
lst = prs.slides._sldIdLst
el = list(lst)[-1]; lst.remove(el); list(lst)[31].addnext(el)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def shape(kind, x, y, w, h, fill, line=None, radius=6):
    sh = s.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    return sh


def write(x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, after=1.2):
    """paras: list of (level, runs); level 0 plain, 1 bullet, 2 sub-bullet."""
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (lvl, runs) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(after)
        if lvl:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int((7 + 10 * (lvl - 1)) * 12700 + 7 * 12700)))
            pPr.set('indent', str(-7 * 12700))
            pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
            pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '•◦'[lvl - 1]}))
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tb


SZ = 7
B = lambda t, c=INK, sz=SZ: (t, sz, True, c)
A = lambda t, c=INK, sz=SZ: (t, sz, False, c)


def tile(x, y, w, h, num, trigger, body, head_h=26):
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, WHITE, LINE)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, head_h, TEAL_PALE)
    badge = shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + 7, y + (head_h - 15) / 2, 54, 15, TEAL, radius=7.5)
    tf = badge.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = f'Ticket {num}'; r.font.size = Pt(7.5); r.font.bold = True; r.font.name = 'Play'
    r.font.color.rgb = RGBColor.from_string(WHITE)
    write(x + 68, y + 2, w - 76, head_h - 4, [(0, [B(trigger, TEAL, 7.5)])], anchor=MSO_ANCHOR.MIDDLE)
    write(x + 10, y + head_h + 5, w - 20, h - head_h - 8, body)


X, W, G = 32.4, 655.2, 6
LW = 330
RX, RW = X + LW + G, W - LW - G
Y0, YB = 100, 372

h8, h9, h11 = 40, 98, 54
y = Y0
tile(X, y, LW, h8, 8, 'C messaged HelpBot containing dispute risk',
     [(0, [A('Follow the '), B('New HDR Ticket Handling Process', AMBER), A('.')])])
y += h8 + G
tile(X, y, LW, h9, 9, 'FC C search contained the word "refund"', [
    (0, [A("This trigger is a Help Center search, not necessarily a message, so first check what the search actually looks like. "
           "Generally, we don't address a Help Center query unless it's related to a message C sent us.")]),
    (1, [B('Search reads like C trying to reach support', TEAL), A(' (not a generic keyword search, but phrased like an actual complaint): treat as BAU and follow the usual Transformation Guidelines.')]),
    (1, [B("Just a search, nothing else", TEAL), A(' showing C is reaching out: simply clear the ticket.')]),
    (2, [A('More info: '), B('Referencing Help Center Search Activity')])])
y += h9 + G
tile(X, y, LW, h11, 11, 'System just prevented C from redeeming <DHJ voucher> as a 2nd voucher & C has a submitted or claimed job', [
    (0, [A('Trained on the Percentage-Based Dispute Assessment: handle as usual. Otherwise, reassign to the '), B('Risk Reduction queue'), A('.')])],
    head_h=30)
y += h11 + G
tile(X, y, LW, YB - y, 12, 'When C cancelled FC, they mentioned the <below> keywords, with historical dispute rates', [
    (0, [A('If C wants a refund for the charges after cancelling FC, follow the usual '), B('Transformation Guidelines'), A('.')])],
    head_h=30)

# ticket 10 (right column)
tile(RX, Y0, RW, YB - Y0, 10, 'Stripe Early Fraud Warning Received', [
    (0, [A('Applies only to Early Fraud Warnings under the action items '), B('Customer.flag_as_dispute_risk'), A(' and '), B('Views.stripegeneral'), A('.')]),
    (0, [B('Check the payment network:', TEAL)]),
    (1, [B('Discover: '), A('trained on the Percentage-Based Dispute Assessment? Handle it directly. Otherwise, reassign to the Risk Reduction queue for now.')]),
    (1, [B('Visa, MasterCard, Amex: '), A('follow the steps below.')]),
    (0, [B('Step 1: How old is the charge?', AMBER), A('  A different clock from the 48-hour rule: this one counts from when the charge was made, not from the HDR flag.', MUTED)]),
    (1, [B('60 days or less: '), A('check the charge against the card-based threshold.')]),
    (2, [A('Above the threshold → leave an internal note, no refund.')]),
    (2, [A('Below the threshold → refund, and take care of any other pain points.')]),
    (1, [B('More than 60 days: '), A("don't refund just because of the Early Fraud Warning flag. The flag alone isn't enough once the charge has aged out.")]),
    (0, [B('Step 2: Is there anything else C wants refunded?', AMBER)]),
    (1, [B('Yes: '), A('follow '), B('C reaches out before or after disputing a charge'), A(' for the rest.')]),
    (1, [B('No: '), A("nothing further to assess. Don't look for other disputable charges to refund proactively.")]),
    (0, [B('One more thing: ', AMBER), A('if the flagged charge is a '), B('used DHJ voucher'), A(', follow that process instead of the steps above.')]),
    (0, [B('Example: ', MUTED), A('Early Fraud Warning on a $65 MasterCard charge made 40 days ago. Within 60 days, so check the $58 threshold: $65 is above it. Leave an internal note, no refund.', MUTED)])])

s.notes_slide.notes_text_frame.text = (
    'HDR tickets 8 to 12.\n'
    'Ticket 8 (C messaged HelpBot containing dispute risk): follow the New HDR Ticket Handling Process.\n'
    'Ticket 9 (FC C search contained the word "refund"): this trigger is a Help Center search, not necessarily a message, so first check what the search looks like. '
    "We generally don't address a Help Center query unless it relates to a message C sent us. If the search reads like C is trying to communicate with support "
    '(phrased like an actual complaint, not a generic keyword search), treat it as BAU and follow the usual Transformation Guidelines. If it is genuinely just a search with nothing else '
    'showing C is reaching out, clear the ticket. See Referencing Help Center Search Activity.\n'
    'Ticket 10 (Stripe Early Fraud Warning received): only for Early Fraud Warnings under the action items Customer.flag_as_dispute_risk and Views.stripegeneral. '
    'Check the payment network: Discover, handle it if trained on the Percentage-Based Dispute Assessment, otherwise reassign to Risk Reduction for now; Visa, MasterCard and Amex follow the steps. '
    "Step 1, how old is the charge? This is a different clock from the 48-hour rule: it counts from the charge date, not the HDR flag. 60 days or less: check against the card-based threshold; "
    'above, leave an internal note and no refund; below, refund and take care of other pain points. More than 60 days: do not refund just because of the Early Fraud Warning flag. '
    'Step 2, is there anything else C wants refunded? Yes: follow C reaches out before or after disputing a charge. No: nothing further to assess; do not hunt for other disputable charges. '
    'If the flagged charge is a used DHJ voucher, it is handled differently: follow that process instead. Example: a $65 MasterCard charge made 40 days ago is within 60 days and above $58, '
    'so leave an internal note and do not refund, even though the Early Fraud Warning flag is present.\n'
    'Ticket 11 (system prevented C from redeeming a DHJ voucher as a 2nd voucher, and C has a submitted or claimed job): if trained on the Percentage-Based Dispute Assessment, handle as usual; '
    'otherwise reassign to the Risk Reduction queue.\n'
    'Ticket 12 (when C cancelled FC they mentioned keywords with historical dispute rates): if C wants a refund for the charges after cancelling FC, follow the usual Transformation Guidelines.')
prs.save(sys.argv[2]); print('saved', sys.argv[2])
