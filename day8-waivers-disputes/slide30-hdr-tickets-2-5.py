"""Day 8 deck slide 30: replace the "HDR tickets 9-12" table with "HDR tickets 2 to 5",
one tile per ticket (2: blocked prepaid card swap, 3: called the card-statement number,
4: cleaningbill.com lookup, 5: BBB complaint). The old tickets 9-12 text moves to the notes.

Usage: python3 slide30-hdr-tickets-2-5.py in.pptx out.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
s = prs.slides[29]
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
assert txt(s.shapes[1]) == 'HDR tickets 9-12', txt(s.shapes[1])
old = next(sh for sh in s.shapes if sh.has_table)
old_rows = [[c.text for c in r.cells] for r in old.table.rows][1:]
old_notes = s.notes_slide.notes_text_frame.text
old._element.getparent().remove(old._element)


def set_text(shape, t):
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = t
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


set_text(s.shapes[1], 'HDR tickets 2 to 5')
set_text(s.shapes[2], 'How to handle each trigger.')


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def shape(kind, x, y, w, h, fill, line=None, radius=6):
    sh = s.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    return sh


def write(x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, after=1.5):
    """paras: list of (level, runs); level 0 plain, 1 bullet, 2 sub-bullet, 3 sub-sub-bullet."""
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (lvl, runs) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(after)
        if lvl:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int((8 + 10 * (lvl - 1)) * 12700 + 8 * 12700)))
            pPr.set('indent', str(-8 * 12700))
            pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
            pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '•◦▪'[lvl - 1]}))
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tb


SZ = 7.5
B = lambda t, c=INK, sz=SZ: (t, sz, True, c)
A = lambda t, c=INK, sz=SZ: (t, sz, False, c)


def tile(x, y, w, h, num, trigger, body, head_h=30):
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, WHITE, LINE)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, head_h, TEAL_PALE)
    badge = shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + 7, y + (head_h - 16) / 2, 52, 16, TEAL, radius=8)
    tf = badge.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = f'Ticket {num}'; r.font.size = Pt(7.5); r.font.bold = True; r.font.name = 'Play'
    r.font.color.rgb = RGBColor.from_string(WHITE)
    write(x + 66, y + 2, w - 74, head_h - 4, [(0, [B(trigger, TEAL, 8)])], anchor=MSO_ANCHOR.MIDDLE)
    write(x + 10, y + head_h + 6, w - 20, h - head_h - 10, body)


X, W, G = 32.4, 655.2, 6
LW = 340
RX, RW = X + LW + G, W - LW - G
Y0, YB = 100, 372

# left column: tickets 2, 4, 5
t2h = 52
tile(X, Y0, LW, t2h, 2, 'C just attempted, and we prevented them from updating the card on file to a prepaid card',
     [(0, [A('This follows the '), B('New HDR Ticket Handling Process', AMBER), A('.')])])
t4y, t4h = Y0 + t2h + G, 100
tile(X, t4y, LW, t4h, 4, 'C just viewed their credit card statement, visited our cleaningbill.com website, and searched for the charge', [
    (0, [A('Check if there are charges C wants refunded:')]),
    (1, [B('Yes: '), A('follow the '), B('New HDR Ticket Handling Process', AMBER), A('.')]),
    (1, [B('No: '), A("ask C for more information about their concern and proactively explain the charges on C's account.")]),
    (2, [B('Note: ', AMBER), A('if the charge is FC-related, explain the FC terms. Explaining the most recent charges lowers the risk of C disputing them.', MUTED)])])
t5y = t4y + t4h + G
tile(X, t5y, LW, YB - t5y, 5, 'BBB Complaint Case received via Support Gmail', [
    (0, [B('CS do:')]),
    (1, [A("C no longer expects a response (e.g. 60+ days since their last message): ignore the ticket.")]),
    (1, [A('Complaint was previously imported and already handled: check for other pain points. If none, clear the ticket.')]),
    (1, [A('Trained on the Percentage-Based Dispute Assessment: handle as usual. Otherwise, reassign to the Risk Reduction queue.')])])

# right column: ticket 3
tile(RX, Y0, RW, YB - Y0, 3, 'C just called our phone number listed on their Credit Card statement', [
    (0, [B('Step 1: Check for a voicemail. ', TEAL), A("Listen closely and note whatever C raised, even if it seems minor, so you don't ask them to repeat it later.")]),
    (0, [B('Step 2: Investigate. ', TEAL), A('Is there more to the story? Has C reached out about anything else recently?')]),
    (0, [B('Step 3: Is there a refund request?', TEAL)]),
    (1, [B('C wants a refund:')]),
    (2, [A('Business hours: try calling C and follow the '), B('New HDR Ticket Handling Process', AMBER), A('.')]),
    (2, [A("Outside business hours: email C and ask when they're available to discuss over the phone.")]),
    (1, [B('No clear concern or refund request:')]),
    (2, [A('Business hours: try calling to gather more details and resolve it over the phone.')]),
    (2, [A('Outside business hours:')]),
    (3, [A('email C asking for more information')]),
    (3, [A('proactively explain any recent charges on the account. FC-related charge: explain the membership terms. Clear explanations reduce the risk of disputes.')]),
    (3, [A("ask when they're available to discuss over the phone")]),
    (0, [B('Watch for: ', AMBER), A('C asking about an unused DHJ voucher? Follow the '), B('Unused DHJ Voucher Refund process'), A(' instead, not a generic refund.')])],
    head_h=30)

s.notes_slide.notes_text_frame.text = (
    'HDR tickets 2 to 5.\n'
    'Ticket 2 (C tried to update the card on file to a prepaid card and we prevented it): follows the New HDR Ticket Handling Process.\n'
    "Ticket 3 (C called the phone number on their credit card statement): Step 1, check for a voicemail; listen closely and jot down whatever C raised, even if minor, so you don't ask them to repeat it. "
    'Step 2, investigate: is there more to the story, and has C reached out about anything else recently? Step 3, is there a refund request? '
    "If C wants a refund: during business hours try calling and follow the New HDR Ticket Handling Process; outside business hours email and ask when they're available to discuss by phone. "
    'If there is no clear concern or refund request: during business hours call to gather details and resolve over the phone; outside business hours email asking for more information, '
    "proactively explain recent charges (explain the membership terms if the charge is ForeverClean-related; clear explanations reduce dispute risk), and ask when they're available to talk. "
    'Watch for C asking about an unused DHJ voucher: follow the Unused DHJ Voucher Refund process instead of treating it as a generic refund.\n'
    'Ticket 4 (C viewed their statement, visited cleaningbill.com and searched for the charge): check whether there are charges C wants refunded. Yes: New HDR Ticket Handling Process. '
    "No: ask C for more information about their concern and proactively explain the charges on the account; if the charge is FC-related, explain the FC terms, since explaining recent charges lowers the risk of a dispute.\n"
    'Ticket 5 (BBB complaint via Support Gmail): if C is no longer expecting a response (for example more than 60 days since their last message), ignore the ticket. '
    'If the complaint was already imported and handled, check for other pain points; if none, clear the ticket. If you are trained on the Percentage-Based Dispute Assessment, handle as usual; otherwise reassign to the Risk Reduction queue.\n'
    'Tickets 9-12 from the previous version of this slide: ' + ' '.join(f'{a}: {b}.' for a, b in old_rows) + '\n' + old_notes)
prs.save(sys.argv[2]); print('saved', sys.argv[2])
