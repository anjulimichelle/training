"""Day 8 deck, ETF Waivers:
  - Slide 43 "Two consecutive serious issue jobs": one-line overview, a table of what counts as a
    serious issue job, and what does NOT count (plus repeat-cancellation counting).
  - Slide 44 "Two consecutive: the process": steps condensed into a left column and the
    retention flow chart (retainable? -> offer / waive + cancel) on the right.

Usage: python3 slides43-44-two-consecutive.py in.pptx out.pptx
"""
import sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE, GREY, ALT = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', '6E6E73', 'F5F6F8'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''


def find(title):
    return next(sl for sl in prs.slides if txt(sl.shapes[1]) == title and txt(sl.shapes[0]) == 'ETF WAIVERS')


def clear(s):
    sl = list(s.shapes)
    keep = {id(x._element) for x in sl[:3]} | {id(x._element) for x in sl if x.is_placeholder}
    for sh in sl:
        if id(sh._element) not in keep:
            sh._element.getparent().remove(sh._element)


def set_text(shape, t):
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = t
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def kit(s):
    def strip(sh):
        st = sh._element.find(qn('p:style'))
        if st is not None:
            sh._element.remove(st)

    def fill(tf, paras, align, anchor, after=1.5):
        tf.word_wrap = True; tf.vertical_anchor = anchor
        for i, item in enumerate(paras):
            lvl, runs = item if isinstance(item, tuple) else (0, item)
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align; p.space_after = Pt(after)
            if lvl:
                pPr = p._p.get_or_add_pPr()
                pPr.set('marL', str(int(15 * 12700))); pPr.set('indent', str(-8 * 12700))
                pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
                pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '•'}))
            for t, size, bold, color in runs:
                r = p.add_run(); r.text = t; f = r.font
                f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
                f.color.rgb = RGBColor.from_string(color)

    def box(x, y, w, h, bg, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=7):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
        sh.adjustments[0] = min(0.5, 6 / min(w, h))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(bg)
        if line:
            sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.75)
        else:
            sh.line.fill.background()
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(2)
        fill(tf, paras, align, anchor)
        return sh

    def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        fill(tf, paras, align, anchor)
        return tb

    def line(x1, y1, x2, y2, head=False):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)); strip(c)
        c.line.color.rgb = RGBColor.from_string(GREY); c.line.width = Pt(1)
        if head:
            ln = c._element.find('.//' + qn('a:ln'))
            etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
    return box, text, line


X, W = 32.4, 655.2

# ======================= slide 43 =======================
s = find('Two consecutive serious issue jobs')
clear(s)
set_text(s.shapes[2], 'Two serious issue jobs in a row (not just the first two, types can differ) can qualify C for an ETF waiver.')
box, text, line = kit(s)
rows = [('False Invoice', "C was charged for a cleaning the CP didn't do: billed for nothing."),
        ('Poor Cleaning Quality', 'CP showed up, but the home was not cleaned to a reasonable standard (missed spots, still visibly dirty).'),
        ('System-cancelled (no CP claim)', 'No CP ever claimed the job. To C it feels exactly like a no-show.'),
        ('Cleaner Cancellation', "Counts only if the job wasn't fulfilled on the original date/time, with more or less than 24h notice. Another CP completes it on time: doesn't count."),
        ('CP no-show / CP late', "No-show: the cleaning didn't happen. Late: C planned around the time window and can lose their day."),
        ('Unexpected job time', "Counts only if the job doesn't happen and it's not because of C's own unavailability at the new time."),
        ('Unauthorized reschedule by CP', 'C agreed to a day and time; the CP moved it without asking.'),
        ('Trust & Safety incidents', 'Theft, damage, unsafe behavior, prohibited substances. Automatic qualifiers.'),
        ('Other CP-related issues', 'Rude or unprofessional behavior, or an undisclosed helper in the home.'),
        ('Overbilling / invoice issues', "Overcharged hours, false invoices, unauthorized added hours: paying for what didn't happen."),
        ('System/CS errors', "Our mistake (wrong cancellation, billing error, retention offer not honored). Counts only with C's complaint attached.")]
TW = 432
heights = [18] + [22.2] * len(rows)
gt = s.shapes.add_table(len(rows) + 1, 2, Pt(X), Pt(100), Pt(TW), Pt(sum(heights)))
tbl = gt.table
tbl.columns[0].width = Pt(118); tbl.columns[1].width = Pt(TW - 118)
tblPr = tbl._tbl.tblPr
for a in ('bandRow', 'firstRow'):
    tblPr.set(a, '0')
sid = tblPr.find(qn('a:tableStyleId'))
if sid is not None:
    tblPr.remove(sid)
for i, row in enumerate([('Counts as a serious issue', 'Why it breaks the promise')] + rows):
    tbl.rows[i].height = Pt(heights[i])
    for j, v in enumerate(row):
        c = tbl.cell(i, j)
        c.margin_left = c.margin_right = Pt(6); c.margin_top = c.margin_bottom = Pt(1)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tcPr = c._tc.get_or_add_tcPr()
        for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
            ln = tcPr.makeelement(qn(tag), {'w': '0' if tag in ('a:lnL', 'a:lnR') else '6350'})
            if ln.get('w') == '0':
                ln.append(ln.makeelement(qn('a:noFill'), {}))
            else:
                sf = ln.makeelement(qn('a:solidFill'), {}); sf.append(sf.makeelement(qn('a:srgbClr'), {'val': LINE})); ln.append(sf)
            tcPr.append(ln)
        c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string(TEAL if i == 0 else (WHITE if i % 2 else ALT))
        fe = tcPr.find(qn('a:solidFill')); tcPr.remove(fe); tcPr.append(fe)
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        r = p.add_run(); r.text = v; f = r.font
        head = i == 0 or j == 0
        f.name = 'Play' if head else 'Arial'; f.bold = head; f.size = Pt(8 if i == 0 else 7)
        f.color.rgb = RGBColor.from_string(WHITE if i == 0 else (TEAL if j == 0 else INK))

RX = X + TW + 10; RW = W - TW - 10
box(RX, 100, RW, 136, AMBER_PALE, AMBER, anchor=MSO_ANCHOR.TOP, margin=9, paras=[
    [('What does NOT count', 9.5, True, AMBER)],
    (1, [('A system-changed start time, by itself. ', 7.5, True, INK), ("If the job still happens, nothing was lost.", 7.5, False, INK)]),
    (1, [('Lockouts ', 7.5, True, INK), ('caused only by an alternate-start-time/RC scheduling change. C home but misses the job because of a system-adjusted schedule.', 7.5, False, INK)]),
    (1, [('Slow replies or poor communication ', 7.5, True, INK), ("from the CP. Annoying, but C still got their cleaning.", 7.5, False, INK)])])
box(RX, 242, RW, 130, TEAL_PALE, TEAL, anchor=MSO_ANCHOR.TOP, margin=9, paras=[
    [('Repeat cancellations (same CP, same job ID)', 9, True, TEAL)],
    (1, [('More than 24h notice, ', 7.5, True, INK), ('even more than once → counts as ', 7.5, False, INK), ('1', 7.5, True, INK), (' issue.', 7.5, False, INK)]),
    (1, [('Less than 24h notice ', 7.5, True, INK), ('→ counts as ', 7.5, False, INK), ('2 separate issues', 7.5, True, INK), ('.', 7.5, False, INK)]),
    (1, [('Different CPs ', 7.5, True, INK), ('failing the same job ID always count separately. CP 1 reschedules then cancels, CP 2 no-shows → 2 issues.', 7.5, False, INK)])])

s.notes_slide.notes_text_frame.text = (
    'Two consecutive serious issue jobs can qualify the customer for an ETF waiver. It is not limited to the first two jobs, and the issue types can differ.\n'
    'What counts as a serious issue job: False invoice: we charged for a cleaning that did not happen, so the customer was billed for nothing. '
    'Poor cleaning quality: the CP showed up but the home was not cleaned to a reasonable standard. System-cancelled (no CP claim): from the customer\'s side this feels exactly like a no-show. '
    'Cleaner cancellation: counts only if the job was not fulfilled on the original date and time; if another CP completes it on the original date it does not count; cancellations with less than '
    '24 hours notice and more than 24 hours notice both count if the job was not fulfilled. Example: CP cancels Aug 10 for an Aug 15 job; if another CP completes it Aug 15, it does not count; '
    'if no one does, it counts. CP no-show or late: a no-show means the cleaning did not happen; late can mess up the customer\'s whole day. '
    'Unexpected job time: if the job still happens, it does not count; if it does not happen because the customer was unavailable at the new time, it is not automatically our failure; '
    'if it does not happen and it is not because of the customer\'s own unavailability, it counts. Unauthorized reschedule by the CP: the agreement was broken without the customer\'s say. '
    'Trust and safety incidents (theft, damage, unsafe behavior, prohibited substances) are automatic qualifiers. Other CP-related issues: rude or unprofessional behavior, or an undisclosed helper. '
    'Overbilling or invoice issues: overcharged hours, false invoices and unauthorized added hours. System or CS errors: our own mistake, such as a wrong cancellation, a billing mistake, or not honoring a retention offer; '
    'it only counts with a customer complaint attached, because a mistake the customer never noticed is not a broken promise to them.\n'
    'What does not count: a system-changed start time by itself; lockouts caused only by an alternate start time or RC scheduling change; slow replies or poor communication from the CP.\n'
    'Repeat cancellations from the same CP and job ID: more than 24 hours notice, even more than once, counts as 1 issue; less than 24 hours notice counts as 2 separate issues. '
    'Different CPs failing the same job ID always count separately.')

# ======================= slide 44 =======================
s = find('Two consecutive: the process')
old_text = {txt(sh): sh for sh in s.shapes}
clear(s)
box, text, line = kit(s)
LW = 268
steps = [('Start here', 'Address what C raised. Then check the account for a second consecutive issue job C didn\'t mention.', TEAL, TEAL_PALE),
         ('Second issue unmentioned', 'Probe openly ("anything else that factored in?") and remove the self-cancel link now. Not a factor? Restore the link and explain the ETF.', TEAL, TEAL_PALE),
         ('Both confirmed', 'Address the new issue, then follow the flow chart: retention first, waive + cancel if C still wants out.', TEAL, TEAL_PALE),
         ('Good job after the two?', 'One good job: still eligible, best chance for real retention. Two or more: no longer eligible. A new issue then needs a fresh pair.', AMBER, AMBER_PALE)]
g = 6; sh_ = (272 - g * 3) / 4
for i, (head, body, col, pale) in enumerate(steps):
    y = 100 + i * (sh_ + g)
    box(X, y, LW, sh_, pale, col, anchor=MSO_ANCHOR.MIDDLE, margin=9,
        paras=[[(f'{i + 1}  ' if i < 3 else '', 8.5, True, col), (head, 8.5, True, col)], [(body, 7.5, False, INK)]])

# flow chart (right)
FX = X + LW + 14; FW = W - LW - 14
cx = FX + FW / 2
lx, rx = FX + 85, FX + FW - 85
B = lambda t, c=WHITE, sz=8: (t, sz, True, c)
A = lambda t, c=INK, sz=7.5: (t, sz, False, c)
box(cx - 115, 100, 230, 24, AMBER, paras=[[B('2 CONSECUTIVE SERIOUS ISSUE JOBS', WHITE, 8.5)]], align=PP_ALIGN.CENTER)
line(cx, 124, cx, 136, head=True)
box(cx - 80, 136, 160, 24, INK, paras=[[A('Is C still retainable?', WHITE, 8.5)]], align=PP_ALIGN.CENTER)
line(cx, 160, cx, 170)
line(lx, 170, rx, 170)
line(lx, 170, lx, 184, head=True); line(rx, 170, rx, 184, head=True)
text(lx + 4, 172, 30, 10, [[B('YES', INK, 7.5)]])
text(rx - 22, 172, 30, 10, [[B('NO', INK, 7.5)]])
box(lx - 82, 184, 164, 50, TEAL, paras=[[A('Offer voucher/credits and free months (if applicable), but ', WHITE),
                                          B('remove the self-cancel link', WHITE, 7.5)]], align=PP_ALIGN.CENTER)
box(rx - 82, 184, 164, 50, AMBER_PALE, AMBER, paras=[[B('Waive the ETF', AMBER, 7.5), A(', if applicable, and '), B('cancel the FC via CRM', INK, 7.5)]],
    align=PP_ALIGN.CENTER)
line(lx, 234, lx, 250, head=True)
box(lx - 65, 250, 130, 24, INK, paras=[[A('C rejected the offer?', WHITE, 8.5)]], align=PP_ALIGN.CENTER)
line(lx + 65, 262, rx, 262)
line(rx, 262, rx, 234, head=True)
text(lx + 80, 251, 30, 10, [[B('YES', INK, 7.5)]])
line(lx, 274, lx, 292, head=True)
text(lx + 4, 277, 30, 10, [[B('NO', INK, 7.5)]])
box(lx - 82, 292, 164, 36, TEAL_PALE, TEAL, paras=[[A('Apply the retention offers and '), B('book a new job', TEAL, 7.5), A('.')]],
    align=PP_ALIGN.CENTER)

nt = s.notes_slide.notes_text_frame
nt.text = nt.text + ('\nFlow chart: with two consecutive serious issue jobs confirmed, ask whether the customer is still retainable. '
                     'Yes: offer voucher or credits and free months if applicable, but remove the self-cancel link. If they reject the offer, waive the ETF '
                     'if applicable and cancel the FC via CRM; if they accept, apply the retention offers and book a new job. '
                     'Not retainable: waive the ETF if applicable and cancel the FC via CRM.')
prs.save(sys.argv[2]); print('saved', sys.argv[2])
