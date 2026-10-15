"""Day 8 deck slides 32-33: sample-ticket buttons.
  - Slide 32 (HDR Ticket 7): add a "View sample ticket" button beside the "Stripe Dispute Received"
    subtitle (pop-up: Stripe dispute internal note), and move the existing Dispute Response ticket
    button right beside the "Related ticket" heading. Clicks: 1 open Stripe dispute sample,
    2 close, 3 open Dispute Response sample, 4 close.
  - Slide 33 (HDR tickets 8 to 12): add a "View sample ticket" button in the Ticket 10 header
    (pop-up: Stripe Early Fraud Warning internal note). Click 1 opens, click 2 closes.

Usage: python3 slides32-33-sample-ticket-buttons.py in.pptx out.pptx
"""
import os, re, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, MUTED, WHITE = '0B7A6F', 'E3F4F1', '6E6E73', 'FFFFFF'
ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets')

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
s32, s33 = prs.slides[31], prs.slides[32]
assert txt(s32.shapes[1]) == 'HDR Ticket 7' and txt(s33.shapes[1]) == 'HDR tickets 8 to 12'


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def box(s, x, y, w, h, fill, line=None, runs=(), align=PP_ALIGN.CENTER, margin=4, lw=0.75):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    sh.adjustments[0] = min(0.5, 6 / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(1)
    p = tf.paragraphs[0]; p.alignment = align
    for t, size, bold, color in runs:
        r = p.add_run(); r.text = t; f = r.font
        f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold; f.color.rgb = RGBColor.from_string(color)
    return sh


def button(s, x, y, w=112, h=18):
    b = box(s, x, y, w, h, TEAL, runs=[('View sample ticket ▸', 8, True, WHITE)])
    b.name = 'SampleTicketButton'
    return b


def popup(s, img, img_size, title, subtitle, iw=580):
    ih = iw * img_size[1] / img_size[0]
    pw, ph = iw + 24, 30 + ih + 30
    px, py = (720 - pw) / 2, max(92, (405 - ph) / 2)
    card = box(s, px, py, pw, ph, WHITE, TEAL, lw=1.5)
    sp = card._element.find(qn('p:spPr'))
    eff = etree.SubElement(sp, qn('a:effectLst'))
    sd = etree.SubElement(eff, qn('a:outerShdw'), blurRad='76200', dist='25400', dir='5400000', algn='t', rotWithShape='0')
    clr = etree.SubElement(sd, qn('a:srgbClr'), val='000000'); etree.SubElement(clr, qn('a:alpha'), val='25000')
    head = box(s, px, py, pw, 26, TEAL, runs=[(title, 11, True, WHITE), ('    ' + subtitle, 8.5, False, TEAL_PALE)],
               align=PP_ALIGN.LEFT, margin=12)
    pic = s.shapes.add_picture(os.path.join(ASSETS, img), Pt(px + 12), Pt(py + 34), Pt(iw), Pt(ih))
    tb = s.shapes.add_textbox(Pt(px + 12), Pt(py + ph - 16), Pt(iw), Pt(10))
    tf = tb.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = 'Click to close'; r.font.size = Pt(7.5); r.font.name = 'Arial'
    r.font.color.rgb = RGBColor.from_string(MUTED)
    return [card.shape_id, head.shape_id, pic.shape_id, tb.shape_id]


def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def leave(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="1" presetClass="exit" presetSubtype="0" fill="hold" grpId="1" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>'
            f'</p:childTnLst></p:cTn></p:par>')


def click(cid, inner):
    return (f'<p:par><p:cTn id="{cid}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{cid+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


def set_timing(s, groups):
    """Each group of shape ids opens on one click and closes on the next, in order."""
    cid, clicks = 3, ''
    for ids in groups:
        start, inner = cid, ''
        cid += 2
        for i, spid in enumerate(ids):
            inner += effect(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
        clicks += click(start, inner)
        start, inner = cid, ''
        cid += 2
        for i, spid in enumerate(ids):
            inner += leave(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 2
        clicks += click(start, inner)
    bld = ''.join(f'<p:bldP spid="{i}" grpId="0"/><p:bldP spid="{i}" grpId="1"/>' for g in groups for i in g)
    NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + clicks +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
    old = s._element.find(qn('p:timing'))
    if old is not None:
        s._element.remove(old)
    t_el = etree.fromstring(tim)
    ext = s._element.find(qn('p:extLst'))
    (ext.addprevious(t_el) if ext is not None else s._element.append(t_el))


def add_note(s, line):
    nt = s.notes_slide.notes_text_frame
    if line not in nt.text:
        nt.text = nt.text + '\n' + line


# ---- slide 32 ----
old_t = s32._element.find(qn('p:timing'))
seen = []
for v in re.findall(r'spTgt spid="(\d+)"', etree.tostring(old_t).decode()):
    if int(v) not in seen:
        seen.append(int(v))
popup_b = seen
assert len(popup_b) == 4, popup_b
sub = s32.shapes[2]
assert txt(sub) == 'Stripe Dispute Received'
sub.width = Pt(170)
button(s32, 35 + 175, 78)
head = next(sh for sh in s32.shapes if txt(sh).startswith('Related ticket:'))
head.width = Pt(345)
old_btn = next(sh for sh in s32.shapes if txt(sh) == 'View sample ticket ▸' and sh.top > Pt(150))
old_btn.left, old_btn.top, old_btn.width, old_btn.height = Pt(32.4 + 340), Pt(217), Pt(112), Pt(18)
popup_a = popup(s32, 'stripe-dispute-received-sample.png', (832, 236), 'Sample ticket', 'Stripe Dispute Received (internal note)')
set_timing(s32, [popup_a, popup_b])
nt = s32.notes_slide.notes_text_frame
nt.text = nt.text.replace('TRAINER: click once to open the sample Dispute Response ticket, click again to close it.',
                          'TRAINER: click 1 opens the Stripe Dispute Received sample, click 2 closes it; click 3 opens the Dispute Response '
                          'ticket sample, click 4 closes it. The Stripe sample shows RDR auto-lost the dispute in Stripe and the system '
                          'estimates no agent action is needed.')

# ---- slide 33 ----
t10 = next(sh for sh in s33.shapes if txt(sh) == 'Stripe Early Fraud Warning Received')
hdr = next(sh for sh in s33.shapes if sh.has_text_frame and txt(sh) == '' and abs(sh.top - t10.top) < Pt(4)
           and sh.width > Pt(250) and sh.left < t10.left < sh.left + sh.width)
bw = 92
t10.width = Pt((hdr.left + hdr.width - t10.left) / 12700 - bw - 10)
b = button(s33, (hdr.left + hdr.width) / 12700 - bw - 6, (hdr.top / 12700) + 5, w=bw, h=16)
b.text_frame.margin_left = b.text_frame.margin_right = Pt(1)
for r in b.text_frame.paragraphs[0].runs:
    r.font.size = Pt(7)
popup_c = popup(s33, 'stripe-early-fraud-warning-sample.png', (819, 241), 'Sample ticket', 'Stripe Early Fraud Warning (internal note)')
set_timing(s33, [popup_c])
add_note(s33, 'TRAINER: click once to open the Stripe Early Fraud Warning sample ticket, click again to close it. In the sample the '
              'dispute was a refundable inquiry, has been fully refunded, and the system estimates no agent action is needed.')
prs.save(sys.argv[2]); print('saved', sys.argv[2])
