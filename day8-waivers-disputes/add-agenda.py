"""Day 8 Waivers, Disputes and Special Cases deck: add a clickable agenda slide after the cover (6 topics, two
columns). Each row jumps to its topic's title slide. Header shapes copied from an existing content slide."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, INK, SOFT, BORDER, WHITE = '0B7A6F', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF'
prs = Presentation(src)
S = list(prs.slides)
ref = S[2]
topics = [s for s in S if s.shapes[0].has_text_frame and s.shapes[0].text_frame.text.strip().upper().startswith('TOPIC 0')]
lines = {'Customer Agreed to Keep MCT': 'They\'ll finish the term, then cancel.',
         'Biohazard or Unsanitary Home': 'A CP reports the home as unsafe.',
         'Assessing Dispute Likelihood': 'Will C escalate to their bank?',
         'HDR Ticket Handling Process': 'Refund proactively, or handle normally?',
         'ETF Waivers': 'When we waive the ETF, and how.',
         'Internal Note Action Items': 'What each system action item means.'}
rows = [(t, t.shapes[1].text_frame.text.strip()) for t in topics]
assert [r[1] for r in rows] == list(lines), [r[1] for r in rows]

s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
num = [x for x in ref.shapes if x.is_placeholder][0]
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
for shape, t in ((s.shapes[0], 'OVERVIEW'), (s.shapes[1], 'Agenda'), (s.shapes[2], 'Six topics. Click a topic to jump to it.')):
    r = shape.text_frame.paragraphs[0].runs; r[0].text = t
    for x in r[1:]: x._r.getparent().remove(x._r)
lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); lst.insert(1, el)

def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
W, H, GAP = (9.1 - 0.15) / 2, 1.05, 0.15
for i, (target, title) in enumerate(rows):
    x0 = 0.45 + (i // 3) * (W + 0.15); y = Inches(1.45 + (i % 3) * (H + GAP))
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x0), y, Inches(W), Inches(H)); strip(card)
    card.adjustments[0] = 0.06; card.fill.solid(); card.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    card.line.color.rgb = RGBColor.from_string(BORDER); card.line.width = Pt(0.75)
    tf = card.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.78); tf.margin_right = Inches(0.45); tf.margin_top = tf.margin_bottom = 0
    for k, (t, font, size, bold, col) in enumerate(((title, 'Play', 13, True, INK), (lines[title], 'Arial', 10, False, SOFT))):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.LEFT; p.space_after = Pt(2)
        rr = p.add_run(); rr.text = t; rr.font.name = font; rr.font.size = Pt(size); rr.font.bold = bold; rr.font.color.rgb = RGBColor.from_string(col)
    for x, w, t, size in ((x0 + 0.1, 0.55, f'{i + 1:02d}', 16), (x0 + W - 0.42, 0.35, '›', 20)):
        b = s.shapes.add_textbox(Inches(x), y, Inches(w), Inches(H))
        tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        rr = p.add_run(); rr.text = t; rr.font.name = 'Play'; rr.font.size = Pt(size); rr.font.bold = True; rr.font.color.rgb = RGBColor.from_string(TEAL)
        b.click_action.target_slide = target
    card.click_action.target_slide = target

s.notes_slide.notes_text_frame.text = 'Agenda. ' + ' '.join(f'Topic {i + 1}: {t} ({lines[t]})' for i, (_, t) in enumerate(rows)) + ' In slideshow mode, click any topic to jump straight to it.'
prs.save(dst)
