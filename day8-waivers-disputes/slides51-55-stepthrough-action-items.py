"""Day 8 deck, Internal Note Action Items Groups A-E (slides 51-55): "flag list + Actions to Take"
layout. Tiles on the left show only the number, trigger title and Sender code; the panel on the
right shows the actions for the highlighted tile. Tile 1 is shown first; each click highlights the
next tile and swaps its actions into the panel (no extra slides). Data: assets/action-items-groups.json.

Usage: python3 slides51-55-stepthrough-action-items.py in.pptx out.pptx
"""
import json, math, os, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

C = {'TEAL': '0B7A6F', 'TEAL_PALE': 'E3F4F1', 'AMBER': '9A6410', 'AMBER_PALE': 'FFF3D6',
     'INK': '1D1D1F', 'MUTED': '6E6E73', 'WHITE': 'FFFFFF', 'LINE': 'DCE1E7', 'PANEL': 'F5F6F8'}
X, W, Y0, YB = 32.4, 655.2, 98, 372
LW = 290
RX = X + LW + 14; RW = W - LW - 14
DATA = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets', 'action-items-groups.json')))

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def shape(s, x, y, w, h, fill, line=None, radius=6, lw=0.75):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(C[fill])
    if line:
        sh.line.color.rgb = RGBColor.from_string(C[line]); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def fill_tf(tf, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, after=2):
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, (num, runs, size) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after); p.alignment = align
        if num:
            pPr = p._p.get_or_add_pPr()
            if num == '•':
                pPr.set('marL', str(int(24 * 12700))); pPr.set('indent', str(-9 * 12700))
                runs = [('• ', True, 'MUTED')] + runs
            else:
                pPr.set('marL', str(int(14 * 12700))); pPr.set('indent', str(-14 * 12700))
                runs = [(num + '  ', True, 'TEAL')] + runs
        for t, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(C[color])


def text(s, x, y, w, h, paras, **k):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, paras, **k)
    return tb


def step_paras(tile, sz):
    paras = []
    if tile['fires']:
        paras.append(('', [('Fires when: ', True, 'MUTED'), (tile['fires'], False, 'MUTED')], sz))
    n = 0
    for st in tile['steps']:
        if st and st[0] == 'sub':
            paras.append(('•', [tuple(r) for r in st[1]], sz))
        else:
            n += 1
            paras.append((str(n), [tuple(r) for r in st], sz))
    return paras


def fits(paras, width, height):
    total = 0
    for num, runs, sz in paras:
        chars = sum(len(r[0]) for r in runs) + 3
        ind = 24 if num == '•' else (14 if num else 0)
        cpl = max(10, (width - ind) / (sz * 0.46))
        total += math.ceil(chars / cpl) * sz * 1.2 + 3
    return total <= height


def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="300"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def leave(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="1" presetClass="exit" presetSubtype="0" fill="hold" grpId="1" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>'
            f'</p:childTnLst></p:cTn></p:par>')


def click(cid, inner):
    return (f'<p:par><p:cTn id="{cid}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{cid+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


def set_timing(s, states):
    """states[0] is visible at start; click k hides states[k-1] and shows states[k]."""
    cid, clicks = 3, ''
    for k in range(1, len(states)):
        start, inner = cid, ''
        cid += 2
        first = True
        for spid in states[k - 1]:
            inner += leave(cid, spid, 'clickEffect' if first else 'withEffect'); cid += 2; first = False
        for spid in states[k]:
            inner += effect(cid, spid, 'withEffect'); cid += 3
        clicks += click(start, inner)
    bld = ''
    for k, ids in enumerate(states):
        for i in ids:
            if k > 0:
                bld += f'<p:bldP spid="{i}" grpId="0"/>'
            if k < len(states) - 1:
                bld += f'<p:bldP spid="{i}" grpId="1"/>'
    old = s._element.find(qn('p:timing'))
    if old is not None:
        s._element.remove(old)
    if not clicks:
        return
    NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + clicks +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
    t_el = etree.fromstring(tim)
    ext = s._element.find(qn('p:extLst'))
    (ext.addprevious(t_el) if ext is not None else s._element.append(t_el))


def tile_paras(i, t, on, tsz):
    return [('', [(f'{i + 1}. ' + t['title'], True, 'TEAL' if on else 'INK')], tsz),
            ('', [(t['sender'], False, 'TEAL' if on else 'MUTED')], 6.5)]


for g in DATA:
    s = next(sl for sl in prs.slides if txt(sl.shapes[0]) == 'INTERNAL NOTE ACTION ITEMS' and txt(sl.shapes[1]) == g['title'])
    sl = list(s.shapes)
    keep = {id(x._element) for x in sl[:3]} | {id(x._element) for x in sl if x.is_placeholder}
    for sh in sl:
        if id(sh._element) not in keep:
            sh._element.getparent().remove(sh._element)
    p = s.shapes[2].text_frame.paragraphs[0]
    p.runs[0].text = 'Click to step through each trigger and its actions.'
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)

    n = len(g['tiles'])
    text(s, X, Y0, LW, 12, [('', [('TRIGGERS', True, 'MUTED')], 7.5)])
    ty0, gap = Y0 + 16, 5
    th = min(46, (YB - ty0 - gap * (n - 1)) / n)
    tsz = 10 if th >= 40 else 9
    # static base tiles
    geo = []
    for i, t in enumerate(g['tiles']):
        y = ty0 + i * (th + gap)
        geo.append(y)
        b = shape(s, X, y, LW, th, 'WHITE', 'LINE')
        tf = b.text_frame; tf.margin_left = Pt(10); tf.margin_right = Pt(8); tf.margin_top = tf.margin_bottom = Pt(1)
        fill_tf(tf, tile_paras(i, t, False, tsz), anchor=MSO_ANCHOR.MIDDLE, after=0)
    # static panel
    shape(s, RX, Y0, RW, YB - Y0, 'PANEL', 'LINE', radius=8)
    text(s, RX + 14, Y0 + 10, RW - 28, 18, [('', [('✓  ', True, 'TEAL'), ('Actions to Take', True, 'INK')], 12)])
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(RX + 14), Pt(Y0 + 34), Pt(RX + RW - 14), Pt(Y0 + 34)); strip(ln)
    ln.line.color.rgb = RGBColor.from_string(C['LINE']); ln.line.width = Pt(0.75)
    # one state per tile: highlighted tile + its actions
    states = []
    for i, t in enumerate(g['tiles']):
        ids = []
        hl = shape(s, X, geo[i], LW, th, 'TEAL_PALE', 'TEAL', lw=1.25)
        tf = hl.text_frame; tf.margin_left = Pt(10); tf.margin_right = Pt(8); tf.margin_top = tf.margin_bottom = Pt(1)
        fill_tf(tf, tile_paras(i, t, True, tsz), anchor=MSO_ANCHOR.MIDDLE, after=0)
        ids.append(hl.shape_id)
        note_h = 28 if t['note'] else 0
        ay = Y0 + 42
        ids.append(text(s, RX + 14, ay, RW - 28, 16, [('', [(t['title'], True, 'TEAL')], 10.5)]).shape_id)
        bw, bh = RW - 28, YB - ay - 22 - note_h - 8
        sz = 10.5
        while sz > 7 and not fits(step_paras(t, sz), bw, bh):
            sz -= 0.5
        ids.append(text(s, RX + 14, ay + 22, bw, bh, step_paras(t, sz), after=3).shape_id)
        if t['note']:
            nb = shape(s, RX + 10, YB - 34, RW - 20, 26, 'AMBER_PALE')
            nt = text(s, RX + 18, YB - 34, RW - 36, 26, [('', [('⚠ ', True, 'AMBER')] + [tuple(r) for r in t['note']], 8.5)],
                      anchor=MSO_ANCHOR.MIDDLE)
            ids += [nb.shape_id, nt.shape_id]
        states.append(ids)
    set_timing(s, states)
    nts = s.notes_slide.notes_text_frame
    if 'TRAINER: tile 1' not in nts.text:
        nts.text = ('TRAINER: tile 1 and its actions show first. Each click highlights the next tile and swaps its actions into the '
                    f'"Actions to Take" panel ({n} tiles, {n - 1} clicks); the click after the last tile moves to the next slide.\n') + nts.text
prs.save(sys.argv[2]); print('saved', sys.argv[2], len(prs.slides), 'slides')
