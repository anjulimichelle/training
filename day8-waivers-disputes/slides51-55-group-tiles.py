"""Day 8 deck, Internal Note Action Items: one slide per group (A-E), one tile per trigger,
each with numbered step-by-step handling (source: Care Knowledge Library, "Internal Note
Action Items - C-Side").
  51 Group A — Fraud & Identity Signals      (was "Group A: fraud and identity")
  52 Group B — Dispute Risk & Chargebacks    (was "Group A: matching past disputes")
  53 Group C — Jobs, Bookings & Scheduling   (was "Group B ... (1 of 2)")
  54 Group D — Comms & Account Flags         (was "Group B ... (2 of 2)")
  55 Group E — Membership & Reminders        (was "Group C: jobs, bookings and scheduling")
  "Group C and D: more triggers" is deleted (its triggers now sit in Groups C and D).

Usage: python3 slides51-55-group-tiles.py in.pptx out.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'
X, W, Y0, YB, G = 32.4, 655.2, 100, 372, 7

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''


def slide(title):
    return next(sl for sl in prs.slides if txt(sl.shapes[0]) == 'INTERNAL NOTE ACTION ITEMS' and txt(sl.shapes[1]) == title)


def reset(s, title, sub):
    sl = list(s.shapes)
    keep = {id(x._element) for x in sl[:3]} | {id(x._element) for x in sl if x.is_placeholder}
    for sh in sl:
        if id(sh._element) not in keep:
            sh._element.getparent().remove(sh._element)
    for k, t in ((1, title), (2, sub)):
        p = s.shapes[k].text_frame.paragraphs[0]
        p.runs[0].text = t
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    s.shapes[1].text_frame.paragraphs[0].runs[0].font.size = Pt(22)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


B = lambda t, c=INK: (t, True, c)
A = lambda t, c=INK: (t, False, c)


def tile(s, x, y, w, h, title, sender, fires, steps, note=None, sz=7.5, head_h=30):
    def shape(x, y, w, h, fill, line=None, radius=6):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
        sh.adjustments[0] = min(0.5, radius / min(w, h))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line:
            sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.75)
        else:
            sh.line.fill.background()
        return sh

    def write(x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, after=1.5):
        tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, (num, runs, size) in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(after)
            if num:
                pPr = p._p.get_or_add_pPr()
                ind = 11 if num != '•' else 20
                pPr.set('marL', str(int(ind * 12700))); pPr.set('indent', str(-9 * 12700 if num == '•' else -11 * 12700))
                runs = [(('•' if num == '•' else num) + ' ', True, MUTED if num == '•' else TEAL)] + runs
            for t, bold, color in runs:
                r = p.add_run(); r.text = t; f = r.font
                f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
                f.color.rgb = RGBColor.from_string(color)
        return tb

    shape(x, y, w, h, WHITE, LINE)
    shape(x, y, w, head_h, TEAL_PALE)
    write(x + 8, y + 2, w - 16, head_h - 4, [('', [B(title, TEAL)], 8.5), ('', [A(sender, MUTED)], 6.5)], anchor=MSO_ANCHOR.MIDDLE, after=0)
    paras = [('', [B('Fires when: ', MUTED), A(fires, MUTED)], sz - 0.5)] if fires else []
    n = 0
    for st in steps:
        if st and st[0] == 'sub':
            paras.append(('•', st[1], sz))
        else:
            n += 1
            paras.append((str(n), st, sz))
    body_h = h - head_h - 8 - (20 if note else 0)
    write(x + 8, y + head_h + 5, w - 16, body_h, paras)
    if note:
        shape(x + 5, y + h - 23, w - 10, 18, AMBER_PALE)
        write(x + 11, y + h - 23, w - 22, 18, [('', [B('⚠ ', AMBER)] + note, sz - 0.5)], anchor=MSO_ANCHOR.MIDDLE, after=0)


def row(s, y, h, items, widths=None, sz=7.5, head_h=30):
    n = len(items)
    widths = widths or [1] * n
    unit = (W - G * (n - 1)) / sum(widths)
    x = X
    for it, wt in zip(items, widths):
        w = unit * wt
        tile(s, x, y, w, h, *it, sz=sz, head_h=head_h)
        x += w + G


SUB = lambda *runs: ('sub', list(runs))
FLAG = 'admin_approved_low_dispute_risk_for_360_days'

# ---------------- Group A ----------------
sA = slide('Group A: fraud and identity')
reset(sA, 'Group A — Fraud & Identity Signals', 'Find the Sender code on the internal note, then follow its tile.')
h1 = 133; h2 = YB - Y0 - h1 - G
row(sA, Y0, h1, [
    ('Repeated address changes', 'api.customerlocation  ·  views.customerprofilelocation (OCW)', 'C updates their address more than once.', [
        [A("One change isn't suspicious (people move, run an Airbnb, book for others). "), B('Repeated'), A(' changes are worth a review.')],
        [A("Follow the instructions on the action item.")],
        [B('Fraud looks likely → '), A('follow '), B('Customer Fraud Handling'), A('.')]]),
    ('Shared address', 'utils.flag_multiple_customers_with_same_addresses', '2+ accounts share the same address.', [
        [A('Usually harmless (roommates, family), but can get around the '), B('one-voucher-per-household'), A(' limit.')],
        [A('Check for fraud, note what you found, and follow the action item.')],
        [A('Active FC on more than one account:')],
        SUB(B('High fraud risk → '), A('cancel FC on all accounts, add '), B('known_customer_fraud'), A('.')),
        SUB(B('No sign of fraud → '), A('add '), B('admin_ignore_address_fraud_flags'), A('.'))]),
    ('"Fake" in name or address', 'utils.flag_risky_customer_jobs', 'C has an upcoming job and "fake" in their name or address.', [
        [A('Another agent already confirmed it → note linking their note, clear the ticket.')],
        [A('Otherwise call: '), B('Primary CRM → Google Voice → Secondary CRM'), A(' (business hours).')],
        [B('Answers → '), A('confirm and update details, remind them about the cleaning.')],
        [B('No answer → '), A('details look real: send a reminder. Look fake: cancel the appointment and RC (correct reason code); remind them FC is still active.')]]),
], sz=7)
row(sA, Y0 + h1 + G, h2, [
    ('Matches a past disputer', 'utils.flag_upcoming_jobs_matching_past_dispute_signals', "C's details look like another customer's who disputed before.", [
        [A('Review the account for fraud.')],
        [B('Matches only on address + last name → '), A('household member, not fraud: add the 360-day flag, note both accounts, keep job and FC, clear.')],
        [B('Fraud likely → '), A('follow Fraud (knowingly intended to dispute and not pay).')],
        [B('Unlikely, FC terms already explained → '), A('cancel upcoming jobs and RC, waive ETF if it applies, refund and invalidate any unused DHJ voucher, tell C why.')],
        [B('Unlikely, not explained yet → '), A('call and explain the FC terms (or the voucher-to-credits conversion).')],
        SUB(B('Answers: '), A('decide on a voucher refund and/or cancel FC + waive ETF. Booking-only issue → 360-day flag, fix the service issue.')),
        SUB(B('No answer: '), A('add the 360-day flag, email the FC terms and a self-service link.'))]),
    ('Identity details keep changing', 'views.customerprofilepassword', 'Name, email or phone changed 3+ times in the last year.', [
        [A('Often normal (new phone, Airbnb, booking for someone else). The '), B('repetition'), A(' is what makes it worth checking.')],
        [A('Follow the instructions on the action item.')]]),
], widths=[2, 1], sz=7)
sA.notes_slide.notes_text_frame.text = (
    'Group A, fraud and identity signals. Each tile is one Sender code: what fires it, then the steps.\n'
    'Repeated address changes (api.customerlocation; views.customerprofilelocation is the OCW version): one change is not suspicious; repeated changes are. Follow the action item; if fraud looks likely, follow Customer Fraud Handling.\n'
    'Shared address: usually roommates or family, but also how someone gets around the one-voucher-per-household limit. Check for fraud and note it. Active FC on more than one account: high risk means cancel FC on all accounts and add known_customer_fraud; no sign of fraud means add admin_ignore_address_fraud_flags.\n'
    '"Fake" in name or address: if another agent already confirmed the account, note it and clear. Otherwise call Primary CRM, then Google Voice, then Secondary CRM. If they answer, confirm details and remind them of the cleaning. If not: real-looking details get a reminder; fake-looking ones get the appointment and RC cancelled, with a reminder that FC is still active.\n'
    'Matches a past disputer: review for fraud. Address and last name only means a household member: add ' + FLAG + ', note both accounts, keep the job and membership, and clear. '
    'Fraud likely: follow Fraud (knowingly intended to dispute and not pay). Unlikely and we already explained FC terms: cancel upcoming jobs and RC, waive the ETF if it applies, refund and invalidate any unused DHJ voucher, and tell them why (e.g. one voucher per household). '
    'Unlikely and not explained yet: call and explain FC terms or the credit conversion. If they answer, decide on a voucher refund or cancelling FC with an ETF waiver; unless they explicitly agreed, assume they may still dispute. Booking-only issue: add the flag and fix the service issue. No answer: add the flag and email with a self-service link.\n'
    'Identity details changed 3+ times in a year: often normal; follow the action item.')

# ---------------- Group B ----------------
sB = slide('Group A: matching past disputes')
reset(sB, 'Group B — Dispute Risk & Chargebacks', "Prevent the dispute, don't invite one.")
hh = (YB - Y0 - G) / 2
row(sB, Y0, hh, [
    ('Flagged as dispute risk', 'Customer.flag_as_dispute_risk', 'Something we did makes a dispute more likely.', [
        [A('Find out '), B('why'), A(' the account was flagged.')],
        [A('Follow the '), B('HDR Ticket Handling Process'), A('.')]]),
    ('FC charge, disputed before', 'ForeverClean.charge', 'C has disputed before and gets charged for FC again.', [
        [A('Run the '), B('dispute-risk assessment'), A('.')],
        [B('Above, C asked → '), A('cancel FC, waive ETF, refund the smaller of disputable charges or the ask. Tell C.')],
        [B('Above, no ask → '), A('cancel FC, waive ETF '), B('without contacting C'), A('.')],
        [B('Below → '), A('call and explain FC terms. No answer: email with the self-cancel link.')]]),
    ('FC charge, no first job', 'ForeverClean.charge  ·  no_first_job', 'MF charge, but C has no submitted, claimed or invoiced jobs.', [
        [A('Run the '), B('dispute-risk assessment'), A('.')],
        [B('Above → '), A('same steps as "disputed before".')],
        [B('Below → '), A("it's about getting C to book their first job. Call in business hours.")]]),
], sz=7)
row(sB, Y0 + hh + G, hh, [
    ('History of disputes, still active', 'utils.flag_dispute_risks', 'C has disputed before and still has FC, RC or an upcoming job.', [
        [B('Contacted us before/after the dispute → '), A('Customer Has Dispute History.')],
        [B('If not → '), A('review the disputed or declined charges (job issue? unaware of FC terms?).')],
        [B('Required: '), A('call to say FC is still active and explain the terms.')],
        [A('No answer or job-only issue → 360-day flag; email with the self-cancel link.')]]),
    ('Pre-disputed charge', 'views.cs_customerprospect', 'C pre-disputes a charge.', [
        [A('Job-related → check for fraud or CP fault; penalize the CP if needed.')],
        [A('Review other jobs on the action item; note findings.')],
        [A('Confirm RC and upcoming appointments are cancelled.')],
        [A('Credits removed? Restore them if needed. Note the risk on other charges.')]],
     [A("Don't refund on this flag alone. C reached out → Customer Has Dispute History.")]),
    ('Dispute or Early Fraud Warning', 'views.stripegeneral', 'C disputes a charge, or Stripe sends an EFW.', [
        [A('Check the internal note for the dispute '), B('type'), A('.')],
        [B('Needs a response → '), A('reassign to the '), B('Dispute Response Team'), A('.')],
        [B('HDR → '), A('follow the HDR Ticket Handling Process.')],
        [B('Otherwise → '), A('follow the action item and note findings.')]]),
], sz=7)
sB.notes_slide.notes_text_frame.text = (
    'Group B, dispute risk and chargebacks. The goal is to prevent the dispute, not invite one.\n'
    'Customer.flag_as_dispute_risk: find out why the account was flagged and follow the HDR Ticket Handling Process.\n'
    'ForeverClean.charge for a customer who disputed before (the highest-risk version): run the dispute-risk assessment. Above the threshold and C asked for a refund or cancellation: cancel FC, waive the ETF if it applies, '
    'refund the smaller of the disputable charges or the ask, and tell them. Above without an ask: cancel FC and waive the ETF without contacting C. Below: call and explain FC terms and charges; '
    'if they answer, decide on cancellations (if their only issue was the job, use the 360-day flag); if not, email with the self-cancellation link.\n'
    'ForeverClean.charge no_first_job: same assessment and above-threshold steps; below the threshold it is about getting them to book their first job, so call in business hours.\n'
    'utils.flag_dispute_risks: if they contacted us before or after the dispute, follow Customer Has Dispute History. If not, review the disputed or declined charges, then call (required) to say FC is still active and explain the terms. '
    'If they answer, decide on cancellations; job-only issue means the 360-day flag. If they do not answer, add the 360-day flag and email with the self-cancellation link.\n'
    'views.cs_customerprospect (pre-dispute): check a job-related charge for fraud or CP fault and penalize if needed, review other jobs on the action item, confirm RC and upcoming appointments are cancelled, restore removed credits if needed, '
    'and note the dispute risk of other charges. Do not refund on this system flag alone; if C also reached out, follow Customer Has Dispute History.\n'
    'views.stripegeneral: check the note for the type. Needs a response: reassign to the Dispute Response Team. HDR: follow the HDR process. Otherwise follow the action item and note findings.')

# ---------------- Group C ----------------
sC = slide('Group B: dispute risk and chargebacks (1 of 2)')
reset(sC, 'Group C — Jobs, Bookings & Scheduling', 'Find the Sender code on the internal note, then follow its tile.')
row(sC, Y0, hh, [
    ('Repeat lockout', 'cleaner_jobs.report_lockout', "CP reports a lockout; not C's first.", [
        [A('Check the '), B('lockout requirements'), A('. Refund from the CP dashboard only if the CP shouldn\'t be paid; tell the CP why.')],
        [A('Premium rate charged → refund it.')],
        [A('Note whether it was valid and what you did.')],
        [A('Message C only if no auto message, C reached out, or FC is active.')]]),
    ('Double-booked customer', 'utils.report_double_booked_clients', '2+ active requests close together (LMC/lockout risk).', [
        [A('Follow the action item. '), B('Call'), A(' in business hours.')],
        [B('Wants both → '), A('silence the flag 30 days (Add Flag > admin_approve_double_book).')],
        [B('Otherwise → '), A('cancel the duplicate(s).')],
        [A('Other issues raised → note, email follow-up signed as Emma.')]]),
    ('DHJ voucher, zip not covered', 'views.choosecp', 'C books with a Legacy DHJ voucher; zip not covered or no CPs.', [
        [A('Check coverage at <brand>.com/coverage.')],
        [A('Refund the voucher via CRM and make sure it\'s '), B('invalidated'), A('.')],
        [A('Cancel FC, RC and upcoming appointments.')]]),
    ('Frequent LMCs or lockouts', 'views.cpjob', 'Invoiced jobs too close together, or frequent LMCs/lockouts.', [
        [A('Check whether the LMC or lockout is '), B('valid'), A('.')],
        [A('Premium Fee charged → refund it.')],
        [A('Contact C only if not auto-notified, C reached out, or FC is active.')],
        [A('Note your findings.')]]),
], sz=6.8)
row(sC, Y0 + hh + G, hh, [
    ('1-star review, CP below 4.0', 'views.cpreviews', "1-star OCW review; CP's recent average is below 4.0 or unrated.", [
        [A('Check for signs of a '), B('false invoice'), A(' or wrong hours.')],
        [B("Don't refund proactively"), A(' without solid evidence the job wasn\'t done.')]]),
    ('Long pending invoice', 'utils.message_or_charge_long_pending_invoice_customers', 'A job stays in pending_invoice too long.', [
        [A('The longer it sits, the higher the risk of an '), B('accidental or false invoice'), A('.')],
        [A('Follow '), B('Long-Pending-Invoice Handling'), A('.')]]),
    ('Refund request flagged', 'views.customer_request_refund', 'Refund request suggests an unfinished job, or damaged/missing items.', [
        [A('Review the specific issue.')],
        [A('Follow the matching policy: '), B('OCH, False Invoice, PCQ'), A(' or related refund policy.')]]),
], sz=7.5)
sC.notes_slide.notes_text_frame.text = (
    'Group C, jobs, bookings and scheduling.\n'
    "cleaner_jobs.report_lockout (not the customer's first lockout): check whether the CP met the lockout requirements; refund from the CP's dashboard only if the CP should not be paid (fraud likely, no C/CP contact, or invalid lockout) and tell the CP why. "
    'Refund a Premium rate if charged, follow the action item, call if needed, note whether the lockout was valid, and message the customer only if they got no automatic message, reached out, or have active FC.\n'
    'utils.report_double_booked_clients: follow the action item and call within business hours. If C obviously wants both, silence the flag for 30 days (C CRM > Do > Add Flag > admin_approve_double_book); otherwise cancel the duplicates. '
    'Other issues raised on the call get an internal note and a follow-up email signed as Emma. CP Gen agents create a tier-3 action item, confirm by email under their own name, and reassign.\n'
    'views.choosecp: Legacy DHJ voucher in an uncovered zip: check coverage, refund and invalidate the voucher via CRM, and cancel FC, RC and upcoming appointments.\n'
    "views.cpjob: check whether the LMC or lockout is valid, refund a Premium Fee if charged, contact the customer only if not auto-notified, they reached out, or FC is active, and note findings. Refund from the CP dashboard only if the CP should not be paid.\n"
    'views.cpreviews: a 1-star review on a CP whose recent average is below 4.0 or unrated: check for a false invoice or wrong hours; do not refund proactively without solid evidence.\n'
    'utils.message_or_charge_long_pending_invoice_customers: follow Long-Pending-Invoice Handling. views.customer_request_refund: follow OCH, False Invoice, PCQ or the related refund policy.')

# ---------------- Group D ----------------
sD = slide('Group B: dispute risk and chargebacks (2 of 2)')
reset(sD, 'Group D — Comms & Account Flags', 'A block-comms flag left on can make C miss reminders and confirmations.')
row(sD, Y0, YB - Y0, [
    ('Block-comms flag on an active C', 'utils.flag_active_block_comms_users', 'An active customer has a block-comms flag.', [
        [A('Find out '), B('why'), A(' the flag is there; resolve any pending issue.')],
        [B('Nothing outstanding → remove the flag'), A(' (preferred): C CRM > Do > Deactivate > Remove block-comms.')],
        [A('Write an internal note explaining why you removed it.')]]),
    ('Upcoming job + block-comms', 'utils.flag_upcoming_jobs_customer_cp_unable_to_receive_sms', "C has an upcoming job but isn't getting notifications.", [
        [B('Flag left on by mistake → '), A('remove it and remind C about the appointment.')],
        [B('On purpose, FC active → '), A('remove the flag. Ask if C wants to keep the appointment:')],
        SUB(B('Yes: '), A('remind them about the appointment and FC.')),
        SUB(B('No: '), A('cancel it (correct reason code); remind them FC (and RC) stays active.')),
        [B('On purpose, FC not active → '), A('ask the same question.')],
        SUB(B('Keep: '), A('remove the flag and remind them.')),
        SUB(B("Don't keep: "), A('cancel the appointment (and RC). Remove the flag unless something requires it.'))]),
    ('Booked while block-comms active', 'views.customerbookagain', 'C books a new appointment while block-comms is active.', [
        [A("C never got a booking confirmation.")],
        [B('Remove the flag'), A('.')],
        [A('Send a message confirming or reminding them about the appointment.')]]),
], widths=[1, 1.35, 1], sz=10, head_h=38)
sD.notes_slide.notes_text_frame.text = (
    'Group D, comms and account flags.\n'
    'utils.flag_active_block_comms_users: an active customer has a block-comms flag. Find out why and resolve any pending issue. If nothing is outstanding, remove the flag (the preferred outcome): C CRM > Do > Deactivate > Remove block-comms, then write an internal note explaining why.\n'
    'utils.flag_upcoming_jobs_customer_cp_unable_to_receive_sms: the customer has an upcoming job and an active block-comms flag, so they are not getting appointment notifications. Was the flag left on purpose? '
    'If not (a previous agent likely forgot to remove it), remove it and remind them about the appointment. If yes and FC is active, remove the flag either way, then ask whether they want to keep the appointment: '
    'yes, remind them about the appointment and FC; no, cancel with the correct reason code and remind them FC (and RC) stays active. If FC is not active, ask the same question: keep it, remove the flag and remind them; '
    "don't, cancel the appointment (and RC) and remove the flag unless something else requires it.\n"
    'views.customerbookagain: the customer booked while block-comms is active, so they never got a booking confirmation. Remove the flag and send a message confirming or reminding them about the appointment.')

# ---------------- Group E ----------------
sE = slide('Group C: jobs, bookings and scheduling')
reset(sE, 'Group E — Membership & Reminders', 'Reminders you set, and High-LTV FC cancellations.')
row(sE, Y0, YB - Y0, [
    ('Timed reminder is due', 'TimedReminder (customer-side)', 'A timed reminder you (or a Project agent) set comes due.', [
        [A('Follow the '), B('instructions in the reminder'), A('.')],
        [B('About cancelling FC after MCT → '), A('see the Membership Retention Playbook (Scenario 8).')],
        [B('Created by a Project agent → '), A('check Project Agent Ticket Reassignment to see if it needs reassigning.')]]),
    ('High-LTV C cancels FC via OCW', 'views.customer_foreverclean', 'A High-LTV customer cancels FC through OCW. Goal: retention, only when it makes sense.', [
        [B('Reassign to FC Reactivations'), A(', except when:')],
        SUB(A('C also contacted us about something unrelated to reactivating')),
        SUB(A('C has disputed a charge before')),
        SUB(A('C is asking for a refund, before or after cancelling')),
        SUB(A('C has a damage_theft_reported flag')),
        SUB(A('there is a BBB, Yelp or social-media threat on file')),
        [A('The same exceptions apply to non-High-LTV customers (see FC Reactivations Routing).')],
        [B('Not reassigned → '), A('handle any other pending messages as usual.')],
        [B('C writes in to reactivate → '), A('AG – ForeverClean Reactivation → A.')]]),
], widths=[1, 1.3], sz=10.5, head_h=40)
sE.notes_slide.notes_text_frame.text = (
    'Group E, membership and reminders.\n'
    'TimedReminder (customer-side): a timed reminder you or a Project agent set earlier comes due. Follow the instructions in the reminder. If it is about the customer wanting to cancel FC after their MCT is complete, see the Membership Retention Playbook (Scenario 8). '
    'If a Project agent created it, check Project Agent Ticket Reassignment to see whether the ticket needs to be reassigned.\n'
    'views.customer_foreverclean: a High-LTV customer cancels FC through OCW. The goal is retention, but only when it makes sense. Reassign the ticket to FC Reactivations, except when the customer also contacted us about something unrelated to reactivating, '
    'has disputed a charge before, is asking to be refunded (before or after cancelling), has a damage_theft_reported flag, or has a BBB, Yelp or social-media threat on file. The same exceptions apply to customers who are not High-LTV (see FC Reactivations Routing). '
    'If you do not reassign, handle any other pending messages as usual. A customer writing in to reactivate: see AG - ForeverClean Reactivation, route A.')

# ---------------- delete "Group C and D: more triggers" ----------------
old = slide('Group C and D: more triggers')
lst = prs.slides._sldIdLst
for sld in list(lst):
    if prs.part.related_part(sld.get(qn('r:id'))) is old.part:
        rid = sld.get(qn('r:id')); lst.remove(sld); prs.part.drop_rel(rid)
        break
prs.save(sys.argv[2]); print('saved', sys.argv[2])
