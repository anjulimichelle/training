"""Day 8 deck: consolidate "Standard call attempt" and "Dispute-risk assessment and the 360-day flag"
into one "Shared steps" slide with two tiles (Standard Call Attempt, Dispute-Risk Assessment
(120-Day Window)) and the 360-day flag as a bottom bar; delete the second slide.

Usage: python3 slides50-51-shared-steps.py in.pptx out.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
S = list(prs.slides)
i50 = next(i for i, sl in enumerate(S) if txt(sl.shapes[1]) == 'Standard call attempt')
s, s51 = S[i50], S[i50 + 1]
assert txt(s51.shapes[1]) == 'Dispute-risk assessment and the 360-day flag'
notes = s.notes_slide.notes_text_frame.text + '\n' + s51.notes_slide.notes_text_frame.text
sl = list(s.shapes)
bar_k = next(i for i, sh in enumerate(sl) if txt(sh).startswith('If they answer'))
keep = {id(x._element) for x in sl[:3]} | {id(x._element) for x in sl if x.is_placeholder} | {id(x._element) for x in sl[bar_k - 2:bar_k + 1]}
for sh in sl:
    if id(sh._element) not in keep:
        sh._element.getparent().remove(sh._element)


def set_text(shape, runs):
    p = shape.text_frame.paragraphs[0]
    first = p.runs[0]
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    first.text = runs[0][0]
    for t, bold in runs[1:]:
        r = p.add_run(); r.text = t; r.font.bold = bold
        r.font.size = first.font.size; r.font.name = 'Play' if bold else first.font.name
        if first.font.color and first.font.color.type:
            r.font.color.rgb = first.font.color.rgb


set_text(s.shapes[1], [('Shared steps', None)])
set_text(s.shapes[2], [('A few steps repeat across many triggers. They are explained once here, and each trigger points back to them.', None)])

# reuse the old bar for the 360-day flag
bar_box, bar_icon, bar_txt = sl[bar_k - 2], sl[bar_k - 1], sl[bar_k]
dy = Pt(336) - bar_box.top
for sh in (bar_box, bar_icon, bar_txt):
    sh.top = sh.top + dy
p = bar_txt.text_frame.paragraphs[0]
runs = p.runs
runs[0].text = '360-day flag:  '
for r in runs[2:]:
    r._r.getparent().remove(r._r)
if len(runs) > 1:
    runs[1].text = ("C's only issue was a specific job? Add admin_approved_low_dispute_risk_for_360_days "
                    '(C CRM > Do > Charges), fix the service issue, note it, and follow up if needed.')
for r in p.runs:
    r.font.size = Pt(8.5)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def shape(x, y, w, h, fill, line=None, radius=6):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    return sh


def write(x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, after=2.5):
    """paras: list of (kind, runs); kind '' plain, '•' bullet, or a number string for a numbered step."""
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (kind, runs) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after)
        if kind:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int(13 * 12700))); pPr.set('indent', str(-13 * 12700 if kind != '•' else -9 * 12700))
            if kind == '•':
                pPr.set('marL', str(int(10 * 12700)))
                pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
                pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '•'}))
            else:
                runs = [(kind + '.  ', runs[0][1], True, TEAL)] + runs
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tb


SZ = 8
B = lambda t, c=INK: (t, SZ, True, c)
A = lambda t, c=INK: (t, SZ, False, c)
X, W, G = 32.4, 655.2, 12
TW = (W - G) / 2
Y0, TH, HH = 100, 228, 26

for n, (title, sub) in enumerate((('Standard Call Attempt', 'Use whenever a trigger tells you to call the customer.'),
                                  ('Dispute-Risk Assessment (120-Day Window)', 'Use whenever a trigger tells you to assess dispute risk.'))):
    x = X + n * (TW + G)
    shape(x, Y0, TW, TH, WHITE, LINE)
    shape(x, Y0, TW, HH, TEAL_PALE)
    badge = shape(x + 8, Y0 + 5, 16, 16, TEAL, radius=8)
    tf = badge.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    r = pp.add_run(); r.text = str(n + 1); r.font.size = Pt(8.5); r.font.bold = True; r.font.name = 'Play'
    r.font.color.rgb = RGBColor.from_string(WHITE)
    write(x + 30, Y0 + 2, TW - 38, HH - 4, [('', [(title, 9.5, True, TEAL)])], anchor=MSO_ANCHOR.MIDDLE)
    write(x + 12, Y0 + HH + 7, TW - 24, 14, [('', [(sub, 7.5, False, MUTED)])])

L = X + 12; LW = TW - 24
write(L, Y0 + HH + 24, LW, TH - HH - 30, [
    ('•', [A('Call only during business hours: '), B('8:00 AM to 8:00 PM in the user\'s local time zone.')]),
    ('•', [A('Ticket comes in '), B('less than 10 minutes before 8:00 AM'), A(': park it and call once business hours start.')]),
    ('•', [A('Ticket comes in '), B('more than 10 minutes before 8:00 AM, or after 8:00 PM'), A(': don\'t wait. Treat it as "customer did not answer".')]),
    ('•', [B('User answers: ', TEAL), A("follow the specific instructions in that trigger's section.")]),
    ('•', [B('User does not answer: ', AMBER), A('in most cases, email them instead. Explain the relevant terms or charges and include a link they can use on their own, if applicable (e.g. a self-cancellation or self-refund link).')])])

R = X + TW + G + 12
write(R, Y0 + HH + 24, LW, 40, [
    ('1', [A("Look at C's FC-related charges from the "), B('last 120 days'), A('.')]),
    ('2', [A('Write an internal note with your assessment, even if the risk is low:')])])
fmt = shape(R + 13, Y0 + HH + 62, LW - 13, 50, AMBER_PALE)
write(R + 22, Y0 + HH + 66, LW - 30, 44, [('', [A(t, AMBER)]) for t in
                                          ('Disputable amount:', 'Dispute threshold:', 'Dispute likelihood:', 'Action:')], after=0.5)
write(R, Y0 + HH + 118, LW, 70, [
    ('3', [B('Above the threshold: '), A("don't wait for the dispute. The trigger's steps walk you through cancelling FC (and waiving the ETF if it applies), and refunding too if C already asked.")]),
    ('4', [B('Below the threshold: '), A('usually a chance to explain the charge, not a reason to refund.')]),
    ('', [A('Full policy: ', MUTED), B('Dispute Risk Assessment', TEAL)])])

s.notes_slide.notes_text_frame.text = (
    'Shared steps: a few steps repeat across many triggers, so they are explained once here and each trigger points back to them.\n' + notes)

lst = prs.slides._sldIdLst
sld = lst[i50 + 1]
rid = sld.get(qn('r:id'))
lst.remove(sld)
prs.part.drop_rel(rid)
prs.save(sys.argv[2]); print('saved', sys.argv[2], 'slide', i50 + 1)
