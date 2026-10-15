"""Day 8 deck: add a slide after "Two consecutive issue jobs: the process":
  - Special case: a good cleaning happened after the two issues (1 good job / 2+ good jobs /
    new issue after eligibility reset).
  - Comms guidance for confirmed 2 consecutive issue jobs (retention-focused, no ETF/MCT mention).
  - Suggested macros (titles).

Usage: python3 add-two-consecutive-special-case.py in.pptx out.pptx
"""
import copy, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
idx, ref = next((i, sl) for i, sl in enumerate(prs.slides)
                if txt(sl.shapes[0]) == 'ETF WAIVERS' and txt(sl.shapes[1]).startswith('Two consecutive') and 'process' in txt(sl.shapes[1]))
num = [x for x in ref.shapes if x.is_placeholder][0]
s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
for k, t in ((1, 'Special case and comms guidance'),
             (2, 'A good cleaning happened after the two issues, and C now cites them to cancel.')):
    r = s.shapes[k].text_frame.paragraphs[0].runs
    r[0].text = t
    for x in r[1:]:
        x._r.getparent().remove(x._r)
lst = prs.slides._sldIdLst
el = list(lst)[-1]; lst.remove(el); list(lst)[idx].addnext(el)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def box(x, y, w, h, bg, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=9, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, 6 / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(bg)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(3)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(2)
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return sh


B = lambda t, c=INK, sz=8: (t, sz, True, c)
A = lambda t, c=INK, sz=8: (t, sz, False, c)
X, W = 32.4, 655.2
LW = 400
box(X, 100, LW, 18, WHITE, None, [[B('SPECIAL CASE: HOW MANY GOOD JOBS HAPPENED AFTER THE TWO ISSUES?', TEAL, 8)]], margin=0).fill.background()
cases = [
    ('1', 'One good job after → still eligible', TEAL, TEAL_PALE,
     [A("One good experience isn't enough to prove trust is rebuilt. This is your best chance for a "), B('real'), A(' retention attempt: '
       'acknowledge the two issues, point to how well the recent job went, and offer to keep them with that CP or a similar one. '
       'Keep the self-cancel link removed meanwhile. Declines → cancel and waive the ETF.')]),
    ('2+', 'Two or more good jobs after → no longer eligible', AMBER, AMBER_PALE,
     [A("Treat it as a standard cancellation: probe for what's driving it now and attempt retention normally. "
        'Still wants out → the ETF applies as usual (unless another waiver reason genuinely exists).')]),
    ('!', 'New issue after eligibility reset (2+ good jobs)', INK, 'F5F6F8',
     [A("That one bad job alone doesn't revive anything. It takes a fresh pair of consecutive issues, same as a brand-new customer, "
        'before this rule applies again.')])]
g = 6; ch = (372 - 122 - g * 2) / 3
for i, (badge, head, col, pale, body) in enumerate(cases):
    y = 122 + i * (ch + g)
    box(X, y, LW, ch, pale, col if col != INK else LINE, anchor=MSO_ANCHOR.MIDDLE, margin=46,
        paras=[[B(head, col, 9)], body])
    box(X + 10, y + (ch - 26) / 2, 26, 26, col, paras=[[B(badge, WHITE, 9 if len(badge) < 2 else 8)]], align=PP_ALIGN.CENTER,
        margin=0, shape=MSO_SHAPE.OVAL)

RX, RW = X + LW + 12, W - LW - 12
box(RX, 100, RW, 116, TEAL, None, anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('COMMS GUIDANCE', TEAL_PALE, 8)],
    [B('Confirmed 2 consecutive serious issue jobs', WHITE, 10)],
    [A('Keep it retention-focused. ', WHITE, 8.5), B('No need to mention the ETF or MCT.', WHITE, 8.5)],
    [A('Acknowledge both issues, show what you will do to fix it, and lead with the retention offer.', TEAL_PALE, 8)]])
box(RX, 222, RW, 18, WHITE, None, [[B('SUGGESTED MACROS', TEAL, 8)]], margin=0).fill.background()
macros = ['ETF waiver: 2 Consecutive Serious Issue Jobs – Probe Before Confirming (Only One Issue Raised)',
          'ETF waiver: 2 consecutive serious issue jobs']
mh = (372 - 244 - g) / 2
for i, m in enumerate(macros):
    y = 244 + i * (mh + g)
    box(RX, y, RW, mh, WHITE, TEAL, anchor=MSO_ANCHOR.MIDDLE, margin=12,
        paras=[[B('▸ Suggested macro', TEAL, 7.5)], [B(m, INK, 8.5)]])

s.notes_slide.notes_text_frame.text = (
    'Special case: a good cleaning happened after the two issues, and the customer is now citing the original two to cancel. Time and a good experience have passed, '
    'so this plays out differently. How many good jobs happened after the two issues changes the right move.\n'
    "One good job after: still eligible. One good experience isn't enough to prove trust is genuinely rebuilt, and this is your best opportunity for a real retention attempt, "
    'because you have something concrete to anchor on. Acknowledge the two issues, point to how well the recent job went, and offer to keep them with that CP or a similar one. '
    'Keep the self-cancel link removed while this is in progress. If they decline, cancel and waive the ETF; they are still eligible.\n'
    "Two or more good jobs after: no longer eligible under this rule. Treat it as a standard cancellation: probe for what's actually driving it now, attempt retention normally, "
    'and if they still want out, the ETF applies as usual (though another waiver reason could still apply if one genuinely exists).\n'
    "A new issue after eligibility has already reset (2+ good jobs): that one bad job alone doesn't revive anything. It takes a fresh pair of consecutive issues, "
    'same as a brand-new customer, before this rule applies again.\n'
    'Comms guidance for confirmed two consecutive serious issue jobs: keep the message retention-focused. There is no need to mention the ETF or MCT.\n'
    'Suggested macros: "ETF waiver: 2 Consecutive Serious Issue Jobs – Probe Before Confirming (Only One Issue Raised)" when C raised only one of the issues, '
    'and "ETF waiver: 2 consecutive serious issue jobs" once both are confirmed.')
prs.save(sys.argv[2]); print('saved', sys.argv[2])
