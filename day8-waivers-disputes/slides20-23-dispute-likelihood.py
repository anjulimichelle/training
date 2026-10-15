"""Day 8 deck (1YxuIWEo...): Assessing Dispute Likelihood.
  - Slide 20 "What is this?": brief overview (judgement call, not a formula) plus the
    process from slide 21 (Investigate, Weigh, Pull the threshold + trainer demo, Judge,
    Write it down). The "Don't mix thresholds" tile becomes a footnote. "Write it down"
    has a button: click 1 opens the internal note format pop-up, click 2 closes it.
  - Slide 21 "The process" is deleted.
  - The two signs slides list every sign: 10 that increase, 7 that decrease.

Usage: python3 slides20-23-dispute-likelihood.py in.pptx out.pptx
"""
import sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE, FIELD = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', 'F5F6F8'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
S = list(prs.slides)
s20, s21, s22, s23 = S[19:23]
assert txt(s20.shapes[1]) == 'What is this?'
assert txt(s21.shapes[1]) == 'The process'
assert txt(s22.shapes[1]).startswith('Signs that increase')
assert txt(s23.shapes[1]).startswith('Signs that decrease')


def clear(s, n_keep):
    sl = list(s.shapes)
    keep = {id(x._element) for x in sl[:n_keep]} | {id(x._element) for x in sl if x.is_placeholder}
    for sh in sl:
        if id(sh._element) not in keep:
            sh._element.getparent().remove(sh._element)


def set_runs(shape, paras):
    """Replace a header shape's text, keeping its first run's formatting."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    proto = p0.runs[0]._r
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    first = True
    for i, runs in enumerate(paras):
        p = p0 if i == 0 else tf.add_paragraph()
        for t, size, bold, color in runs:
            if first:
                r = p0.runs[0]; first = False
            else:
                r = p.add_run()
            r.text = t; f = r.font
            if size: f.size = Pt(size)
            if bold is not None: f.bold = bold; f.name = 'Play' if bold else 'Arial'
            if color: f.color.rgb = RGBColor.from_string(color)


def kit(s):
    def strip(sh):
        st = sh._element.find(qn('p:style'))
        if st is not None:
            sh._element.remove(st)

    def fill_text(tf, paras, align, anchor, after=1.5):
        tf.word_wrap = True; tf.vertical_anchor = anchor
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align; p.space_after = Pt(after)
            for t, size, bold, color in runs:
                r = p.add_run(); r.text = t; f = r.font
                f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
                f.color.rgb = RGBColor.from_string(color)

    def box(x, y, w, h, fill, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=6, margin=6, lw=0.75):
        sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = min(0.5, radius / min(w, h))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line:
            sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
        else:
            sh.line.fill.background()
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(1)
        fill_text(tf, paras, align, anchor)
        return sh

    def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=1.5):
        tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        fill_text(tf, paras, align, anchor, after)
        return tb

    def dot(x, y, d, label, fill, color=WHITE, size=8.5):
        return box(x, y, d, d, fill, paras=[[(label, size, True, color)]], align=PP_ALIGN.CENTER,
                   shape=MSO_SHAPE.OVAL, margin=0)
    return box, text, dot


def set_notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ---------------- slide 20: overview + process ----------------
clear(s20, 3)
set_runs(s20.shapes[2], [[('When a customer wants money back: how likely are they to file a dispute with their bank?', 11, False, MUTED)],
                        [("It's a judgement call, not a formula.", 11, True, TEAL)]])
s20.shapes[2].height = Pt(30)
box, text, dot = kit(s20)
L, W = 32.4, 655.2

# three context tiles
tiles = [('What it is', 'Beyond "is the ask valid", you read how likely C is to dispute with their bank. It is one input into your resolution path.', TEAL, TEAL_PALE),
         ('Under-reading', 'You refuse to budge, they dispute. We lose the amount, pay a chargeback fee, and likely lose the customer.', AMBER, AMBER_PALE),
         ('Over-reading', 'Refunding reflexively gives away money and teaches the account that asking loudly gets results.', AMBER, AMBER_PALE)]
tw = (W - 2 * 8) / 3
for i, (head, body, col, fill) in enumerate(tiles):
    x = L + i * (tw + 8)
    box(x, 110, tw, 50, fill, paras=[[(head, 9.5, True, col)], [(body, 8, False, INK)]], margin=8)
box(L, 165, W, 18, WHITE, TEAL, paras=[[('Goal:  ', 8.5, True, TEAL),
                                        ('an accurate read, not a cautious one. Only charges within 120 days of the charge date count as disputable.', 8.5, False, INK)]],
    margin=8)

# process
text(L, 190, 200, 11, [[('THE PROCESS', 8, True, TEAL)]])
steps = [('Investigate', "Pull job history, comms and CS/CP notes. Don't stop at the ticket in front of you."),
         ("Weigh, don't score", 'One sign rarely tells the story. Read the signs together, like a retainability read.'),
         ('Pull the threshold', 'CRM > Agent Navigation Bar > Refund if Above Dispute Threshold. Use the live table per card network.'),
         ('Judge', "Does your read clear that network's threshold? Clears: resolve proactively. Doesn't: the standard offer is enough."),
         ('Write it down', 'Leave an internal note of your assessment either way.')]
cw, gap, cy, ch = (W - 4 * 7) / 5, 7, 204, 151
for i, (head, body) in enumerate(steps):
    x = L + i * (cw + gap)
    box(x, cy, cw, ch, TEAL_PALE)
    dot(x + 8, cy + 8, 17, str(i + 1), TEAL)
    text(x + 30, cy + 8, cw - 36, 17, [[(head, 9.5, True, TEAL)]], anchor=MSO_ANCHOR.MIDDLE)
    text(x + 8, cy + 32, cw - 16, 70, [[(body, 8, False, INK)]])
    if i == 0:
        box(x + 6, cy + ch - 42, cw - 12, 36, WHITE, TEAL,
            paras=[[('Check the signs ▸', 8, True, TEAL)], [('next slides: 10 that increase, 7 that decrease', 7, False, INK)]])
    if i == 2:
        box(x + 6, cy + ch - 42, cw - 12, 36, AMBER_PALE, AMBER,
            paras=[[('▶ Trainer demo', 8, True, AMBER)], [('Open the live threshold table in CRM.', 7, False, INK)]])
    if i == 4:
        btn = box(x + 6, cy + ch - 42, cw - 12, 36, TEAL,
                  paras=[[('Internal note format', 8, True, WHITE)], [('click to open ▸', 7, False, TEAL_PALE)]],
                  align=PP_ALIGN.CENTER)
        btn.name = 'NoteFormatButton'
box(L, 362, W, 15, AMBER_PALE, paras=[[('⚠  Not the card-based assessment:  ', 8, True, AMBER),
                                         ('this is a percentage threshold, not the HDR dollar thresholds per card (Amex, Mastercard, Visa).', 8, False, INK)]],
    margin=8)

# pop-up: internal note format
pop = []
px, py, pw, ph = 170, 92, 380, 262
card = box(px, py, pw, ph, WHITE, TEAL, lw=1.5)
sp = card._element.find(qn('p:spPr'))
eff = etree.SubElement(sp, qn('a:effectLst'))
sh_ = etree.SubElement(eff, qn('a:outerShdw'), blurRad='76200', dist='25400', dir='5400000', algn='t', rotWithShape='0')
clr = etree.SubElement(sh_, qn('a:srgbClr'), val='000000'); etree.SubElement(clr, qn('a:alpha'), val='25000')
pop.append(card)
pop.append(box(px, py, pw, 30, TEAL, paras=[[('Internal note format', 11, True, WHITE),
                                             ('    leave it on every assessment, either way', 8.5, False, TEAL_PALE)]], margin=12))
fields = ['Disputable charge:', 'Dispute threshold:', 'Dispute likelihood:', 'Action:',
          'Signs that increase dispute likelihood:', 'Signs that decrease dispute likelihood:']
fy, fh = py + 40, 31
for i, f in enumerate(fields):
    y = fy + i * fh
    pop.append(box(px + 12, y, pw - 24, fh - 5, FIELD, LINE, paras=[[(f, 9, True, INK)]], margin=8))
pop.append(text(px + 12, py + ph - 14, pw - 24, 10, [[('Click to close', 7.5, False, MUTED)]], align=PP_ALIGN.RIGHT))


def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def leave(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="1" presetClass="exit" presetSubtype="0" fill="hold" grpId="1" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>'
            f'</p:childTnLst></p:cTn></p:par>')


def click(cid, inner):
    return (f'<p:par><p:cTn id="{cid}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{cid+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


ids = [sh.shape_id for sh in pop]
cid, c1 = 3, ''
for i, spid in enumerate(ids):
    c1 += effect(cid + 2, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
click1 = click(3, c1); cid += 2
start2, c2 = cid, ''
cid += 2
for i, spid in enumerate(ids):
    c2 += leave(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 2
click2 = click(start2, c2)
bld = ''.join(f'<p:bldP spid="{i}" grpId="0"/><p:bldP spid="{i}" grpId="1"/>' for i in ids)
NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
       '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
       + click1 + click2 +
       '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
       '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
       '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
old = s20._element.find(qn('p:timing'))
if old is not None:
    s20._element.remove(old)
t_el = etree.fromstring(tim)
ext = s20._element.find(qn('p:extLst'))
(ext.addprevious(t_el) if ext is not None else s20._element.append(t_el))

set_notes(s20, (
    "When a customer wants money back, beyond whether the ask is valid, you read how likely they are to file a dispute with their bank. "
    "It is a judgement call, not a formula. Under-reading means we lose the amount plus a chargeback fee and likely the customer. "
    "Over-reading means we refund reflexively and teach the account that asking loudly works. The goal is accuracy, not caution. "
    "Only charges within 120 days of the charge date count as disputable.\n"
    "The process: 1) Investigate the account: job history, comms, CS and CP notes, and check the signs on the next two slides. "
    "2) Weigh the signs together rather than scoring them. "
    "3) Pull the live threshold for the customer's card network: CRM > Agent Navigation Bar > Refund if Above Dispute Threshold. "
    "TRAINER DEMO: open CRM, go to the Agent Navigation Bar > Refund if Above Dispute Threshold, and show the live table per card network. "
    "4) Judge whether your read clears that network's threshold: if it does, resolve proactively; if not, the standard offer is enough. "
    "5) Leave an internal note of your assessment either way.\n"
    "Don't confuse this with the card-based assessment: this is a percentage threshold, not the HDR dollar thresholds per card (Amex, Mastercard, Visa).\n"
    "TRAINER: click once to open the internal note format pop-up, click again to close it. The note lists: Disputable charge, Dispute threshold, "
    "Dispute likelihood, Action, Signs that increase dispute likelihood, Signs that decrease dispute likelihood."))


# ---------------- signs slides ----------------
def signs_slide(s, title, items, accent, accent_pale, tip):
    clear(s, 2)
    set_runs(s.shapes[1], [[(title, None, None, None)]])
    box, text, dot = kit(s)
    rows = 5
    cw, gap, y0, rh, rg = (W - 10) / 2, 10, 88, 50, 6
    cells = list(items) + ([tip] if tip else [])
    for k, it in enumerate(cells):
        col, row = divmod(k, rows)
        x, y = L + col * (cw + gap), y0 + row * (rh + rg)
        if it is tip:
            box(x, y, cw, rh, accent_pale, accent, paras=[[(tip[0], 9.5, True, accent)], [(tip[1], 8.5, False, INK)]], margin=10)
            continue
        head, body = it
        box(x, y, cw, rh, WHITE, LINE)
        dot(x + 9, y + (rh - 20) / 2, 20, str(k + 1), accent)
        text(x + 38, y + 4, cw - 46, rh - 8, [[(head, 10, True, accent)], [(body, 8.5, False, INK)]], anchor=MSO_ANCHOR.MIDDLE)


inc = [('Early in FC', 'Only on their 1st or 2nd paid month, or the ETF was charged right after it.'),
       ('Said "cancel" or "refund"', 'Used either word in a message or in a Help Center search.'),
       ('Called the billing number', 'Called the billing phone number about the charge.'),
       ('Visited cleaningbill.com', 'Went to cleaningbill.com and looked up the charge.'),
       ('Bad experience', "Couldn't get matched, or low ratings plus few invoiced jobs. Low ratings alone are not the signal."),
       ('Tried to swap the card', 'Tried to change the card on file to avoid paying.'),
       ('Says they disputed', 'Check the tense: "I already called my bank" is stronger than "I\'m going to".'),
       ('Public threats', 'Threatens to take the complaint public.'),
       ('BBB case', 'Opened or mentions a Better Business Bureau case.'),
       ('Escalated tone', '"Scam", "fraud", expletives.')]
dec = [('Satisfied jobs', '2+ satisfied invoiced jobs, high ratings or a repeat cleaner.'),
       ('Wants an upcoming job', 'Self-booked it, or confirmed it with CS or the CP.'),
       ('Wants to keep FC', 'Says or implies they want to keep the membership.'),
       ('Understands the terms', 'Says or implies they understand the FC terms.'),
       ('3+ paid FC months', 'If they were going to dispute those charges, they would have.'),
       ('Asks for more', 'Asks for credits or free months. They are invested.'),
       ('Wants a charge explained', 'Asks us to explain a charge rather than reverse it. They are still engaged.')]
signs_slide(s22, 'Signs that increase dispute likelihood', inc, AMBER, AMBER_PALE, None)
signs_slide(s23, 'Signs that decrease dispute likelihood', dec, TEAL, TEAL_PALE,
            ("Weigh, don't score", 'One sign rarely tells the story. Read these together with the signs that increase likelihood.'))
set_notes(s22, (
    'Signs that increase dispute likelihood (all 10): 1) early in FC: only on the 1st or 2nd paid month, or the ETF charged right after it; '
    '2) used the words cancel or refund in a message or Help Center search; 3) called the billing number; 4) visited cleaningbill.com and looked up the charge; '
    "5) a genuinely bad experience: couldn't get matched, or low ratings plus few invoiced jobs; 6) tried to swap the card to avoid paying; "
    '7) says they disputed; 8) public threats; 9) a BBB case; 10) an escalated tone such as "scam", "fraud" or expletives. '
    "Two reading traps: low ratings alone aren't the signal (it's low ratings plus few jobs), and check the tense, because "
    "'I already called my bank' is stronger than 'I'm going to'."))
set_notes(s23, (
    'Signs that decrease dispute likelihood (all 7): 1) 2+ satisfied invoiced jobs, high ratings or a repeat cleaner; 2) an upcoming job they want, '
    'self-booked or confirmed with CS or the CP; 3) wanting to keep FC; 4) understanding the FC terms; 5) 3+ paid FC months; '
    '6) asking for more from us, such as credits or free months; 7) wanting a charge explained. '
    "Weigh them together with the signs that increase likelihood: one sign rarely tells the story."))

# ---------------- delete slide 21 ----------------
lst = prs.slides._sldIdLst
sld = lst[20]
rid = sld.get(qn('r:id'))
lst.remove(sld)
prs.part.drop_rel(rid)

prs.save(sys.argv[2]); print('saved', sys.argv[2])
