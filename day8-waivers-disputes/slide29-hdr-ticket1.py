"""Day 8 deck slide 29: replace the "HDR tickets 5-8" table with "HDR Ticket 1: C's
RefundRequest expired & received a courtesy voucher", laid out as a 1st Response / 2nd
Response table (label column + content), with the partial-refund note and $105 example.
The old tickets 5-8 handling text moves to the speaker notes.

Usage: python3 slide29-hdr-ticket1.py in.pptx out.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'

prs = Presentation(sys.argv[1])
s = prs.slides[28]
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
assert txt(s.shapes[1]) == 'HDR tickets 5-8', txt(s.shapes[1])
old = next(sh for sh in s.shapes if sh.has_table)
old_rows = [[c.text for c in r.cells] for r in old.table.rows][1:]
old_notes = s.notes_slide.notes_text_frame.text
old._element.getparent().remove(old._element)


def set_text(shape, t, size=None):
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = t
    if size: p.runs[0].font.size = Pt(size)
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


set_text(s.shapes[1], 'HDR Ticket 1')
set_text(s.shapes[2], "C's RefundRequest expired & received a courtesy voucher")
r = s.shapes[2].text_frame.paragraphs[0].runs[0]
r.font.bold = True; r.font.name = 'Play'; r.font.color.rgb = RGBColor.from_string(TEAL)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def rect(x, y, w, h, fill, line=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    return sh


def write(x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: list of (level, runs); level 0 = plain, 1 = bullet, 2 = sub-bullet."""
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (lvl, runs) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(2.5)
        if lvl:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int((12 + 14 * (lvl - 1)) * 12700 + 9 * 12700)))
            pPr.set('indent', str(-9 * 12700))
            bf = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}); pPr.append(bf)
            bc = pPr.makeelement(qn('a:buChar'), {'char': '•' if lvl == 1 else '◦'}); pPr.append(bc)
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return tb


B = lambda t, c=INK, sz=8.5: (t, sz, True, c)
A = lambda t, c=INK, sz=8.5: (t, sz, False, c)

X, W, LW = 32.4, 655.2, 92
CX, CW = X + LW, W - LW
r1y, r1h = 100, 98
r2y, r2h = r1y + r1h, 272 - r1h

# row 1: 1st response
rect(X, r1y, LW, r1h, AMBER, LINE)
write(X, r1y, LW, r1h, [(0, [B('1st Response', WHITE, 10)])], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
rect(CX, r1y, CW, r1h, AMBER_PALE, LINE)
write(CX + 12, r1y + 8, CW - 24, r1h - 12, [
    (0, [A("Before anything else, find out what actually happened with this job: pull up both C's and the CP's accounts and check the cleaning really took place.")]),
    (1, [A('Looks like a '), B('false invoice'), A('? Follow the usual False Invoice process.')]),
    (1, [A('CP '), B('overcharged'), A(' C for the job? Follow the usual OCH process.')]),
    (1, [B('Job happened as described?')]),
    (2, [A('Remind C about the courtesy voucher they already have. Sometimes this alone resolves it.')]),
    (2, [A('Offer to set them up with a different cleaner using that voucher.')])])

# row 2: 2nd response
rect(X, r2y, LW, r2h, AMBER_PALE, LINE)
write(X, r2y, LW, r2h, [(0, [B('2nd Response', AMBER, 10)])], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
rect(CX, r2y, CW, r2h, WHITE, LINE)
write(CX + 12, r2y + 8, CW - 24, 110, [
    (0, [A('If C comes back still asking for a refund, first check the timing: how long has it been since C was flagged HDR?')]),
    (1, [B('Within 48 hours: ', TEAL), A('leave an internal note of your card-based threshold assessment, then:')]),
    (2, [B('Above the threshold → '), A("don't refund, unless something else makes C eligible (usually a material legal threat or a genuine hardship case). Nothing applies? Handle it as usual.")]),
    (2, [B('Below the threshold → '), A('refund, and address anything else C is concerned about.')]),
    (1, [B('Past 48 hours: ', TEAL), A('the urgency window has closed. Follow the usual Transformation Guidelines (48-hour rule).')])])
nb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(CX + 10), Pt(r2y + r2h - 54), Pt(CW - 20), Pt(46)); strip(nb)
nb.adjustments[0] = 0.12
nb.fill.solid(); nb.fill.fore_color.rgb = RGBColor.from_string(TEAL_PALE); nb.line.fill.background()
write(CX + 20, r2y + r2h - 50, CW - 40, 40, [
    (0, [B('Partial refund?  ', TEAL), A('Assess against the full original Stripe amount, not just the part C asks back.')]),
    (0, [B('Example: ', MUTED, 8), A('C was charged $105 and asks for $20 back. Assess the full $105. It is above the card-based threshold, so no refund, even though $20 alone would have cleared it.', MUTED, 8)])],
    anchor=MSO_ANCHOR.MIDDLE)

s.notes_slide.notes_text_frame.text = (
    "HDR Ticket 1: C's RefundRequest expired and C received a courtesy voucher.\n"
    "1st response: before anything else, figure out what actually happened with the job by pulling up both C's and the CP's accounts and checking the cleaning took place. "
    'If it looks like a false invoice, follow the usual False Invoice process. If the CP overcharged C, follow the usual OCH process. '
    'If the job happened as described, remind C about the courtesy voucher they already have (sometimes this alone resolves it) and offer to set them up with a different cleaner using that voucher.\n'
    '2nd response: if C comes back still pushing for a refund, check how long it has been since they were flagged HDR. '
    'Within 48 hours, write a quick internal note on the card-based threshold, then: above the threshold, hold off on the refund unless something else makes C eligible anyway '
    '(a material legal threat or genuine hardship case are the usual exceptions), otherwise handle as usual; below the threshold, go ahead and refund and take care of anything else bothering them. '
    'Past 48 hours, the urgency window is closed: follow the usual Transformation Guidelines.\n'
    'Note: even if C only wants a partial refund on a job charge, base the assessment on the full original amount from Stripe. Example: C was charged $105 and asks for $20 back. '
    'Assess against $105; since $105 is above the card-based threshold, no refund goes out, even though $20 would have cleared the threshold on its own.\n'
    'Tickets 5-8 from the previous version of this slide: ' + ' '.join(f'{a}: {b}.' for a, b in old_rows) + '\n' + old_notes)
prs.save(sys.argv[2]); print('saved', sys.argv[2])
