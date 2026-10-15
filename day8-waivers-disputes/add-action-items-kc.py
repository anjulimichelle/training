"""Day 8 deck (helpers copied from add-practice-and-hdr-kc.py):
  - After "Group E — Membership & Reminders": add a knowledge check covering Internal Note
    Action Items (one answer revealed per click).

Usage: python3 add-action-items-kc.py in.pptx out.pptx
"""
import copy, os, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7'
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets', 'cancel-after-mct-customer-message.png')

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''


def find(title, eyebrow_start=None):
    for i, sl in enumerate(prs.slides):
        if txt(sl.shapes[1]) == title and (eyebrow_start is None or txt(sl.shapes[0]).startswith(eyebrow_start)):
            return i, sl
    raise SystemExit('not found: ' + title)


def new_slide(ref, n_head, after_idx, texts):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    num = [x for x in ref.shapes if x.is_placeholder][0]
    for el in [ref.shapes[k]._element for k in range(n_head)] + [num._element]:
        s.shapes._spTree.append(copy.deepcopy(el))
    for k, t in enumerate(texts):
        r = s.shapes[k].text_frame.paragraphs[0].runs
        r[0].text = t
        for x in r[1:]:
            x._r.getparent().remove(x._r)
    lst = prs.slides._sldIdLst
    el = list(lst)[-1]; lst.remove(el); list(lst)[after_idx].addnext(el)
    return s


def kit(s):
    def strip(sh):
        st = sh._element.find(qn('p:style'))
        if st is not None:
            sh._element.remove(st)

    def fill(tf, paras, align, anchor, after=3):
        tf.word_wrap = True; tf.vertical_anchor = anchor
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align; p.space_after = Pt(after)
            for t, size, bold, color in runs:
                r = p.add_run(); r.text = t; f = r.font
                f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
                f.color.rgb = RGBColor.from_string(color)

    def box(x, y, w, h, bg, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=8,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=False, name=None):
        sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = min(0.5, 6 / min(w, h))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(bg)
        if line:
            sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(1 if dash else 0.75)
            if dash:
                sh.line.dash_style = MSO_LINE.DASH
        else:
            sh.line.fill.background()
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(3)
        fill(tf, paras, align, anchor)
        if name:
            sh.name = name
        return sh

    def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
        tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        fill(tf, paras, align, anchor)
        if name:
            tb.name = name
        return tb
    return box, text


def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def click(cid, inner):
    return (f'<p:par><p:cTn id="{cid}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{cid+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


def reveal_timing(s, groups):
    """Each group of shapes appears on its own click, in order."""
    cid, clicks = 3, ''
    for ids in groups:
        start, inner = cid, ''
        cid += 2
        for i, spid in enumerate(ids):
            inner += effect(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
        clicks += click(start, inner)
    bld = ''.join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for g in groups for i in g)
    NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + clicks +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
    old = s._element.find(qn('p:timing'))
    if old is not None:
        s._element.remove(old)
    t_el = etree.fromstring(tim)
    ext = s._element.find(qn('p:extLst'))
    (ext.addprevious(t_el) if ext is not None else s._element.append(t_el))



# ---------------- Action items knowledge check (after Group E) ----------------
proc_i = next(i for i, sl in enumerate(prs.slides) if txt(sl.shapes[0]) == 'INTERNAL NOTE ACTION ITEMS' and txt(sl.shapes[1]).startswith('Group E'))
ref = next(sl for sl in prs.slides if txt(sl.shapes[0]).endswith('KNOWLEDGE CHECK'))
kc = new_slide(ref, 2, proc_i, ['INTERNAL NOTE ACTION ITEMS  ·  KNOWLEDGE CHECK', 'Let’s See What You’ve Got!'])
box, text = kit(kc)
X, W = 32.4, 655.2
qa = [("Where do you find what triggered an action item?",
       "The Sender field at the top right of the internal note. It's a system code: look it up, then follow its steps."),
      ("A trigger says to call C, and the ticket came in at 7:55 AM C's local time. What do you do?",
       "It's under 10 minutes before 8 AM: park it and call once business hours start."),
      ("ForeverClean.charge fires for a C who disputed before. Risk is above the threshold, and C hasn't asked for anything.",
       "Cancel FC and waive the ETF without contacting C. The goal is to prevent the dispute, not invite one."),
      ("Shared address flag: two accounts with active FC, and no sign of fraud. What do you add?",
       "admin_ignore_address_fraud_flags. (High fraud risk would mean cancel FC on all accounts and add known_customer_fraud.)"),
      ("C has an upcoming job, and a previous agent forgot to remove their block-comms flag. What do you do?",
       "Remove the flag and remind C about the appointment."),
      ("A High-LTV C cancels FC through OCW but has disputed a charge before. Reassign to FC Reactivations?",
       "No. A past dispute is one of the exceptions, so handle the ticket yourself.")]
n = len(qa); gap = 5; h = (272 - gap * (n - 1)) / n
groups = []
for i, (q, a) in enumerate(qa):
    y = 100 + i * (h + gap)
    box(X, y, 364, h, WHITE, LINE)
    box(X + 10, y + (h - 20) / 2, 20, 20, TEAL_PALE, paras=[[(str(i + 1), 9, True, TEAL)]], align=PP_ALIGN.CENTER, margin=0,
        shape=MSO_SHAPE.OVAL)
    text(X + 40, y, 316, h, [[(q, 8.5, False, INK)]], anchor=MSO_ANCHOR.MIDDLE)
    b = box(X + 371, y, W - 371, h, TEAL_PALE, name=f'step{i + 1}Box')
    t = text(X + 381, y, W - 391, h, [[(a, 8.5, False, INK)]], anchor=MSO_ANCHOR.MIDDLE, name=f'step{i + 1}Ans')
    groups.append([b.shape_id, t.shape_id])
reveal_timing(kc, groups)
kc.notes_slide.notes_text_frame.text = (
    'Quick knowledge check: Internal Note Action Items. Ask each question and let trainees answer before revealing. Each click reveals the next answer.\n'
    '1) The Sender field in the top-right corner of the internal note holds a system code, such as utils.report_double_booked_clients. Look it up, then follow its steps.\n'
    "2) Standard Call Attempt: call 8:00 AM to 8:00 PM in the user's local time. Under 10 minutes before 8 AM, park the ticket and call at 8; more than 10 minutes before 8 AM or after 8 PM, don't wait and treat it as no answer.\n"
    '3) ForeverClean.charge for a customer who disputed before: run the dispute-risk assessment. Above the threshold with no request, cancel FC and waive the ETF without contacting the customer. If they asked, also refund the smaller of the disputable charges or the ask.\n'
    '4) utils.flag_multiple_customers_with_same_addresses: with active FC on more than one account, high fraud risk means cancel FC on all accounts and add known_customer_fraud; no sign of fraud means add admin_ignore_address_fraud_flags.\n'
    '5) utils.flag_upcoming_jobs_customer_cp_unable_to_receive_sms: if the block-comms flag was left on by mistake, remove it and remind the customer about the appointment.\n'
    '6) views.customer_foreverclean: reassign High-LTV OCW cancellations to FC Reactivations, except when the customer contacted us about something unrelated, has disputed before, is asking for a refund, has a damage_theft_reported flag, or has a BBB, Yelp or social-media threat on file.')
prs.save(sys.argv[2]); print('saved', sys.argv[2])
