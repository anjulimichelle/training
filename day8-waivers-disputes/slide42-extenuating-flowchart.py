"""Day 8 deck slide 42 (ETF Waivers: Extenuating circumstances): replace the table with a flow
chart. Root -> Compassion / Fairness / Compliance / Affordability -> one card per circumstance,
with "waive + cancel" in amber and "retain / pause" in teal. The table text moves to the notes.

Usage: python3 slide42-extenuating-flowchart.py in.pptx out.pptx
"""
import sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE, GREY = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', '6E6E73'

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
idx = next(i for i, sl in enumerate(prs.slides) if txt(sl.shapes[1]) == 'Extenuating circumstances' and txt(sl.shapes[0]) == 'ETF WAIVERS')
s = prs.slides[idx]
old = next(sh for sh in s.shapes if sh.has_table)
rows = [[c.text for c in r.cells] for r in old.table.rows][1:]
old_notes = s.notes_slide.notes_text_frame.text
old._element.getparent().remove(old._element)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def box(x, y, w, h, bg, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=7):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    sh.adjustments[0] = min(0.5, 7 / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(bg)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(3)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(1.5)
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)
    return sh


def line(x1, y1, x2, y2, head=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)); strip(c)
    c.line.color.rgb = RGBColor.from_string(GREY); c.line.width = Pt(1)
    if head:
        ln = c._element.find('.//' + qn('a:ln'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')


SZ = 7.5
H = lambda t, c: (t, 8.5, True, c)
A = lambda t: (t, SZ, False, INK)
WAIVE = ('→ waive + cancel', SZ, True, AMBER)
R = lambda t: ('→ ' + t, SZ, True, TEAL)


def card(x, y, w, h, title, lines, color, pale):
    paras = [[H(title, color)]]
    for cond, action in lines:
        paras.append([A(cond)])
        paras.append([action])
    return box(x, y, w, h, pale, color, paras, anchor=MSO_ANCHOR.TOP)


X, W = 32.4, 655.2
n, g = 4, 10
cw = (W - g * (n - 1)) / n
cx = [X + i * (cw + g) for i in range(n)]
mid = [x + cw / 2 for x in cx]

box(X + W / 2 - 95, 98, 190, 22, INK, paras=[[('Extenuating Circumstances', 10, True, WHITE)]], align=PP_ALIGN.CENTER)
line(X + W / 2, 120, X + W / 2, 127)
line(mid[0], 127, mid[-1], 127)
cats = [('Compassion', 'Severe illness, death', AMBER, AMBER_PALE),
        ('Fairness', 'Military deployment, severe weather', TEAL, TEAL_PALE),
        ('Compliance', 'Disability need', AMBER, AMBER_PALE),
        ('Affordability', 'Extreme financial hardship', TEAL, TEAL_PALE)]
cy, ch = 134, 30
for i, (name, sub, col, pale) in enumerate(cats):
    line(mid[i], 127, mid[i], cy, head=True)
    box(cx[i], cy, cw, ch, col, paras=[[(name, 9.5, True, WHITE)], [(sub, 7.5, False, WHITE)]], align=PP_ALIGN.CENTER, margin=4)

r1y, r1h = cy + ch + 10, 80
r2y = r1y + r1h + 10
r2h = 372 - r2y
for i in range(n):
    line(mid[i], cy + ch, mid[i], r1y, head=True)

card(cx[0], r1y, cw, r1h, 'Severe illness', [
    ('Wants to cancel because of it', WAIVE),
    ('Illness is why they use the service', R('retain first'))], AMBER, AMBER_PALE)
card(cx[0], r2y, cw, r2h, 'Death or bereavement', [
    ('FC member has passed away', WAIVE),
    ('Family member passed, so C no longer needs the service', WAIVE),
    ('Family member passed, C still needs the service', R('retain first, offer to pause'))], AMBER, AMBER_PALE)
line(mid[0], r1y + r1h, mid[0], r2y, head=True)

card(cx[1], r1y, cw, r1h, 'Military deployment', [
    ('Deploying now, wants to cancel', WAIVE),
    ('Just a heads up, still using the service', R('retain first, set a reminder for later'))], TEAL, TEAL_PALE)
card(cx[1], r2y, cw, r2h, 'Severe weather | Disaster', [
    ("Can't use the service anymore", WAIVE),
    ('Temporarily disrupted, will come back', R('pause instead'))], TEAL, TEAL_PALE)
line(mid[1], r1y + r1h, mid[1], r2y, head=True)

card(cx[2], r1y, cw, r1h, 'ADA accommodation', [
    ('Disability means they can no longer use the service', WAIVE),
    ('Can still use it with an adjustment', R('retain first'))], AMBER, AMBER_PALE)
card(cx[3], r1y, cw, r1h, 'Financial hardship', [
    ('Extreme financial hardship is clear', ('→ 50% off ETF; full waiver if declined', SZ, True, AMBER)),
    ('Hardship is vague, no specifics given', R('retain first; MF reduction to address affordability'))], TEAL, TEAL_PALE)

wx = cx[2]; ww = cw * 2 + g
box(wx, r2y, ww, r2h, WHITE, LINE, paras=[
    [('Also remember', 8.5, True, TEAL)],
    [('Moving to a care facility or assisted living ', SZ, False, INK), ('→ waive + cancel', SZ, True, AMBER),
     (' (no health decline statement or TL approval needed).', SZ, False, INK)],
    [('Refund windows: ', SZ, True, INK), ('severe illness, family member passed, deployment: unused MFs within 120 days. FC member passed: up to 1 year.', SZ, False, INK)],
    [('Qualifying makes the waiver ', SZ, False, MUTED), ('available', SZ, True, MUTED), (', not automatic.', SZ, False, MUTED)]],
    anchor=MSO_ANCHOR.TOP, margin=9)

s.notes_slide.notes_text_frame.text = (
    'Extenuating circumstances fall into four groups: compassion (severe illness, death), fairness (military deployment, severe weather), '
    'compliance (disability need) and affordability (extreme financial hardship).\n'
    'Severe illness: cancelling because of it, waive and cancel; if the illness is why they use the service, retain first. '
    'Death or bereavement: FC member passed, waive and cancel; family member passed and C no longer needs the service, waive and cancel; family member passed but C still needs the service, retain first and offer to pause. '
    'Military deployment: deploying now and wants to cancel, waive and cancel; just a heads up and still using the service, retain first and set a reminder for later. '
    "Severe weather or disaster: can't use the service anymore, waive and cancel; temporarily disrupted and will come back, pause instead. "
    'ADA accommodation: the disability means they can no longer use the service, waive and cancel; can still use it with an adjustment, retain first. '
    'Financial hardship: extreme hardship is clear, 50% off the ETF, full waiver if they decline; hardship is vague with no specifics, retain first with an MF reduction to address affordability.\n'
    'Detail from the previous table: ' + ' '.join(f'{a}: {b}.' for a, b in rows) + '\n' + old_notes)
prs.save(sys.argv[2]); print('saved slide', idx + 1, sys.argv[2])
