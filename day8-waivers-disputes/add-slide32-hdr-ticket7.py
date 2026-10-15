"""Day 8 deck: add slide 32 "HDR Ticket 7: Stripe Dispute Received" after HDR Ticket 6.
  - Flow chart: 1 check auto-accepted disputed charges -> 2 check other charges C wants refunded.
  - Related ticket (Dispute Response Ticket with Dispute Risk): steps 1-3, with a
    "View sample ticket" button (click 1 opens the screenshot pop-up, click 2 closes it).
  - Active experiment note (fc_dispute_cancel_on_second_dispute).

Usage: python3 add-slide32-hdr-ticket7.py in.pptx out.pptx
"""
import copy, os, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEAL, TEAL_PALE, AMBER, AMBER_PALE = '0B7A6F', 'E3F4F1', '9A6410', 'FFF3D6'
INK, MUTED, WHITE, LINE, GREY = '1D1D1F', '6E6E73', 'FFFFFF', 'DCE1E7', '6E6E73'
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets', 'dispute-response-ticket-sample.png')

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
S = list(prs.slides)
ref = S[30]
assert txt(ref.shapes[1]) == 'HDR Ticket 6', txt(ref.shapes[1])
num = [x for x in ref.shapes if x.is_placeholder][0]

s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
for shape, t in ((s.shapes[1], 'HDR Ticket 7'), (s.shapes[2], 'Stripe Dispute Received')):
    r = shape.text_frame.paragraphs[0].runs
    r[0].text = t
    for x in r[1:]:
        x._r.getparent().remove(x._r)
lst = prs.slides._sldIdLst
el = list(lst)[-1]; lst.remove(el); list(lst)[30].addnext(el)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill(tf, paras, align, anchor, after=1.5):
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, item in enumerate(paras):
        lvl, runs = item if isinstance(item, tuple) else (0, item)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(after)
        if lvl:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int((8 + 12 * (lvl - 1)) * 12700 + 8 * 12700)))
            pPr.set('indent', str(-8 * 12700))
            pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
            pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '•◦'[lvl - 1]}))
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(color)


def shape(kind, x, y, w, h, bg, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=8, radius=6, lw=0.75):
    sh = s.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(bg)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(2)
    fill(tf, paras, align, anchor)
    return sh


def text(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill(tf, paras, align, anchor)
    return tb


def dot(x, y, label, bg=TEAL, d=16):
    return shape(MSO_SHAPE.OVAL, x, y, d, d, bg, paras=[[(label, 8.5, True, WHITE)]], align=PP_ALIGN.CENTER, margin=0)


def arrow(x1, y1, x2, y2, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)); strip(c)
    c.line.color.rgb = RGBColor.from_string(GREY); c.line.width = Pt(1)
    if head:
        ln = c._element.find('.//' + qn('a:ln'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
    return c


SZ = 7.5
B = lambda t, c=INK, sz=SZ: (t, sz, True, c)
A = lambda t, c=INK, sz=SZ: (t, sz, False, c)
X, W = 32.4, 655.2

# ---- flow chart (screenshot 1) ----
fy, fh = 104, 104
p1 = shape(MSO_SHAPE.PENTAGON, X, fy, 345, fh, AMBER_PALE, AMBER)
dot(X + 10, fy + 9, '1', AMBER)
text(X + 32, fy + 9, 280, fh - 14, [
    [B('Check auto-accepted disputed charges.', AMBER, 8.5)],
    (1, [A('Charge related to a '), B('job'), A(': create a CP HB (holdback) if we have sufficient evidence (e.g. C/CP comms) the job was not completed, and follow the '),
         B('False Invoice'), A(' instructions.')]),
    (1, [A('Charge for the '), B("CP's tip"), A(': create a CP HB with Holdback type "Client <C name> requested refund of tip given for <job date> cleaning".')])])
p2 = shape(MSO_SHAPE.CHEVRON, X + 332, fy, W - 332, fh, TEAL_PALE, TEAL)
p2.adjustments[0] = 0.22
dot(X + 368, fy + 9, '2', TEAL)
text(X + 390, fy + 9, W - 332 - 92, fh - 14, [
    [B('Check if there are other charges C wants refunded:', TEAL, 8.5)],
    (1, [B('Yes: '), A('follow '), B('Customer Has Dispute History', TEAL), A('.')]),
    (1, [B('No: '), A("C hasn't reached out for additional refunds, so there's no need to assess or issue refunds for the other disputable charges.")])])

# ---- related ticket ----
ry = 218
text(X, ry, 470, 16, [[B('Related ticket:  ', MUTED, 9), B('Dispute Response Ticket with Dispute Risk', TEAL, 9.5)]], anchor=MSO_ANCHOR.MIDDLE)
btn = shape(MSO_SHAPE.ROUNDED_RECTANGLE, X + W - 128, ry - 2, 128, 20, TEAL,
            paras=[[B('View sample ticket ▸', WHITE, 8)]], align=PP_ALIGN.CENTER, margin=4)
btn.name = 'SampleTicketButton'
sy, sh_ = 240, 98
shape(MSO_SHAPE.ROUNDED_RECTANGLE, X, sy, 330, sh_, WHITE, TEAL)
dot(X + 9, sy + 8, '1')
text(X + 30, sy + 8, 290, sh_ - 12, [
    [A('Check the '), B('CRM'), A(' and '), B('Stripe'), A(' to see whether the dispute response has already been submitted and is awaiting results.')],
    (1, [B('CRM: '), A('an internal note says the dispute response has been submitted.')]),
    (1, [B('Stripe: '), A('shows '), B('<#> days till results'), A(' if responded to, or '), B('<#> days to respond'), A(' if there is no response yet.')])])
bx = X + 348
arrow(X + 330, sy + sh_ / 2, X + 339, sy + sh_ / 2, head=False)
arrow(X + 339, sy + 23, X + 339, sy + sh_ - 23, head=False)
arrow(X + 339, sy + 23, bx, sy + 23); arrow(X + 339, sy + sh_ - 23, bx, sy + sh_ - 23)
bw = W - 348
shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, sy, bw, 46, TEAL_PALE, TEAL,
      paras=[[B('Submitted: ', TEAL, 8.5), A('handle BAU (business as usual).', INK, 8.5)]], margin=30)
dot(bx + 9, sy + 15, '2')
shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, sy + sh_ - 46, bw, 46, AMBER_PALE, AMBER,
      paras=[[B('NOT submitted: ', AMBER, 8.5), B("don't resolve the ticket", INK, 8.5), A(' and change the queue to ', INK, 8.5),
              B('Dispute Response Team', INK, 8.5), A('.', INK, 8.5)]], margin=30)
dot(bx + 9, sy + sh_ - 31, '3', AMBER)

# ---- active experiment note ----
shape(MSO_SHAPE.ROUNDED_RECTANGLE, X, 345, W, 29, AMBER_PALE, paras=[[
    B('⚑ Active experiment:  ', AMBER, 8), A('if the ', INK, 8), B('Customer.flag_as_dispute_risk', INK, 8),
    A(' note says the FC was kept on purpose for the ', INK, 8), B('fc_dispute_cancel_on_second_dispute', INK, 8),
    A(" experiment, don't cancel the FC or refund unless C asks. If C asks, follow ", INK, 8), B('Customer Has Dispute History', INK, 8), A('.', INK, 8)]],
    margin=10)

# ---- pop-up: sample ticket ----
pop = []
iw = 600; ih = iw * 355 / 976
px, py, pw = 60, 96, iw + 24
ph = 30 + ih + 34
card = shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, py, pw, ph, WHITE, TEAL, lw=1.5)
sp = card._element.find(qn('p:spPr'))
eff = etree.SubElement(sp, qn('a:effectLst'))
sd = etree.SubElement(eff, qn('a:outerShdw'), blurRad='76200', dist='25400', dir='5400000', algn='t', rotWithShape='0')
clr = etree.SubElement(sd, qn('a:srgbClr'), val='000000'); etree.SubElement(clr, qn('a:alpha'), val='25000')
pop.append(card)
pop.append(shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, py, pw, 26, TEAL,
                 paras=[[B('Sample ticket', WHITE, 11), A('    Dispute Response Ticket with Dispute Risk (internal note)', TEAL_PALE, 8.5)]], margin=12))
pop.append(s.shapes.add_picture(IMG, Pt(px + 12), Pt(py + 34), Pt(iw), Pt(ih)))
pop.append(text(px + 12, py + ph - 16, iw, 10, [[A('Click to close', MUTED, 7.5)]], align=PP_ALIGN.RIGHT))


def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def leave(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="1" presetClass="exit" presetSubtype="0" fill="hold" grpId="1" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>'
            f'</p:childTnLst></p:cTn></p:par>')


def click(cid, inner):
    return (f'<p:par><p:cTn id="{cid}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{cid+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')


ids = [x.shape_id for x in pop]
cid, c1 = 3, ''
for i, spid in enumerate(ids):
    c1 += effect(cid + 2, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
click1 = click(3, c1); cid += 2
start2, c2 = cid, ''
cid += 2
for i, spid in enumerate(ids):
    c2 += leave(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 2
click2 = click(start2, c2)
bld = ''.join(f'<p:bldP spid="{i}" grpId="0"/><p:bldP spid="{i}" grpId="1"/>' for i in ids if True)
NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
       '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
       + click1 + click2 +
       '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
       '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
       '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
t_el = etree.fromstring(tim)
ext = s._element.find(qn('p:extLst'))
(ext.addprevious(t_el) if ext is not None else s._element.append(t_el))

s.notes_slide.notes_text_frame.text = (
    'HDR Ticket 7: Stripe Dispute Received.\n'
    'Step 1 of the flow: check the auto-accepted disputed charges. If the charge is related to a job, create a CP holdback (HB) if there is sufficient evidence, such as C/CP comms, '
    'that the job was not completed, and follow the False Invoice instructions. If the charge is for the CP\'s tip, create a CP HB and set the Holdback type to '
    '"Client <C name> requested refund of tip given for <job date> cleaning".\n'
    'Step 2: check whether there are other charges C wants refunded. Yes: follow Customer Has Dispute History. No: if C has not reached out for additional refunds, '
    'there is no need to assess or issue additional refunds for the other disputable charges.\n'
    'Related ticket, Dispute Response Ticket with Dispute Risk: 1) check the CRM and Stripe to see whether the dispute response has already been submitted and is awaiting results. '
    'In the CRM, look for an internal note that the dispute response was submitted. In Stripe, a responded dispute shows <#> days till results; one with no response yet shows <#> days to respond. '
    '2) If the dispute response has been submitted, handle BAU. 3) If it has NOT been submitted, do not resolve the ticket, and change the queue to Dispute Response Team.\n'
    'Active experiment: if the Customer.flag_as_dispute_risk note says the FC was kept on purpose for the fc_dispute_cancel_on_second_dispute experiment, '
    "don't cancel the FC or refund unless C asks. If C asks, follow Customer Has Dispute History.\n"
    'TRAINER: click once to open the sample Dispute Response ticket, click again to close it.')
prs.save(sys.argv[2]); print('saved', sys.argv[2])
