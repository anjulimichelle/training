"""Build the "Capital One Card on File" training deck in the New Hire Care Training theme.
Template: an export of the Day 8 deck (cover, topic divider, content header and Questions slides
are reused for the theme). Content: Care Knowledge Library, "Capital One Card on File".

Usage: python3 build-deck.py day8-template.pptx out.pptx
"""
import copy, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

C = {'TEAL': '0B7A6F', 'TEAL_PALE': 'E3F4F1', 'AMBER': '9A6410', 'AMBER_PALE': 'FFF3D6',
     'INK': '1D1D1F', 'MUTED': '6E6E73', 'WHITE': 'FFFFFF', 'LINE': 'DCE1E7', 'GREY': '6E6E73', 'ALT': 'F5F6F8'}
X, W = 32.4, 655.2

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
S = list(prs.slides)
lst = prs.slides._sldIdLst
el_of = {id(s): e for s, e in zip(S, list(lst))}
cover = S[0]
divider = next(s for s in S if txt(s.shapes[0]).startswith('TOPIC') and len(s.shapes) > 20)
questions = next(s for s in S if txt(s.shapes[1]) == 'Questions?')
content_ref = next(s for s in S if txt(s.shapes[0]) == 'ASSESSING DISPUTE LIKELIHOOD' and txt(s.shapes[1]) == 'What is this?')
keep = {id(cover), id(divider), id(questions)}


def set_runs(shape, t):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    lines = t.split('\n')
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    p0.runs[0].text = lines[0]
    for ln in lines[1:]:
        np_ = copy.deepcopy(p0._p); p0._p.getparent().append(np_)
        rr = np_.findall(qn('a:r'))
        rr[0].find(qn('a:t')).text = ln
    return shape


# cover
set_runs(cover.shapes[0], 'HOMEAGLOW  ·  NEW HIRE CARE TRAINING')
set_runs(cover.shapes[1], 'Capital One\nCard on File')
set_runs(cover.shapes[2], 'Why membership charges are skipped, how to confirm it, and how to handle each case.')
cover.notes_slide.notes_text_frame.text = 'Capital One Card on File (Retention and Membership). Source: Care Knowledge Library, "Capital One Card on File".'

# divider
set_runs(divider.shapes[0], 'RETENTION AND MEMBERSHIP')
set_runs(divider.shapes[1], 'Capital One Card on File')
set_runs(divider.shapes[2], "When every card on the account is Capital One, we can't charge MFs or ETFs.")
for k, t in zip((6, 9, 12, 15), ('What it is and how to spot it', 'The rules and what to say', 'Confirm the case, pick the branch',
                                   'Handle each branch with the macros')):
    set_runs(divider.shapes[k], t)
divider.notes_slide.notes_text_frame.text = ('This module covers customers whose only cards on file are Capital One. By the end, trainees can spot the case, '
                                             'explain it without blaming the card, and handle each branch with the right macro.')

order = [cover, divider]


def new_slide(eyebrow, title, sub):
    s = prs.slides.add_slide(content_ref.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    num = [x for x in content_ref.shapes if x.is_placeholder][0]
    for e in (content_ref.shapes[0]._element, content_ref.shapes[1]._element, content_ref.shapes[2]._element, num._element):
        s.shapes._spTree.append(copy.deepcopy(e))
    set_runs(s.shapes[0], eyebrow); set_runs(s.shapes[1], title); set_runs(s.shapes[2], sub)
    order.append(s)
    return s


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill_tf(tf, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, after=2):
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, item in enumerate(paras):
        num, runs = item if isinstance(item, tuple) else ('', item)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after); p.alignment = align
        if num:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int(12 * 12700))); pPr.set('indent', str(-10 * 12700))
            runs = [(num + ' ', runs[0][1], True, 'TEAL' if num != '•' else 'MUTED')] + runs
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(C[color])


def box(s, x, y, w, h, bg, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=9,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=6, lw=0.75):
    sh = s.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(C[bg])
    if line:
        sh.line.color.rgb = RGBColor.from_string(C[line]); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(3)
    fill_tf(tf, paras, anchor, align)
    return sh


def text(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, after=2, name=None):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, paras, anchor, align, after)
    if name:
        tb.name = name
    return tb


def arrow(s, x1, y1, x2, y2, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)); strip(c)
    c.line.color.rgb = RGBColor.from_string(C['GREY']); c.line.width = Pt(1)
    if head:
        ln = c._element.find('.//' + qn('a:ln'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')


def dot(s, x, y, label, bg='TEAL', d=20, sz=9):
    return box(s, x, y, d, d, bg, paras=[[(label, sz, True, 'WHITE')]], align=PP_ALIGN.CENTER, margin=0, shape=MSO_SHAPE.OVAL)


B = lambda t, c='INK', sz=8.5: (t, sz, True, c)
A = lambda t, c='INK', sz=8.5: (t, sz, False, c)


def tile(s, x, y, w, h, n, head, body, col='TEAL', pale='TEAL_PALE', sz=8.5):
    box(s, x, y, w, h, 'WHITE', 'LINE')
    dot(s, x + 10, y + 10, n, col, 22)
    text(s, x + 40, y + 12, w - 50, 18, [[B(head, col, 10.5)]])
    text(s, x + 12, y + 40, w - 24, h - 46, body, after=3)


EB = 'CAPITAL ONE CARD ON FILE'

# ---- What is this? ----
s = new_slide(EB, 'What is this?', "Homeaglow doesn't accept Capital One cards (credit or debit) for membership payments.")
tw = (W - 16) / 3
tile(s, X, 102, tw, 166, '1', 'What happens', [
    [A('When '), B('every'), A(' card on the account is Capital One, MFs and ETFs are '), B('skipped automatically'), A('.')],
    [A('There is '), B('no failed charge'), A(' in Stripe.')]])
tile(s, X + tw + 8, 102, tw, 166, '2', "How you'll spot it", [
    ('•', [A('Overdue MFs with nothing failed in Stripe')]),
    ('•', [A('C asks why they keep being told to update their card')]),
    ('•', [A('An FCF first cleaning cancelled for "no valid payment method"')])])
tile(s, X + 2 * (tw + 8), 102, tw, 166, '3', 'Why it matters', [
    [A('C almost always thinks something is wrong with their card or account. '), B('Nothing is.')],
    [A('The card works elsewhere. If you go looking for a decline, you won\'t find one, and saying it was declined isn\'t true.')]],
    col='AMBER', pale='AMBER_PALE')
box(s, X, 278, W, 50, 'AMBER_PALE', paras=[
    [B('This is our decision, so C doesn\'t pay to leave.  ', 'AMBER', 9.5)],
    [A('No cleaning yet → refund them. ETF already applies → waive it. Your job: offer the one fix that works (a card from another bank) and make leaving clean if they don\'t want it.', 'INK', 8.5)]], margin=12)
s.notes_slide.notes_text_frame.text = (
    "Homeaglow doesn't accept Capital One cards, credit or debit, for membership payments. When every card on an account is Capital One, MFs and ETFs are skipped "
    'automatically and there is no failed charge in Stripe. You will usually spot it one of three ways: overdue MFs with nothing failed in Stripe, a customer asking why they '
    'keep being told to update their card, or an FCF first cleaning cancelled for "no valid payment method".\n'
    'The customer almost always thinks something is wrong with their card or account. Nothing is: the card works fine elsewhere; we just don\'t accept it for memberships. '
    "If you go looking for a decline you won't find one, and telling the customer their card was declined is not true.\n"
    "Because this is our decision, not something the customer did, we don't make them pay to leave: if they haven't had a cleaning yet, we refund them; if an ETF already "
    'applies, we waive it. The job is to offer the one fix that works, a card from another bank, and make leaving clean if they don\'t want it.')

# ---- The rules ----
s = new_slide(EB, 'The rules first', 'Retention here means fixing the payment problem.')
LW = 318
rules = [('No way to charge it', [A('No manual charges, no CRM retries, no alternate descriptor. Another Capital One card won\'t work either: '), B('the block is by bank, not by card.')]),
         ('Retention = fixing the payment', [A("Don't offer a lower MF or a reduced ETF; neither helps if we still can't charge C. Capital One is an "),
                                             B('ETF waiver reason'), A(', so it overrides Graceful Closure like the others.')]),
         ("Don't give a reason", [A('Describe '), B('what we accept'), A(', not why.')])]
rh = (270 - 2 * 7) / 3
for i, (h, body) in enumerate(rules):
    y = 102 + i * (rh + 7)
    box(s, X, y, LW, rh, 'TEAL_PALE')
    dot(s, X + 10, y + 10, str(i + 1))
    text(s, X + 38, y + 10, LW - 48, rh - 14, [[B(h, 'TEAL', 10)], body], after=3)
RX, RW = X + LW + 12, W - LW - 12
rows = [('Say', 'Skip'),
        ('"We aren\'t able to accept Capital One cards for membership payments."', '"Your card was declined" / "insufficient funds"'),
        ('"This isn\'t about your card or your account."', '"Capital One is blocking us" / "there were disputes"'),
        ('"A card from a different bank will work."', '"It\'s temporary" / "we\'re working on a fix"')]
heights = [24, 82, 82, 82]
gt = s.shapes.add_table(4, 2, Pt(RX), Pt(102), Pt(RW), Pt(sum(heights)))
tbl = gt.table
tbl.columns[0].width = Pt(RW / 2); tbl.columns[1].width = Pt(RW / 2)
tblPr = tbl._tbl.tblPr
for a in ('bandRow', 'firstRow'):
    tblPr.set(a, '0')
sid = tblPr.find(qn('a:tableStyleId'))
if sid is not None:
    tblPr.remove(sid)
for i, row in enumerate(rows):
    tbl.rows[i].height = Pt(heights[i])
    for j, v in enumerate(row):
        c = tbl.cell(i, j)
        c.margin_left = c.margin_right = Pt(10); c.vertical_anchor = MSO_ANCHOR.MIDDLE
        bg = ('TEAL' if j == 0 else 'AMBER') if i == 0 else ('TEAL_PALE' if j == 0 else 'AMBER_PALE')
        c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string(C[bg])
        tcPr = c._tc.get_or_add_tcPr()
        for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
            ln = tcPr.makeelement(qn(tag), {'w': '12700'})
            sf = ln.makeelement(qn('a:solidFill'), {}); sf.append(sf.makeelement(qn('a:srgbClr'), {'val': 'FFFFFF'})); ln.append(sf)
            tcPr.append(ln)
        fe = tcPr.find(qn('a:solidFill')); tcPr.remove(fe); tcPr.insert(0, fe)
        p = c.text_frame.paragraphs[0]; c.text_frame.word_wrap = True
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        r = p.add_run(); r.text = ('✓  ' if j == 0 else '✕  ') + v if i == 0 else v
        r.font.size = Pt(10 if i == 0 else 9.5); r.font.bold = i == 0; r.font.name = 'Play' if i == 0 else 'Arial'
        r.font.color.rgb = RGBColor.from_string(C['WHITE'] if i == 0 else C['INK'])
s.notes_slide.notes_text_frame.text = (
    "The rules first. We can't charge Capital One cards for MFs or ETFs: no manual charges, no CRM retries, no alternate descriptor. Another Capital One card won't work either, "
    "because the block is by bank, not by card. Retention here means fixing the payment problem: don't offer a lower MF or a reduced ETF, because neither helps if we still can't "
    'charge the customer. Capital One is an ETF waiver reason, so it overrides Graceful Closure the same way the others do (see ETF Waiver Reasons Override Graceful Closure). '
    "Don't give a reason for the change: describe what we accept, not why.\n"
    'Say: "We aren\'t able to accept Capital One cards for membership payments." / "This isn\'t about your card or your account." / "A card from a different bank will work." '
    'Skip: "Your card was declined" or "insufficient funds" / "Capital One is blocking us" or "there were disputes" / "It\'s temporary" or "we\'re working on a fix".')

# ---- Step 1 ----
s = new_slide(EB, 'Step 1: Confirm it\'s a Capital One case', 'Check the issuer before you assume it\'s a bug.')
box(s, X, 102, 400, 166, 'WHITE', 'LINE')
text(s, X + 14, 112, 372, 150, [
    ('1', [A('In ', 'INK', 10), B('Stripe', 'INK', 10), A(', check ', 'INK', 10), B('every', 'INK', 10), A(' card on the account.', 'INK', 10)]),
    ('2', [A('Capital One shows as ', 'INK', 10), B('CAPITAL ONE, NATIONAL ASSOCIATION', 'TEAL', 10), A('.', 'INK', 10)]),
    ('3', [A('Every card is Capital One → ', 'INK', 10), B('this is the case', 'INK', 10), A('. Go to Step 2.', 'INK', 10)]),
    ('4', [A('Even ', 'INK', 10), B('one', 'INK', 10), A(' card from another bank → the system charges it and ', 'INK', 10),
           B('this article doesn\'t apply', 'INK', 10), A('.', 'INK', 10)])], after=7)
box(s, X + 412, 102, W - 412, 166, 'AMBER_PALE', anchor=MSO_ANCHOR.MIDDLE, margin=14, paras=[
    [B('💡 Quick tip', 'AMBER', 10.5)],
    [A('Overdue MFs with ', 'INK', 9.5), B('no failed charge', 'INK', 9.5), A(' in Stripe is the giveaway.', 'INK', 9.5)]])
box(s, X, 278, W, 44, 'TEAL_PALE', margin=14, paras=[
    [B('Example:  ', 'TEAL', 9.5), A('Capital One card plus a Chase card on file → not a Capital One case. The Chase card is charged, and standard rules apply.', 'INK', 9.5)]])
s.notes_slide.notes_text_frame.text = (
    'Step 1: confirm it is a Capital One case. In Stripe, check every card on the account. Capital One shows as CAPITAL ONE, NATIONAL ASSOCIATION. '
    "If even one card is from another bank, the system charges it and this article doesn't apply. Quick tip: overdue MFs with no failed charge in Stripe is the giveaway, "
    "so check the issuer before you assume it's a bug. Example: a Capital One card plus a Chase card on file is not a Capital One case; the Chase card is charged and standard rules apply.")

# ---- Step 2: which branch ----
s = new_slide(EB, 'Step 2: Does an ETF apply?', 'Each branch stands on its own.')
cx = X + W / 2
box(s, cx - 120, 100, 240, 26, 'INK', paras=[[B('Every card is Capital One', 'WHITE', 9.5)]], align=PP_ALIGN.CENTER)
arrow(s, cx, 126, cx, 136)
box(s, cx - 120, 136, 240, 26, 'TEAL', paras=[[B('Check: does an ETF apply?', 'WHITE', 9.5)]], align=PP_ALIGN.CENTER)
bw = (W - 20) / 3
mids = [X + i * (bw + 10) + bw / 2 for i in range(3)]
arrow(s, cx, 162, cx, 172, head=False)
arrow(s, mids[0], 172, mids[2], 172, head=False)
for m in mids:
    arrow(s, m, 172, m, 184)
branches = [('Branch 1: No ETF yet', 'FCF, no completed cleaning', [A('Ask for a card from a different bank. No credit offer.')], 'TEAL', 'TEAL_PALE'),
            ('Branch 2: ETF applies', 'Completed cleaning, within MCT', [A('Ask for a card from a different bank (+ $20 credit once added). Won\'t add one → you cancel and waive the ETF.')], 'AMBER', 'AMBER_PALE'),
            ('Beyond MCT', 'No ETF applies', [A('Standard cancellation.')], 'INK', 'ALT')]
for i, (h, sub, body, col, pale) in enumerate(branches):
    x = X + i * (bw + 10)
    box(s, x, 184, bw, 96, pale, col if col != 'INK' else 'LINE', anchor=MSO_ANCHOR.TOP, margin=11, paras=[
        [B(h, col, 10.5)], [A(sub, 'MUTED', 8)], body])
box(s, X, 292, W, 50, 'WHITE', 'LINE', margin=14, paras=[
    [B('Overdue MFs (both branches): ', 'TEAL', 9), A('same as any other overdue MF. Nothing changes because the card is Capital One.', 'INK', 9)],
    [A('Whether to raise it: ', 'MUTED', 8), B('Failed Job Charge', 'INK', 8), A(' or ', 'MUTED', 8), B('Overdue Membership Fee', 'INK', 8),
     A('. How many to bring current: ', 'MUTED', 8), B('Understanding Offer Logic → Free Months (Job 2)', 'INK', 8), A('.', 'MUTED', 8)]])
s.notes_slide.notes_text_frame.text = (
    'Step 2: check whether an ETF applies, then follow that branch. Each branch stands on its own. Branch 1, no ETF yet: FCF with no completed cleaning. '
    'Branch 2, ETF applies: a completed cleaning, within MCT. Beyond MCT, no ETF applies, so it is a standard cancellation.\n'
    'Overdue MFs apply to both branches and are handled the same as any other overdue MF: use Failed Job Charge or Overdue Membership Fee for whether to raise it, '
    'and Understanding Offer Logic → Free Months (Job 2) for how many to bring current.')

# ---- Branch 1 ----
s = new_slide(EB, 'Branch 1: No ETF yet', 'FCF, no completed cleaning.')
box(s, X, 100, W, 44, 'TEAL_PALE', margin=14, paras=[
    [A('Explain that we aren\'t able to accept Capital One cards for membership payments, and ', 'INK', 9.5), B('ask for a card from a different bank', 'INK', 9.5),
     A('. No credit offer here: the first cleaning is already free or heavily discounted.', 'INK', 9.5)]])
tw = (W - 16) / 3
outs = [('They add one', [A('Keep the upcoming job, or '), B('rebook'), A(' it if it was already cancelled.')], 'TEAL'),
        ('Can\'t or won\'t, want to cancel', [B('Cancel and refund everything paid'), A(': the $19 MF plus any add-on, even past the 24-hour window.')], 'AMBER'),
        ('Won\'t cancel, but only want Capital One', [A('Offer both paths plainly: add a card from a different bank, or cancel and refund. '), B('Let them choose.')], 'TEAL')]
for i, (h, body, col) in enumerate(outs):
    x = X + i * (tw + 8)
    box(s, x, 152, tw, 112, 'WHITE', 'LINE')
    text(s, x + 12, 162, tw - 24, 18, [[B(h, col, 10)]])
    text(s, x + 12, 192, tw - 24, 68, [body])
box(s, X, 272, W, 30, 'ALT', margin=14, paras=[
    [B('Example:  ', 'TEAL', 9), A('FCF C paid $19 plus the $19 4-hour add-on, first cleaning cancelled, no other card → cancel and refund $38.', 'INK', 9)]])
box(s, X, 310, W, 34, 'WHITE', 'TEAL', margin=14, paras=[
    [B('Macros:  ', 'TEAL', 9), A('Capital One – Add a Different Card (No Cleaning Yet)', 'INK', 9), A('   ·   ', 'MUTED', 9),
     A('Capital One – Cancel and Refund (No Cleaning Yet)', 'INK', 9)]])
s.notes_slide.notes_text_frame.text = (
    "Branch 1, no ETF yet (FCF, no completed cleaning). Explain that we aren't able to accept Capital One cards for membership payments, and ask the customer to add a card "
    'from a different bank. There is no credit offer here, because the first cleaning is already free or heavily discounted. If they add one, keep the upcoming job, or rebook '
    "it if it was already cancelled. If they can't or won't and want to cancel, cancel and refund everything paid: the $19 MF plus any add-on, even past the 24-hour window. "
    "If they don't want to cancel but only want to use Capital One, offer both paths plainly (add a card from a different bank, or cancel and refund) and let them choose. "
    'Example: an FCF customer paid $19 plus the $19 4-hour add-on, the first cleaning was cancelled, and there is no other card: cancel and refund $38. '
    'Macros: Capital One – Add a Different Card (No Cleaning Yet); Capital One – Cancel and Refund (No Cleaning Yet).')

# ---- Branch 2 ----
s = new_slide(EB, 'Branch 2: ETF applies', 'Completed cleaning, within MCT.')
LW = 330
box(s, X, 100, LW, 92, 'WHITE', 'LINE', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('C wants to cancel? Is the card the only reason?', 'TEAL', 10)],
    [A('Something else too (e.g. a bad cleaning) → '), B('handle both in the same reply'), A(': resolve the other issue with its own playbook, and ask for a card from a different bank.')]])
box(s, X, 200, LW, 96, 'AMBER_PALE', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('$20 credit for adding a card', 'AMBER', 10)],
    [A('Once C adds a card from a different bank, add '), B('$20 in credits'), A(' toward their next cleaning. Offer it once; it doesn\'t go up if they say no.')],
    [A('The other playbook already gives a credit or voucher (e.g. a PCQ voucher)? That '), B('replaces'), A(' the $20. Don\'t stack them.')]])
box(s, X, 304, LW, 40, 'ALT', margin=12, paras=[[B('Macro:  ', 'TEAL', 8.5), A('Capital One – Add a Card + $20 Credit')]])
RX, RW = X + LW + 12, W - LW - 12
box(s, RX, 100, RW, 74, 'TEAL_PALE', 'TEAL', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('They add a card', 'TEAL', 10)],
    [A('Now a normal case. Still want to cancel → standard '), B('FC Retention and ETF'), A(' guidance.')]])
box(s, RX, 182, RW, 92, 'AMBER_PALE', 'AMBER', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('They won\'t add a card', 'AMBER', 10)],
    [B('Cancel the membership yourself'), A(' and '), B('waive the ETF'), A('.')],
    [B('Don\'t send the self-cancel link: ', 'AMBER'), A('self-cancel tries to charge the ETF and can leave the account stuck.')],
    [A('Macro: Capital One – Cancellation, No Fee', 'MUTED', 8)]])
box(s, RX, 282, RW, 62, 'WHITE', 'LINE', anchor=MSO_ANCHOR.TOP, margin=12, paras=[
    [B('Examples', 'MUTED', 8.5)],
    [A('PCQ + Capital One, wants to cancel → PCQ playbook offer (e.g. a voucher) + ask for a card. No extra $20.', 'INK', 8)],
    [A('Legacy DHJ, 3 paid months, MCT not done, only Capital One, won\'t add one → you cancel, no ETF.', 'INK', 8)]])
s.notes_slide.notes_text_frame.text = (
    'Branch 2, ETF applies (completed cleaning, within MCT). If the customer wants to cancel, check whether the card is the only reason. If there is something else too, like a bad '
    'cleaning, handle both in the same reply: resolve the other issue with its own playbook, and ask for a card from a different bank.\n'
    "For the card: once they add a card from a different bank, add $20 in credits toward their next cleaning. Offer it once; it doesn't go up if they say no. If the other issue's "
    "playbook already gives a credit or voucher (for example a PCQ voucher), that replaces the $20. Don't stack them: both exist to get the customer to book again.\n"
    "If they add a card, this is now a normal case: if they still want to cancel, use standard FC Retention and ETF guidance. If they won't add a card, cancel the membership "
    "yourself and waive the ETF. Don't send the self-cancel link, because self-cancel tries to charge the ETF and can leave the account stuck.\n"
    'Examples: PCQ plus Capital One and wants to cancel: the PCQ playbook offer (e.g. a voucher) plus the ask for a card from another bank, no extra $20. '
    "A Legacy DHJ member with 3 paid months, MCT not done, only a Capital One card, who won't add one: you cancel, no ETF.")

# ---- Macros ----
s = new_slide(EB, 'Suggested macros', 'Pick the macro for the branch and what C decides.')
macros = [('Capital One – Add a Different Card (No Cleaning Yet)', 'Branch 1: ask for another card', 'TEAL',
           'Thanks for reaching out! Your card itself is fine, and this isn\'t anything you did. We just aren\'t able to accept Capital One cards for membership payments. '
           'To keep your membership and your cleaning, you can add a card from a different bank here: <<update card link>>. If that doesn\'t work for you, I can cancel your '
           'membership and refund what you\'ve paid instead. Just let me know which you\'d prefer!'),
          ('Capital One – Cancel and Refund (No Cleaning Yet)', 'Branch 1: C chose to cancel', 'TEAL',
           'Thanks for letting me know. I\'ve cancelled your membership and refunded the <<amount paid>> you paid. You\'ll see it back within 5–10 business days. '
           'If you\'d ever like to try again with a card from a different bank, we\'d love to have you.'),
          ('Capital One – Add a Card + $20 Credit', 'Branch 2: ask for another card', 'AMBER',
           'Thanks for reaching out about this. Your card is fine, and this isn\'t anything you did. We just aren\'t able to accept Capital One cards for membership payments. '
           'If you\'d like to keep your cleanings going, you can add a card from a different bank here: <<update card link>>. Once it\'s added, I\'ll put $20 in credits toward your next cleaning.'),
          ('Capital One – Cancellation, No Fee', 'Branch 2: C won\'t add a card', 'AMBER',
           'Thanks for letting me know. Since we aren\'t able to accept the card on your account, I\'ve cancelled your membership, and there\'s no cancellation fee. '
           'You won\'t be charged anything further. If you\'d ever like to come back, just add a card from a different bank and we\'d be happy to get you booked.')]
mw = (W - 10) / 2; mh = 104
for i, (t, when, col, body) in enumerate(macros):
    r, c = divmod(i, 2)
    x = X + c * (mw + 10); y = 100 + r * (mh + 8)
    box(s, x, y, mw, mh, 'WHITE', 'LINE')
    box(s, x, y, mw, 34, 'TEAL_PALE' if col == 'TEAL' else 'AMBER_PALE', margin=10, paras=[[B(t, col, 8.5)], [A(when, 'MUTED', 7)]])
    text(s, x + 10, y + 40, mw - 20, mh - 44, [[A(body, 'INK', 7.2)]])
box(s, X, 324, W, 48, 'ALT', margin=12, anchor=MSO_ANCHOR.MIDDLE, paras=[
    [B('If they ask why: ', 'TEAL', 7.8), A('"That\'s a decision on our side about which cards we accept for memberships. It isn\'t anything to do with your card."', 'INK', 7.8)],
    [B('"It worked for the voucher": ', 'TEAL', 7.8), A('"It can go through for the initial purchase, but we can\'t use it for ongoing membership payments."', 'INK', 7.8)],
    [B('Don\'t use: ', 'AMBER', 7.8), A('the "Credit Card failed on FC Charge" macros. They say the card "may have expired or is no longer valid", which isn\'t the case here.', 'INK', 7.8)]])
s.notes_slide.notes_text_frame.text = (
    'Four macros. Branch 1: Capital One – Add a Different Card (No Cleaning Yet) to ask for another card, and Capital One – Cancel and Refund (No Cleaning Yet) once the customer '
    'chooses to cancel. Branch 2: Capital One – Add a Card + $20 Credit to ask for another card, and Capital One – Cancellation, No Fee if they won\'t add one.\n'
    'If they ask why: "That\'s a decision on our side about which cards we accept for memberships. It isn\'t anything to do with your card." If they say it worked for the voucher: '
    '"It can go through for the initial purchase, but we can\'t use it for ongoing membership payments." Don\'t use the "Credit Card failed on FC Charge" macros for Capital One, '
    'because they say the card "may have expired or is no longer valid", which isn\'t the case here.')

# ---- Knowledge check ----
s = new_slide('CAPITAL ONE CARD ON FILE  ·  KNOWLEDGE CHECK', 'Let’s See What You’ve Got!', 'Ask each question, then click to reveal the answer.')
qa = [('C has overdue MFs but nothing failed in Stripe. What do you check first?',
       'Every card on the account in Stripe. All CAPITAL ONE, NATIONAL ASSOCIATION → Capital One case.'),
      ('C has a Capital One card and a Chase card on file. Is this a Capital One case?',
       'No. The Chase card gets charged, so standard rules apply.'),
      ('C asks to try another Capital One card. Will that work?',
       "No. The block is by bank, not by card. Ask for a card from a different bank."),
      ('FCF C, no completed cleaning, won\'t add another card and wants to cancel. Paid $19 + $19 add-on.',
       'Cancel and refund everything paid: $38, even past the 24-hour window.'),
      ('Within MCT, C won\'t add a card and wants to cancel. Do you send the self-cancel link?',
       "No. Cancel it yourself and waive the ETF. Self-cancel tries to charge the ETF and can leave the account stuck."),
      ('PCQ + Capital One, within MCT. C gets a PCQ voucher. Do you also add the $20 credit?',
       "No. The PCQ voucher replaces the $20. Don't stack them.")]
n = len(qa); gap = 5; h = (272 - gap * (n - 1)) / n
groups = []
for i, (q, a) in enumerate(qa):
    y = 100 + i * (h + gap)
    box(s, X, y, 364, h, 'WHITE', 'LINE')
    box(s, X + 10, y + (h - 20) / 2, 20, 20, 'TEAL_PALE', paras=[[(str(i + 1), 9, True, 'TEAL')]], align=PP_ALIGN.CENTER, margin=0, shape=MSO_SHAPE.OVAL)
    text(s, X + 40, y, 316, h, [[A(q, 'INK', 8.5)]], anchor=MSO_ANCHOR.MIDDLE)
    b = box(s, X + 371, y, W - 371, h, 'TEAL_PALE')
    t = text(s, X + 381, y, W - 391, h, [[A(a, 'INK', 8.5)]], anchor=MSO_ANCHOR.MIDDLE)
    groups.append([b.shape_id, t.shape_id])


def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="400"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


cid, clicks = 3, ''
for ids in groups:
    start, inner = cid, ''
    cid += 2
    for i, spid in enumerate(ids):
        inner += effect(cid, spid, 'clickEffect' if i == 0 else 'withEffect'); cid += 3
    clicks += (f'<p:par><p:cTn id="{start}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
               f'<p:par><p:cTn id="{start+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
               f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
bld = ''.join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for g in groups for i in g)
NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
       '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + clicks +
       '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
       '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
       '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
ext = s._element.find(qn('p:extLst'))
t_el = etree.fromstring(tim)
(ext.addprevious(t_el) if ext is not None else s._element.append(t_el))
s.notes_slide.notes_text_frame.text = (
    'Quick knowledge check: Capital One Card on File. Ask each question and let trainees answer before revealing. Each click reveals the next answer.\n'
    '1) Overdue MFs with no failed charge is the giveaway: check every card in Stripe for CAPITAL ONE, NATIONAL ASSOCIATION. '
    '2) Even one card from another bank means the system charges it, so this article does not apply. 3) The block is by bank, not by card. '
    '4) Branch 1: cancel and refund everything paid, the $19 MF plus any add-on, even past the 24-hour window. '
    '5) Branch 2: cancel the membership yourself and waive the ETF; never send the self-cancel link. 6) A credit or voucher from the other playbook replaces the $20; do not stack them.')

order.append(questions)
questions.notes_slide.notes_text_frame.text = 'Open the floor for questions. When in doubt, check the Knowledge Library (Capital One Card on File) or ask your Trainer or Team Lead.'

# keep only the slides we built, in order
want = [id(s) for s in order]
all_slides = list(prs.slides)
els = list(lst)
pairs = dict((id(sl), e) for sl, e in zip(all_slides, els))
for sl, e in zip(all_slides, els):
    if id(sl) not in want:
        rid = e.get(qn('r:id')); lst.remove(e); prs.part.drop_rel(rid)
for e in list(lst):
    lst.remove(e)
for sl in order:
    lst.append(pairs[id(sl)])
prs.save(sys.argv[2]); print('saved', sys.argv[2], len(prs.slides), 'slides')
