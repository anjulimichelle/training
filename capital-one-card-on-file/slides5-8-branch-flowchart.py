"""Capital One deck: merge slides 5-8 (Step 1, Step 2, Branch 1, Branch 2) into one flow-chart
slide: Step 1 -> Step 2 -> three branches (Branch 1 No ETF yet, Branch 2 ETF applies, Beyond MCT).
Branch details are hidden at first; each click opens the next branch in the panel (one at a time).
Related macros listed along the bottom.

Usage: python3 slides5-8-branch-flowchart.py in.pptx out.pptx
"""
import sys
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

C = {'TEAL': '0B7A6F', 'TEAL_PALE': 'E3F4F1', 'AMBER': '9A6410', 'AMBER_PALE': 'FFF3D6',
     'INK': '1D1D1F', 'MUTED': '6E6E73', 'WHITE': 'FFFFFF', 'LINE': 'DCE1E7', 'GREY': '6E6E73', 'ALT': 'F5F6F8'}
X, W = 32.4, 655.2

prs = Presentation(sys.argv[1])
txt = lambda sh: sh.text_frame.text if sh.has_text_frame else ''
lst = prs.slides._sldIdLst
S = list(prs.slides)
titles = ["Step 1: Confirm it's a Capital One case", 'Step 2: Does an ETF apply?', 'Branch 1: No ETF yet', 'Branch 2: ETF applies']
idx = [next(i for i, s in enumerate(S) if txt(s.shapes[1]) == t) for t in titles]
s = S[idx[0]]
notes = '\n'.join(S[i].notes_slide.notes_text_frame.text for i in idx)
for i in idx[1:]:
    e = list(lst)[i - (idx[1:].index(i))]
for i in sorted(idx[1:], reverse=True):
    e = list(lst)[i]
    rid = e.get(qn('r:id')); lst.remove(e); prs.part.drop_rel(rid)

sl = list(s.shapes)
keep = {id(x._element) for x in sl[:3]} | {id(x._element) for x in sl if x.is_placeholder}
for sh in sl:
    if id(sh._element) not in keep:
        sh._element.getparent().remove(sh._element)
for k, t in ((1, 'Which branch is it?'), (2, 'Confirm the case, check for an ETF, then follow one branch. Click to open each branch.')):
    p = s.shapes[k].text_frame.paragraphs[0]
    p.runs[0].text = t
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
t_old = s._element.find(qn('p:timing'))
if t_old is not None:
    s._element.remove(t_old)


def strip(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None:
        sh._element.remove(st)


def fill_tf(tf, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, after=2):
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, item in enumerate(paras):
        num, runs = item if isinstance(item, tuple) else ('', item)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after); p.alignment = align
        if num:
            pPr = p._p.get_or_add_pPr()
            pPr.set('marL', str(int(11 * 12700))); pPr.set('indent', str(-9 * 12700))
            runs = [(num + ' ', runs[0][1], True, 'MUTED')] + runs
        for t, size, bold, color in runs:
            r = p.add_run(); r.text = t; f = r.font
            f.name = 'Play' if bold else 'Arial'; f.size = Pt(size); f.bold = bold
            f.color.rgb = RGBColor.from_string(C[color])


def box(x, y, w, h, bg, line=None, paras=(), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=9, radius=6, lw=0.75):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)); strip(sh)
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(C[bg])
    if line:
        sh.line.color.rgb = RGBColor.from_string(C[line]); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(3)
    fill_tf(tf, paras, anchor, align)
    return sh


def text(x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, after=2):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, paras, anchor, align, after)
    return tb


def arrow(x1, y1, x2, y2, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)); strip(c)
    c.line.color.rgb = RGBColor.from_string(C['GREY']); c.line.width = Pt(1)
    if head:
        ln = c._element.find('.//' + qn('a:ln'))
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')


B = lambda t, c='INK', sz=8: (t, sz, True, c)
A = lambda t, c='INK', sz=8: (t, sz, False, c)

# steps
sw = 312
box(X, 98, sw, 54, 'WHITE', 'TEAL', anchor=MSO_ANCHOR.MIDDLE, margin=10, paras=[
    [B('Step 1: Confirm it\'s a Capital One case', 'TEAL', 9.5)],
    [A('In Stripe, every card shows '), B('CAPITAL ONE, NATIONAL ASSOCIATION'), A('. Even one card from another bank → not this case.')]])
arrow(X + sw, 125, X + sw + 18, 125)
s2x = X + sw + 18; s2w = W - sw - 18
box(s2x, 98, s2w, 54, 'TEAL', anchor=MSO_ANCHOR.MIDDLE, margin=10, paras=[
    [B('Step 2: Does an ETF apply?', 'WHITE', 9.5)],
    [A('Each branch stands on its own. Overdue MFs: handle like any other overdue MF (both branches).', 'TEAL_PALE')]])
# branch headers
bw = (W - 20) / 3
mids = [X + i * (bw + 10) + bw / 2 for i in range(3)]
cx2 = s2x + s2w / 2
arrow(cx2, 152, cx2, 160, head=False)
arrow(mids[0], 160, mids[2], 160, head=False)
for m in mids:
    arrow(m, 160, m, 168)
heads = [('1  Branch 1: No ETF yet', 'FCF, no completed cleaning', 'TEAL', 'TEAL_PALE'),
         ('2  Branch 2: ETF applies', 'Completed cleaning, within MCT', 'AMBER', 'AMBER_PALE'),
         ('3  Beyond MCT', 'No ETF applies', 'INK', 'ALT')]
hy, hh = 168, 32
for i, (h, sub, col, pale) in enumerate(heads):
    box(X + i * (bw + 10), hy, bw, hh, 'WHITE', 'LINE', align=PP_ALIGN.CENTER, margin=6,
        paras=[[B(h, col if col != 'INK' else 'INK', 9)], [A(sub, 'MUTED', 7)]])
py, ph = hy + hh + 8, 128
# placeholder (state 0)
st0 = [box(X, py, W, ph, 'ALT', 'LINE', align=PP_ALIGN.CENTER, paras=[
    [B('Click to open each branch', 'TEAL', 11)], [A('Branch 1  →  Branch 2  →  Beyond MCT, one at a time.', 'MUTED', 8.5)]]).shape_id]
states = [st0]
half = (W - 30) / 2


def panel(i, col, pale, left, right):
    ids = []
    hl = box(X + i * (bw + 10), hy, bw, hh, pale, col if col != 'INK' else 'TEAL', align=PP_ALIGN.CENTER, margin=6, lw=1.5,
             paras=[[B(heads[i][0], col, 9)], [A(heads[i][1], 'MUTED', 7)]])
    pn = box(X, py, W, ph, pale, col if col != 'INK' else 'LINE', lw=1)
    ids += [hl.shape_id, pn.shape_id]
    ids.append(text(X + 12, py + 9, half, ph - 16, left, after=2.5).shape_id)
    if right:
        ids.append(text(X + 22 + half, py + 9, half, ph - 16, right, after=2.5).shape_id)
    states.append(ids)


panel(0, 'TEAL', 'TEAL_PALE', [
    [B('Ask for a card from a different bank.', 'TEAL', 9)],
    [A('Explain we aren\'t able to accept Capital One cards for membership payments. '), B('No credit offer'),
     A(': the first cleaning is already free or heavily discounted.')],
    [B('Macros: ', 'MUTED', 7.5), A('Capital One – Add a Different Card (No Cleaning Yet) · Capital One – Cancel and Refund (No Cleaning Yet)', 'MUTED', 7.5)]], [
    ('•', [B('They add one → '), A('keep the upcoming job, or rebook it if it was already cancelled.')]),
    ('•', [B("Can't or won't, want to cancel → "), A('cancel and refund everything paid: the $19 MF plus any add-on, even past the 24-hour window.')]),
    ('•', [B("Won't cancel, but only want Capital One → "), A('offer both paths plainly (another bank\'s card, or cancel and refund). Let them choose.')]),
    [B('Example: ', 'TEAL', 7.5), A('paid $19 + $19 4-hour add-on, cleaning cancelled, no other card → cancel and refund $38.', 'INK', 7.5)]])
panel(1, 'AMBER', 'AMBER_PALE', [
    [B('Wants to cancel? Is the card the only reason?', 'AMBER', 9)],
    [A('Something else too (e.g. a bad cleaning) → '), B('handle both in the same reply'), A(': its own playbook + ask for a card.')],
    [B('$20 credit: ', 'AMBER'), A('once C adds a card from another bank. Offer it once; it doesn\'t go up. A credit or voucher from the other playbook (e.g. PCQ) '),
     B('replaces'), A(' it. Don\'t stack.')],
    [B('Macro: ', 'MUTED', 7.5), A('Capital One – Add a Card + $20 Credit', 'MUTED', 7.5)]], [
    ('•', [B('They add a card → '), A('normal case. Still want to cancel → standard FC Retention and ETF guidance.')]),
    ('•', [B("They won't add a card → "), A('cancel the membership yourself and waive the ETF. '), B("Don't send the self-cancel link", 'AMBER'),
           A(': it tries to charge the ETF and can leave the account stuck.')]),
    [B('Macro: ', 'MUTED', 7.5), A('Capital One – Cancellation, No Fee', 'MUTED', 7.5)],
    [B('Examples: ', 'AMBER', 7.5), A('PCQ + Capital One → PCQ offer + card ask, no extra $20. Legacy DHJ, 3 paid months, only Capital One, won\'t add one → you cancel, no ETF.', 'INK', 7.5)]])
panel(2, 'INK', 'ALT', [
    [B('Beyond MCT: no ETF applies.', 'INK', 10)],
    [A('Standard cancellation.', 'INK', 9.5)]], [
    [B('Overdue MFs (both branches)', 'TEAL', 9)],
    [A('Same as any other overdue MF; nothing changes because the card is Capital One.')],
    [A('Whether to raise it: '), B('Failed Job Charge'), A(' or '), B('Overdue Membership Fee'), A('.')],
    [A('How many to bring current: '), B('Understanding Offer Logic → Free Months (Job 2)'), A('.')]])

# related macros
my = py + ph + 8
text(X, my, 120, 12, [[B('RELATED MACROS', 'MUTED', 7.5)]])
mw = (W - 18) / 4
for i, m in enumerate(['Capital One – Add a Different Card (No Cleaning Yet)', 'Capital One – Cancel and Refund (No Cleaning Yet)',
                       'Capital One – Add a Card + $20 Credit', 'Capital One – Cancellation, No Fee']):
    box(X + i * (mw + 6), my + 13, mw, 26, 'WHITE', 'TEAL', align=PP_ALIGN.CENTER, margin=4,
        paras=[[B(m, 'TEAL', 7), A('  new', 'AMBER', 6.5)]])


def effect(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{cid+2}" dur="300"/>{tgt}</p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def leave(cid, spid, node):
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    return (f'<p:par><p:cTn id="{cid}" presetID="1" presetClass="exit" presetSubtype="0" fill="hold" grpId="1" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>'
            f'</p:childTnLst></p:cTn></p:par>')


cid, clicks = 3, ''
for k in range(1, len(states)):
    start, inner = cid, ''
    cid += 2
    for j, spid in enumerate(states[k - 1]):
        inner += leave(cid, spid, 'clickEffect' if j == 0 else 'withEffect'); cid += 2
    for spid in states[k]:
        inner += effect(cid, spid, 'withEffect'); cid += 3
    clicks += (f'<p:par><p:cTn id="{start}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
               f'<p:par><p:cTn id="{start+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
               f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
bld = ''
for k, ids in enumerate(states):
    for i in ids:
        if k > 0:
            bld += f'<p:bldP spid="{i}" grpId="0"/>'
        if k < len(states) - 1:
            bld += f'<p:bldP spid="{i}" grpId="1"/>'
NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
tim = ('<p:timing ' + NS + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
       '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + clicks +
       '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
       '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
       '</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + bld + '</p:bldLst></p:timing>')
ext = s._element.find(qn('p:extLst'))
t_el = etree.fromstring(tim)
(ext.addprevious(t_el) if ext is not None else s._element.append(t_el))
s.notes_slide.notes_text_frame.text = (
    'TRAINER: the branch details start hidden. Click 1 opens Branch 1 (no ETF yet), click 2 swaps in Branch 2 (ETF applies), click 3 swaps in Beyond MCT; '
    'the next click moves to the next slide.\n' + notes)
prs.save(sys.argv[2]); print('saved', sys.argv[2], len(prs.slides), 'slides')
