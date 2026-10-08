"""Retention deck slide 59 (Unused DHJ Voucher, Case 4): add the leftover-voucher CRM screenshot and the
related macros. Patches the live Google export so Anjuli's edits elsewhere stay exactly as they are."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst, img = sys.argv[1:4]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, RED = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', 'B3261E'
prs = Presentation(src)
s = prs.slides[58]
assert s.shapes[1].text_frame.text == 'Case 4: Already self-refunded'

def box(x, y, w, h, fill, line=None):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    st = b._element.find(P_NS + 'style')
    if st is not None: b._element.remove(st)
    b.adjustments[0] = 0.06; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
    else: b.line.fill.background()
    return b
def text(x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, after=3):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(after)
        for t_, size, bold, col in runs:
            r = p.add_run(); r.text = t_; r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
    return t

Y = 3.88
# Left: the leftover voucher as it shows in the CRM.
box(0.45, Y, 4.6, 1.3, WHITE, BORDER)
text(0.6, Y + 0.08, 4.3, 0.2, [[('LEFTOVER VOUCHER IN THE CRM', 8, True, SOFT), ('   look for “Refunded but still valid”', 8, False, RED)]])
w = 3.55; h = w * 90 / 349
pic = s.shapes.add_picture(img, Inches(0.6), Inches(Y + 0.32), Inches(w), Inches(h))
pic.line.color.rgb = RGBColor.from_string(BORDER); pic.line.width = Pt(0.75)
# Right: related macros.
box(5.2, Y, 4.35, 1.3, TEALSOFT, TEAL)
text(5.38, Y + 0.08, 4.0, 0.2, [[('RELATED MACROS', 8, True, TEAL)]])
macros = ['Unused DHJ Voucher — Refund Request', 'Unused DHJ Voucher — Wants to ‘Try it Out’ (Trial Cleaning)',
          'Unused DHJ Voucher — Wants One Time Cleaning', 'Deactivated (send before deactivating)']
text(5.38, Y + 0.32, 4.05, 0.95, [[('•  ' + m, 9.5, False, INK)] for m in macros], after=2)

n = s.notes_slide.notes_text_frame
n.text = n.text.rstrip() + (' ON THE SLIDE: how a leftover voucher looks in the CRM after a self-refund: the DHJ voucher row is tagged '
    '"Refunded but still valid". That is the system-generated voucher to invalidate. Related macros: Unused DHJ Voucher — Refund Request; '
    'Unused DHJ Voucher — Wants to \'Try it Out\' (Trial Cleaning); Unused DHJ Voucher — Wants One Time Cleaning; Deactivated (send before deactivating).')
prs.save(dst)
