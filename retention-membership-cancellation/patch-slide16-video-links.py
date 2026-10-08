"""Retention deck slide 16 (Vouchers & credits): two 'watch' buttons linking to the how-to videos.
Patches the live Google export so Anjuli's own edits stay exactly as they are."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

src, dst = sys.argv[1:3]
prs = Presentation(src)
s = prs.slides[15]
assert s.shapes[1].text_frame.text == 'Vouchers & credits'
links = [('▶  Watch: Issuing credits', 'https://drive.google.com/file/d/1rKOi9--Q23bl4Ybf7m7bFXzy4p9xjSHc/view?usp=drive_link'),
         ('▶  Watch: Issuing a voucher', 'https://drive.google.com/file/d/14NVsd3tL6L5rAQEvbCbj0QSR-7-rZN7D/view?usp=drive_link')]
w, gap = 2.9, 0.15
for i, (label, url) in enumerate(links):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45 + i * (w + gap)), Inches(4.65), Inches(w), Inches(0.4))
    st = b._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: b._element.remove(st)
    b.name = f'videoLink{i + 1}'; b.adjustments[0] = 0.2
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string('0B7A6F'); b.line.color.rgb = RGBColor.from_string('0B7A6F')
    b.click_action.hyperlink.address = url
    tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.size = Pt(10.5); r.font.bold = True
    r.font.color.rgb = RGBColor.from_string('FFFFFF'); r.font.name = 'SF Pro Display'
n = s.notes_slide.notes_text_frame
n.text = n.text.rstrip() + ' TRAINER: in slideshow mode, click "Watch: Issuing credits" or "Watch: Issuing a voucher" to play that how-to video.'
prs.save(dst)
