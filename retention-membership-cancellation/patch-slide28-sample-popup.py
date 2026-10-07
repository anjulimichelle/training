"""Slide 28 (Cancellation Intent: Get the opening right): add a 'see a sample' button and a
pop-up with the triggered ticket and the CP's message (appears on click 1, closes on click 2).
Patches the live Google export so Anjuli's own edits stay exactly as they are."""
import re, sys, zipfile, io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

src, dst, img1, img2 = sys.argv[1:5]
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE = (RGBColor.from_string(c) for c in ('0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF'))
prs = Presentation(src)
s = prs.slides[27]
assert 'Get the opening right' in ''.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)

def text(shape, runs, size, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = align
    for t, bold, color in runs:
        r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = 'SF Pro Display'

def rect(x, y, w, h, fill, line, name, kind=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h)); sh.name = name
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line; sh.line.width = Pt(0.75); sh.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = 0.04
    st = sh._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: sh._element.remove(st)  # no theme shadow
    return sh

# Cue button (always visible).
b = rect(0.45, 3.85, 4.4, 0.45, TEAL, TEAL, 'sampleButton'); b.adjustments[0] = 0.18
text(b, [('▶  See a sample: triggered ticket + CP’s message', True, WHITE)], 10.5, PP_ALIGN.CENTER)
T = s.shapes.add_textbox(Inches(4.95), Inches(3.85), Inches(4.6), Inches(0.45)); T.name = 'sampleHint'
text(T, [('Click once to show both, click again to close.', False, SOFT)], 9)

# Pop-up: panel, two labelled screenshots.
rect(0.38, 1.3, 9.24, 3.92, WHITE, BORDER, 'pop1_2Panel')
def label(n, t, y, name):
    L = s.shapes.add_textbox(Inches(0.7), Inches(y), Inches(8.6), Inches(0.25)); L.name = name
    text(L, [(f'{n}  ', True, TEAL), (t[0], True, INK), (t[1], False, SOFT)], 10)
label(1, ('Triggered ticket', '  — what Care takes action on'), 1.42, 'pop1_2Label1')
p1 = s.shapes.add_picture(img1, Inches(0.7), Inches(1.7), width=Inches(4.9)); p1.name = 'pop1_2Img1'
y2 = 1.7 + p1.height / 914400 + 0.12
label(2, ('CP’s message', '  — the cleaner reporting that the customer wants to cancel'), y2, 'pop1_2Label2')
p2 = s.shapes.add_picture(img2, Inches(0.7), Inches(y2 + 0.28), width=Inches(6.8)); p2.name = 'pop1_2Img2'
for p in (p1, p2):
    p.line.color.rgb = BORDER; p.line.width = Pt(0.75)
print('pop-up bottom', round(y2 + 0.28 + p2.height / 914400, 2))
buf = io.BytesIO(); prs.save(buf)

# Click animation for slide 28 only (the rest of the deck keeps Google's own timing untouched).
ns = {}
exec(open(__file__.replace('patch28.py', 'add_story_anim.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == 'ppt/slides/slide28.xml':
        x = data.decode('utf8'); assert '<p:timing' not in x
        ids = re.findall(r'<p:cNvPr id="(\d+)" name="pop1_2[^"]*"', x)
        cid, ins, outs, bld = 3, '', '', []
        for i, sid in enumerate(ids):
            ins += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if i == 0 else 'withEffect'); cid += 3
        c1 = ns['click'](cid, ins); cid += 2
        for i, sid in enumerate(ids):
            outs += ns['leave'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if i == 0 else 'withEffect'); cid += 3
        c2 = ns['click'](cid, outs)
        for sid in ids:
            if re.search(rf'<p:sp>(?:(?!</p:sp>).)*?<p:cNvPr id="{sid}"', x, re.S):
                bld += [f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>', f'<p:bldP spid="{sid}" grpId="1" animBg="1"/>']
        x = x.replace('</p:sld>', ns['timing']([c1, c2], bld) + '</p:sld>')
        data = x.encode('utf8'); print('slide28 pop-up shapes', len(ids))
    zout.writestr(item, data)
zout.close()
