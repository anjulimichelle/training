"""Retention deck: insert 'If C declines the resolution' flowchart right after slide 56.
Built on the live Google export so Anjuli's edits elsewhere stay exactly as they are."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
INK, SOFT, WHITE = '1D1D1F', '6E6E73', 'FFFFFF'
GREEN, GREENSOFT = '2E7D32', 'E4F2E1'
PURPLE, PURPLESOFT = '5E4A9C', 'EEEAF8'
ROSE, ROSESOFT = 'A8456B', 'F8E6EE'
prs = Presentation(src)
ref = prs.slides[55]
assert ref.shapes[1].text_frame.text == 'Case 2: Scenario challenge'
s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
num = [x for x in ref.shapes if x.is_placeholder][0]
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
def retext(shape, t):
    r = shape.text_frame.paragraphs[0].runs; r[0].text = t
    for x in r[1:]: x._r.getparent().remove(x._r)
sh = list(s.shapes)
retext(sh[1], 'Case 2: They decline the resolution'); retext(sh[2], 'They still want a refund or a cancellation. What they ask for decides the steps.')

def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
def box(x, y, w, h, fill, line=None, adj=0.08):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
    b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
    else: b.line.fill.background()
    return b
def text(shape, paras, size=9.5, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, after=3, margin=0.0):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
        for t, bold, col, *sz in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(sz[0] if sz else size); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
def tb(x, y, w, h): return s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
def line(x1, y1, x2, y2, arrow=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); strip(c)
    c.line.color.rgb = RGBColor.from_string(SOFT); c.line.width = Pt(1)
    if arrow:
        ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
B = lambda t, col=INK: ('•  ' + t, False, col)

# Start
top = box(3.45, 1.4, 3.1, 0.4, INK); text(top, [[('If C declines the resolution', True, WHITE, 11)]], align=PP_ALIGN.CENTER)
line(5.0, 1.8, 5.0, 1.98, False); line(1.75, 1.98, 5.1, 1.98, False)
line(1.75, 1.98, 1.75, 2.14); line(5.1, 1.98, 5.1, 2.14)
# Left: refund
lh = box(0.45, 2.14, 2.6, 0.5, GREEN); text(lh, [[('Cancel and refund', True, WHITE, 10.5)], [('OR refund only', True, WHITE, 10.5)]], align=PP_ALIGN.CENTER, after=0)
lb = box(0.45, 2.74, 2.6, 2.38, GREENSOFT, GREEN)
text(tb(0.6, 2.84, 2.32, 2.2), [[B('Refund the voucher through CRM ')], [('   CRM → voucher → refund DHJ/FC', False, SOFT, 8.5)],
                               [B('Invalidate it right away')],
                               [B('Advise C it’s refunded: 5–10 business days')],
                               [B('Deactivate the FC table, if there’s one')]], size=9.5, anchor=MSO_ANCHOR.TOP, after=5)
# Middle: cancel only?
md = box(3.3, 2.14, 3.6, 0.62, PURPLE)
text(md, [[('Cancel only.', True, WHITE, 10.5), (' Do they also want their personal information deleted?', False, WHITE, 9.5)]], align=PP_ALIGN.CENTER, margin=0.1)
# YES
line(5.1, 2.76, 5.1, 2.94); text(tb(5.18, 2.76, 0.5, 0.18), [[('YES', True, SOFT, 8)]])
yb = box(3.3, 2.94, 3.6, 2.18, PURPLESOFT, PURPLE)
text(tb(3.45, 3.02, 3.32, 2.06), [[('Credit card removal', True, INK), ('  → escalate to a TL', False, INK)],
                                  [('Personal data', True, INK), ('  → log it to the Data right to know/delete request tracker', False, INK)],
                                  [B('Invalidate the voucher through the CRM')],
                                  [B('Deactivate the FC table, if there’s one')],
                                  [B('Send a closure email (before deactivating)')],
                                  [B('Deactivate via CRM')]], size=9.5, anchor=MSO_ANCHOR.TOP, after=4)
# NO
line(6.9, 2.45, 7.15, 2.45); text(tb(6.93, 2.24, 0.3, 0.18), [[('NO', True, SOFT, 7.5)]])
nb = box(7.15, 2.14, 2.4, 2.98, ROSESOFT, ROSE)
text(tb(7.28, 2.24, 2.16, 2.8), [[B('Invalidate the voucher through the CRM')],
                                 [B('Deactivate the FC table, if there’s one')],
                                 [('Also wants the account closed?', True, ROSE)],
                                 [B('Send a closure email (before deactivating)')],
                                 [B('Deactivate via CRM')]], size=9.5, anchor=MSO_ANCHOR.TOP, after=5)

s.notes_slide.notes_text_frame.text = (
    'If the customer declines the resolution and insists on a refund or a cancellation. '
    'Cancel and refund, or refund only: refund the voucher through CRM (CRM > voucher > refund DHJ/FC) and invalidate it right away; advise the customer the voucher has been refunded and give the 5–10 business day timeframe; deactivate the ForeverClean (FC) table if there is one. '
    'Cancel only: ask whether they also want their personal information deleted. '
    'YES: credit card removal, escalate to a TL to remove the card from Stripe; personal data, log it to the Data right to know/delete request tracker; invalidate the voucher through the CRM; deactivate the FC table if there is one; send a closure email; deactivate via CRM. '
    'NO: invalidate the voucher through the CRM and deactivate the FC table if there is one. If the customer also wants to close or cancel the account itself, send a closure email and deactivate via CRM. '
    'Always send the closure email before deactivating: once the account is deactivated, there is no way to reach them by email.')
lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(56, el)
prs.save(dst)
