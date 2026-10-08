"""Retention deck slide 78 ('They accept: set it up'): split into a DHJ slide and an FCF slide, each with
side-by-side OTC and TC step lists (from the KB table 'If C Accepts the Trial Cleaning or One-Time Cleaning
offered by CARE'). The DHJ slide reuses slide 78 itself; the FCF slide is inserted right after it."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', '9A6410', 'FFF3D6'
prs = Presentation(src)
dhj = prs.slides[77]
assert dhj.shapes[1].text_frame.text.strip().startswith('They accept') and any(x.has_table for x in dhj.shapes)
num = [x for x in dhj.shapes if x.is_placeholder][0]
head = [dhj.shapes[0], dhj.shapes[1], dhj.shapes[2]]
for x in list(dhj.shapes):
    if x not in head and x.shape_id != num.shape_id: x._element.getparent().remove(x._element)

fcf = prs.slides.add_slide(dhj.slide_layout)
for ph in list(fcf.placeholders): ph._element.getparent().remove(ph._element)
for el in (head[0]._element, head[1]._element, head[2]._element, num._element):
    fcf.shapes._spTree.append(copy.deepcopy(el))
lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); lst.insert(78, el)

def retext(shape, t):
    r = shape.text_frame.paragraphs[0].runs; r[0].text = t
    for x in r[1:]: x._r.getparent().remove(x._r)

def build(s, model, otc, tc):
    retext(s.shapes[1], f'They accept ({model}): set it up')
    retext(s.shapes[2], 'Same flow for OTC and TC, except the steps in gold.')
    def strip(x):
        st = x._element.find(P_NS + 'style')
        if st is not None: x._element.remove(st)
    def box(x, y, w, h, fill, line=None, adj=0.08):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
        b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
        else: b.line.fill.background()
        return b
    def text(shape, runs, size, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        p = tf.paragraphs[0]; p.alignment = align
        for t, bold, col in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
    w = 4.45
    for i, (label, steps) in enumerate((('One-Time Cleaning (OTC)', otc), ('Trial Cleaning (TC)', tc))):
        x = 0.45 + i * (w + 0.2)
        hb = box(x, 1.4, w, 0.34, TEAL, adj=0.2)
        text(hb, [(label, True, WHITE)], 10.5, align=PP_ALIGN.CENTER)
        for j, (lead, rest, diff) in enumerate(steps):
            y = 1.82 + j * 0.415
            box(x, y, w, 0.37, GOLDSOFT if diff else WHITE, GOLD if diff else BORDER)
            bd = box(x + 0.07, y + 0.055, 0.26, 0.26, GOLD if diff else TEALSOFT, adj=0.5)
            text(bd, [(str(j + 1), True, WHITE if diff else TEAL)], 8.5, align=PP_ALIGN.CENTER)
            t = s.shapes.add_textbox(Inches(x + 0.42), Inches(y), Inches(w - 0.5), Inches(0.37))
            text(t, [(lead, True, GOLD if diff else INK), (rest, False, INK)], 8.5)

CONFIRM = ('Confirm account info: ', 'address, card on file, date, time, duration', False)
CHARGE = ('Charge upfront ', '(FC Sales Calculator), then book it (CRM or C dashboard)', False)
DOUBLE = ('Prevent a double charge: ', 'admin_courtesy voucher, don’t notify C', False)
DONE = ('Confirm to C: ', 'charged upfront, booked, notified once the CP confirms', False)
build(dhj, 'DHJ',
      [CONFIRM,
       ('Still wants a refund? ', 'Refund the DHJ voucher via CRM, invalidate it', False),
       ('No refund? ', 'Deduct the voucher price from the full OTC price, invalidate the voucher', False),
       CHARGE,
       ('FC table: deactivate it. ', 'C only wants a one-time cleaning', True),
       DOUBLE, DONE],
      [CONFIRM,
       ('Still wants a refund? ', 'Refund the DHJ voucher via CRM, invalidate it', False),
       ('No refund? ', 'Deduct the voucher price from the TC price, invalidate the voucher', False),
       CHARGE,
       ('FC table: leave it as is. ', 'DHJ MFs only start after a completed job', True),
       DOUBLE,
       ('Legacy CRM voucher: ', '“ForeverClean force” + “1 free month”; update the ETF to $99', True),
       DONE])
FREF = ('Still wants a refund? ', 'Log the 1st charge (1st MF + add-on) in the FCF Refund log; a TL refunds via Stripe', False)
build(fcf, 'FCF',
      [CONFIRM, FREF,
       ('No refund? ', 'Deduct all MFs paid to date (+ add-on) from the full OTC price, invalidate the voucher', False),
       CHARGE,
       ('FC table: deactivate it. ', 'C only wants a one-time cleaning', False),
       DOUBLE, DONE],
      [CONFIRM, FREF,
       ('No refund? ', 'Deduct all MFs paid to date (+ add-on) from the TC price, invalidate the voucher', False),
       CHARGE,
       ('FC table: deactivate it. ', 'FCF starts at sign-up and auto-renews', False),
       DOUBLE,
       ('Legacy CRM voucher: ', '“ForeverClean force” + “1 free month”, add C to the FC table, ETF to $99', True),
       DONE])

dhj.notes_slide.notes_text_frame.text = (
    'DHJ: the customer accepts the TC or OTC offered by Care. Steps 1–4, 6 and the last step are the same for OTC and TC; the gold steps differ. '
    '1) Confirm account info: complete address, a valid card on file, appointment date, time and duration. '
    '2) Still wants a refund: refund the DHJ voucher via CRM and invalidate it immediately. '
    '3) No longer wants a refund: deduct the voucher amount from the full cleaning price (OTC) or the trial cleaning price (TC), and invalidate the DHJ voucher. '
    '4) Charge upfront using the FC Sales Calculator, then book (CRM if applicable, otherwise the customer dashboard). '
    '5) FC table: OTC, invalidate the voucher and deactivate the FC table, since C only wants a one-time cleaning, not a membership. TC, leave the FC table as is: in the DHJ model, MFs only start after a completed job. '
    '6) Prevent a double charge: issue an admin_courtesy voucher (don\'t notify the customer). '
    'TC only: click the voucher in Legacy CRM and select "ForeverClean force" (FC triggers once the job is completed) and "1 free month" (delays the next MF by 30 days instead of charging right after the cleaning); update the ETF to $99. '
    'Last: confirm to the customer: charged upfront, appointment booked, notified once the CP confirms.')
fcf.notes_slide.notes_text_frame.text = (
    'FCF: the customer accepts the TC or OTC offered by Care. Same steps for OTC and TC; only the gold TC step is extra. '
    '1) Confirm account info: complete address, a valid card on file, appointment date, time and duration. '
    '2) Still wants a refund: log the 1st charge (1st MF + add-on, if applicable) in the FCF Refund log so a TL can refund it via Stripe. '
    '3) No longer wants a refund: deduct all MFs paid to date (+ add-on, if applicable) from the full OTC price or the TC price, then invalidate the FCF voucher. '
    'Example: 2 MFs + $59 add-on = $137 paid; a 6-hour clean at $378 means C pays the remaining $241. If what C paid exceeds the price, C pays $0 and the difference is not refunded (e.g. $150 in MFs vs a $120 cleaning: $0 due, $30 not refunded). '
    '4) Charge upfront using the FC Sales Calculator, then book (CRM if applicable, otherwise the customer dashboard). '
    '5) FC table: deactivate it in both cases. OTC, C only wants a one-time cleaning. TC, the FCF membership starts at sign-up and auto-renews even without a completed job. '
    '6) Prevent a double charge: issue an admin_courtesy voucher (don\'t notify the customer). '
    'TC only: click the voucher in Legacy CRM and select "ForeverClean force" (FC triggers once the job is completed) and "1 free month" (delays the next MF by 30 days); add C to the FC table, since the membership becomes active after the 30-day free period; update the ETF to $99. '
    'Last: confirm to the customer: charged upfront, appointment booked, notified once the CP confirms.')
prs.save(dst)
