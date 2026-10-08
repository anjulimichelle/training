"""Retention deck: insert an FCF pre-invoice scenario challenge right after 'Retention strategy' (slide 64),
and on 'Exit strategy' drop the cheat-sheet button and reword the refund note. Patches the live Google export."""
import copy, io, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, BG = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', 'F6F7F9'
prs = Presentation(src)
ref = prs.slides[63]
assert ref.shapes[1].text_frame.text.strip() == 'Retention strategy'
ex = prs.slides[64]
assert ex.shapes[1].text_frame.text.strip() == 'Exit strategy'

# ---------- Exit strategy: remove the cheat-sheet button, reword and widen the refund note ----------
btn = [x for x in ex.shapes if x.has_text_frame and 'Refund eligibility cheat sheet' in x.text_frame.text and x.text_frame.text.startswith('▶')]
assert len(btn) == 1
btn[0]._element.getparent().remove(btn[0]._element)
note = [x for x in ex.shapes if x.has_text_frame and x.text_frame.text.startswith('Wants a refund?')][0]
p2 = note.text_frame.paragraphs[1]; r = p2.runs; r[0].text = 'Review the 24-hour refund window, then check the Refund eligibility cheat sheet.'
for x in r[1:]: x._r.getparent().remove(x._r)
note.width = Inches(8.7)
ex.notes_slide.notes_text_frame.text = ex.notes_slide.notes_text_frame.text.replace(
    'After attempting retention, check the Refund eligibility cheat sheet', 'Review the 24-hour refund window, then check the Refund eligibility cheat sheet')

# ---------- New slide: scenario challenge ----------
s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
num = [x for x in ref.shapes if x.is_placeholder][0]
for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
    s.shapes._spTree.append(copy.deepcopy(el))
for shape, t in ((s.shapes[1], 'Scenario challenge'), (s.shapes[2], 'Read the customer’s message. How would you respond?')):
    rr = shape.text_frame.paragraphs[0].runs; rr[0].text = t
    for x in rr[1:]: x._r.getparent().remove(x._r)
lst = prs.slides._sldIdLst; ids = list(lst); el = ids[-1]; lst.remove(el); lst.insert(64, el)

def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
def box(x, y, w, h, fill, line=None, adj=0.06, name=None, dash=False):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
    b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
        if dash:
            ln = b._element.spPr.find(qn('a:ln')); ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))
    else: b.line.fill.background()
    if name: b.name = name
    return b
def text(shape, paras, size=9, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=4):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
        for t, bold, col, *sz in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(sz[0] if sz else size); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
def tb(x, y, w, h, name=None):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: t.name = name
    return t

scen = [
    ('SCENARIO 1', 'CUSTOMER',
     'Stop this! You’ve charged me 5x now! Refund everything or I will report fraud to my bank!',
     'Context: C hasn’t booked their free 3-hr first cleaning and has now been charged 5 MFs.',
     ['Thanks for reaching out! I’m happy to clarify the charges on your account and go over your options. If you don’t mind me asking, was there a reason you haven’t had a chance to book your first cleaning yet? I’d love to help if there’s anything getting in the way.',
      'The charges you’re seeing are your ForeverClean monthly fees. ForeverClean is our membership program which included your free 3-hour first cleaning during sign up. The membership fees continue monthly for $59 unless canceled, even if you haven’t used the free cleaning yet.',
      'If you’d still like to use your free 3-hour cleaning, I’d be happy to help get that set up. I can also add 3 free months to your membership, so your next monthly fee wouldn’t be charged until <date>. Quick note: the free months don’t count toward your 6-month term.',
      'If you’d rather cancel, you can do that here:',
      'Let me know what works best for you, and I’ll be happy to help.']),
    ('SCENARIO 2', 'CUSTOMER',
     'I will not be needing your service. I’m moving out of state.', None,
     ['Thanks for letting me know! I hope the move goes smoothly.',
      'Just so you know, we’re available across all 50 U.S. states, so if you’d like to continue with us at your new address, I’d be happy to help get your free first cleaning set up once you’re settled.',
      'If you need a little more time before booking, I can also add a free month so your next monthly fee wouldn’t be charged until <date>. Just a quick note: the free month doesn’t count toward your 6-month term.',
      'If you’d rather cancel, you can do that here: [cancellation link].',
      'Let me know what works best for you, and I’ll be happy to help.']),
]
w = (9.1 - 0.2) / 2
for i, (lab, who, msg, ctx, resp) in enumerate(scen):
    x = 0.45 + i * (w + 0.2)
    text(tb(x, 1.38, w, 0.2), [[(lab, True, TEAL, 8)]])
    box(x, 1.6, w, 0.9, BG, BORDER, adj=0.1)
    paras = [[(who, True, SOFT, 7.5)], [('“' + msg + '”', False, INK, 10)]]
    if ctx: paras.append([(ctx, False, SOFT, 8.5)])
    text(tb(x + 0.15, 1.6, w - 0.3, 0.9), paras, anchor=MSO_ANCHOR.MIDDLE, after=2)
    box(x, 2.6, w, 2.6, WHITE, BORDER, dash=True)
    text(tb(x, 3.35, w, 0.25), [[('Write your reply first.', False, SOFT, 9)]], align=PP_ALIGN.CENTER)
    cue = box(x + (w - 2.3) / 2, 3.65, 2.3, 0.38, TEAL, adj=0.3)
    text(cue, [[('▶  Show CS response', True, WHITE, 10.5)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(x, 2.6, w, 2.6, TEALSOFT, TEAL, name=f'step{i + 1}Box')
    text(tb(x + 0.14, 2.67, w - 0.28, 2.48, name=f'step{i + 1}Text'),
         [[('SAMPLE CS RESPONSE', True, TEAL, 7.5)]] + [[(p, False, INK, 8.5)] for p in resp], after=4)

s.notes_slide.notes_text_frame.text = (
    'SCENARIO CHALLENGE (FCF pre-invoice: unused free first cleaning, monthly fees charged). For each scenario, read the customer\'s message aloud and give trainees a few minutes to write a reply. '
    'Read a few aloud, then click to reveal the sample CS response (click 1: scenario 1, click 2: scenario 2). '
    'Scenario 1: "Stop this! You\'ve charged me 5x now! Refund everything or I will report fraud to my bank!" Context: C hasn\'t booked the free 3-hr first cleaning and has been charged 5 MFs. '
    'Stay calm, clarify the charges (ForeverClean monthly fees, $59/month until canceled, even if the free cleaning is unused), probe for the root cause, and offer to book the free cleaning plus up to 3 free months (multiple unused paid MFs). Free months don\'t count toward the 6-month term. Give the self-cancel option. '
    'Scenario 2: "I will not be needing your service. I\'m moving out of state." Personal/service limit (moving): remind them Homeaglow is available in all 50 states, offer to set up the free first cleaning at the new address, or a free month if they need more time. Give the self-cancel link. '
    'Sample responses are on the slide; <date> = the next MF date after the free months.')
buf = io.BytesIO(); prs.save(buf)

ns = {}
exec(open(__file__.replace('add-slide65-fcf-scenario-challenge.py', 'add-story-animation.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == s.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        groups = {}
        for sid, n in re.findall(r'<p:cNvPr id="(\d+)" name="step(\d+)[^"]*"', x): groups.setdefault(int(n), []).append(sid)
        cid, clicks, bld = 3, [], []
        for n in sorted(groups):
            inner = ''
            for j, sid in enumerate(groups[n]):
                inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if j == 0 else 'withEffect'); cid += 3
                bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
            clicks.append(ns['click'](cid, inner)); cid += 2
        x = x.replace('</p:sld>', ns['timing'](clicks, bld) + '</p:sld>')
        data = x.encode('utf8'); print('clicks:', len(clicks))
    zout.writestr(item, data)
zout.close()
