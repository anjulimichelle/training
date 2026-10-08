"""Retention deck: slide 23 keeps only the 'Which tool?' tile and gains a worked scenario; a new
'ETF reduction' slide goes in right after it (before the untouched '50% cap' slide).
Built on the live Google export so Anjuli's edits elsewhere stay exactly as they are."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF'
RED, PINK, GOLD, GOLDSOFT = 'B3261E', 'FCEBEA', '9A6410', 'FFF3D6'
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
prs = Presentation(src)

def mk(s):
    def box(x, y, w, h, fill, line=None, rect=False):
        b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE if rect else MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        st = b._element.find(P_NS + 'style')
        if st is not None: b._element.remove(st)
        if not rect: b.adjustments[0] = 0.08
        b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
        else: b.line.fill.background()
        return b
    def text(x, y, w, h, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT, after=2):
        t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame
        tf.word_wrap = True; tf.vertical_anchor = anchor
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(after); p.alignment = align
            for txt, size, bold, col in runs:
                r = p.add_run(); r.text = txt; r.font.size = Pt(size); r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
        return t
    def head(x, y, w, label, col):
        b = box(x, y, w, 0.34, col); tf = b.text_frame
        for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label; r.font.size = Pt(10.5); r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE); r.font.name = 'SF Pro Display'
    return box, text, head

# ---------- Slide 23: Which tool + sample scenario ----------
s = prs.slides[22]
sh = list(s.shapes)
assert sh[1].text_frame.text.startswith('Exit offers') and sh[6].text_frame.text == 'Which tool?'
for i in (3, 4, 8):
    sh[i]._element.getparent().remove(sh[i]._element)
def place(shape, x, y, w, h):
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)
place(sh[5], 0.45, 1.42, 9.1, 1.1); place(sh[6], 0.45, 1.42, 9.1, 0.34); place(sh[7], 0.63, 1.84, 8.75, 0.62)
box, text, head = mk(s)
text(0.45, 2.66, 6, 0.2, [[('SAMPLE SCENARIO', 8, True, SOFT)]])
box(0.45, 2.9, 9.1, 0.46, TEALSOFT)
text(0.65, 2.9, 8.7, 0.46, [[('C has ', 10.5, False, INK), ('1 paid $59 MF', 10.5, True, INK), (', a ', 10.5, False, INK), ('6-month MCT', 10.5, True, INK),
                            (' from sign-up, and cancelling now costs a ', 10.5, False, INK), ('$140 ETF', 10.5, True, INK), ('.', 10.5, False, INK)]])
w = (9.1 - 0.15) / 2
box(0.45, 3.48, w, 0.92, WHITE, BORDER)
text(0.65, 3.48, w - 0.4, 0.92, [[('Reduce the MCT by 1 month', 10.5, True, INK), ('  (4 months remaining)', 9.5, False, SOFT)],
                                [('$59 × 4 months = ', 12, False, INK), ('$236', 14, True, INK)]])
box(0.45 + w + 0.15, 3.48, w, 0.92, TEALSOFT, TEAL)
text(0.65 + w + 0.15, 3.48, w - 0.4, 0.92, [[('Reduce the ETF', 10.5, True, TEAL), ('  (max 50%)', 9.5, False, SOFT)],
                                           [('$140 × 50% = ', 12, False, INK), ('$70', 14, True, TEAL), ('   ✓ lower', 10, True, TEAL)]])
box(0.45, 4.52, 9.1, 0.44, TEAL)
text(0.65, 4.52, 8.7, 0.44, [[('Decision: ', 10.5, True, WHITE), ('reducing the ETF is the most cost-effective exit offer for this customer.', 10.5, False, WHITE)]])
n = s.notes_slide.notes_text_frame
n.text = n.text.rstrip() + (' SAMPLE SCENARIO: the customer has 1 paid $59 MF and a 6-month MCT from sign-up, and cancelling now costs a $140 ETF. '
    'Reducing the MCT by 1 month leaves 4 months: $59 x 4 = $236. Reducing the ETF by the maximum 50%: $140 x 50% = $70. '
    'Decision: reducing the ETF is the most cost-effective exit offer for this customer. The first-job hours check and when to offer an ETF reduction are on the next slide.')

# ---------- New slide: ETF reduction ----------
ref = prs.slides[23]
assert ref.shapes[1].text_frame.text.startswith('ETF reduction: the 50% cap')
ns = prs.slides.add_slide(ref.slide_layout)
for ph in list(ns.placeholders): ph._element.getparent().remove(ph._element)
for i in (0, 1, 2, len(ref.shapes) - 1):
    ns.shapes._spTree.append(copy.deepcopy(ref.shapes[i]._element))
def retext(shape, t):
    r = shape.text_frame.paragraphs[0].runs; r[0].text = t
    for x in r[1:]: x._r.getparent().remove(x._r)
c = list(ns.shapes)
retext(c[1], 'ETF reduction'); retext(c[2], 'It depends entirely on why the customer is cancelling.')
box, text, head = mk(ns)
# 1. Fix the first job first
box(0.45, 1.42, 3.55, 2.2, WHITE, BORDER); head(0.45, 1.42, 3.55, '1  Check the first job’s hours', TEAL)
text(0.63, 1.84, 3.2, 1.72, [[('ETF = $35/hr × first-cleaning hours', 10, True, INK)],
                             [('An overcharge on that job inflates the ETF. Resolve the OCH first.', 9, False, SOFT)],
                             [('Invoiced 4 hrs, CP’s messages show 2.5 → refund 1.5 hrs', 9, False, INK)],
                             [('ETF $140 → $87.50', 14, True, TEAL)],
                             [('That’s the ETF you present.', 9, True, INK)]], anchor=MSO_ANCHOR.TOP, after=3)
# 2. When to offer it
text(4.15, 1.42, 5.4, 0.2, [[('2  WHEN TO OFFER IT', 8, True, SOFT)]])
rows = [(GOLD, GOLDSOFT, 'Unawareness only', '(or “I only wanted a one-time cleaning”)', 'Don’t offer it proactively. Wait for a hint they’d pay a reduced ETF.'),
        (TEAL, TEALSOFT, 'Any other reason mixed in', '(service issue, cost complaint…)', 'Standard exit tool once retention is exhausted. MCT cheaper? Use that.'),
        (RED, PINK, 'Extreme financial hardship', '(skip retention attempts)', 'Go straight to 50% off. Declined? Waive the ETF entirely.')]
for i, (col, soft, a, a2, b) in enumerate(rows):
    y = 1.66 + i * 0.66
    box(4.15, y, 5.4, 0.58, soft, col)
    text(4.3, y, 2.1, 0.58, [[(a, 9.5, True, col)], [(a2, 8, False, SOFT)]], after=0)
    text(6.5, y, 2.95, 0.58, [[(b, 9, False, INK)]])
# 3. Deductions, cap, documentation
tiles = [('3  What can be deducted', [[('MF and/or voucher price', 9, True, INK), (': don’t stack automatically, only when justified', 9, False, INK)],
                                      [('Additional-hours charge', 9, True, INK), (': if the FC-triggering job ran past the voucher’s hours', 9, False, INK)]]),
         ('Cap: 50% of the original ETF', [[('Whatever is deducted. More line items don’t raise the ceiling, they just change how you get there.', 9, False, INK)]]),
         ('Before you apply it', [[('Note the customer’s ', 9, False, INK), ('LTNR', 9, True, INK), (' in the internal note (context, not a limit).', 9, False, INK)],
                                  [('Existing reduction? Always calculate from the ', 9, False, INK), ('original', 9, True, INK), (' ETF.', 9, False, INK)]])]
tw = (9.1 - 0.3) / 3
for i, (t, body) in enumerate(tiles):
    x = 0.45 + i * (tw + 0.15)
    box(x, 3.76, tw, 1.3, WHITE, BORDER); head(x, 3.76, tw, t, [INK, TEAL, GOLD][i])
    text(x + 0.15, 4.16, tw - 0.3, 0.86, body, anchor=MSO_ANCHOR.TOP, after=3)
ns.notes_slide.notes_text_frame.text = (
    'ETF reduction. Check the first job\'s hours before you calculate any exit offer. The ETF is $35/hr x first-cleaning hours, so an overcharge on that job inflates the ETF. '
    'Resolve the OCH first, then present the corrected ETF and work the exit offer from that number. First job invoiced 4 hrs, the CP\'s own messages show 2.5: refund 1.5 hrs, and the ETF drops from $140 to $87.50. That\'s the ETF you present. '
    'This one depends entirely on why the customer is cancelling, specifically whether unawareness is the whole story or just part of it. '
    'If unawareness (or "I only wanted a one-time cleaning") is the only reason (you probed, and there\'s genuinely nothing else going on), don\'t offer an ETF reduction proactively. Wait for a hint that the customer would actually pay a reduced ETF before putting it on the table. Unawareness alone doesn\'t earn a discount; some willingness from the customer to pay is what opens that door. '
    'If there\'s any other reason mixed in (a service issue, a cost complaint, anything beyond plain unawareness), ETF reduction is just part of your standard exit strategy once retention\'s been exhausted. The only thing that changes your choice is cost: if reducing the MCT would actually be cheaper for the customer than reducing the ETF, use that instead. Pick whichever tool gets the customer out for less, not whichever feels more generous. '
    'Exception, extreme financial hardship: there\'s no point running through the usual retention attempts first. Go straight to exit: offer 50% off the ETF. If they decline even that, the next step is waiving the ETF entirely. '
    'What can be deducted, and in what order: MF and/or voucher price, the standard deduction; don\'t stack these automatically, only when justified (the customer contests paying the ETF on top of what they\'ve already paid, or explicitly asks for both to be deducted). Additional-hours charge: if the job that triggered FC ran longer than the voucher\'s covered hours, and the customer was charged for that overage on their card, that charge can also be deducted from the ETF. '
    'The cap is always 50% off the original ETF, regardless of which deductions apply. More deductible line items don\'t raise the ceiling, they just change how you get there. '
    'Note the customer\'s Lifetime Net Revenue (LTNR) in your internal note when reducing the ETF, the same way you would for LMC/Lockout. It\'s a documentation step, not a limiter, and it doesn\'t override the rule that the 50% reduction must be maxed out before moving to Graceful Closure. '
    'Existing offer? Check whether a reduction was already applied before granting another. Always calculate any reduction from the original ETF, never an already-discounted amount.')
lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(23, el)
prs.save(dst)
