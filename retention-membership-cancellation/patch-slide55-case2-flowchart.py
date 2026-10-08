"""Retention deck slide 55 (Unused DHJ Voucher, Case 2): replace the step list with Anjuli's flowchart.
Service and Membership offers appear one per click. Patches the live Google export."""
import io, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', '9A6410', 'FFF3D6'
PURPLE, PURPLESOFT = '5E4A9C', 'EEEAF8'
prs = Presentation(src)
s = prs.slides[54]
sh = list(s.shapes)
assert sh[1].text_frame.text.startswith('Case 2: Refund or cancel intent') and sh[33].text_frame.text.startswith('Only if they want a refund')
for x in sh[3:-1]: x._element.getparent().remove(x._element)
s.shapes[2].text_frame.paragraphs[0].runs[0].text = 'Acknowledge first, then follow the reason.'
for r in s.shapes[2].text_frame.paragraphs[0].runs[1:]: r._r.getparent().remove(r._r)

def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
def box(x, y, w, h, fill, line=None, adj=0.08, name=None):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
    b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
    else: b.line.fill.background()
    if name: b.name = name
    return b
def fill_text(shape, paras, size=9.5, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, after=2, margin=0.0):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right'): setattr(tf, m, Inches(margin))
    tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
        for t, bold, col, *sz in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(sz[0] if sz else size); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
def tb(x, y, w, h, name=None):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: t.name = name
    return t
def line(x1, y1, x2, y2, arrow=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); strip(c)
    c.line.color.rgb = RGBColor.from_string(SOFT); c.line.width = Pt(1)
    if arrow:
        ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))

# Decision
d = box(2.55, 1.42, 4.1, 0.48, INK); fill_text(d, [[('Acknowledge the request. Did the C provide a reason?', True, WHITE, 10.5)]], align=PP_ALIGN.CENTER)
# NO branch
line(6.65, 1.66, 7.4, 1.66)
fill_text(tb(6.78, 1.44, 0.5, 0.2), [[('NO', True, SOFT, 8)]])
pb = box(7.4, 1.4, 2.15, 0.82, TEALSOFT, TEAL)
fill_text(pb, [[('Probe for the reason', True, TEAL, 10.5)], [('Unsubscribe if C has a cancel intent only', False, INK, 8.5)]], align=PP_ALIGN.CENTER, margin=0.08)
# YES branch
line(4.6, 1.9, 4.6, 2.08, False); line(1.7, 2.08, 5.3, 2.08, False)
line(1.7, 2.08, 1.7, 2.26); line(5.3, 2.08, 5.3, 2.26)
fill_text(tb(4.68, 1.92, 0.5, 0.16), [[('YES', True, SOFT, 8)]])
# Service and Membership tiles
for x, w, col, soft, title, sub in ((0.45, 2.5, TEAL, TEALSOFT, 'Service', 'address the actual issue'), (3.1, 4.4, PURPLE, PURPLESOFT, 'Membership', 'membership-terms objection')):
    h = box(x, 2.26, w, 0.38, col); fill_text(h, [[(title, True, WHITE, 11), ('   ' + sub, False, WHITE, 8.5)]], align=PP_ALIGN.CENTER)
    box(x, 2.7, w, 1.58, soft, col)
svc = [[('Service-related issues', True, INK), ('  (CP no-show, CP cancellation, no CP claim)', False, SOFT, 8)], [('→ Offer $10–$25 credits', True, TEAL)]], \
      [[('Service limitations / personal', True, INK)], [('→ No offers necessary', True, TEAL)]]
for i, paras in enumerate(svc):
    fill_text(tb(0.6, 2.78 + i * 0.72, 2.25, 0.66, name=f'step{i + 1}Svc'), paras, size=9, anchor=MSO_ANCHOR.TOP)
mem = [([('MF expensive', True, INK), ('  → reduce by $10–$15 off the original amount', False, INK)], 0.27),
       ([('MCT too long', True, INK), ('  → reduce by 2 months from the original MCT', False, INK)], 0.27),
       ([('Doesn’t want a membership', True, INK), ('  → no offers; advise them to contact us for more flexible terms. OTC can be offered', False, INK)], 0.4),
       ([('OTC', True, INK), ('  → full price, charged upfront', False, INK)], 0.27),
       ([('TC', True, INK), ('  → regular price, charged upfront', False, INK)], 0.27)]
y = 2.76
for i, (runs, h) in enumerate(mem):
    fill_text(tb(3.25, y, 4.15, h, name=f'step{i + 3}Mem'), [[('•  ', True, PURPLE)] + runs], size=9, anchor=MSO_ANCHOR.TOP)
    y += h + 0.02
# Into the closing box
line(1.7, 4.28, 1.7, 4.42); line(5.3, 4.28, 5.3, 4.42); line(8.47, 2.22, 8.47, 4.42)
cb = box(0.45, 4.42, 9.1, 0.7, GOLDSOFT, GOLD)
fill_text(cb, [[('•  Remind them the voucher comes with a membership and is valid for 1 year', False, INK)],
               [('•  Cancel RC and upcoming appointments', False, INK)],
               [('•  Provide the self-refund link ', False, INK), ('(refund intent only)', True, GOLD)]], size=9, margin=0.25, after=1)

n = s.notes_slide.notes_text_frame
n.text = ('Unused DHJ Voucher with refund or cancel intent. Acknowledge the request, then check: did the customer give a reason? '
          'NO: probe for the reason (unsubscribe them if they have a cancel intent only). '
          'YES, service: address the actual issue (block the C/CP pairing, coach or penalize the CP as needed). Service-related issues (CP no-show, CP cancellation, no CP claim): offer $10–$25 credits. Service limitations or personal reasons (found another cleaner, wants same-day, moved, no phone support, will self-clean): no offers necessary. '
          'YES, membership: MF expensive, offer to reduce by $10–$15 off the original amount; MCT too long, offer to reduce by 2 months from the original MCT; does not want a membership in general, no offers, advise them to contact us if they want more flexible terms, and OTC can be offered; OTC at full price, charged upfront; TC at regular price, charged upfront. '
          'Then, in every case: remind them the voucher comes with a membership and is valid for 1 year; cancel the RC and upcoming appointments (if there is no sign they want to keep them, with the correct cancellation reason code); provide the self-refund link for refund intent only (enter the voucher code before clicking "Refund Voucher", and double-check the code). '
          'If they decline and insist on a refund: refund the voucher through CRM (CRM > voucher > refund DHJ/FC) and invalidate it right away, advise the 5–10 business day timeframe, and deactivate the FC table if there is one. '
          'TRAINER: clicks 1–2 reveal the Service offers, clicks 3–7 the Membership offers, one at a time.')
buf = io.BytesIO(); prs.save(buf)

ns = {}
exec(open(__file__.replace('patch55_flow.py', 'add_story_anim.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == s.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        steps = sorted(((int(n), sid) for sid, n in re.findall(r'<p:cNvPr id="(\d+)" name="step(\d+)[^"]*"', x)))
        cid, clicks, bld = 3, [], []
        for n_, sid in steps:
            clicks.append(ns['click'](cid, ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect'))); cid += 5
            bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
        x = x.replace('</p:sld>', ns['timing'](clicks, bld) + '</p:sld>')
        data = x.encode('utf8'); print('clicks:', len(clicks))
    zout.writestr(item, data)
zout.close()
