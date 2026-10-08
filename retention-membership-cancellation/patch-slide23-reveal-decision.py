"""Retention deck slide 23: the decision is hidden by default. Trainees see a neutral pair of options
and a question; one click reveals the decision bar, a highlight on the ETF option and a 'Lower cost' tag.
Patches the live Google export so Anjuli's edits elsewhere stay exactly as they are."""
import io, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', '9A6410', 'FFF3D6'
prs = Presentation(src)
s = prs.slides[22]
sh = list(s.shapes)
assert sh[1].text_frame.text.startswith('Exit offers') and sh[15].text_frame.text.startswith('Decision:')
etf_box, etf_txt, dec_box, dec_txt = sh[12], sh[13], sh[14], sh[15]

# Neutral ETF option (same look as the MCT one) and no '✓ lower' giveaway.
etf_box.fill.solid(); etf_box.fill.fore_color.rgb = RGBColor.from_string(WHITE)
etf_box.line.color.rgb = RGBColor.from_string(BORDER); etf_box.line.width = Pt(0.75)
p0, p1 = etf_txt.text_frame.paragraphs[0], etf_txt.text_frame.paragraphs[1]
p0.runs[0].font.color.rgb = RGBColor.from_string(INK)
for r in p1.runs:
    if '✓' in r.text: r._r.getparent().remove(r._r)
    elif r.text.strip() == '$70': r.font.color.rgb = RGBColor.from_string(INK)

def strip_style(shape):
    st = shape._element.find(P_NS + 'style')
    if st is not None: shape._element.remove(st)
def rbox(x, y, w, h, fill, line, lw=0.75, name=None):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip_style(b)
    b.adjustments[0] = 0.08
    if fill: b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    else: b.fill.background()
    b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(lw)
    if name: b.name = name
    return b
def write(shape, runs, size, align=PP_ALIGN.LEFT):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = align
    for t, bold, col in runs:
        r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'

# Default state: the question, where the decision bar sits (the bar covers it once revealed).
q = rbox(0.45, 4.52, 9.1, 0.44, GOLDSOFT, GOLDSOFT, name='askQuestion')
qt = s.shapes.add_textbox(Inches(0.65), Inches(4.52), Inches(6.4), Inches(0.44)); qt.name = 'askText'
write(qt, [('Your turn: ', True, GOLD), ('which exit offer would you choose, and why?', False, INK)], 10.5)
cue = rbox(7.25, 4.58, 2.2, 0.32, GOLD, GOLD, name='askCue'); cue.adjustments[0] = 0.3
write(cue, [('▶  Show the decision', True, WHITE)], 9.5, PP_ALIGN.CENTER)
# Keep the question underneath the decision bar.
tree = dec_box._element.getparent()
for el in (q._element, qt._element, cue._element):
    tree.remove(el); dec_box._element.addprevious(el)

# Revealed on click: highlight + tag on the ETF option, and the decision bar.
hl = rbox(5.08, 3.48, 4.47, 0.92, None, TEAL, 2.25, name='step1Highlight')
tag = rbox(8.1, 3.58, 1.32, 0.28, TEALSOFT, TEAL, name='step1Tag'); tag.adjustments[0] = 0.3
write(tag, [('✓ Lower cost', True, TEAL)], 9, PP_ALIGN.CENTER)
dec_box.name = 'step1DecisionBar'; dec_txt.name = 'step1DecisionText'
for el in (dec_box._element, dec_txt._element):   # decision on top of everything
    tree.remove(el); tree.append(el)

n = s.notes_slide.notes_text_frame
n.text = n.text.rstrip() + (' TRAINER: the decision is hidden at first. Ask trainees which exit offer they would choose and why, '
                            'then click to reveal the answer: reducing the ETF ($70) costs the customer less than reducing the MCT ($236).')
buf = io.BytesIO(); prs.save(buf)

# One click: everything named step1* appears together.
ns = {}
exec(open(__file__.replace('patch23_reveal.py', 'add_story_anim.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == s.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        ids = re.findall(r'<p:cNvPr id="(\d+)" name="step1[^"]*"', x)
        cid, inner, bld = 3, '', []
        for i, sid in enumerate(ids):
            inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if i == 0 else 'withEffect'); cid += 3
            bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
        x = x.replace('</p:sld>', ns['timing']([ns['click'](cid, inner)], bld) + '</p:sld>')
        data = x.encode('utf8'); print('revealed on click:', len(ids))
    zout.writestr(item, data)
zout.close()
