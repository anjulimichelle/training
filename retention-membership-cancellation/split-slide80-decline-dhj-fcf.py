"""Retention deck slide 80 ('They decline the offer'): split into a DHJ slide (slide 80 itself) and an FCF
slide inserted right after it, from the KB table 'If C Declines the TC/OTC Offer from CARE'.
Steps are the same for OTC and TC within each model."""
import copy, sys
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
prs = Presentation(src)
dhj = prs.slides[79]
assert dhj.shapes[1].text_frame.text.strip() == 'They decline the offer'

fcf = prs.slides.add_slide(dhj.slide_layout)
for ph in list(fcf.placeholders): ph._element.getparent().remove(ph._element)
for x in dhj.shapes:
    el = copy.deepcopy(x._element)
    for blip in el.iter(qn('a:blip')):
        img = dhj.part.related_part(blip.get(qn('r:embed')))
        blip.set(qn('r:embed'), fcf.part.relate_to(img, RT.IMAGE))
    fcf.shapes._spTree.append(el)
lst = prs.slides._sldIdLst; el = list(lst)[-1]; lst.remove(el); lst.insert(80, el)

def by_text(s, start):
    m = [x for x in s.shapes if x.has_text_frame and x.text_frame.text.startswith(start)]
    assert len(m) == 1, start
    return m[0]
def retext(shape, *texts):
    runs = shape.text_frame.paragraphs[0].runs
    for r, t in zip(runs, texts): r.text = t
    for r in runs[len(texts):]: r._r.getparent().remove(r._r)

for s, model, refund, lead, advise in (
        (dhj, 'DHJ', 'Within 1 year of purchase: refund the DHJ voucher via CRM and invalidate it.',
         'Advise C:  ', 'the refund is processed and takes 5–10 business days.'),
        (fcf, 'FCF', 'Check the Refund eligibility cheat sheet: No completed job yet (FCF Pre-Invoice Cancellation).',
         'Advise C:  ', 'what happens next depends on the Refund eligibility cheat sheet.')):
    retext(by_text(s, 'They decline the offer'), f'They decline the offer ({model})')
    retext(by_text(s, 'Refund what they'), 'Same steps for OTC and TC. Refund what they’re owed and close things out.')
    retext(by_text(s, 'DHJ: within 1 year'), refund)
    retext(by_text(s, 'In every case'), 'Always, for both OTC and TC.')
    retext(by_text(s, 'DHJ refund:'), lead, advise)

dhj.notes_slide.notes_text_frame.text = (
    'DHJ: the customer declines the TC/OTC offer. Same steps for OTC and TC. '
    'Refund: if C requests a refund within 1 year of purchase, refund the DHJ voucher through CRM and invalidate the voucher. '
    'FC table: deactivate it. Advise the customer: the refund is processed, 5–10 business day timeframe. '
    'CRM account: if C wants their account closed, deactivate the CRM account, but send comms before deactivating.')
fcf.notes_slide.notes_text_frame.text = (
    'FCF: the customer declines the TC/OTC offer. Same steps for OTC and TC. '
    'Refund: check the Refund Eligibility Cheat Sheet (No Completed Job Yet) in the FCF Pre-Invoice Cancellation article. '
    'FC table: deactivate it. Advise the customer: depends on the cheat sheet outcome. '
    'CRM account: if C wants their account closed, deactivate the CRM account, but send comms before deactivating.')
prs.save(dst)
