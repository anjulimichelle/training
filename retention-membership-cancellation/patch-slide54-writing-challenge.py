"""Retention deck slide 54 (Unused DHJ Voucher, Case 1): replace the 'This is enough' quote with a
5-minute writing challenge; the sample CS response stays hidden until a click.
Patches the live Google export so Anjuli's edits elsewhere stay exactly as they are."""
import io, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, GOLD, GOLDSOFT, BG = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', '9A6410', 'FFF3D6', 'F6F7F9'
prs = Presentation(src)
s = prs.slides[53]
sh = list(s.shapes)
assert sh[1].text_frame.text.startswith('Case 1') and sh[20].text_frame.text.startswith('THIS IS ENOUGH')
for i in (18, 19, 20): sh[i]._element.getparent().remove(sh[i]._element)
for i in (21, 22, 23): sh[i].top = sh[i].top - Inches(0.93)   # tip bar moves up into the freed space

def strip_style(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
def rbox(x, y, w, h, fill, line=None, lw=0.75, name=None, adj=0.08, dash=False):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip_style(b)
    b.adjustments[0] = adj
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(lw)
        if dash:
            ln = b._element.spPr.find(qn('a:ln')); d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'}); ln.append(d)
    else: b.line.fill.background()
    if name: b.name = name
    return b
def text(shape_or_box, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT, after=3):
    tf = shape_or_box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(after)
        for t, size, bold, col in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
def tbox(x, y, w, h, name=None):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: t.name = name
    return t

Y = 3.43
# Header strip
hd = rbox(0.45, Y, 9.1, 0.32, TEAL)
text(hd, [[('✍  5-min writing challenge', 11, True, WHITE), ('   Read the customer’s message and write your reply.', 9.5, False, WHITE)]])
hd.text_frame.margin_left = Inches(0.18)
# Left: the customer's message
rbox(0.45, Y + 0.4, 4.45, 1.3, WHITE, BORDER)
text(tbox(0.65, Y + 0.48, 4.0, 0.2), [[('CUSTOMER MESSAGE', 8, True, SOFT)]])
bub = rbox(0.65, Y + 0.7, 4.05, 0.62, BG, BORDER, adj=0.15)
text(tbox(0.8, Y + 0.7, 3.75, 0.62), [[('“Can you make sure I’m scheduled for Nov 15, 1PM cleaning?”', 11, False, INK)]])
text(tbox(0.65, Y + 1.36, 4.05, 0.28), [[('Write your reply. Then compare it with the sample.', 9, False, SOFT)]])
# Right: hidden sample response with a cue
rbox(5.05, Y + 0.4, 4.5, 1.3, WHITE, BORDER, dash=True)
cue = rbox(6.05, Y + 0.86, 2.5, 0.38, TEAL, adj=0.3, name='cueButton')
text(cue, [[('▶  Show CS response', 10.5, True, WHITE)]], align=PP_ALIGN.CENTER)
resp = rbox(5.05, Y + 0.4, 4.5, 1.3, TEALSOFT, TEAL, name='step1Response')
rt = tbox(5.25, Y + 0.48, 4.1, 1.18, name='step1ResponseText')
text(rt, [[('SAMPLE CS RESPONSE', 8, True, TEAL)],
          [('I’ve got you! You’re all set for Nov 15 at 1PM.', 10.5, False, INK)],
          [('Just a quick reminder, your voucher comes with a membership which will start once it’s used.', 10.5, False, INK)]],
     anchor=MSO_ANCHOR.TOP, after=5)

n = s.notes_slide.notes_text_frame
n.text = n.text.rstrip() + (' 5-MINUTE WRITING CHALLENGE: the customer writes "Can you make sure I\'m scheduled for Nov 15, 1PM cleaning?" '
    'Give trainees five minutes to write their reply, read a few aloud, then click to show the sample CS response: '
    '"I\'ve got you! You\'re all set for Nov 15 at 1PM. Just a quick reminder, your voucher comes with a membership which will start once it\'s used." '
    'Point out that it answers their actual question first, then adds the one-line membership reminder.')
buf = io.BytesIO(); prs.save(buf)

ns = {}
exec(open(__file__.replace('patch54_challenge.py', 'add_story_anim.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == s.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        ids = re.findall(r'<p:cNvPr id="(\d+)" name="step1[^"]*"', x)
        cid, inner, bld = 3, '', []
        for i, sid in enumerate(ids):
            inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if i == 0 else 'withEffect'); cid += 3
            bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
        x = x.replace('</p:sld>', ns['timing']([ns['click'](cid, inner)], bld) + '</p:sld>')
        data = x.encode('utf8'); print('revealed on click:', len(ids))
    zout.writestr(item, data)
zout.close()
