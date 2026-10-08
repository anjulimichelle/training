"""Retention deck slide 56: replace the 'Which offer?' table with a two-scenario writing challenge.
Each sample CS response stays hidden until its click. Patches the live Google export."""
import io, re, sys, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, BG = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', 'F6F7F9'
prs = Presentation(src)
s = prs.slides[55]
sh = list(s.shapes)
assert sh[1].text_frame.text == 'Case 2: Which offer?' and sh[3].shape_type == 19
sh[3]._element.getparent().remove(sh[3]._element)
def retext(shape, t):
    r = shape.text_frame.paragraphs[0].runs; r[0].text = t
    for x in r[1:]: x._r.getparent().remove(x._r)
retext(sh[1], 'Case 2: Scenario challenge'); retext(sh[2], 'Read the customer’s message. How would you respond?')

def strip(x):
    st = x._element.find(P_NS + 'style')
    if st is not None: x._element.remove(st)
def box(x, y, w, h, fill, line=None, adj=0.06, name=None, dash=False):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
    b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
        if dash:
            ln = b._element.spPr.find(qn('a:ln')); ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))
    else: b.line.fill.background()
    if name: b.name = name
    return b
def text(shape, paras, size=9, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=4):
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
        for t, bold, col, *sz in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(sz[0] if sz else size); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
def tb(x, y, w, h, name=None):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: t.name = name
    return t

scen = [
    ('I don’t want to use your service. Give me my money back!',
     ['Sorry to hear that! Would you mind sharing why? I’d be happy to see if there’s anything I can do that might work better for you.',
      'If now just isn’t the right time, your voucher is good for a year, so you can always use it later. Just a quick reminder, your voucher includes a 6-month membership which will start once it’s used. If you’d rather go ahead with the refund, you can use this link: <refund link>',
      'Let me know what works best for you. I’m here to help!']),
    ('I can’t afford this membership. Please just cancel.',
     ['I completely understand. If the membership isn’t within your budget right now, I’d be happy to help.',
      'Before you decide, I can lower your monthly membership from $59 to $44 if that would make it more manageable. Your voucher is also valid for a year, so you can hold onto it and use it later if the timing works better. The 6-month membership won’t start until you use the voucher for your first cleaning.',
      'Let me know what works best for you. I’m happy to help either way!']),
]
w = (9.1 - 0.2) / 2
for i, (msg, resp) in enumerate(scen):
    x = 0.45 + i * (w + 0.2)
    text(tb(x, 1.4, w, 0.2), [[(f'SCENARIO {i + 1}', True, TEAL, 8)]])
    box(x, 1.62, w, 0.7, BG, BORDER, adj=0.12)
    text(tb(x + 0.15, 1.62, w - 0.3, 0.7), [[('CUSTOMER', True, SOFT, 7.5)], [('“' + msg + '”', False, INK, 10.5)]], anchor=MSO_ANCHOR.MIDDLE, after=2)
    box(x, 2.44, w, 2.7, WHITE, BORDER, dash=True)
    text(tb(x, 3.2, w, 0.25), [[('Write your reply first.', False, SOFT, 9)]], align=PP_ALIGN.CENTER)
    cue = box(x + (w - 2.3) / 2, 3.5, 2.3, 0.38, TEAL, adj=0.3)
    text(cue, [[('▶  Show CS response', True, WHITE, 10.5)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(x, 2.44, w, 2.7, TEALSOFT, TEAL, name=f'step{i + 1}Box')
    text(tb(x + 0.16, 2.52, w - 0.32, 2.56, name=f'step{i + 1}Text'),
         [[('SAMPLE CS RESPONSE', True, TEAL, 7.5)]] + [[(p, False, INK, 10)] for p in resp], after=6)

s.notes_slide.notes_text_frame.text = (
    'SCENARIO CHALLENGE (unused DHJ voucher, refund or cancel intent). For each scenario, read the customer\'s message aloud and give trainees a few minutes to write a reply. '
    'Read a few aloud, then click to reveal the sample CS response (click 1: scenario 1, click 2: scenario 2). '
    'Scenario 1, "I don\'t want to use your service. Give me my money back!" No reason given, so probe first; remind them the voucher is valid for a year and comes with a 6-month membership that starts once it\'s used; refund intent, so send the self-refund link. '
    'Sample response: "Sorry to hear that! Would you mind sharing why? I\'d be happy to see if there\'s anything I can do that might work better for you. If now just isn\'t the right time, your voucher is good for a year, so you can always use it later. Just a quick reminder, your voucher includes a 6-month membership which will start once it\'s used. If you\'d rather go ahead with the refund, you can use this link: <refund link>. Let me know what works best for you. I\'m here to help!" '
    'Scenario 2, "I can\'t afford this membership. Please just cancel." Membership-related (MF expensive), so offer a $10–$15 MF reduction off the original $59 (here $59 to $44); remind them the voucher is valid for a year and the membership only starts once it\'s used. '
    'Sample response: "I completely understand. If the membership isn\'t within your budget right now, I\'d be happy to help. Before you decide, I can lower your monthly membership from $59 to $44 if that would make it more manageable. Your voucher is also valid for a year, so you can hold onto it and use it later if the timing works better. The 6-month membership won\'t start until you use the voucher for your first cleaning. Let me know what works best for you. I\'m happy to help either way!"')
buf = io.BytesIO(); prs.save(buf)

ns = {}
exec(open(__file__.replace('patch56_scenarios.py', 'add_story_anim.py')).read().split('\nzin = ')[0].split('src, dst = sys.argv[1], sys.argv[2]')[1], ns)
zin = zipfile.ZipFile(buf); zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == s.part.partname.lstrip('/'):
        x = data.decode('utf8'); assert '<p:timing' not in x
        groups = {}
        for sid, n in re.findall(r'<p:cNvPr id="(\d+)" name="step(\d+)[^"]*"', x): groups.setdefault(int(n), []).append(sid)
        cid, clicks, bld = 3, [], []
        for n in sorted(groups):
            inner = ''
            for j, sid in enumerate(groups[n]):
                inner += ns['effect'](cid + 2, f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>', 'clickEffect' if j == 0 else 'withEffect'); cid += 3
                bld.append(f'<p:bldP spid="{sid}" grpId="0" animBg="1"/>')
            clicks.append(ns['click'](cid, inner)); cid += 2
        x = x.replace('</p:sld>', ns['timing'](clicks, bld) + '</p:sld>')
        data = x.encode('utf8'); print('clicks:', len(clicks))
    zout.writestr(item, data)
zout.close()
