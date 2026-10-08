"""Retention deck: insert a 'Guardrail: asked to pause again' slide right after slide 20.
Built on the live Google export so Anjuli's edits elsewhere stay exactly as they are."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

src, dst = sys.argv[1:3]
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE = 'B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF'
TEAL = '0B7A6F'; RED, PINK, GOLD, GOLDSOFT = 'B3261E', 'FCEBEA', '9A6410', 'FFF3D6'
prs = Presentation(src)
ref = prs.slides[19]
assert 'repeat pauses' in ref.shapes[1].text_frame.text
s = prs.slides.add_slide(ref.slide_layout)
for ph in list(s.placeholders):
    if ph.placeholder_format.type != 13 and 'Slide Number' not in ph.name: ph._element.getparent().remove(ph._element)

# Eyebrow, title, subtitle and slide number: copies of slide 20's, so fonts and positions match.
for i in (0, 1, 2, 28):
    s.shapes._spTree.append(copy.deepcopy(ref.shapes[i]._element))
def retext(shape, text):
    r = shape.text_frame.paragraphs[0].runs
    r[0].text = text
    for x in r[1:]: x._r.getparent().remove(x._r)
sh = list(s.shapes)
retext(sh[-4], 'OFFER LOGIC'); retext(sh[-3], 'Guardrail: asked to pause again')
retext(sh[-2], 'When a customer asks for a new pause right after a previous one.')

def box(x, y, w, h, fill, line=None):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    st = b._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: b._element.remove(st)
    b.adjustments[0] = 0.08
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
    b.line.color.rgb = RGBColor.from_string(line or fill); b.line.width = Pt(0.75)
    return b
def text(x, y, w, h, paras, anchor=MSO_ANCHOR.MIDDLE):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        for txt, size, bold, col in runs:
            r = p.add_run(); r.text = txt; r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
    return t

text(0.45, 1.42, 5.2, 0.2, [[('SITUATION', 8, True, SOFT)]])
text(5.8, 1.42, 3.75, 0.2, [[('WHAT TO DO', 8, True, SOFT)]])
rows = [
    (0.62, [[('Fewer than ', 10.5, False, INK), ('2 MFs', 10.5, True, INK), (' charged and paid since the last pause ended', 10.5, False, INK)]],
     RED, PINK, [[('\U0001F6A9  Support approval first', 10.5, True, RED)], [('No matter how long they’re asking for', 9, False, INK)]]),
    (1.0, [[('All three', 10.5, True, INK), (' are true:', 10.5, False, INK)],
           [('•  At least 2 MFs charged and paid since the last pause ended', 9.5, False, INK)],
           [('•  Actually ', 9.5, False, INK), ('using the service', 9.5, True, INK), (' since then (booked or completed cleanings)', 9.5, False, INK)],
           [('•  ', 9.5, False, INK), ('Not', 9.5, True, INK), (' also asking for a refund right now', 9.5, False, INK)]],
     TEAL, TEALSOFT, [[('✅  Give the pause normally', 10.5, True, TEAL)], [('Within the caps (up to 6 months)', 9, False, INK)]]),
    (0.62, [[('Any one', 10.5, True, INK), (' of those three is ', 10.5, False, INK), ('not', 10.5, True, INK), (' true', 10.5, False, INK)]],
     RED, PINK, [[('\U0001F6A9  Support approval first', 10.5, True, RED)], [('It could be abuse, not a real need for a pause', 9, False, INK)]]),
]
y = 1.66
for h, sit, col, soft, act in rows:
    box(0.45, y, 5.2, h, WHITE, BORDER)
    s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(y + 0.08), Inches(0.05), Inches(h - 0.16))
    bar = list(s.shapes)[-1]; st = bar._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: bar._element.remove(st)
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(col); bar.line.fill.background()
    text(0.65, y, 4.9, h, sit)
    box(5.8, y, 3.75, h, soft, col)
    text(5.98, y, 3.45, h, act)
    y += h + 0.1
box(0.45, y + 0.06, 9.1, 0.86, GOLDSOFT)
text(0.65, y + 0.06, 8.7, 0.86, [[('Why this matters: ', 9.5, True, GOLD), ('pausing again with no MF charged in between looks like avoiding the MF (and the ETF that comes with cancelling), not needing a break. ', 9.5, False, INK),
                                  ('Payment alone doesn’t prove it’s genuine. Payment + actual usage + no refund request together closes those gaps.', 9.5, False, INK)]])
print('bottom', round(y + 0.92, 2))

s.notes_slide.notes_text_frame.text = (
    'Guardrail: when a customer asks to pause again. This applies when a customer asks for a new pause right after a previous pause. '
    'Fewer than 2 MFs have been charged and paid since the last pause ended: needs Support approval before you give another pause, no matter how long they are asking for. '
    'All three are true (at least 2 MFs charged and paid since the last pause ended; the customer has actually been using the service, booked or completed cleanings, since then; '
    'and the customer is not also asking for a refund right now): you can give the pause normally, within the caps (pausing is capped at 6 months; more needs Support approval). '
    'Any one of those three is not true: needs Support approval; this could be a sign of abuse rather than a real need for a pause. '
    'Why this matters: if a customer pauses, then immediately pauses again with no MF ever charged in between, it looks more like they are trying to avoid paying the MF '
    '(and the ETF that comes with cancelling) than genuinely needing a break. Payment alone does not prove it is genuine either: someone could pay without using the service, '
    'or ask to pause right when they are also asking for a refund. Checking for payment + actual usage + no refund request together closes those gaps.')

# Move the new slide to sit right after slide 20.
lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(20, el)
prs.save(dst)
