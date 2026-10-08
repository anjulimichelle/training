"""Retention deck: insert 'Retention strategy' and 'Exit strategy' (FCF pre-invoice) right after slide 63.
Built on the live Google export so Anjuli's edits elsewhere stay exactly as they are."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

src, dst = sys.argv[1:3]
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TEAL, TEALSOFT, INK, SOFT, BORDER, WHITE, GOLD, GOLDSOFT = '0B7A6F', 'E3F4F1', '1D1D1F', '6E6E73', 'DCE1E7', 'FFFFFF', '9A6410', 'FFF3D6'
PEACH, PEACHSOFT, BLUE, BLUESOFT, PURPLE, PURPLESOFT, RED, PINK = 'C0612B', 'FCEADF', '2F6FA3', 'E3EEF7', '5E4A9C', 'EEEAF8', 'B3261E', 'FCEBEA'
prs = Presentation(src)
ref = prs.slides[62]
assert ref.shapes[1].text_frame.text == 'What is this?' and 'FCF' in ref.shapes[0].text_frame.text.upper()
cheat = [s for s in prs.slides if s.shapes[1].has_text_frame and s.shapes[1].text_frame.text.startswith('Refund cheat sheet: no service issue')][0]

def new_slide(title, sub):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    num = [x for x in ref.shapes if x.is_placeholder][0]
    for el in (ref.shapes[0]._element, ref.shapes[1]._element, ref.shapes[2]._element, num._element):
        s.shapes._spTree.append(copy.deepcopy(el))
    for shape, t in ((s.shapes[1], title), (s.shapes[2], sub)):
        r = shape.text_frame.paragraphs[0].runs; r[0].text = t
        for x in r[1:]: x._r.getparent().remove(x._r)
    return s

def kit(s):
    def strip(x):
        st = x._element.find(P_NS + 'style')
        if st is not None: x._element.remove(st)
    def box(x, y, w, h, fill, line=None, adj=0.06):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); strip(b)
        b.adjustments[0] = adj; b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line: b.line.color.rgb = RGBColor.from_string(line); b.line.width = Pt(0.75)
        else: b.line.fill.background()
        return b
    def text(shape, paras, size=9, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=2, margin=0.0):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = 0
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.space_after = Pt(after)
            for t, bold, col, *sz in runs:
                r = p.add_run(); r.text = t; r.font.size = Pt(sz[0] if sz else size); r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string(col); r.font.name = 'SF Pro Display' if bold else 'SF Pro Text'
        return shape
    def tb(x, y, w, h): return s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    def line(x1, y1, x2, y2):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); strip(c)
        c.line.color.rgb = RGBColor.from_string(SOFT); c.line.width = Pt(1)
        ln = c.line._get_or_add_ln(); ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    return box, text, tb, line

def badge(box, text, x, y, n, col):
    b = box(x, y, 0.28, 0.28, col, adj=0.5); text(b, [[(str(n), True, WHITE, 10)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---------- Retention strategy ----------
s1 = new_slide('Retention strategy', 'The customer asks to cancel before their first cleaning. Retention comes first.')
box, text, tb, line = kit(s1)
top = box(0.45, 1.38, 9.1, 0.36, PURPLESOFT, PURPLE)
text(top, [[('Customer requests cancellation:  ', True, PURPLE, 9.5), ('acknowledge  •  send the self-cancel link or offer to cancel for them  •  don’t say “voucher”', False, INK, 9)]],
     anchor=MSO_ANCHOR.MIDDLE, margin=0.15)
badge(box, text, 0.45, 1.82, 1, TEAL)
text(tb(0.82, 1.82, 8.7, 0.28), [[('Identify the root cause & address it', True, INK, 11), ('   (not provided? probe)', False, SOFT, 9.5)]], anchor=MSO_ANCHOR.MIDDLE)
tiles = [
    (PEACH, PEACHSOFT, 'Personal / service limits', 'Off-platform cleaner, moved, self-cleaning, no credit card, same-day cleaning needs',
     ['Address the concern', 'Highlight the value of the FC membership', 'Encourage booking the free or heavily discounted cleaning, if it fits']),
    (GOLD, GOLDSOFT, 'Service failure', 'Invalid LO, CP cancellation, CP LMC, no CP claim, false invoice, CP no-show',
     ['Fix it (block / coach / penalize the CP) AND resolve it', 'Offer priority booking for the next clean', 'No credits: the cleaning is already free or heavily discounted', 'Next MF within 2 weeks? Offer a free month']),
    (BLUE, BLUESOFT, 'Customer-fault failure', 'Valid LO submission, customer LMC',
     ['C still gets the free / heavily discounted cleaning from their original purchase', 'Refund any LO or LMC charge, whichever applies']),
    (PURPLE, PURPLESOFT, 'Membership terms', 'MF too expensive, MCT too long, unaware of membership, one-time cleaning only',
     ['MF too expensive: $10–$15 off the original $59', 'MCT too long: 1 month off the original 6-month MCT', 'Unaware: address it, then probe the service experience', 'One-time only: OTC at full price, upfront']),
]
w = (9.1 - 0.3) / 4
for i, (col, soft, title, ex, items) in enumerate(tiles):
    x = 0.45 + i * (w + 0.1)
    box(x, 2.16, w, 2.0, soft, col)
    h = box(x, 2.16, w, 0.32, col); text(h, [[(title, True, WHITE, 10)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(tb(x + 0.1, 2.54, w - 0.2, 1.58), [[('Ex: ' + ex, False, SOFT, 7.5)]] + [[('• ' + it, False, INK, 8)] for it in items], after=3)
note = box(0.45, 4.24, 9.1, 0.34, GOLDSOFT, GOLD)
text(note, [[('NOTE: ', True, GOLD, 9.5), ('multiple unused, paid monthly fees? Offer up to ', False, INK, 9.5), ('3 free months', True, INK, 9.5), (' on top of the standard retention offers.', False, INK, 9.5)]], anchor=MSO_ANCHOR.MIDDLE, margin=0.15)
badge(box, text, 0.45, 4.7, 2, TEAL)
st2 = box(0.82, 4.64, 8.73, 0.44, TEALSOFT, TEAL)
text(st2, [[('\U0001F4C5  Set a Timed Reminder (TR) ', True, TEAL, 10), ('at least 3 days before the next MF', True, INK, 10), (' to cancel FC if the customer doesn’t respond to the retention attempt.', False, INK, 9.5)]], anchor=MSO_ANCHOR.MIDDLE, margin=0.15)
s1.notes_slide.notes_text_frame.text = (
    'RETENTION STRATEGY (FCF, no completed cleaning yet). Customer requests cancellation: acknowledge the request; provide the self-cancellation link or offer to cancel on their behalf; do not use "voucher" terminology. '
    'Step 1, identify the root cause and address it (if not provided, probe). '
    'Personal / service limits (e.g. off-platform cleaner, C moved, self-cleaning, no credit card, same-day cleaning needs): address the concern appropriately, highlight the value of the FC membership, and encourage booking using the free or heavily discounted cleaning if appropriate. '
    'Service failure (e.g. invalid LO, CP cancellation, CP LMC, no CP claim, false invoice, CP no-show): address the service failure (block/coach/penalize the CP, etc.) AND provide the necessary resolution; offer priority booking for the next clean; do not offer credits since the cleaning is already free or heavily discounted; offer a free month if the next MF will be charged within 2 weeks. '
    'Customer-fault failure (e.g. valid LO submission, customer LMC): the customer should still get the free/heavily discounted cleaning based on their original purchase; refund any LO or LMC charges, whichever is applicable. '
    'Membership terms: MF too expensive, $10–$15 reduction from the original $59 MF; MCT too long, 1-month reduction from the original 6-month MCT; unaware of membership, address the unawareness and probe for the service experience (if applicable) and any other underlying reason for cancelling; one-time cleaning, offer a one-time cleaning at full price, charged upfront. '
    'NOTE: if the customer has multiple unused paid monthly fees, offer up to 3 free months on top of the standard retention offers. '
    'Step 2, set a Timed Reminder (TR) at least 3 days before the next monthly fee to cancel FC if the customer does not respond to the retention attempt.')

# ---------- Exit strategy ----------
s2 = new_slide('Exit strategy', 'Retention didn’t work. Close it out cleanly.')
box, text, tb, line = kit(s2)
wb = box(0.45, 1.42, 9.1, 0.8, PINK, RED)
text(wb, [[('WHEN', True, RED, 8)],
          [('The TR triggers and C hasn’t responded', True, INK, 11), ('   or   ', False, SOFT, 10), ('C responds and still insists on cancelling after the retention attempt', True, INK, 11)]],
     anchor=MSO_ANCHOR.MIDDLE, margin=0.2, after=2)
steps = [('Cancel FC via CRM', ''), ('Invalidate the voucher', ''), ('Invalidate the TR', 'if it hasn’t triggered yet')]
sw, gap = 2.8, 0.35
for i, (t, sub) in enumerate(steps):
    x = 0.45 + i * (sw + gap)
    b = box(x, 2.5, sw, 1.0, WHITE, BORDER)
    badge(box, text, x + 0.18, 2.86, i + 1, TEAL)
    text(tb(x + 0.6, 2.5, sw - 0.75, 1.0), [[(t, True, INK, 12.5)]] + ([[(sub, False, SOFT, 9.5)]] if sub else []), anchor=MSO_ANCHOR.MIDDLE)
    if i < 2: line(x + sw, 3.0, x + sw + gap, 3.0)
rb = box(0.45, 3.8, 9.1, 0.9, GOLDSOFT, GOLD)
text(tb(0.65, 3.8, 5.9, 0.9), [[('Wants a refund?', True, GOLD, 12)], [('After attempting retention, check the Refund eligibility cheat sheet.', False, INK, 10)]], anchor=MSO_ANCHOR.MIDDLE)
btn = box(6.65, 4.04, 2.7, 0.42, GOLD, adj=0.3)
text(btn, [[('▶  Refund eligibility cheat sheet', True, WHITE, 10)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
btn.click_action.target_slide = cheat
s2.notes_slide.notes_text_frame.text = (
    'EXIT STRATEGY (FCF, no completed cleaning yet). When the TR triggers and the customer has not responded, or the customer responds and insists on cancelling after the retention attempt: '
    '1) cancel FC via CRM; 2) invalidate the voucher; 3) invalidate the TR (if it has not triggered yet). '
    'Wants a refund? Refer to the Refund eligibility cheat sheet after attempting retention. TRAINER: in slideshow mode, click "Refund eligibility cheat sheet" to jump to it.')

lst = prs.slides._sldIdLst
a, b = lst[-2], lst[-1]
lst.remove(a); lst.remove(b); lst.insert(63, a); lst.insert(64, b)
prs.save(dst)
